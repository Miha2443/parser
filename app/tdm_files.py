"""Helpers for choosing dashboard datasets to send to TDM."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re
import warnings
import zipfile


TDM_FILE_SUFFIXES = {".xlsx", ".json", ".csv", ".pdf", ".png"}
TEMP_DOWNLOAD_SUFFIXES = {".crdownload", ".download", ".part", ".tmp"}


@dataclass(frozen=True)
class TdmFileOption:
    path: Path
    rel: str
    mtime: float
    size_bytes: int

    @property
    def label(self) -> str:
        mtime_text = datetime.fromtimestamp(self.mtime).strftime("%d.%m.%Y %H:%M")
        size_kb = self.size_bytes / 1024
        size_text = f"{size_kb:.0f} KB" if size_kb < 1024 else f"{size_kb / 1024:.1f} MB"
        return f"{self.rel}  ({mtime_text}, {size_text})"


@dataclass(frozen=True)
class TdmDatasetSpec:
    key: str
    section: str
    title: str
    patterns: tuple[str, ...]
    bundle_name: str


@dataclass(frozen=True)
class TdmDatasetOption:
    spec: TdmDatasetSpec
    files: tuple[Path, ...]

    @property
    def newest_mtime(self) -> float:
        return max((path.stat().st_mtime for path in self.files), default=0)


TDM_DATASETS: tuple[TdmDatasetSpec, ...] = (
    TdmDatasetSpec(
        "ipc", "Данные Мосстата / Росстата", "ИПЦ — все графики",
        ("downloads/*Индексы потребительских цен*часть1*.xls*",
         "downloads/*Индексы потребительских цен*часть2*.xls*"),
        "ИПЦ",
    ),
    TdmDatasetSpec(
        "salary", "Данные Мосстата / Росстата", "Заработная плата — все графики",
        ("downloads/*Среднемесячная номинальная начисленная заработная плата*экономике п.xls*",
         "downloads/*Среднемесячная номинальная начисленная заработная плата*экономике с.xls*"),
        "Заработная_плата",
    ),
    TdmDatasetSpec(
        "national_accounts", "Данные Мосстата / Росстата", "ВВП, ВРП и ВДС",
        ("downloads/*VVP_god*.xls*", "downloads/*VVP_na_dushu*.xls*",
         "downloads/*ВРП с 1998*.xls*", "downloads/*ВДС годы ОКВЭД2*2016*.xls*",
         "downloads/*VDS_god_OKVED2*.xls*", "downloads/*VRP_OKVED2007*.xls*"),
        "ВВП_ВРП_ВДС",
    ),
    TdmDatasetSpec(
        "current_construction", "Рынок недвижимости", "Текущее строительство — оперативные данные",
        ("data/raw/realty/nashdom/construction_operational_*.json",
         "data/raw/realty/nashdom/monitoring_2_0_*.xlsx"),
        "Текущее_строительство",
    ),
    TdmDatasetSpec(
        "kvartirografia", "Рынок недвижимости", "Квартирография",
        ("data/raw/realty/nashdom/kvartirografia_*",),
        "Квартирография",
    ),
    TdmDatasetSpec(
        "rasprodannost", "Рынок недвижимости", "Распроданность и стройготовность",
        ("data/raw/realty/nashdom/rasprodannost_*.xlsx",),
        "Распроданность_и_стройготовность",
    ),
    TdmDatasetSpec(
        "vvod", "Рынок недвижимости", "Ввод недвижимости — годовые данные",
        ("data/raw/realty/vvod/vvod.xlsx", "data/raw/realty/vvod/Stroi_111_*.xls*"),
        "Ввод_недвижимости",
    ),
    TdmDatasetSpec(
        "developer_construction", "Профиль застройщика", "Объём текущего строительства ЕРЗ",
        ("data/raw/realty/erzrf/top_obyem_stroitelstva_rf_*.xlsx",
         "data/raw/realty/erzrf/top_obyem_stroitelstva_msk_*.xlsx"),
        "ЕРЗ_текущее_строительство",
    ),
    TdmDatasetSpec(
        "developer_accumulated", "Профиль застройщика", "Накопленный ввод жилья ЕРЗ",
        ("data/raw/realty/erzrf/top_nakopl_vvod_rf_*.xlsx",
         "data/raw/realty/erzrf/top_nakopl_vvod_msk_*.xlsx"),
        "ЕРЗ_накопленный_ввод",
    ),
    TdmDatasetSpec(
        "escrow", "Профиль застройщика", "Наполненность счетов эскроу",
        ("data/raw/realty/escrow_manual/*.xlsx",),
        "Счета_эскроу",
    ),
)


def _latest_match(root: Path, pattern: str) -> Path | None:
    candidates = [
        path for path in root.glob(pattern)
        if path.is_file()
        and "_archive" not in path.parts
        and path.suffix.lower() not in TEMP_DOWNLOAD_SUFFIXES
    ]
    return max(candidates, key=lambda path: path.stat().st_mtime, default=None)


def _is_confirmed_accumulated_erz(path: Path) -> bool:
    """Reject old ERZ files that were another rating under this filename."""
    try:
        from openpyxl import load_workbook

        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            headers = {
                str(value).replace("\xa0", " ").strip()
                for row in sheet.iter_rows(min_row=1, max_row=3, values_only=True)
                for value in row if value is not None
            }
            return "Введено, м²" in headers
        finally:
            workbook.close()
    except Exception:  # noqa: BLE001
        return False


def list_tdm_datasets(root: Path) -> list[TdmDatasetOption]:
    """Resolve friendly dashboard choices to the newest source file(s)."""
    options: list[TdmDatasetOption] = []
    for spec in TDM_DATASETS:
        files: list[Path] = []
        for pattern in spec.patterns:
            match = _latest_match(root, pattern)
            if match is not None and match not in files:
                files.append(match)
        if len(files) != len(spec.patterns):
            continue
        if spec.key == "developer_accumulated" and not all(
            _is_confirmed_accumulated_erz(path) for path in files
        ):
            continue
        if files:
            options.append(TdmDatasetOption(spec=spec, files=tuple(files)))
    return options


def prepare_tdm_dataset(option: TdmDatasetOption, temp_dir: Path) -> tuple[Path, bool]:
    """Return a sendable file; combine multi-source datasets into one ZIP."""
    if len(option.files) == 1:
        return option.files[0], False
    temp_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^0-9A-Za-zА-Яа-яЁё_-]+", "_", option.spec.bundle_name).strip("_")
    target = temp_dir / f"{safe_name}_{datetime.now():%Y%m%d}.zip"
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in option.files:
            archive.write(path, arcname=path.name)
    return target, True


def list_tdm_realty_files(realty_root: Path) -> list[TdmFileOption]:
    """Return sendable active realty files sorted newest first."""
    if not realty_root.exists():
        return []

    options: list[TdmFileOption] = []
    for path in realty_root.rglob("*"):
        if not path.is_file() or "_archive" in path.parts:
            continue
        suffix = path.suffix.lower()
        if suffix in TEMP_DOWNLOAD_SUFFIXES or suffix not in TDM_FILE_SUFFIXES:
            continue
        try:
            stat = path.stat()
            rel = str(path.relative_to(realty_root))
        except OSError:
            continue
        options.append(
            TdmFileOption(
                path=path,
                rel=rel,
                mtime=stat.st_mtime,
                size_bytes=stat.st_size,
            )
        )
    options.sort(key=lambda item: item.mtime, reverse=True)
    return options
