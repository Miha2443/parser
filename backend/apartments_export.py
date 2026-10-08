"""XLSX of the apartment API payload: full rows, original units and provenance."""
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_workbook(payload):
    book = Workbook()
    book.remove(book.active)

    def safe(value):
        if isinstance(value, (list, dict)):
            value = json.dumps(value, ensure_ascii=False, allow_nan=False)
        return "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value

    def sheet(name, headers, rows):
        tab = book.create_sheet(name)
        tab.append(headers)
        for row in rows:
            tab.append([safe(value) for value in row])
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1C3444")
        for row in tab.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
        for index, header in enumerate(headers, 1):
            tab.column_dimensions[get_column_letter(index)].width = min(55, max(18, len(header) + 2))

    def table(name, rows):
        sheet(name, ["Наименование", "Квартиры, тыс. шт.", "Площадь, тыс. м²", "Место",
                     "1 комн, %", "2 комн, %", "3 комн, %", "4+ комн, %"],
              [[row["name"], row["apartmentThousandCount"], row["areaThousandM2"], row.get("place"),
                *[room["sharePercent"] for room in row["rooms"]]] for row in rows])

    sheet("Отчёт", ["Показатель", "Значение"], [
        ["Регион", payload["region"]], ["Дата отчёта", payload["reportDate"]],
        ["Сформировано (UTC)", payload["generatedAt"]], ["Версия источников", payload["version"]],
        ["Версия схемы", payload["schemaVersion"]],
    ])
    if "developer" in payload:
        table("Девелопер", [payload["developer"]])
        summary = payload["summary"]
        labels = [("countThousand", "Квартиры, тыс. шт."), ("areaThousandM2", "Площадь, тыс. м²"),
                  ("averageAreaM2", "Средняя площадь квартиры, м²"), ("marketSharePercent", "Доля рынка, %"),
                  ("marketBaseAreaThousandM2", "База рынка: сумма площадей девелоперов, тыс. м²"),
                  ("place", "Место"), ("totalDevelopers", "Всего девелоперов")]
        sheet("Показатели", ["Показатель", "Значение"], [[label, summary[key]] for key, label in labels])
        sheet("Средние по регионам", ["Регион", "Средняя площадь квартиры, м²"],
              [[row["region"], row["averageAreaM2"]] for row in payload["referenceAverages"]])
        sheet("Комнатность", ["Тип", "Доля, %"], [[row["type"], row["sharePercent"]] for row in payload["rooms"]])
        table("Сравнение", payload["comparison"])
    else:
        sheet("Типы квартир", ["Тип", "Количество квартир, шт.", "Площадь, тыс. м²"],
              [[row["type"], row["count"], row["areaThousandM2"]] for row in payload["apartments"]])
        sheet("Распределение", ["Площадь, м²", "Доля, %"],
              [[row["range"], row["sharePercent"]] for row in payload["distribution"]])
        table("Девелоперы", payload["developers"])
        table("Регионы", payload["regions"])
    source = payload["source"]
    sheet("Источники", ["Показатель", "Значение"], [
        ["Дата источников", source["date"]], ["Файлы-кандидаты", source["files"]],
        ["Примечание", "Кандидаты реестра DataAccess, не доказательство участия; выбор mart/raw по существующим правилам."],
        ["Диагностика загрузчика", source["issues"]],
    ])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
