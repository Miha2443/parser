import copy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from unittest.mock import MagicMock

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import create_app
from backend.profile_export import export_workbook
from backend.profile_service import ProfileService
from pipeline.data_access import DataContext

ROOT = Path(__file__).resolve().parents[2]


class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.service = ProfileService(DataContext(ROOT), check_interval=0)
        cls.baseline = json.loads((ROOT / "frontend/public/profile-snapshot.json").read_text(encoding="utf-8"))
        cls.client = TestClient(create_app(cls.service))

    def test_catalog_not_limited_to_demo_and_metadata(self):
        response = self.client.get("/api/v1/catalog")
        self.assertEqual(response.status_code, 200)
        catalog = response.json()
        self.assertGreater(len(catalog["controls"]["developers"]), 3)
        self.assertFalse(catalog["controls"]["frozen"])
        self.assertEqual(len(catalog["version"]), 64)
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertNotIn("profiles", catalog)

    def test_all_five_baseline_profiles_exact_values(self):
        for expected in self.baseline["profiles"]:
            with self.subTest(profile=expected["id"]):
                response = self.client.get("/api/v1/profile", params={"developer": expected["developerKey"], "region": expected["region"]})
                self.assertEqual(response.status_code, 200)
                actual = response.json()["profiles"][0]
                # Selector display spelling can be canonicalized; calculations cannot.
                actual["developer"] = expected["developer"]
                actual["controls"]["developer"] = expected["controls"]["developer"]
                self.assertEqual(actual, expected)

    def test_invalid_filters(self):
        for params, status in [({}, 422), ({"developer": "пик", "region": "oops"}, 422),
                               ({"developer": "nonexistent", "region": "msk"}, 404),
                               ({"developer": "capital group", "region": "msk"}, 404)]:
            with self.subTest(params=params):
                self.assertEqual(self.client.get("/api/v1/profile", params=params).status_code, status)

    def test_excel_matches_registries_and_nulls(self):
        payload = self.service.profile("пик", "msk")
        response = self.client.get("/api/v1/profile/export", params={"developer": "пик", "region": "msk"})
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content), data_only=False)
        for key, name in (("commissioned", "Реестр РВ"), ("permitted", "Реестр ОКС")):
            table = payload["profiles"][0]["objects"][key]
            self.assertEqual(book[name].max_row - 1, table["rowCount"])
            self.assertEqual([c.value for c in book[name][1]], table["columns"])
        self.assertEqual(book["Реестр РВ"].freeze_panes, "A2")
        self.assertEqual(response.headers["x-data-version"], payload["version"])

    def test_excel_source_text_is_not_formula(self):
        payload = copy.deepcopy(self.service.profile("пик", "msk"))
        payload["profiles"][0]["developer"] = '=HYPERLINK("evil")'
        book = load_workbook(BytesIO(export_workbook(payload)), data_only=False)
        self.assertEqual(book["Профиль"]["B2"].data_type, "s")

    def test_export_rejects_changed_generation(self):
        response = self.client.get("/api/v1/profile/export", params={"developer": "пик", "region": "msk", "required_version": "old-version"})
        self.assertEqual(response.status_code, 409)
        current = self.service.profile("пик", "msk")["version"]
        response = self.client.get("/api/v1/profile/export", params={"developer": "пик", "region": "msk", "required_version": current})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["x-data-version"], current)

    def test_schema_key_error_is_source_failure_not_filter_failure(self):
        broken = MagicMock()
        broken.profile.side_effect = KeyError("rv")
        response = TestClient(create_app(broken)).get("/api/v1/profile", params={"developer": "пик", "region": "msk"})
        self.assertEqual(response.status_code, 503)

    def test_reported_loader_error_cannot_become_empty_observation(self):
        broken = MagicMock()
        broken.issues = [("error", "broken workbook")]
        for name in self.service._data:
            getattr(broken, f"load_{name}").return_value = self.service._data[name]
        service = ProfileService(DataContext(ROOT), check_interval=0)
        with patch("backend.profile_service.DataAccess", return_value=broken), patch.object(service, "_versions", return_value=("test",)):
            client = TestClient(create_app(service))
            self.assertEqual(client.get("/api/v1/catalog").status_code, 503)
        self.assertIsNone(service._metadata)

    def test_sources_changed_during_load_are_not_published(self):
        source = MagicMock()
        source.issues = []
        for name in self.service._data:
            getattr(source, f"load_{name}").return_value = self.service._data[name]
        service = ProfileService(DataContext(ROOT), check_interval=0)
        with patch("backend.profile_service.DataAccess", return_value=source), patch.object(service, "_versions", side_effect=[("before",), ("after",)]):
            self.assertEqual(TestClient(create_app(service)).get("/api/v1/catalog").status_code, 503)
        self.assertIsNone(service._metadata)

    def test_missing_data_is_503_not_demo(self):
        with tempfile.TemporaryDirectory() as directory:
            service = ProfileService(DataContext(Path(directory)))
            client = TestClient(create_app(service))
            self.assertEqual(client.get("/api/v1/health").status_code, 503)
            self.assertEqual(client.get("/api/v1/catalog").status_code, 503)

    def test_generation_changes_invalidate_profile_cache(self):
        service = ProfileService(DataContext(ROOT), check_interval=0)
        initial = service.profile("пик", "msk")
        signature = service._signature
        original = service._versions
        with patch.object(service, "_versions", side_effect=lambda access: ("changed", original(access))):
            updated = service.profile("пик", "msk")
        self.assertNotEqual(initial["version"], updated["version"])
        self.assertNotEqual(service._signature, signature)
        self.assertEqual(initial["profiles"], updated["profiles"])


if __name__ == "__main__":
    unittest.main()
