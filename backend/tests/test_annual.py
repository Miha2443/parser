"""Real-source parity with execution of all original annual charts and exports."""
import ast
import copy
import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest

from fastapi import FastAPI
from fastapi.testclient import TestClient
import pandas as pd
import plotly.graph_objects as go
from openpyxl import load_workbook

from backend.annual_routes import LOADERS, catalog, create_router, create_service, export_workbook, report
from pipeline.annual_calculations import (
    _housing_composition, _moscow_housing_composition, annual_charts, annual_fact_period, annual_forecast,
)
from pipeline.data_access import DataAccess, DataContext, monitoring_by_year, month_label, quarter_label
from pipeline.report_values import records
from scripts.check_apartments_parity import Capture

ROOT = Path(__file__).resolve().parents[2]


class StopPage(Exception):
    pass


class AnnualCapture(Capture):
    def stop(self):
        raise StopPage()


class AnnualTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.context = DataContext(ROOT)
        access = DataAccess(cls.context)
        cls.before = tuple(access.data_version(name) for name in LOADERS)
        cls.data = {name: getattr(access, name)() for name in LOADERS}

    def capture(self, region):
        path = ROOT / "app/pages/8_Ввод_недвижимости.py"
        tree = ast.parse(path.read_text(encoding="utf-8"))
        tree.body = [node for node in tree.body if not isinstance(node, (ast.Import, ast.ImportFrom))]
        ui = AnnualCapture(region)
        env = {"st": ui, "pd": pd, "go": go,
               "COLORS": {key: "#333333" for key in ("red", "green", "blue", "amber", "teal", "ink")},
               "apply_theme": lambda: None, "page_header": lambda *args: None,
               "style_plotly": lambda *args, **kwargs: None, "chart_download_button": lambda *args, **kwargs: None,
               "table_download_buttons": lambda frame, **kwargs: ui.exports.append(frame.copy()),
               "monitoring_by_year": monitoring_by_year, "month_label": month_label, "quarter_label": quarter_label,
               **{name: lambda value=value: value for name, value in self.data.items()}}
        try:
            exec(compile(tree, str(path), "exec"), env)
        except StopPage:
            pass
        return ui

    def test_all_legacy_charts_export_rows_and_summary_values(self):
        for region, count in (("Москва", 7), ("РФ", 2)):
            with self.subTest(region=region):
                ui = self.capture(region)
                charts = annual_charts(self.data, region)
                self.assertEqual(len(charts), count)
                self.assertEqual(len(ui.figures), count)
                self.assertEqual(len(ui.exports), count)
                for chart, figure, frame in zip(charts, ui.figures, ui.exports):
                    # Moscow's agreed classification and RF's source taxonomy
                    # have separate category/null/provenance tests.
                    if chart["id"] == "b2_1":
                        self.assertEqual(chart["controlTotal"], "TOTAL")
                        self.assertNotIn("TOTAL", [series["id"] for series in chart["series"]])
                        self.assertTrue(chart["rawCategoryValues"])
                        continue
                    bars = [trace for trace in figure.data if trace.type == "bar"]
                    self.assertEqual([row["name"] for row in chart["series"]], [trace.name for trace in bars])
                    for series, trace in zip(chart["series"], bars):
                        self.assertEqual([p["x"] for p in series["points"]], list(trace.x))
                        expected = [None if region == "РФ" and point["x"] == "2026" else value
                                    for point, value in zip(series["points"], trace.y)]
                        self.assertEqual([p["y"] for p in series["points"]], expected)
                    adapted = [{"year": row["year"], **{col["label"]: row[col["id"]] for col in chart["columns"][1:]}}
                               for row in chart["rows"]]
                    expected_rows = records(frame)
                    if region == "РФ":
                        expected_rows[-1] = {key: value if key == "year" else None for key, value in expected_rows[-1].items()}
                    self.assertEqual(adapted, expected_rows)
                    for summary in chart["summaries"]:
                        subset = frame[frame["year"].between(summary["from"], summary["to"])]
                        for value in summary["values"]:
                            if region == "РФ" and summary["to"] == 2026:
                                self.assertIsNone(value["value"])
                            else:
                                self.assertAlmostEqual(value["value"], float(subset[value["label"]].sum()), places=9)
        access = DataAccess(self.context)
        self.assertEqual(self.before, tuple(access.data_version(name) for name in LOADERS))

    def test_full_workbook_notes_units_and_region_scope(self):
        service = create_service(self.context, check_interval=0)
        msk, rf = service.report(region="msk"), service.report(region="rf")
        self.assertEqual(len(msk["charts"]), 7)
        self.assertEqual(len(rf["charts"]), 2)
        self.assertTrue(any("нет годового итога" in note for note in rf["notes"]))
        book = load_workbook(BytesIO(export_workbook(msk)))
        self.assertEqual(len(book.sheetnames), 15 + sum(bool(chart.get("forecast")) for chart in msk["charts"]))
        self.assertNotIn("forecastSource", rf)
        self.assertTrue(all("forecast" not in chart for chart in rf["charts"]))
        for chart in msk["charts"]:
            self.assertEqual(chart["unit"], "млн м²")
            self.assertEqual(book[chart["id"]].max_row, len(chart["rows"]) + 1)

    def test_real_monitoring_annual_ihs_2011_2014_reaches_chart_and_provenance(self):
        history = self.data["load_monitoring_annual_ihs"]
        if history.empty:
            self.skipTest("Saved monitoring has no annual IHS block")
        chart = next(chart for chart in annual_charts(self.data, "Москва") if chart["id"] == "b2_1")
        rows = {row["year"]: row for row in chart["rows"]}
        provenance = {row["year"]: row for row in chart["categoryProvenance"]}
        raw = {row["year"]: row for row in chart["rawCategoryValues"]}
        self.assertEqual([series["id"] for series in chart["series"]], ["МКД", "ИЖС"])
        self.assertNotIn("Остальное_жильё", [column["id"] for column in chart["columns"]])
        for year, cell, total, value, mkd in ((2011, "S24", 2.2, .0258, 2.1742),
                                            (2012, "S25", 3.05, .2218, 2.8282),
                                            (2013, "S26", 3.146, .3124, 2.8336),
                                            (2014, "S27", 3.342, .4974, 2.8446)):
            self.assertAlmostEqual(rows[year]["ИЖС"], value)
            self.assertAlmostEqual(rows[year]["TOTAL"], total)
            self.assertAlmostEqual(rows[year]["МКД"], mkd)
            self.assertIsNone(raw[year]["МКД"])
            self.assertEqual(provenance[year]["ИЖС_source_cell"], cell)
            self.assertEqual(provenance[year]["ИЖС_source_sheet"], "Данные с 2011 года")
        # The old positive grey remainder is now included in Moscow's MKD.
        self.assertAlmostEqual(raw[2015]["МКД"], 3.2591)
        self.assertAlmostEqual(rows[2015]["МКД"], 3.3455)
        for row in chart["rows"]:
            with self.subTest(year=row["year"]):
                self.assertEqual(row["TOTAL"], raw[row["year"]]["TOTAL"])
                self.assertEqual(row["ИЖС"], raw[row["year"]]["ИЖС"])
                if row["TOTAL"] is not None and row["ИЖС"] is not None and row["TOTAL"] >= row["ИЖС"]:
                    self.assertAlmostEqual(row["TOTAL"], row["МКД"] + row["ИЖС"], places=12)
                    proof = provenance[row["year"]]
                    self.assertEqual(proof["МКД_source_kind"], "user_agreed_calculation")
                    self.assertEqual(proof["МКД_source_formula"], "МКД = TOTAL - ИЖС")
                    self.assertIn("Пользовательское согласование", proof["МКД_source_label"])
                    for suffix in ("file", "sheet", "cell", "row", "rows", "value"):
                        self.assertIsNone(proof.get("МКД_source_" + suffix))
        payload = dict(version="test", generatedAt="test", source={}, **report(self.data, catalog(self.data)))
        book = load_workbook(BytesIO(export_workbook(payload)))
        for excel_row, row in zip(list(book["b2_1"].values)[1:], chart["rows"]):
            for actual, column in zip(excel_row, chart["columns"]):
                expected = row[column["id"]]
                if expected is None:
                    self.assertIsNone(actual)
                else:
                    self.assertAlmostEqual(actual, expected, places=12)
        metadata = dict(book["Report"].values)
        for field in ("rawCategoryValues", "rawCategoryProvenance", "categoryProvenance"):
            self.assertEqual(json.loads(metadata["b2_1/" + field]), chart[field])


class MoscowHousingCompositionTests(unittest.TestCase):
    def test_agreed_difference_includes_old_grey_without_claiming_a_source_cell(self):
        original = pd.DataFrame([
            {"year": 2011, "TOTAL": 3, "МКД": None, "ИЖС": .25},
            {"year": 2015, "TOTAL": 4, "МКД": 3, "ИЖС": .5},
            {"year": 2020, "TOTAL": 3, "МКД": 4, "ИЖС": .5},
            {"year": 2024, "TOTAL": 3, "МКД": None, "ИЖС": 0},
            {"year": 2025, "TOTAL": 0, "МКД": None, "ИЖС": 0},
        ])
        original["МКД_source_file"] = "old.xlsx"
        original["МКД_source_cell"] = "B12"
        original["МКД_source_rows"] = [[12]] * len(original)
        original["МКД_source_value"] = 3000
        original["TOTAL_source_precision_million"] = .001
        original["ИЖС_source_precision_million"] = .0001
        snapshot = original.copy(deep=True)
        composition = _housing_composition(original, confirmed_ihs_years=(2011,))
        self.assertAlmostEqual(composition.loc[1, "Остальное_жильё"], .5)
        shown = _moscow_housing_composition(composition)
        self.assertEqual(shown["МКД"].tolist(), [2.75, 3.5, 2.5, 3, 0])
        pd.testing.assert_series_equal(shown["TOTAL"], original["TOTAL"])
        pd.testing.assert_series_equal(shown["ИЖС"], original["ИЖС"])
        self.assertTrue(shown["Остальное_жильё"].isna().all())
        for column in ("file", "cell", "rows", "value"):
            self.assertTrue(shown["МКД_source_" + column].isna().all())
        self.assertTrue(shown["МКД_source_kind"].eq("user_agreed_calculation").all())
        self.assertTrue(shown["МКД_source_formula"].eq("МКД = TOTAL - ИЖС").all())
        self.assertTrue(shown["МКД_source_precision_million"].eq(.0011).all())
        pd.testing.assert_frame_equal(original, snapshot)

    def test_missing_or_conflicting_controls_never_imply_zero_or_a_confirmed_remainder(self):
        original = pd.DataFrame([
            {"year": 2015, "TOTAL": None, "МКД": 2, "ИЖС": 1},
            {"year": 2016, "TOTAL": 3, "МКД": 2, "ИЖС": None},
            {"year": 2017, "TOTAL": 3, "МКД": None, "ИЖС": None},
            {"year": 2018, "TOTAL": None, "МКД": None, "ИЖС": .5},
            {"year": 2019, "TOTAL": .5, "МКД": 2, "ИЖС": 1},
            {"year": 2020, "TOTAL": 3, "МКД": 2, "ИЖС": -1},
            {"year": 2021, "TOTAL": float("inf"), "МКД": 2, "ИЖС": 1},
            {"year": 2022, "TOTAL": 3, "МКД": 2, "ИЖС": float("inf")},
            {"year": 2023, "TOTAL": -1, "МКД": 2, "ИЖС": 0},
            {"year": 2024, "TOTAL": None, "МКД": -1, "ИЖС": None},
            {"year": 2025, "TOTAL": 1, "МКД": 2, "ИЖС": None},
        ])
        original["МКД_source_cell"] = "B12"
        composition = _housing_composition(original)
        shown = _moscow_housing_composition(composition)
        self.assertEqual(shown.loc[:1, "МКД"].tolist(), [2, 2])
        self.assertEqual(shown.loc[:1, "МКД_source_cell"].tolist(), ["B12", "B12"])
        self.assertTrue(shown.loc[2:, "МКД"].isna().all())
        self.assertTrue(shown.loc[2:, "МКД_source_cell"].isna().all())
        self.assertTrue(shown.loc[4:8, "МКД_source_kind"].eq("unresolved_conflict").all())
        self.assertEqual(shown.loc[10, "МКД_source_kind"], "unresolved_conflict")
        self.assertTrue(shown["МКД_source_formula"].isna().all())
        self.assertTrue(shown["МКД_source_reason"].notna().all())
        self.assertTrue(shown["Остальное_жильё"].isna().all())
        pd.testing.assert_series_equal(shown["TOTAL"], original["TOTAL"])
        pd.testing.assert_series_equal(shown["ИЖС"], original["ИЖС"])


class AnnualForecastTests(unittest.TestCase):
    def setUp(self):
        years = list(range(2011, 2026))
        housing = pd.DataFrame({"year": years, "жильё": [3.0, None] + [4.0] * 13, "нежильё": [2.0] * 15})
        nonres = pd.DataFrame({"year": years, "нежильё": [2.0] * 15, "нежилые_в_жилье": [.5] * 15})
        budget = pd.DataFrame({"year": years, "Бюджет": [1.0] * 15, "Небюджет": [3.0] * 15})
        rv = pd.DataFrame([{"Год ввода по Мосстату": year, "Месяц ввода по Мосстату": month,
                            "Группировка": "Жилье", "Отрасли": "Жилые объекты", "Подтип объекта": "МКД",
                            "Источник финансирования": "Внебюджет", "Инвестор для Реновации": "да",
                            "Жилая площадь": 1_000_000, "category_жилое": 1_000_000,
                            "category_моп": 500_000, "category_нежилое_отдельное": 200_000,
                            "category_нежилое_в_жилом": 100_000}
                           for year, month in [(2025, "декабрь"), (2026, "сентябрь"), (2027, "декабрь")]])
        self.data = {"load_vvod_static": {"msk_total": housing, "rf_total": housing.copy(),
                      "msk_residential": pd.DataFrame({"year": years, "МКД": [3.0] * 15, "ИЖС": [1.0] * 15}),
                      "msk_nonres": nonres, "msk_nonres_branches": pd.DataFrame({"year": years, "офисы": [1.0] * 15})},
                     "load_emiss_34118": pd.DataFrame(columns=["year"]), "load_emiss_34118_periods": pd.DataFrame(),
                     "load_monitoring_2_0": {"rv": rv},
                     "load_monitoring_2011_2026_static": {"residential_budget_split": budget, "nonres_budget_split": budget.copy()},
                     "load_commissioning_forecasts": {"source_file": "monitoring_2_0_20261021.xlsx", "source_date": "2026-10-21",
                         "unitBasis": "assumed_thousand_m2", "notes": ["Площадь принята в тыс. м²."], "rows": [
                             {"year": 2025, "housing_total": 999999},
                             {"year": 2026, "housing_living": 6000, "housing_total": 10000, "renovation_living": 2500,
                              "nonres_standalone": 4000, "nonres_embedded": 2000, "nonres_total": 6000, "offices": 123456},
                             {"year": 2027, "housing_living": 0, "housing_total": None, "nonres_standalone": 0},
                             {"year": 2029, "roads": 100},
                             {"year": 2031, "housing_living": 8000, "housing_total": 12000, "nonres_standalone": 5000}]}}

    def test_only_matching_moscow_totals_convert_units_and_preserve_all_facts(self):
        before = annual_charts(self.data, "Москва")
        source = copy.deepcopy(self.data["load_commissioning_forecasts"])
        payload = report(self.data, catalog(self.data))
        expected = {"b1": 14.0, "b2_1": 6.0, "b2_2": 6.0, "b2_3": 2.5, "b3_1": 6.0, "b3_2": 4.0, "b3_3": 4.0}
        for original, chart in zip(before, payload["charts"]):
            self.assertEqual({key: value for key, value in chart.items() if key not in ("factPeriod", "forecast")}, original)
            self.assertEqual(chart["factPeriod"], {"year": 2026, "month": 9})
            self.assertEqual(chart["forecast"]["points"][0], {"year": 2026, "value": expected[chart["id"]]})
            self.assertNotIn("offices", chart["forecast"]["sourceFields"])
            self.assertNotIn(2029, [point["year"] for point in chart["forecast"]["points"]])
            self.assertNotIn(2025, [point["year"] for point in chart["forecast"]["points"]])
        charts = {chart["id"]: chart for chart in payload["charts"]}
        self.assertEqual([point["year"] for point in charts["b1"]["forecast"]["points"]], [2026, 2031])
        self.assertEqual(charts["b2_1"]["forecast"]["points"][1], {"year": 2027, "value": 0.0})
        self.assertIsNone(charts["b1"]["summaries"][0]["values"][0]["value"])
        self.assertEqual(payload["forecastSource"]["unitBasis"], "assumed_thousand_m2")
        self.assertEqual(self.data["load_commissioning_forecasts"], source)
        rf = report(self.data, catalog(self.data), region="rf")
        self.assertNotIn("forecastSource", rf)
        self.assertTrue(all("forecast" not in chart for chart in rf["charts"]))
        self.assertIsNone(rf["charts"][0]["rows"][-1]["жильё"])

    def test_no_interpolation_decomposition_or_unmatched_only_bars(self):
        for rows in ([], [{"year": 2026, "roads": 5}], [{"year": 2026, "housing_total": 10000}],
                     [{"year": 2026, "housing_total": 10000, "nonres_standalone": None}]):
            self.data["load_commissioning_forecasts"]["rows"] = rows
            payload = report(self.data, catalog(self.data))
            self.assertTrue(all("forecast" not in chart for chart in payload["charts"]))
        self.assertIsNone(annual_forecast("unmatched", {"rows": [{"year": 2026, "housing_living": 4000}]}, 2026))
        for value in (None, float("nan"), float("inf"), -10, True):
            self.assertIsNone(annual_forecast("b2_1", {"rows": [{"year": 2026, "housing_living": value}]}, 2026))

    def test_fact_period_uses_rv_year_month_not_candidate_or_forecast_dates(self):
        self.assertEqual(annual_fact_period(self.data, 2026, "Москва"), {"year": 2026, "month": 9})
        self.assertEqual(annual_fact_period(self.data, 2026, "РФ"), {"year": 2026, "month": None})
        self.data["load_monitoring_2_0"]["rv"]["Месяц ввода по Мосстату"] = self.data["load_monitoring_2_0"]["rv"]["Месяц ввода по Мосстату"].astype(object)
        self.data["load_monitoring_2_0"]["rv"].loc[1, "Месяц ввода по Мосстату"] = 8
        self.assertEqual(annual_fact_period(self.data, 2026, "Москва")["month"], 8)
        self.data["load_monitoring_2_0"]["rv"].loc[1, "Месяц ввода по Мосстату"] = "неизвестно"
        self.assertIsNone(annual_fact_period(self.data, 2026, "Москва")["month"])

    def test_excel_keeps_forecasts_out_of_fact_and_summary_sheets(self):
        payload = dict(version="test", generatedAt="test", source={}, **report(self.data, catalog(self.data)))
        book = load_workbook(BytesIO(export_workbook(payload)))
        for chart in payload["charts"]:
            rows = list(book[chart["id"]].values)
            self.assertEqual(rows[1:], [tuple(row[column["id"]] for column in chart["columns"]) for row in chart["rows"]])
            forecast_rows = list(book[chart["id"] + "_forecast"].values)
            self.assertEqual([row[-1] for row in forecast_rows[1:]], [point["value"] for point in chart["forecast"]["points"]])
            self.assertTrue(all(row[1] == "Прогноз" for row in forecast_rows[1:]))
        self.assertIn("load_commissioning_forecasts", LOADERS)
        self.assertIn("commissioning_forecasts", create_service(DataContext(ROOT)).families)

    def test_annual_ihs_is_moscow_only_with_missing_years_and_preserved_source_categories(self):
        rf_before = annual_charts(self.data, "РФ")
        previous = self.data["load_vvod_static"]["msk_residential"]
        previous["ИЖС_source_file"] = "old.xlsx"
        previous["ИЖС_source_rows"] = [[999]] * len(previous)
        previous["ИЖС_сельское"] = .01
        history = pd.DataFrame([
            {"region": "Москва", "year": year, "month": 12, "period_type": "year", "ИЖС": value,
             "ИЖС_source_file": "monitoring.xlsx", "ИЖС_source_sheet": "Данные с 2011 года",
             "ИЖС_source_cell": cell, "ИЖС_source_unit": "тыс. м²", "ИЖС_source_value": value * 1000}
            for year, value, cell in ((2011, .0258, "S24"), (2013, .3124, "S26"),
                                       (2014, .4974, "S27"), (2025, .6015, "S38"), (2026, 999, "S39"))])
        history = pd.concat([history, pd.DataFrame([
            {"region": "РФ", "year": 2012, "month": 12, "period_type": "year", "ИЖС": 999},
            {"region": "Москва", "year": 2016, "month": 12, "period_type": "ytd", "ИЖС": 999},
            {"region": "Москва", "year": 2017, "month": 8, "period_type": "year", "ИЖС": 999},
        ])], ignore_index=True)
        self.data["load_monitoring_annual_ihs"] = history
        snapshot = copy.deepcopy(self.data)
        payload = report(self.data, catalog(self.data))
        chart = next(chart for chart in payload["charts"] if chart["id"] == "b2_1")
        rows = {row["year"]: row for row in chart["rows"]}
        for year, value in ((2011, .0258), (2013, .3124), (2014, .4974), (2025, .6015)):
            self.assertAlmostEqual(rows[year]["ИЖС"], value)
            self.assertAlmostEqual(rows[year]["МКД"], rows[year]["TOTAL"] - value)
        self.assertEqual(rows[2011]["TOTAL"], 3)
        self.assertIsNone(rows[2012]["ИЖС"])
        self.assertIsNone(rows[2026]["ИЖС"])
        self.assertEqual(rows[2016]["ИЖС"], 1)
        self.assertEqual(rows[2017]["ИЖС"], 1)
        self.assertIsNone(next(value["value"] for value in chart["summaries"][0]["values"] if value["id"] == "ИЖС"))
        provenance = {row["year"]: row for row in chart["categoryProvenance"]}
        self.assertEqual(provenance[2011]["ИЖС_source_cell"], "S24")
        self.assertEqual(provenance[2011]["ИЖС_source_file"], "monitoring.xlsx")
        self.assertIsNone(provenance[2011]["ИЖС_source_rows"])
        raw = {row["year"]: row for row in chart["rawCategoryValues"]}
        self.assertEqual(raw[2011]["МКД"], 3)
        self.assertEqual(raw[2011]["ИЖС_сельское"], .01)
        self.assertEqual([series["id"] for series in chart["series"]], ["МКД", "ИЖС"])
        self.assertFalse(any("до 2015" in note for note in payload["notes"]))
        self.assertFalse(any("не называется МКД" in note or "Серый сегмент" in note for note in payload["notes"]))
        self.assertEqual(annual_charts(self.data, "РФ"), rf_before)
        rf = report(self.data, catalog(self.data), region="rf")
        self.assertTrue(any("не называется МКД" in note for note in rf["notes"]))
        self.assertTrue(any("Серый сегмент" in note for note in rf["notes"]))
        self.assertNotIn("rawCategoryProvenance", rf["charts"][1])
        pd.testing.assert_frame_equal(self.data["load_monitoring_annual_ihs"], snapshot["load_monitoring_annual_ihs"])
        pd.testing.assert_frame_equal(previous, snapshot["load_vvod_static"]["msk_residential"])

        book = load_workbook(BytesIO(export_workbook(dict(version="test", generatedAt="test", source={}, **payload))))
        ihs_column = next(index for index, column in enumerate(chart["columns"], 1) if column["id"] == "ИЖС")
        self.assertAlmostEqual(book["b2_1"].cell(2, ihs_column).value, .0258)
        self.assertIsNone(book["b2_1"].cell(3, ihs_column).value)
        metadata = dict(book["Report"].values)
        exported = json.loads(metadata["b2_1/categoryProvenance"])
        self.assertEqual(next(row for row in exported if row["year"] == 2011)["ИЖС_source_cell"], "S24")
        self.assertIn("load_monitoring_annual_ihs", LOADERS)
        self.assertIn("monitoring_annual_ihs", create_service(DataContext(ROOT)).families)

    def test_api_chart_summaries_and_excel_share_moscow_classification_and_raw_evidence(self):
        self.data["load_vvod_static"]["msk_total"].loc[1, "жильё"] = 4
        self.data["load_vvod_static"]["msk_residential"]["МКД_source_file"] = "explicit.xlsx"
        self.data["load_vvod_static"]["msk_residential"]["МКД_source_cell"] = "B12"
        self.data["load_vvod_static"]["msk_residential"]["МКД_source_value"] = 3000
        self.data["load_monitoring_annual_ihs"] = pd.DataFrame([
            {"region": "Москва", "year": year, "month": 12, "period_type": "year", "ИЖС": .25,
             "ИЖС_source_file": "annual.xlsx", "ИЖС_source_cell": f"S{year - 1987}"}
            for year in range(2011, 2026)])
        options = catalog(self.data)
        service = SimpleNamespace(report=lambda region: dict(
            version="test", generatedAt="test", source={}, **report(self.data, options, region)))
        app = FastAPI()
        app.include_router(create_router(None, lambda action, **kwargs: action(), service=service))
        with TestClient(app) as client:
            response = client.get("/api/v1/commissioning/annual?region=msk")
            self.assertEqual(response.status_code, 200)
            payload = response.json()
            json.dumps(payload, allow_nan=False)
            response = client.get("/api/v1/commissioning/annual/export?region=msk&required_version=test")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["X-Data-Version"], payload["version"])
        chart = next(chart for chart in payload["charts"] if chart["id"] == "b2_1")
        self.assertEqual([column["id"] for column in chart["columns"]], ["year", "МКД", "ИЖС", "TOTAL"])
        self.assertEqual(chart["rows"][0], {"year": 2011, "МКД": 2.75, "ИЖС": .25, "TOTAL": 3})
        self.assertIsNone(chart["rows"][-1]["ИЖС"])
        book = load_workbook(BytesIO(response.content))
        excel_rows = list(book["b2_1"].values)
        self.assertEqual(excel_rows[0], tuple(column["label"] for column in chart["columns"]))
        self.assertEqual(excel_rows[1:], [tuple(row[column["id"]] for column in chart["columns"]) for row in chart["rows"]])
        for series in chart["series"]:
            self.assertEqual(series["points"], [{"x": str(row["year"]), "y": row[series["id"]]} for row in chart["rows"]])
        expected_totals = []
        for period in chart["summaries"]:
            for value in period["values"]:
                measurements = [row[value["id"]] for row in chart["rows"] if period["from"] <= row["year"] <= period["to"]]
                expected = None if None in measurements else sum(measurements)
                self.assertEqual(value["value"], expected)
                expected_totals.append((period["from"], period["to"], value["label"], expected))
        self.assertEqual(list(book["b2_1_totals"].values)[1:], expected_totals)
        self.assertEqual(chart["summaries"][0]["values"][0]["value"], 55.25)
        self.assertIsNone(chart["summaries"][1]["values"][1]["value"])
        metadata = dict(book["Report"].values)
        for field in ("rawCategoryValues", "rawCategoryProvenance", "categoryProvenance"):
            self.assertEqual(json.loads(metadata["b2_1/" + field]), chart[field])
        original = chart["rawCategoryProvenance"][0]
        self.assertEqual(original["МКД_source_file"], "explicit.xlsx")
        self.assertEqual(original["МКД_source_cell"], "B12")
        self.assertEqual(original["МКД_source_value"], 3000)
        self.assertEqual(chart["rawCategoryValues"][0]["МКД"], 3)
        calculated = chart["categoryProvenance"][0]
        self.assertIsNone(calculated["МКД_source_cell"])
        self.assertEqual(calculated["МКД_source_kind"], "user_agreed_calculation")


if __name__ == "__main__":
    unittest.main()
