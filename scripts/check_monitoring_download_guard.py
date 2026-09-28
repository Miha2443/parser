"""Real-XLSX offline regression checks for monitoring prepublication validation."""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from contextlib import ExitStack, redirect_stdout
import json
import os
import socket
import sys
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl
import nashdom_checker as nc


HEADERS = {
    "Реестр РВ": ["УИН", "Группа компаний", "Год ввода по Мосстату", "Отрасли", "Группировка", "Общая площадь", "Жилая площадь", "Количество квартир"],
    "Реестр ОКС": ["УИН", "Группа компаний", "Назначение", "Общая площадь", "Жилая площадь", "Количество квартир"],
}


def workbook_bytes(*, omit_sheet=None, omit_header=None, empty=False, area=100,
                   year=2026, count=10, formula_headers=False, cached=True):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, columns in HEADERS.items():
        if title == omit_sheet:
            continue
        sheet = book.create_sheet(title)
        columns = [c for c in columns if c != omit_header]
        sheet.append([f'="{c}"' if formula_headers else c for c in columns])
        row = {"УИН": "shared-uin", "Группа компаний": "Developer", "Год ввода по Мосстату": year,
               "Отрасли": "Жилые объекты", "Группировка": "Жилье", "Назначение": "Жилье",
               "Общая площадь": area, "Жилая площадь": 70 if area else 0, "Количество квартир": count}
        if not empty:
            # Equal rows must remain distinct source components.
            sheet.append([row[c] for c in columns])
            sheet.append([row[c] for c in columns])
    target = BytesIO()
    book.save(target)
    book.close()
    if not formula_headers or not cached:
        return target.getvalue()
    # openpyxl cannot write formula caches; reproduce valid Google OOXML cached
    # string headers directly, leaving <f> formulas present for data_only=False.
    original = BytesIO(target.getvalue())
    rewritten = BytesIO()
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(original) as source, ZipFile(rewritten, "w", ZIP_DEFLATED) as destination:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename.startswith("xl/worksheets/sheet"):
                tree = ET.fromstring(content)
                for cell in tree.findall("m:sheetData/m:row[@r='1']/m:c", ns):
                    formula = cell.find("m:f", ns)
                    if formula is not None:
                        cell.set("t", "str")
                        cell.find("m:v", ns).text = formula.text.strip('"')
                content = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            destination.writestr(entry, content)
    return rewritten.getvalue()


class FixedTime:
    @staticmethod
    def now():
        return datetime(2026, 9, 28, 14, 0, 0)


class Response:
    status_code = 200
    url = "https://docs.google.com/export"
    headers = {"content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}

    def __init__(self, content):
        self.content = content


class MonitoringDownloadGuard(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.raw = self.root / "raw"
        self.raw.mkdir()
        self.state_path = self.root / "state.json"
        self.target = self.raw / "monitoring_2_0_20260928.xlsx"
        self.good = workbook_bytes()
        self.target.write_bytes(self.good)
        self.state = {"monitoring_2_0": {"filename": self.target.name,
                      "sha256": nc._sha256_bytes(self.good), "size_bytes": len(self.good),
                      "downloaded_at": "2026-09-28T12:00:00"}, "other_source": {"keep": True}}
        self.state_path.write_text(json.dumps(self.state, ensure_ascii=False), encoding="utf-8")
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.stack.enter_context(patch.object(nc, "DOWNLOAD_DIR", self.raw))
        self.stack.enter_context(patch.object(nc, "STATE_FILE", self.state_path))
        self.stack.enter_context(patch.object(nc, "datetime", FixedTime))
        self.stack.enter_context(patch.dict(os.environ, {"NASHDOM_FORCE": "0", "TDM_DISABLED": "1"}))
        self.stack.enter_context(patch.object(socket.socket, "connect", side_effect=AssertionError("Network forbidden in test")))
        self.stack.enter_context(patch.object(socket.socket, "connect_ex", side_effect=AssertionError("Network forbidden in test")))

    def fetch(self, content):
        with patch.object(nc.requests, "get", return_value=Response(content)), redirect_stdout(StringIO()):
            return nc.fetch_monitoring_2_0(self.state)

    def assert_rejected_preserving_last_good(self, content):
        before_state = deepcopy(self.state)
        before_disk_state = self.state_path.read_bytes()
        before_time = self.target.stat().st_mtime_ns
        self.assertEqual(self.fetch(content), ([], False))
        self.assertEqual(self.target.read_bytes(), self.good)
        self.assertEqual(self.target.stat().st_mtime_ns, before_time)
        self.assertEqual(self.state, before_state)
        self.assertEqual(self.state_path.read_bytes(), before_disk_state)
        self.assertEqual([p.name for p in self.raw.iterdir()], [self.target.name])

    def test_rejects_empty_fake_zip_and_wrong_or_incomplete_real_workbooks(self):
        cases = [b"", b"PK\x03\x04not a workbook", workbook_bytes(omit_sheet="Реестр ОКС"),
                 workbook_bytes(omit_header="Жилая площадь"), workbook_bytes(empty=True),
                 workbook_bytes(area=""), workbook_bytes(area="unknown"), workbook_bytes(year="invalid")]
        for index, content in enumerate(cases):
            with self.subTest(case=index):
                self.assert_rejected_preserving_last_good(content)

    def test_nonfinite_and_negative_values_do_not_publish(self):
        cases = [workbook_bytes(area="inf"), workbook_bytes(area="-inf"), workbook_bytes(area="NaN"),
                 workbook_bytes(area=-1), workbook_bytes(count="inf"), workbook_bytes(count="NaN")]
        for index, content in enumerate(cases):
            with self.subTest(case=index):
                self.assert_rejected_preserving_last_good(content)

    def test_valid_export_and_equal_components_publish_without_deduplication(self):
        content = workbook_bytes(area=123)
        self.assertEqual(self.fetch(content), ([self.target], True))
        self.assertEqual(self.target.read_bytes(), content)
        self.assertEqual(self.state["monitoring_2_0"]["sha256"], nc._sha256_bytes(content))
        book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            self.assertEqual(book["Реестр РВ"].max_row, 3)
        finally:
            book.close()
        self.assertEqual(self.state["other_source"], {"keep": True})

    def test_zero_is_valid_and_same_day_verified_file_keeps_mtime(self):
        content = workbook_bytes(area=0, count=0)
        self.assertEqual(self.fetch(content), ([self.target], True))
        before_time = self.target.stat().st_mtime_ns
        self.assertEqual(self.fetch(content), ([], True))
        self.assertEqual(self.target.stat().st_mtime_ns, before_time)

    def test_formatted_numeric_area_uses_shared_normalization(self):
        self.assertEqual(self.fetch(workbook_bytes(area="1\xa0234,5")), ([self.target], True))

    def test_google_formula_headers_use_cached_values_and_missing_cache_fails(self):
        content = workbook_bytes(formula_headers=True)
        book = openpyxl.load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(book["Реестр РВ"]["A1"].value, '="УИН"')
        book.close()
        self.assertEqual(self.fetch(content), ([self.target], True))
        self.good = content
        self.assert_rejected_preserving_last_good(workbook_bytes(formula_headers=True, cached=False))

    def test_same_day_state_cannot_bypass_semantic_validation(self):
        bad = workbook_bytes(omit_header="Назначение")
        self.target.write_bytes(bad)
        self.state["monitoring_2_0"].update(sha256=nc._sha256_bytes(bad), size_bytes=len(bad))
        before = deepcopy(self.state)
        self.assertEqual(self.fetch(bad), ([], False))
        self.assertEqual(self.state, before)

    def test_corrupt_same_day_target_is_replaced_by_valid_matching_response(self):
        self.target.write_bytes(b"externally corrupted")
        self.assertEqual(self.fetch(self.good), ([self.target], True))
        self.assertEqual(self.target.read_bytes(), self.good)


if __name__ == "__main__":
    unittest.main(verbosity=2)
