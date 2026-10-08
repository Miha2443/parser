"""Execute unchanged economics pages; compare chart points, pivots and exports."""
import ast
import copy
import hashlib
from io import BytesIO
import itertools
import json
from pathlib import Path
from pickle import UnpicklingError
import tempfile
import unittest
from unittest.mock import patch
from zipfile import BadZipFile

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import create_app
from backend.economics_export import export_workbook
from backend.economics_service import EconomicsService, LOADERS, build_catalog
from backend.profile_service import FilterUnavailable, SourceUnavailable
from pipeline.data_access import (
    DataAccess, DataContext, MONTH_NAMES_RU, QUARTER_NAMES_RU, file_signature,
    format_thousands, month_label, quarter_label,
)
from pipeline.economics_calculations import MSK, RF, VIEWS, accounts_report, ipc_report, salary_report
from pipeline.profile_calculations import clean, records
from pipeline.download_provenance import source_provenance, download_summary

ROOT = Path(__file__).resolve().parents[2]
PAGES = {"salary": "1_Заработная_плата.py", "ipc": "2_ИПЦ.py", "accounts": "3_ВРП_и_ВВП.py"}


class Capture:
    def __init__(self, family, selection):
        self.family, self.selection = family, selection
        self.figures, self.exports, self.frames = [], [], []
        self.session_state = {}

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def columns(self, sizes):
        return [self] * (sizes if isinstance(sizes, int) else len(sizes))

    def container(self, *args, **kwargs):
        return self

    def expander(self, *args, **kwargs):
        return self

    def radio(self, label, options, **kwargs):
        key = kwargs.get("key")
        if key in ("b5_region", "b6_region"):
            value = self.selection["structure_region" if key == "b5_region" else "index_region"]
        elif key == "b5_mode":
            value = "В рублях" if self.selection["structure_mode"] == "value" else "Доля, %"
        elif label == "Период":
            value = {"year": "Год", "quarter": "Квартал", "month": "Месяц"}[self.selection["period"]]
        elif label == "Регион":
            value = self.selection["region"]
        elif label == "Значение":
            value = "С начала года" if self.selection["ytd"] else options[0]
        else:
            value = {"month_to_month": "К предыдущему месяцу", "ytd_to_yago": "С начала года к АППГ"}[self.selection["index_base"]]
        assert value in options, (label, value)
        return value

    def multiselect(self, label, options, **kwargs):
        key = kwargs.get("key")
        if key and key.startswith("na_block"):
            value = self.selection[key.replace("na_", "")]
        elif key in ("b5_industries", "b6_industries"):
            value = self.selection["structure_industries" if key == "b5_industries" else "index_industries"]
        else:
            name = {"Отрасль": "views", "Регионы": "regions", "Кварталы": "quarters", "Месяцы": "months"}[label]
            value = self.selection[name]
        assert all(item in options for item in value), (label, value, options)
        return value

    def checkbox(self, *args, **kwargs):
        return self.selection["show_total"]

    def plotly_chart(self, figure, **kwargs):
        self.figures.append(figure)

    def dataframe(self, frame, **kwargs):
        self.frames.append((frame if isinstance(frame, pd.DataFrame) else frame.data).copy())

    def __getattr__(self, name):
        return lambda *args, **kwargs: None


def capture(family, data, selection):
    path = ROOT / "app/pages" / PAGES[family]
    tree = ast.parse(path.read_text(encoding="utf-8"))
    tree.body = [node for node in tree.body if not isinstance(node, (ast.Import, ast.ImportFrom))]
    ui = Capture(family, selection)
    env = {"st": ui, "pd": pd, "px": px, "go": go,
           LOADERS[family]: lambda: data.copy(deep=True),
           "COLORS": {key: "#333333" for key in ("red", "blue", "neutral")}, "SERIES": ["#333333", "#444444"],
           "MONTH_NAMES_RU": MONTH_NAMES_RU, "QUARTER_NAMES_RU": QUARTER_NAMES_RU,
           "month_label": month_label, "quarter_label": quarter_label, "format_thousands": format_thousands,
           "moscow_first": lambda values: sorted(list(values), key=lambda value: 0 if "моск" in str(value).casefold() else 1),
           "apply_theme": lambda: None, "page_header": lambda *args: None,
           "style_plotly": lambda *args, **kwargs: None,
           "chart_download_button": lambda *args, **kwargs: None,
           "table_download_buttons": lambda frame, **kwargs: ui.exports.append(frame.copy()),
           "dataset_download_summary": lambda *args: "", "show_dataset_sources": lambda *args: None}
    exec(compile(tree, str(path), "exec"), env)
    return ui


def pivot_records(frame):
    frame = frame.copy()
    if frame.columns.name == "year" or all(isinstance(col, (int, float)) for col in frame.columns):
        frame.columns = [str(int(col)) for col in frame.columns]
    return records(frame.reset_index())


def input_fingerprints(context):
    access = DataAccess(context)
    files = {Path(item[0]) for loader in LOADERS.values() for item in access.data_version(loader) if item[1] is not None}
    for loader in LOADERS.values():
        frame = getattr(access, loader)()
        for name in frame.source_file.dropna().unique():
            path = context.downloads / str(name)
            if Path(str(name)).name == str(name) and path.is_file():
                files.add(path)
    return {str(path): hashlib.sha256(path.read_bytes()).hexdigest() for path in files}


class LegacyParityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.context = DataContext(ROOT)
        cls.hashes = input_fingerprints(cls.context)
        cls.raw_marts = tuple(file_signature(path) for base in (ROOT / "data/raw", ROOT / "data/marts")
                              for path in sorted(base.rglob("*")) if path.is_file())
        access = DataAccess(cls.context)
        cls.data = {family: getattr(access, loader)() for family, loader in LOADERS.items()}
        cls.services = {family: EconomicsService(cls.context, family) for family in LOADERS}

    def compare(self, family, payload, data=None):
        ui = capture(family, self.data[family] if data is None else data, payload["selection"])
        self.assertEqual(len(payload["charts"]), len(ui.figures))
        self.assertEqual(len(payload["charts"]), len(ui.exports))
        self.assertEqual(len(payload["tables"]), len(ui.frames))
        for chart, figure, exported in zip(payload["charts"], ui.figures, ui.exports):
            self.assertEqual([series["name"] for series in chart["series"]], [trace.name for trace in figure.data])
            for series, trace in zip(chart["series"], figure.data):
                self.assertEqual([point["x"] for point in series["points"]], clean(list(trace.x)))
                self.assertEqual([point["y"] for point in series["points"]], clean(list(trace.y)))
            self.assertEqual([col["label"] for col in chart["columns"]], list(exported.columns))
            self.assertEqual(chart["rows"], records(exported))
        for table, frame in zip(payload["tables"], ui.frames):
            self.assertEqual(table["rows"], pivot_records(frame))
        return ui

    def test_salary_all_periods_modes_regions_industries_and_subsets(self):
        for region, views, (period, ytd), subset in itertools.product(
                (MSK, RF), (VIEWS, VIEWS[:1], VIEWS[1:]),
                (("year", False), ("month", False), ("month", True), ("quarter", False), ("quarter", True)), (False, True)):
            filters = dict(region=region, views=views, period=period, ytd=ytd)
            if subset:
                filters.update(months=[1, 5, 12], quarters=[1, 4])
            with self.subTest(**filters):
                self.compare("salary", self.services["salary"].report(**filters))

    def test_ipc_all_periods_bases_regions_and_subsets(self):
        for regions, (period, index_base), subset in itertools.product(
                ([MSK], [RF], [MSK, RF]),
                (("year", "month_to_month"), ("month", "month_to_month"), ("month", "ytd_to_yago"),
                 ("quarter", "month_to_month"), ("quarter", "ytd_to_yago")), (False, True)):
            filters = dict(regions=regions, period=period, index_base=index_base)
            if subset:
                filters.update(months=[2, 3, 12], quarters=[2, 4])
            with self.subTest(**filters):
                self.compare("ipc", self.services["ipc"].report(**filters))

    def test_accounts_six_blocks_independent_regions_modes_industries_total(self):
        catalog = self.services["accounts"].catalog()
        for structure_region, index_region, mode, show_total, many in itertools.product(
                (MSK, RF), (MSK, RF), ("value", "share"), (False, True), (False, True)):
            filters = dict(structure_region=structure_region, index_region=index_region,
                           structure_mode=mode, show_total=show_total,
                           block1_regions=[MSK, RF], block2_regions=[RF], block3_regions=[MSK], block4_regions=[RF, MSK])
            if many:
                filters.update(structure_industries=[item["id"] for item in catalog["structureIndustries"][structure_region][mode]],
                               index_industries=[item["id"] for item in catalog["indexIndustries"][index_region]])
            with self.subTest(**filters):
                payload = self.services["accounts"].report(**filters)
                self.compare("accounts", payload)
                self.assertEqual(len(payload["charts"]), 6)
                self.assertTrue(all(point["x"] >= 2011 for chart in payload["charts"] for series in chart["series"] for point in series["points"]))
                last = payload["charts"][-1]
                self.assertEqual(any(series["id"] == "total" for series in last["series"]), show_total)
                self.assertTrue(any(row["view"] == "Всего" for row in last["rows"]))

    def test_empty_selection_preserves_page_no_chart_semantics(self):
        for family, filters in (("salary", {"views": []}), ("salary", {"months": []}),
                                ("salary", {"period": "quarter", "quarters": []}), ("ipc", {"regions": []}),
                                ("ipc", {"months": []}), ("ipc", {"period": "quarter", "quarters": []}),
                                ("accounts", {**{f"block{i}_regions": [] for i in range(1, 5)},
                                              "structure_industries": [], "index_industries": []})):
            payload = self.services[family].report(**filters)
            self.compare(family, payload)
            self.assertEqual(payload["charts"], [])

    def test_duplicate_source_rows_nulls_labels_and_first_pivots(self):
        for family in ("salary", "ipc", "accounts"):
            data = self.data[family].copy(deep=True)
            # Additional conflicting source rows stay in series/downloads; pivots use first.
            eligible = data[data.year.ge(2016)]
            keys = ["region", "view", "metric"] if family == "accounts" else ["region", "view", "period_type"]
            duplicates = eligible.groupby(keys, sort=False).head(2).copy()
            duplicates["value"] *= 2
            duplicates["source_file"] = "another-source.xlsx"
            data = pd.concat([data, duplicates], ignore_index=True)
            candidates = data.index[data.year.ge(2016)]
            data.loc[candidates[:8], "value"] = float("nan")
            if family == "salary":
                for period, ytd in (("year", False), ("month", False), ("month", True), ("quarter", False), ("quarter", True)):
                    selection = self.services[family].report(period=period, ytd=ytd, views=VIEWS)["selection"]
                    payload = {"selection": selection, **salary_report(data, **selection)}
                    self.compare(family, payload, data)
            elif family == "ipc":
                for period, base in itertools.product(("year", "quarter", "month"), ("month_to_month", "ytd_to_yago")):
                    selection = self.services[family].report(period=period, index_base=base, regions=[MSK, RF])["selection"]
                    payload = {"selection": selection, **ipc_report(data, **selection)}
                    self.compare(family, payload, data)
            else:
                for region, mode, total in itertools.product((MSK, RF), ("value", "share"), (False, True)):
                    selection = self.services[family].report(block1_regions=[MSK, RF], block2_regions=[MSK, RF],
                                                             block3_regions=[MSK, RF], block4_regions=[MSK, RF],
                                                             structure_region=region, index_region=region,
                                                             structure_mode=mode, show_total=total)["selection"]
                    payload = {"selection": selection, **accounts_report(data, **selection)}
                    self.compare(family, payload, data)

    def test_exact_accounts_scaling_modern_options_and_legacy_rest(self):
        payload = self.services["accounts"].report(block1_regions=[MSK, RF], block2_regions=[MSK, RF])
        self.assertEqual([chart["unit"] for chart in payload["charts"]],
                         ["трлн руб", "млн руб/чел", "%", "%", "трлн руб", "%"])
        structure = payload["charts"][4]
        self.assertTrue(any(point["x"] < 2016 and point["y"] > 0 for point in structure["series"][-1]["points"]))
        # Unit selection is deliberately from the first row, not per-row normalization.
        data = self.data["accounts"].copy(deep=True)
        data.loc[data.metric.eq("vds_value") & data.region.eq(MSK), "unit"] = "млрд руб"
        changed = {"selection": payload["selection"], **accounts_report(data, **payload["selection"])}
        self.compare("accounts", changed, data)
        original = structure["series"][0]["points"][0]["y"]
        self.assertAlmostEqual(changed["charts"][4]["series"][0]["points"][0]["y"], original * 1000)

    def test_full_workbooks_all_rows_points_pivots_units_and_null(self):
        for family in LOADERS:
            payload = self.services[family].report()
            book = load_workbook(BytesIO(export_workbook(payload)))
            for chart in payload["charts"]:
                tab = book[chart["id"] + "_series"]
                expected = [[series["id"], series["name"], series["unit"], p["x"], p["y"]]
                            for series in chart["series"] for p in series["points"]]
                actual = [list(row) for row in tab.iter_rows(min_row=2, values_only=True)]
                self.assertEqual(len(actual), len(expected))
                for found, wanted in zip(actual, expected):
                    self.assertEqual(found[:4], wanted[:4])
                    if wanted[-1] is None:
                        self.assertIsNone(found[-1])
                    else:
                        self.assertAlmostEqual(found[-1], wanted[-1], places=8)
            for item in [*payload["charts"], *payload["tables"]]:
                self.assertEqual(book[item["id"]].max_row, len(item["rows"]) + 1)
                self.assertEqual(list(next(book[item["id"]].values)), [col["label"] for col in item["columns"]])
            self.assertEqual(book["Provenance"].max_row, len(payload["source"]["provenance"]["rows"]) + 1)
            self.assertIn(payload["version"], [row[1] for row in book["Report"].values])
        payload = self.services["salary"].report()
        payload["charts"][0]["series"][0]["points"][0]["y"] = None
        payload["charts"][0]["rows"][0]["view"] = "=1+1"
        payload["charts"][0]["columns"][0]["label"] = "+header"
        payload["selection"]["region"] = "@evil"
        book = load_workbook(BytesIO(export_workbook(payload)))
        tab = book["salary_month"]
        self.assertEqual(tab["A1"].value, "+header")
        self.assertEqual(tab["A1"].data_type, "s")
        self.assertTrue(any(cell.value == "=1+1" and cell.data_type == "s" for row in tab for cell in row))
        self.assertIsNone(book["salary_month_series"]["E2"].value)

    def test_catalog_defaults_copy_and_provenance_matches_helper(self):
        for family in LOADERS:
            catalog = self.services[family].catalog()
            payload = self.services[family].report()
            self.assertEqual(catalog["version"], payload["version"])
            self.assertEqual(payload["schemaVersion"], 1)
            self.assertEqual(catalog["source"]["fileEvidence"], "candidates")
            self.assertTrue(catalog["source"]["recordedSources"])
            self.assertTrue(catalog["source"]["provenance"]["rows"])
            self.assertEqual(len(catalog["controls"]), 10 if family == "accounts" else 6 if family == "salary" else 5)
            source_data = self.data[family]
            if family == "accounts":
                source_data = source_data[source_data.year.ge(2011)]
            expected = source_provenance(source_data, self.context.downloads)
            self.assertEqual(catalog["source"]["downloadSummary"], download_summary(expected))
            expected_rows = records(expected)
            for row in expected_rows:
                row["Исходный файл"] = self.services[family]._relative(row["Исходный файл"])
            self.assertEqual(catalog["source"]["provenance"]["rows"], expected_rows)
            catalog["defaults"].clear()
            payload["charts"].clear()
            payload["source"]["files"].clear()
            self.assertTrue(self.services[family].catalog()["defaults"])
            self.assertTrue(self.services[family].report()["charts"])
            self.assertTrue(self.services[family].report()["source"]["files"])

    def test_http_real_three_families_full_contract_and_pins(self):
        client = TestClient(create_app())
        requests = {
            "salary": [("period", "quarter"), ("ytd", "true"), ("views", VIEWS[0]), ("views", VIEWS[1]),
                       ("quarters", 1), ("quarters", 4)],
            "ipc": [("period", "month"), ("index_base", "ytd_to_yago"), ("regions", MSK), ("regions", RF),
                    ("months", 1), ("months", 12)],
            "accounts": [("block1_regions", MSK), ("block1_regions", RF), ("block2_regions", RF),
                         ("block3_regions", ""), ("block4_regions", MSK), ("structure_region", RF),
                         ("structure_mode", "share"), ("index_region", MSK), ("show_total", "false")],
        }
        for family, params in requests.items():
            base = f"/api/v1/economics/{family}"
            with self.subTest(family=family):
                catalog = client.get(base + "/catalog")
                self.assertEqual(catalog.status_code, 200, catalog.text)
                report = client.get(base + "/report", params=params)
                self.assertEqual(report.status_code, 200, report.text)
                payload = report.json()
                self.assertEqual(payload["version"], catalog.json()["version"])
                self.assertEqual(client.get(base + "/export", params=params).status_code, 422)
                bad = client.get(base + "/export", params=[*params, ("required_version", "stale")])
                self.assertEqual(bad.status_code, 409)
                self.assertEqual(bad.headers["X-Data-Version"], payload["version"])
                good = client.get(base + "/export", params=[*params, ("required_version", payload["version"])])
                self.assertEqual(good.status_code, 200, good.text if good.status_code != 200 else "")
                book = load_workbook(BytesIO(good.content))
                for chart in payload["charts"]:
                    self.assertEqual(book[chart["id"]].max_row, len(chart["rows"]) + 1)
                self.assertEqual(client.get(base, params={"structure_industries" if family == "accounts" else
                                                         "views" if family == "salary" else "regions": "unknown"}).status_code, 404)

    def test_zz_all_inputs_unchanged(self):
        self.assertEqual(self.hashes, input_fingerprints(self.context))
        self.assertEqual(self.raw_marts, tuple(file_signature(path) for base in (ROOT / "data/raw", ROOT / "data/marts")
                                              for path in sorted(base.rglob("*")) if path.is_file()))


class FakeAccess:
    version = 1
    reads = 0
    fail = False
    change = False
    error = None
    data = None

    def __init__(self, context):
        self.context = context
        self.issues = [("warning", "Synthetic source warning")]

    def data_version(self, loader):
        return ((str(self.context.root / "source.pkl"), self.version),)

    def load_salary(self):
        type(self).reads += 1
        if self.error:
            raise self.error
        if self.fail:
            self.issues.append(("error", "Synthetic read failure"))
        if self.change:
            type(self).version += 1
        return self.data.copy(deep=True)


class ServiceAndHttpTests(unittest.TestCase):
    def setUp(self):
        FakeAccess.version, FakeAccess.reads, FakeAccess.fail, FakeAccess.change, FakeAccess.error = 1, 0, False, False, None
        FakeAccess.data = pd.DataFrame([dict(year=2024, month=1, quarter=1, region=MSK, view=VIEWS[0],
                                           period_type="month", value=123.4, unit="руб", source_file="old.xls")])
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.context = DataContext(Path(self.temp.name))
        self.patcher = patch("backend.economics_service.DataAccess", FakeAccess)
        self.patcher.start()
        self.addCleanup(self.patcher.stop)
        self.service = EconomicsService(self.context, "salary", check_interval=0)

    def test_throttle_copy_refresh_warning_and_immutable_generation(self):
        self.service.check_interval = 5
        with patch("backend.economics_service.time.monotonic", return_value=10):
            before = self.service.report()
        self.assertEqual(before["source"]["issues"], [["warning", "Synthetic source warning"]])
        before["charts"][0]["rows"].clear()
        FakeAccess.version = 2
        with patch("backend.economics_service.time.monotonic", return_value=14):
            self.assertEqual(self.service.catalog()["version"], before["version"])
        self.assertEqual(FakeAccess.reads, 1)
        with patch("backend.economics_service.time.monotonic", return_value=15):
            after = self.service.report()
        self.assertNotEqual(before["version"], after["version"])
        self.assertTrue(after["charts"][0]["rows"])
        self.assertEqual(FakeAccess.reads, 2)
        self.service.report()
        self.assertEqual(FakeAccess.reads, 2)

    def test_failed_refresh_never_publishes_or_serves_previous_generation(self):
        before = self.service.catalog()
        FakeAccess.version, FakeAccess.fail = 2, True
        for _ in range(2):
            with self.assertRaises(SourceUnavailable):
                self.service.report()
        self.assertEqual(self.service._metadata["version"], before["version"])
        FakeAccess.fail = False
        self.assertNotEqual(before["version"], self.service.report()["version"])

    def test_changed_inputs_mid_load_are_not_published(self):
        FakeAccess.change = True
        with self.assertRaises(SourceUnavailable):
            self.service.catalog()
        self.assertIsNone(self.service._metadata)
        FakeAccess.change = False
        self.assertTrue(self.service.report()["charts"])

    def test_changes_during_provenance_never_publish(self):
        from pipeline.download_provenance import source_provenance

        def changing(*args):
            FakeAccess.version += 1
            return source_provenance(*args)

        with patch("backend.economics_service.source_provenance", side_effect=changing):
            with self.assertRaises(SourceUnavailable):
                self.service.catalog()
        self.assertIsNone(self.service._metadata)

    def test_referenced_file_and_receipt_changes_refresh_provenance_not_values(self):
        from pipeline.download_provenance import receipt_path

        before = self.service.report()
        downloads = self.context.downloads
        downloads.mkdir()
        source = downloads / "old.xls"
        source.write_bytes(b"local source fixture")
        after = self.service.report()
        self.assertNotEqual(before["version"], after["version"])
        self.assertEqual(before["charts"], after["charts"])
        receipt = receipt_path(source)
        receipt.parent.mkdir()
        receipt.write_text(json.dumps({"downloaded_at": "2024-01-01T00:00:00+00:00",
                                      "sha256": hashlib.sha256(source.read_bytes()).hexdigest()}), encoding="utf-8")
        confirmed = self.service.catalog()
        self.assertNotEqual(after["version"], confirmed["version"])
        self.assertEqual(confirmed["source"]["provenance"]["rows"][0]["Статус"], "Скачивание подтверждено")
        self.assertEqual(self.service.report()["charts"], before["charts"])

    def test_null_units_values_and_unavailable_regions_stay_explicit(self):
        FakeAccess.data.loc[0, "unit"] = None
        FakeAccess.data.loc[0, "value"] = float("nan")
        payload = self.service.report()
        self.assertIsNone(payload["charts"][0]["series"][0]["points"][0]["y"])
        self.assertIsNone(payload["charts"][0]["rows"][0]["значение_руб"])
        self.assertEqual(payload["charts"][0]["unit"], "руб.")
        book = load_workbook(BytesIO(export_workbook(payload)))
        self.assertIsNone(book["salary_month_series"]["E2"].value)
        with self.assertRaises(FilterUnavailable):
            self.service.report(region=RF)

    def test_pinned_export_rejects_new_generation_instead_of_stale_cache(self):
        client = TestClient(create_app())
        with patch("backend.economics_service.time.monotonic", return_value=10):
            old = client.get("/api/v1/economics/salary").json()["version"]
        FakeAccess.version += 1
        with patch("backend.economics_service.time.monotonic", return_value=15):
            response = client.get("/api/v1/economics/salary/export", params={"required_version": old})
        self.assertEqual(response.status_code, 409)
        self.assertNotEqual(response.headers["X-Data-Version"], old)

    def test_non_numeric_empty_and_missing_schema_are_503_not_filter_errors(self):
        for data in (pd.DataFrame(), FakeAccess.data.drop(columns="value"), FakeAccess.data.assign(value="bad")):
            with self.subTest(data=data.shape):
                FakeAccess.data = data
                with self.assertRaises(SourceUnavailable):
                    self.service.catalog()
        with self.assertRaises(SourceUnavailable):
            build_catalog("accounts", pd.DataFrame([dict(year=2010, region=MSK, view="Всего", value=1.0,
                                                       unit="руб", metric="vrp_total", source_file="a.xlsx")]))

    def test_filter_failure_not_cached_and_partial_source_is_not_fabricated(self):
        for filters in ({"region": RF}, {"views": ["bad"]}, {"months": [13]}, {"period": "quarter", "quarters": [5]}):
            with self.assertRaises(FilterUnavailable):
                self.service.report(**filters)
        self.assertEqual(self.service.report(period="year")["charts"], [])
        self.assertTrue(self.service.report()["charts"])
        self.assertEqual(FakeAccess.reads, 1)

    def test_http_registration_repeated_typed_lists_empty_and_mandatory_pin(self):
        client = TestClient(create_app())
        response = client.get("/api/v1/economics/salary/report", params=[("months", "1"), ("views", VIEWS[0])])
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["selection"]["months"], [1])
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.assertEqual(client.get("/api/v1/economics/salary/export").status_code, 422)
        self.assertEqual(client.get("/api/v1/economics/salary/export", params={"required_version": "wrong"}).status_code, 409)
        good = client.get("/api/v1/economics/salary/export", params={"required_version": payload["version"], "months": 1, "views": VIEWS[0]})
        self.assertEqual(good.status_code, 200)
        self.assertEqual(good.headers["X-Data-Version"], payload["version"])
        self.assertEqual(load_workbook(BytesIO(good.content))["salary_month"].max_row, 2)
        for key in ("views", "months"):
            empty = client.get("/api/v1/economics/salary", params={key: ""})
            self.assertEqual(empty.status_code, 200, empty.text)
            self.assertEqual(empty.json()["charts"], [])
        for params, status in (({"months": "1,2"}, 422), ({"period": "week"}, 422), ({"ytd": "bad"}, 422),
                               ({"region": RF}, 404), ({"months": 13}, 404), ({"region": ""}, 422)):
            self.assertEqual(client.get("/api/v1/economics/salary", params=params).status_code, status)
        paths = client.get("/openapi.json").json()["paths"]
        for family in LOADERS:
            for suffix in ("", "/catalog", "/report", "/export"):
                self.assertIn(f"/api/v1/economics/{family}{suffix}", paths)
        for path in ("annual", "operational"):
            self.assertTrue(any(f"commissioning/{path}" in item for item in paths))
        for path in ("/api/v1/linear", "/api/v1/map", "/api/v1/updates"):
            self.assertIn(path, paths)

    def test_http_corrupt_pickle_workbook_and_diagnostics_503_no_fallback(self):
        for error in (UnpicklingError("bad"), BadZipFile("bad"), EOFError("bad")):
            FakeAccess.error = error
            client = TestClient(create_app())
            with self.assertLogs("backend.app", level="ERROR"):
                self.assertEqual(client.get("/api/v1/economics/salary/catalog").status_code, 503)
            FakeAccess.error = None
            self.assertEqual(client.get("/api/v1/economics/salary/catalog").status_code, 200)
        FakeAccess.fail = True
        client = TestClient(create_app())
        with self.assertLogs("backend.app", level="ERROR"):
            self.assertEqual(client.get("/api/v1/economics/salary").status_code, 503)

    def test_export_uses_one_payload_when_version_changes_after_copy(self):
        from backend.economics_export import export_workbook as actual_export

        def changing(payload):
            FakeAccess.version += 1
            return actual_export(payload)

        client = TestClient(create_app())
        version = client.get("/api/v1/economics/salary").json()["version"]
        with patch("backend.economics_routes.export_workbook", side_effect=changing):
            response = client.get("/api/v1/economics/salary/export", params={"required_version": version})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["X-Data-Version"], version)
        book = load_workbook(BytesIO(response.content))
        self.assertIn(version, [row[1] for row in book["Report"].values])

    def test_backend_and_calculations_have_no_streamlit_imports(self):
        for path in [ROOT / "pipeline/economics_calculations.py", *ROOT.glob("backend/economics_*.py")]:
            for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))):
                if isinstance(node, ast.ImportFrom):
                    self.assertFalse((node.module or "").startswith(("streamlit", "app.data_access")))
                elif isinstance(node, ast.Import):
                    self.assertTrue(all(not alias.name.startswith("streamlit") for alias in node.names))


if __name__ == "__main__":
    unittest.main()
