"""Read-only service generations, public filters and pinned workbook exports."""
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.preferential_mortgage_routes import create_router, export_workbook
from backend.preferential_mortgage_service import PreferentialMortgageService
from backend.profile_service import FilterUnavailable, SourceUnavailable
from pipeline.data_access import DataContext, file_signature
from pipeline.preferential_mortgage import CHAR_METRICS, SOURCE_FILE
from tests.test_preferential_mortgage import characteristic, flow, partial_source, source, weighted_source


class PreferentialMortgageServiceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.context = DataContext(Path(self.tmp.name))
        self.path = self.context.root / SOURCE_FILE
        self.path.parent.mkdir(parents=True)
        self.service = PreferentialMortgageService(self.context)

    def write(self, document=None):
        document = source() if document is None else document
        self.path.write_text(json.dumps(document, ensure_ascii=False), encoding="utf-8")
        return document

    def client(self):
        def call(function, *args, **kwargs):
            try:
                return function(*args)
            except SourceUnavailable as exc:
                raise HTTPException(503, "Source unavailable") from exc
            except FilterUnavailable as exc:
                raise HTTPException(404, "Filter unavailable") from exc
        app = FastAPI()
        app.include_router(create_router(self.context, call, self.service))
        return TestClient(app)

    def test_catalog_region_specific_defaults_and_independent_characteristic_selection(self):
        self.write(source([flow("2026-01"), flow("2026-10"), flow("2026-02", region="rf"),
                           flow("2026-12", region="rf", count=None, volume=None)],
                          [characteristic("2026-07", averageAmountMln=5, ratePct=6),
                           characteristic("2026-10", averageAmountMln=5),
                           characteristic("2026-01", region="rf", averageAmountMln=5, ratePct=6)]))
        catalog = self.service.catalog()
        self.assertEqual(catalog["defaultRegion"], "msk")
        self.assertEqual(catalog["regions"][0]["periods"], ["2026-01", "2026-10"])
        self.assertEqual(catalog["regions"][0]["characteristicPeriods"], ["2026-07"])
        self.assertEqual([r["id"] for r in catalog["months"]], list(range(1, 13)))
        self.assertEqual([r["id"] for r in catalog["programs"]], ["preferential", "family", "it"])
        self.assertEqual(self.service.report()["selection"], {"region": "msk", "month": 10, "characteristicMonth": 7})
        self.assertEqual(self.service.report("rf")["selection"], {"region": "rf", "month": 2, "characteristicMonth": 1})
        self.assertEqual(self.service.report()["version"], catalog["version"])
        self.assertEqual(catalog["source"]["files"], [SOURCE_FILE])
        client = self.client()
        for month in range(1, 13):
            report = client.get("/api/v1/preferential-mortgage", params={"month": month, "characteristicMonth": 13 - month})
            self.assertEqual(report.status_code, 200)
            self.assertEqual(report.json()["selection"]["characteristicMonth"], 13 - month)

    def test_zero_region_available_null_only_region_unavailable_and_default_is_available_region(self):
        self.write(source([flow("2026-01", count=None, volume=None), flow("2026-02", region="rf", count=0, volume=0)]))
        self.assertEqual(self.service.catalog()["defaultRegion"], "rf")
        self.assertEqual(len(self.service.catalog()["regions"]), 1)
        self.assertEqual(self.service.report()["selection"]["characteristicMonth"], 2)
        self.assertEqual(self.client().get("/api/v1/preferential-mortgage?region=msk").status_code, 404)

    def test_public_errors_invalid_query_filters_missing_data_and_no_unrelated_fallback(self):
        client = self.client()
        self.path.with_name("other.json").write_text(json.dumps(source()), encoding="utf-8")
        endpoints = ("/api/v1/preferential-mortgage/catalog", "/api/v1/preferential-mortgage",
                     "/api/v1/preferential-mortgage/export?required_version=old")
        for value in (None, "{", "[]", json.dumps({**source(), "sourceUrl": "wrong"}), json.dumps({**source(), "flows": []})):
            with self.subTest(value=value):
                if value is not None:
                    self.path.write_text(value, encoding="utf-8")
                for endpoint in endpoints:
                    self.assertEqual(client.get(endpoint).status_code, 503)
        self.write()
        for key, value in (("region", "other"), ("month", 0), ("month", 13), ("characteristicMonth", 0), ("characteristicMonth", 13)):
            self.assertEqual(client.get("/api/v1/preferential-mortgage", params={key: value}).status_code, 404)
        for key in ("month", "characteristicMonth"):
            self.assertEqual(client.get("/api/v1/preferential-mortgage", params={key: "bad"}).status_code, 422)
        self.assertEqual(client.get("/api/v1/preferential-mortgage/export").status_code, 422)

    def test_fresh_generations_digest_changes_and_no_stale_data_after_failure(self):
        self.write()
        signature = file_signature(self.path)
        with patch("backend.preferential_mortgage_service.file_signature", return_value=signature):
            first = self.service.report()
            self.write(source([flow("2026-01", count=99)]))
            second = self.service.report()
        self.assertNotEqual(first["version"], second["version"])
        self.assertEqual(second["cumulative"][-1]["programs"]["family"]["count"], 99)
        self.path.write_text("broken", encoding="utf-8")
        with self.assertRaises(SourceUnavailable):
            self.service.report()
        self.assertIsNone(self.service._generation)
        self.write()
        self.service.catalog()
        self.path.unlink()
        with self.assertRaises(SourceUnavailable):
            self.service.catalog()
        self.assertIsNone(self.service._generation)

    def test_source_replacement_during_load_validation_or_cached_return_is_rejected(self):
        self.write()
        signature = file_signature(self.path)
        for sequence in ((signature, ("different",)), (signature, signature, ("different",))):
            self.service._signature = None
            with self.subTest(sequence=sequence), patch("backend.preferential_mortgage_service.file_signature", side_effect=sequence):
                with self.assertRaises(SourceUnavailable):
                    self.service.catalog()
                self.assertIsNone(self.service._generation)
        self.service.catalog()
        with patch("backend.preferential_mortgage_service.file_signature", side_effect=(signature, signature, ("different",))):
            with self.assertRaises(SourceUnavailable):
                self.service.catalog()
        self.assertIsNone(self.service._generation)

    def test_payload_and_raw_snapshots_are_detached_and_read_only(self):
        original = self.write()
        before = self.path.read_bytes()
        signature = file_signature(self.path)
        catalog = self.service.catalog()
        catalog["regions"].clear()
        catalog["source"]["files"].clear()
        report, raw = self.service.export_snapshot()
        report["cumulative"][-1]["programs"]["family"]["count"] = 999
        raw["flows"].clear()
        self.assertEqual(self.service.export_snapshot()[1], original)
        self.assertEqual(self.service.report()["cumulative"][-1]["programs"]["family"]["count"], 2)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(file_signature(self.path), signature)

    def test_version_pinned_export_exact_raw_values_horizons_and_formula_safety(self):
        document = weighted_source()
        document["flows"][0].update(volumeMln=1234567.891234, count=0.5, sourceCells={"count": "M10"}, comment=" =1+2")
        document["notes"] = ["=HYPERLINK(\"bad\")"]
        self.write(document)
        client = self.client()
        report = client.get("/api/v1/preferential-mortgage", params={"month": 2, "characteristicMonth": 3}).json()
        self.assertEqual(client.get("/api/v1/preferential-mortgage/export?required_version=old").status_code, 409)
        response = client.get("/api/v1/preferential-mortgage/export", params={"required_version": report["version"], "month": 2, "characteristicMonth": 3})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["x-data-version"], report["version"])
        self.assertIn("preferential-mortgage.xlsx", response.headers["content-disposition"])
        self.assertIn("spreadsheetml.sheet", response.headers["content-type"])
        book = load_workbook(BytesIO(response.content), data_only=False)
        metadata = dict(book["Report"].values)
        self.assertEqual(metadata["month"], 2)
        self.assertEqual(metadata["characteristicMonth"], 3)
        sheet = book["Source flows"]
        observed = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
        self.assertEqual(observed["volumeMln"], 1234567.891234)
        self.assertEqual(observed["count"], 0.5)
        self.assertEqual(json.loads(observed["sourceCells"]), {"count": "M10"})
        self.assertEqual(observed["comment"], "' =1+2")
        raw_chars = book["Source characteristics"]
        headings = [c.value for c in raw_chars[1]]
        self.assertIn("propertyAmountMln", headings)
        self.assertEqual(book["Cumulative characteristics"].cell(2, 5).value, None)
        self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
        self.assertTrue(any(c.value == "'=HYPERLINK(\"bad\")" for row in book["Sources and method"] for c in row))

    def test_exact_source_window_is_preserved_in_public_report_and_pinned_export(self):
        self.write(partial_source())
        client = self.client()
        for month, expected in ((7, 30.0552221846817), (6, None)):
            with self.subTest(month=month):
                params = {"region": "msk", "characteristicMonth": month}
                response = client.get("/api/v1/preferential-mortgage", params=params)
                self.assertEqual(response.status_code, 200)
                report = response.json()
                self.assertEqual(report["characteristicsCumulative"][-1]["programs"]["family"]["ownFundsPct"], expected)
                exported = client.get("/api/v1/preferential-mortgage/export", params={**params, "required_version": report["version"]})
                self.assertEqual(exported.status_code, 200)
                book = load_workbook(BytesIO(exported.content))
                observed = next(row for row in book["Cumulative characteristics"].iter_rows(min_row=2, values_only=True)
                                if row[0] == 2026 and row[2] == "family" and row[3] == "ownFundsPct")
                self.assertEqual(observed[4], month)
                self.assertEqual(observed[5], expected)

    def test_export_uses_one_pinned_snapshot_even_during_source_replacement(self):
        original = self.write()
        version = self.service.catalog()["version"]
        def replaced(payload, snapshot):
            self.write(source([flow("2026-01", count=999)]))
            self.assertEqual(snapshot, original)
            self.assertEqual(payload["version"], version)
            return export_workbook(payload, snapshot)
        with patch("backend.preferential_mortgage_routes.export_workbook", side_effect=replaced):
            response = self.client().get("/api/v1/preferential-mortgage/export", params={"required_version": version})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["x-data-version"], version)
        current = self.service.catalog()["version"]
        stale = self.client().get("/api/v1/preferential-mortgage/export", params={"required_version": version})
        self.assertEqual(stale.status_code, 409)
        self.assertEqual(stale.headers["x-data-version"], current)

    def test_export_preserves_long_evidence_without_cell_truncation(self):
        document = source()
        document["flows"][0]["evidence"] = {"text": "a" * 80000}
        document["extraEvidence"] = {"text": "b" * 90000}
        self.write(document)
        report, raw = self.service.export_snapshot()
        book = load_workbook(BytesIO(export_workbook(report, raw)))
        fragments = {}
        for location, part, value in book["Long source fields"].iter_rows(min_row=2, values_only=True):
            fragments.setdefault(location, []).append((part, value))
        for location, expected in (("flows/0/evidence", document["flows"][0]["evidence"]),
                                   ("metadata/extraEvidence", document["extraEvidence"])):
            restored = "".join(value for _, value in sorted(fragments[location]))
            self.assertEqual(json.loads(restored), expected)

    def test_app_registration_uses_profile_pinned_root_and_no_store(self):
        self.write()
        from backend.app import create_app
        class ProfileStub:
            context = self.context
        client = TestClient(create_app(service=ProfileStub()))
        response = client.get("/api/v1/preferential-mortgage/catalog")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertEqual(response.json()["source"]["files"], [SOURCE_FILE])
        report = client.get("/api/v1/preferential-mortgage").json()
        for collection in ("characteristicsAnnual", "characteristicsCumulative"):
            self.assertEqual(set(report[collection][-1]["programs"]["family"]), set(CHAR_METRICS))
            self.assertEqual(set(report[collection][-1]["throughMonths"]), set(CHAR_METRICS))


if __name__ == "__main__":
    unittest.main()
