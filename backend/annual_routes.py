"""Annual commissioning API; historical rules remain in the extracted pure module."""
from io import BytesIO
import json

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook

from backend.generation_service import GenerationService
from backend.profile_service import FilterUnavailable, SourceUnavailable
from pipeline.annual_calculations import annual_charts

LOADERS = ("load_vvod_static", "load_emiss_34118", "load_emiss_34118_periods", "load_monitoring_2_0", "load_monitoring_2011_2026_static")


def catalog(data):
    if all(frame.empty for frame in data["load_vvod_static"].values()) and data["load_emiss_34118"].empty:
        raise SourceUnavailable("Annual commissioning sources are missing")
    return {"regions": [{"id": "msk", "label": "Москва"}, {"id": "rf", "label": "РФ"}]}


def report(data, options, region="msk"):
    chosen = next((item for item in options["regions"] if item["id"] == region), None)
    if chosen is None:
        raise FilterUnavailable("Annual region is unavailable")
    charts = annual_charts(data, chosen["label"])
    if not charts:
        raise SourceUnavailable("Annual report has no source data")
    notes = ["Единицы всех графиков: млн м²; период 2011–2026, реновация с 2017 года.",
             "Заполнение пропусков нулём в годовых графиках сохранено по прежнему правилу; это не подтверждение полноты исходных данных.",
             "ИЖС до 2015 года не показывается: сопоставимый подтверждённый ряд начинается с 2015 года."]
    if region == "rf":
        notes.append("2026 год по РФ — нет данных; нулевой столбец не означает нулевой фактический ввод.")
    else:
        notes.extend(["2026 год по Москве дополняется мониторингом 2.0; исторические значения МОП и реновации сохранены из утверждённой страницы.",
                      "Дополнительные отрасли показываются при покрытии мониторингом не менее 95% годового объёма; иначе площадь остаётся в «Прочем».",
                      "Состав отраслей фиксирован для всех лет: доля более 10% хотя бы в одном году и не менее шести крупнейших групп."])
    return {"region": chosen, "charts": charts, "notes": notes}


def create_service(context, check_interval=5):
    return GenerationService(context, LOADERS, tuple(name.removeprefix("load_") for name in LOADERS),
                             catalog, report, check_interval)


def export_workbook(payload):
    book = Workbook()
    tab = book.active
    tab.title = "Report"
    tab.append(["Field", "Value"])
    for key in ("version", "generatedAt", "region", "source", "notes"):
        value = payload[key]
        tab.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    for chart in payload["charts"]:
        tab = book.create_sheet(chart["id"])
        tab.append([column["label"] for column in chart["columns"]])
        for row in chart["rows"]:
            tab.append([row[column["id"]] for column in chart["columns"]])
        summary = book.create_sheet(f"{chart['id']}_totals")
        summary.append(["From", "To", "Indicator", "million m2"])
        for period in chart["summaries"]:
            for row in period["values"]:
                summary.append([period["from"], period["to"], row["label"], row["value"]])
    for tab in book:
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = 30
        for row in tab:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.lstrip().startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def create_router(context, call, service=None):
    router = APIRouter(prefix="/api/v1/commissioning/annual")
    service = service or create_service(context)

    @router.get("/catalog")
    def options():
        return call(service.catalog, scope="Annual commissioning")

    @router.get("")
    def view(region: str = Query(default="msk", pattern="^(msk|rf)$")):
        return call(lambda: service.report(region=region), scope="Annual commissioning")

    @router.get("/export")
    def export(region: str = Query(default="msk", pattern="^(msk|rf)$"), required_version: str = Query(min_length=1, max_length=128)):
        payload = view(region)
        if required_version != payload["version"]:
            raise HTTPException(409, "Source changed; refresh the displayed report before exporting")
        return Response(export_workbook(payload), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="annual-commissioning.xlsx"',
                                 "X-Data-Version": payload["version"]})

    return router
