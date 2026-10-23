"""DOM.RF program workbook extraction and population-checked calendar summaries."""
from __future__ import annotations

import copy
from calendar import monthrange
from datetime import date, datetime, timezone
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from pipeline.mortgage_calculations import MONTH_LABELS


SOURCE_FILE = "data/raw/realty/domrf/preferential_mortgage.json"
SOURCE_URL = "https://дом.рф/programmy-gosudarstvennoj-podderzhki/operational-reporting/"
REGIONS = {"msk": "Москва", "rf": "Российская Федерация"}
PROGRAMS = {"preferential": "Льготная ипотека", "family": "Семейная ипотека", "it": "IT ипотека"}
FLOW_METRICS = ("volumeMln", "count")
CHAR_METRICS = ("averageAmountMln", "ratePct", "termMonths", "ownFundsPct", "areaM2", "priceThousandM2")
CHAR_SOURCE_METRICS = (*CHAR_METRICS, "propertyAmountMln")
FLOW_SHEET = "01_02_01"
CHAR_SHEET = "01_02_05"
_REGION_LABELS = {"г. Москва": "msk", "Москва": "msk", "Российская Федерация": "rf"}
_PROGRAM_LABELS = {label: key for key, label in PROGRAMS.items()}
_CHAR_LABELS = {
    "Средняя сумма кредита, млн руб.": "averageAmountMln",
    "Средневзвешенный размер текущей процентной ставки, %": "ratePct",
    "Средний срок кредита, мес.": "termMonths",
    "Размер собственных средств заемщика, %": "ownFundsPct",
    "Средняя площадь, м2": "areaM2",
    "Средняя стоимость 1 м2, тыс. руб.": "priceThousandM2",
    "Средняя стоимость помещения, млн руб.": "propertyAmountMln",
}
_WEIGHTS = {"ratePct": "Сумма кредита, срок кредита", "termMonths": "Сумма кредита",
            "ownFundsPct": "Стоимость квартиры"}
_METHOD_LABELS = {"Процентная ставка": "ratePct", "Срок кредита": "termMonths",
                  "Доля собственных средств": "ownFundsPct"}
_MONTHS_GENITIVE = ("января", "февраля", "марта", "апреля", "мая", "июня",
                   "июля", "августа", "сентября", "октября", "ноября", "декабря")
_NOTES = [
    "Объем выдач указан в млн руб., количество в кредитах. Дробные количества из оперативных данных сохранены.",
    "Итого включает только Льготную, Семейную и IT ипотеку, а не строку источника «Все программы».",
    "Годовые значения взяты из явных годовых итогов источника. Последний год может быть неполным.",
    "Накопительные выдачи суммируются с января. Пропуски не заменяются нулями.",
    "Характеристики не усредняются арифметически. Веса проверяются по методике и сверкой с годовым итогом.",
    "При совпадении периода с подтвержденным периодом показателя используется опубликованный итог без пересчета.",
    "Для ставки вес равен объему выдач × сроку, для срока — объему, для собственных средств — стоимости жилья.",
    "Средняя сумма и площадь взвешиваются по числу кредитов; цена м² равна общей стоимости жилья / общей площади.",
    "Если охват или веса не подтверждаются, накопительная характеристика недоступна (null).",
]


def _text(value):
    return " ".join(value.replace("\u00a0", " ").split()) if isinstance(value, str) else ""


def _number(value):
    if value is None or (isinstance(value, str) and _text(value) in ("", "-", "–", "—", "н/д")):
        return None
    if type(value) not in (int, float):
        raise ValueError(f"Non-numeric workbook value: {value!r}")
    try:
        valid = math.isfinite(value) and value >= 0
    except OverflowError:
        valid = False
    if not valid:
        raise ValueError("Invalid workbook numeric value")
    return value


def _columns(header, first_value):
    columns, seen = [], set()
    for index, value in enumerate(header[first_value:], first_value):
        label = _text(value)
        if isinstance(value, (datetime, date)):
            if value.day != 1 or value.year < 2018:
                raise ValueError("Unexpected monthly header")
            basis, period = "month", value.strftime("%Y-%m")
        elif re.fullmatch(r"Итого\s+\d{4}\s+год", label, re.IGNORECASE):
            basis, period = "annual", re.search(r"\d{4}", label).group()
        elif label.upper() == "ИТОГО (ЗА ВЕСЬ ПЕРИОД)":
            basis, period = "allTime", "all"
        elif value is None:
            continue
        else:
            raise ValueError(f"Unexpected workbook period header: {value!r}")
        if (basis, period) in seen:
            raise ValueError("Duplicate workbook period column")
        seen.add((basis, period))
        columns.append((index, basis, period))
    if not any(basis == "month" for _, basis, _ in columns) or not any(basis == "annual" for _, basis, _ in columns):
        raise ValueError("Workbook requires monthly columns and explicit annual totals")
    return columns


def parse_workbook(path):
    """Stable downloader entry point. Read, but never modify, the supplied XLSX."""
    path = Path(path)
    flows, characteristics, region_labels, notes, methodology = {}, {}, set(), [], {}
    content = path.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if FLOW_SHEET not in book.sheetnames or CHAR_SHEET not in book.sheetnames:
            raise ValueError("DOM.RF program workbook sheets are missing")
        source_date = None
        if "01_02_00" in book.sheetnames:
            for row_number, row in enumerate(book["01_02_00"].iter_rows(values_only=True), 1):
                for column, value in enumerate(row):
                    label = _text(value)
                    if "Данные актуальны по состоянию на" in label:
                        match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", label)
                        if match:
                            source_date = date(*map(int, match.group(3, 2, 1))).isoformat()
                    if label and column == 1 and 13 <= row_number < 30:
                        notes.append(label)
                if len(row) > 2 and _text(row[1]) in _METHOD_LABELS:
                    metric = _METHOD_LABELS[_text(row[1])]
                    methodology[metric] = {"weight": _text(row[2]), "sheet": "01_02_00", "row": row_number}

        def record(store, region, program, basis, period, metrics, sheet):
            key = region, program, basis, period
            return store.setdefault(key, {"region": region, "program": program, "basis": basis,
                                          "period": period, **dict.fromkeys(metrics), "sourceSheet": sheet,
                                          "sourceCells": {}})

        sheet = book[FLOW_SHEET]
        rows = sheet.iter_rows(min_row=4, values_only=True)
        header = next(rows, ())
        if len(header) < 3 or _text(header[0]) != "Программы" or _text(header[1]) != "Размерность величин":
            raise ValueError("Unexpected flow header layout")
        columns = _columns(header, 2)
        region, seen_flow = None, set()
        for row_number, row in enumerate(rows, 5):
            label, unit = _text(row[0]), _text(row[1])
            if label and not unit:
                region_labels.add(label)
                region = _REGION_LABELS.get(label)
                continue
            program = _PROGRAM_LABELS.get(label)
            if region is None or program is None:
                continue
            metric = {"Всего, шт.": "count", "Всего, млн руб.": "volumeMln"}.get(unit)
            if metric is None:
                raise ValueError(f"Unexpected source units: {unit!r}")
            if (region, program, metric) in seen_flow:
                raise ValueError("Duplicate program flow row")
            seen_flow.add((region, program, metric))
            for index, basis, period in columns:
                target = record(flows, region, program, basis, period, FLOW_METRICS, FLOW_SHEET)
                target[metric] = _number(row[index])
                target["sourceCells"][metric] = f"{get_column_letter(index + 1)}{row_number}"

        sheet = book[CHAR_SHEET]
        rows = sheet.iter_rows(min_row=4, values_only=True)
        header = next(rows, ())
        if len(header) < 2 or _text(header[0]) != "Программы":
            raise ValueError("Unexpected characteristics header layout")
        columns = _columns(header, 1)
        region, program, seen_char = None, None, set()
        for row_number, row in enumerate(rows, 5):
            label = _text(row[0])
            if label in region_labels or label in _REGION_LABELS:
                region, program = _REGION_LABELS.get(label), None
                continue
            metric = _CHAR_LABELS.get(label)
            if not metric:
                if label:
                    program = _PROGRAM_LABELS.get(label)
                continue
            if region is None or program is None:
                continue
            if (region, program, metric) in seen_char:
                raise ValueError("Duplicate program characteristic row")
            seen_char.add((region, program, metric))
            for index, basis, period in columns:
                target = record(characteristics, region, program, basis, period, CHAR_SOURCE_METRICS, CHAR_SHEET)
                target[metric] = _number(row[index])
                target["sourceCells"][metric] = f"{get_column_letter(index + 1)}{row_number}"
        for region in REGIONS:
            for program in PROGRAMS:
                if any((region, program, metric) not in seen_flow for metric in FLOW_METRICS):
                    raise ValueError("Required region/program flow rows are missing")
                if any((region, program, metric) not in seen_char for metric in CHAR_METRICS):
                    raise ValueError("Required region/program characteristic rows are missing")
    finally:
        book.close()
    return validate_source({"schemaVersion": 1, "sourceUrl": SOURCE_URL, "sourceDate": source_date,
                            "collectedAt": datetime.now(timezone.utc).isoformat(),
                            "sourceWorkbook": {"filename": path.name, "sha256": digest},
                            "flows": list(flows.values()), "characteristics": list(characteristics.values()),
                            "methodology": methodology, "notes": notes})


def validate_source(document):
    """Reject duplicate dimensions, ambiguous units and non-finite observations."""
    if not isinstance(document, dict) or type(document.get("schemaVersion")) is not int or document["schemaVersion"] != 1:
        raise ValueError("Unsupported preferential mortgage source schema")
    if document.get("sourceUrl") != SOURCE_URL:
        raise ValueError("Unexpected preferential mortgage source URL")
    collected = document.get("collectedAt")
    if not isinstance(collected, str) or datetime.fromisoformat(collected.replace("Z", "+00:00")).tzinfo is None:
        raise ValueError("Collection timestamp must include a timezone")
    source_date = document.get("sourceDate")
    if source_date is not None:
        if not isinstance(source_date, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", source_date):
            raise ValueError("Invalid source date")
        date.fromisoformat(source_date)
    for collection, metrics in (("flows", FLOW_METRICS), ("characteristics", CHAR_SOURCE_METRICS)):
        rows = document.get(collection)
        if not isinstance(rows, list) or not rows:
            raise ValueError(f"Source {collection} are missing")
        keys, populated = set(), False
        for row in rows:
            if not isinstance(row, dict):
                raise ValueError("Invalid source observation")
            if (not isinstance(row.get("region"), str) or row["region"] not in REGIONS
                    or not isinstance(row.get("program"), str) or row["program"] not in PROGRAMS):
                raise ValueError("Invalid source region or program")
            basis, period = row.get("basis"), row.get("period")
            patterns = {"month": r"(?:201[8-9]|20[2-9]\d|[3-9]\d{3})-(?:0[1-9]|1[0-2])",
                        "annual": r"(?:201[8-9]|20[2-9]\d|[3-9]\d{3})", "allTime": r"all"}
            if not isinstance(basis, str) or basis not in patterns or not isinstance(period, str) or not re.fullmatch(patterns[basis], period):
                raise ValueError("Invalid source basis or period")
            key = row["region"], row["program"], basis, period
            if key in keys:
                raise ValueError("Duplicate source observation")
            keys.add(key)
            for metric in metrics:
                if metric not in row:
                    raise ValueError(f"Missing source metric: {metric}")
                value = row[metric]
                if value is not None:
                    if _number(value) is None:
                        raise ValueError("Source metrics must be numbers or null")
                    populated = True
        if collection == "flows" and not populated:
            raise ValueError(f"Source {collection} contain no observed values")
    if not isinstance(document.get("notes", []), list) or any(not isinstance(note, str) for note in document.get("notes", [])):
        raise ValueError("Invalid source notes")
    methodology = document.get("methodology", {})
    if not isinstance(methodology, dict):
        raise ValueError("Invalid source methodology")
    for metric, evidence in methodology.items():
        if metric not in _WEIGHTS or not isinstance(evidence, dict) or not isinstance(evidence.get("weight"), str):
            raise ValueError("Invalid weighting evidence")
    json.dumps(document, allow_nan=False)
    return copy.deepcopy(document)


def region_periods(document, characteristics=False):
    collection = "characteristics" if characteristics else "flows"
    metrics = CHAR_METRICS[1:] if characteristics else FLOW_METRICS
    return {region: sorted({row["period"] for row in document[collection]
                            if row["region"] == region and row["basis"] == "month"
                            and any(row[metric] is not None for metric in metrics)}) for region in REGIONS}


def _sum(values):
    values = list(values)
    if not values or any(value is None for value in values):
        return None
    try:
        result = sum(values) if all(type(value) is int for value in values) else math.fsum(values)
        return result if math.isfinite(result) else None
    except OverflowError:
        return None


def _divide(numerator, denominator):
    if numerator is None or denominator in (None, 0):
        return None
    try:
        value = numerator / denominator
        return value if math.isfinite(value) else None
    except OverflowError:
        return None


def _close(left, right):
    return left is not None and right is not None and math.isclose(left, right, rel_tol=1e-10, abs_tol=1e-9)


def _weighted(rows, metric):
    products, weights = [], []
    for flow, char in rows:
        count, volume = flow.get("count"), flow.get("volumeMln")
        if count == 0 and volume == 0:
            continue
        if (count is None or volume is None or count <= 0
                or not _close(_divide(volume, count), char.get("averageAmountMln"))):
            return None
        value = char.get(metric)
        if value is None:
            return None
        if metric in ("averageAmountMln", "areaM2"):
            weight = count
        elif metric == "termMonths":
            weight = volume
        elif metric == "ratePct":
            term = char.get("termMonths")
            if term is None:
                return None
            weight = volume * term
        else:
            property_amount = char.get("propertyAmountMln")
            if property_amount is None:
                return None
            if metric == "ownFundsPct":
                weight = count * property_amount
            else:
                area = char.get("areaM2")
                if area is None or area <= 0 or not _close(value, property_amount / area * 1000):
                    return None
                weight = count * area
        if metric == "averageAmountMln":
            products.append(volume)
        elif metric == "priceThousandM2":
            products.append(1000 * char["propertyAmountMln"] * count)
        else:
            products.append(value * weight)
        weights.append(weight)
    return _divide(_sum(products), _sum(weights))


def _metric_horizon(index, program, year, metric):
    return max((month for month in range(1, 13)
                if index.get((program, "month", f"{year}-{month:02d}"), {}).get(metric) is not None), default=None)


def _reimbursement_cutoff(document, program, year):
    notes = [_text(note) for note in document.get("notes", [])]
    parameters_only = any(re.search(r"параметр\w* кредит\w*.*отчет\w* о возмещени\w*", note, re.IGNORECASE)
                          and re.search(r"оперативн\w* статистик\w* отсутств", note, re.IGNORECASE) for note in notes)
    source_date = document.get("sourceDate")
    if not parameters_only or not source_date or int(source_date[:4]) != year:
        return None
    pattern = (r"\bс\s+(" + "|".join(_MONTHS_GENITIVE) + r")\s+(\d{4})\s+г\.\s*[-–—]\s*"
               r"(?:ежемесячные опросы|еженедельные отчеты)")
    cutoffs = set()
    for note in notes:
        if f"{PROGRAMS[program]}:" not in note:
            continue
        for match in re.finditer(pattern, note, re.IGNORECASE):
            transition_month = _MONTHS_GENITIVE.index(match[1].lower()) + 1
            transition_year = int(match[2])
            if source_date < f"{transition_year}-{transition_month:02d}-01":
                continue
            cutoff_year, cutoff_zero_month = divmod(transition_year * 12 + transition_month - 2, 12)
            if cutoff_year == year:
                cutoffs.add(cutoff_zero_month + 1)
    return cutoffs.pop() if len(cutoffs) == 1 else None


def _published_window_matches(index, flow_index, document, program, year, through, metric):
    if through != _metric_horizon(index, program, year, metric):
        return False
    flow_horizon = max((month for month in range(1, 13)
                        if any(flow_index.get((program, "month", f"{year}-{month:02d}"), {}).get(key) is not None
                               for key in FLOW_METRICS)), default=0)
    if flow_horizon < through:
        return False
    if metric != "averageAmountMln":
        confirmed_cutoff = _reimbursement_cutoff(document, program, year)
        if confirmed_cutoff is not None:
            return through == confirmed_cutoff
    remaining = [flow_index.get((program, "month", f"{year}-{month:02d}"), {})
                 for month in range(through + 1, 13)]
    if all(row.get("count") == 0 and row.get("volumeMln") == 0 for row in remaining):
        return True
    source_date = document.get("sourceDate")
    if source_date and int(source_date[:4]) == year and int(source_date[5:7]) == through and flow_horizon == through:
        return True
    # A lagged aggregate without dated population evidence remains unavailable.
    return False


def _cumulative_characteristic(index, flow_index, document, program, year, through, metric):
    def pairs(end):
        return [(flow_index.get((program, "month", f"{year}-{month:02d}"), {}),
                 index.get((program, "month", f"{year}-{month:02d}"), {})) for month in range(1, end + 1)]

    annual = index.get((program, "annual", str(year)), {}).get(metric)
    if annual is not None and _published_window_matches(index, flow_index, document, program, year, through, metric):
        return annual
    rows = pairs(through)
    if through == 1:
        return rows[0][1].get(metric)
    if metric in _WEIGHTS and document.get("methodology", {}).get(metric, {}).get("weight") != _WEIGHTS[metric]:
        return None
    source_horizon = _metric_horizon(index, program, year, metric)
    if source_horizon is None or annual is None:
        return None
    confirmed_cutoff = _reimbursement_cutoff(document, program, year) if metric != "averageAmountMln" else None
    if confirmed_cutoff is not None and source_horizon != confirmed_cutoff:
        return None
    # A known source annual total verifies both the weight and its population.
    if not _close(_weighted(pairs(source_horizon), metric), annual):
        return None
    return _weighted(rows, metric)


def build_report(document, region, month, characteristic_month):
    if (not isinstance(region, str) or region not in REGIONS
            or any(type(value) is not int or not 1 <= value <= 12 for value in (month, characteristic_month))):
        raise ValueError("Invalid preferential mortgage report filters")
    periods = region_periods(document)[region]
    if not periods:
        raise ValueError("Region has no observed flows")
    latest = periods[-1]
    years = range(2018, int(latest[:4]) + 1)
    flows = {(r["program"], r["basis"], r["period"]): r for r in document["flows"] if r["region"] == region}
    chars = {(r["program"], r["basis"], r["period"]): r for r in document["characteristics"] if r["region"] == region}
    annual, cumulative, char_annual, char_cumulative = [], [], [], []
    notes = [*_NOTES, *document.get("notes", [])]
    unverified = {}
    for year in years:
        observed = [int(period[5:]) for period in periods if period.startswith(f"{year}-")]
        horizon = max(observed, default=12 if year < int(latest[:4]) else 0)
        through = min(month, horizon)
        annual_programs, cumulative_programs = {}, {}
        for program in PROGRAMS:
            annual_programs[program] = {key: flows.get((program, "annual", str(year)), {}).get(key) for key in FLOW_METRICS}
            cumulative_programs[program] = {key: _sum(flows.get((program, "month", f"{year}-{m:02d}"), {}).get(key)
                                                        for m in range(1, through + 1)) for key in FLOW_METRICS}
        for target, programs, end in ((annual, annual_programs, horizon), (cumulative, cumulative_programs, through)):
            target.append({"year": year, "throughMonth": end, "programs": programs,
                           "total": {metric: _sum(p[metric] for p in programs.values()) for metric in FLOW_METRICS}})

        program_horizons = {program: {metric: _metric_horizon(chars, program, year, metric)
                                     for metric in CHAR_METRICS} for program in PROGRAMS}
        metric_horizons = {metric: max((p[metric] for p in program_horizons.values() if p[metric] is not None),
                                      default=None) for metric in CHAR_METRICS}
        detailed_horizon = max((metric_horizons[metric] for metric in CHAR_METRICS[1:]
                                if metric_horizons[metric] is not None), default=0)
        annual_values, cumulative_values, cumulative_horizons = {}, {}, {}
        for program in PROGRAMS:
            annual_values[program] = {metric: chars.get((program, "annual", str(year)), {}).get(metric) for metric in CHAR_METRICS}
            cumulative_values[program], cumulative_horizons[program] = {}, {}
            for metric in CHAR_METRICS:
                available = program_horizons[program][metric]
                end = min(characteristic_month, available) if available is not None else None
                cumulative_horizons[program][metric] = end
                if characteristic_month == 12 and horizon == 12:
                    value = annual_values[program][metric]
                elif end is not None:
                    value = _cumulative_characteristic(chars, flows, document, program, year, end, metric)
                    if value is None and annual_values[program][metric] is not None:
                        unverified.setdefault(metric, []).append(f"{PROGRAMS[program]} {year}")
                else:
                    value = None
                cumulative_values[program][metric] = value
        annual_end = max((m for m in metric_horizons.values() if m is not None), default=horizon)
        cumulative_metric_horizons = {metric: min(characteristic_month, end) if end is not None else None
                                      for metric, end in metric_horizons.items()}
        char_end = max((m for m in cumulative_metric_horizons.values() if m is not None), default=min(characteristic_month, horizon))
        char_annual.append({"year": year, "throughMonth": annual_end, "throughMonths": metric_horizons,
                            "programThroughMonths": program_horizons, "programs": annual_values})
        char_cumulative.append({"year": year, "throughMonth": char_end, "throughMonths": cumulative_metric_horizons,
                                "programThroughMonths": cumulative_horizons, "programs": cumulative_values})
        if year == int(latest[:4]) and detailed_horizon and detailed_horizon < horizon:
            notes.append(f"{year}: выдачи и средняя сумма доступны по месяц {horizon:02d}; подробные характеристики — по месяц {detailed_horizon:02d}. Годовой столбец источника имеет разные горизонты по показателям.")
    for metric, affected in unverified.items():
        notes.append(f"Накопительный {metric} не рассчитан: веса или охват не подтверждены годовым итогом ({'; '.join(affected)}).")
    source_date = document.get("sourceDate")
    latest_partial = None
    if source_date and source_date[:7] > latest:
        latest_partial = False
    elif source_date and source_date[:7] == latest:
        instant = date.fromisoformat(source_date)
        latest_partial = instant.day < monthrange(instant.year, instant.month)[1]
    if latest_partial:
        notes.append(f"Последний месяц {latest} неполный: данные по состоянию на {source_date}.")
    characteristic_periods = region_periods(document, characteristics=True)[region]
    return {"selection": {"region": region, "month": month, "characteristicMonth": characteristic_month},
            "regionLabel": REGIONS[region], "latestPeriod": latest, "latestPeriodPartial": latest_partial,
            "characteristicLatestPeriod": characteristic_periods[-1] if characteristic_periods else None,
            "annual": annual, "cumulative": cumulative,
            "characteristicsAnnual": char_annual, "characteristicsCumulative": char_cumulative, "notes": notes}
