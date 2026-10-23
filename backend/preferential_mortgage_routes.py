"""Preferential mortgage API and version-pinned, full-precision XLSX exports."""
from io import BytesIO
import json

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from backend.preferential_mortgage_service import PreferentialMortgageService
from pipeline.preferential_mortgage import CHAR_METRICS, CHAR_SOURCE_METRICS, FLOW_METRICS, PROGRAMS


def export_workbook(payload, source):
    book = Workbook()
    overflow = None

    def cell_value(value, location):
        nonlocal overflow
        result = json.dumps(value, ensure_ascii=False, allow_nan=False) if isinstance(value, (dict, list)) else value
        if not isinstance(result, str) or len(result) <= 30000:
            return result
        if overflow is None:
            overflow = book.create_sheet("Long source fields")
            overflow.append(["Field", "Part", "Text"])
        for start in range(0, len(result), 30000):
            overflow.append([location, start // 30000 + 1, result[start:start + 30000]])
        return "Long source fields: " + location

    metadata = book.active
    metadata.title = "Report"
    metadata.append(["Field", "Value"])
    for key, value in {"version": payload["version"], "generatedAt": payload["generatedAt"],
                       **payload["selection"], "latestPeriod": payload["latestPeriod"],
                       "latestPeriodPartial": payload["latestPeriodPartial"],
                       "characteristicLatestPeriod": payload["characteristicLatestPeriod"],
                       **payload["source"]}.items():
        metadata.append([key, cell_value(value, f"report/{key}")])
    for collection, name, metrics in (("flows", "Source flows", FLOW_METRICS),
                                       ("characteristics", "Source characteristics", CHAR_SOURCE_METRICS)):
        sheet = book.create_sheet(name)
        headers = ["region", "program", "basis", "period", *metrics]
        headers += sorted({key for row in source[collection] for key in row} - set(headers))
        sheet.append(headers)
        for index, row in enumerate(source[collection]):
            sheet.append([cell_value(row.get(key), f"{collection}/{index}/{key}") for key in headers])
    for collection, title in (("annual", "Annual flows"), ("cumulative", "Cumulative flows")):
        sheet = book.create_sheet(title)
        sheet.append(["year", "throughMonth", "program", *FLOW_METRICS])
        for row in payload[collection]:
            for program in (*PROGRAMS, "total"):
                values = row["total"] if program == "total" else row["programs"][program]
                sheet.append([row["year"], row["throughMonth"], program, *[values[key] for key in FLOW_METRICS]])
    for collection, title in (("characteristicsAnnual", "Annual characteristics"),
                               ("characteristicsCumulative", "Cumulative characteristics")):
        sheet = book.create_sheet(title)
        sheet.append(["year", "throughMonth", "program", "metric", "metricThroughMonth", "value"])
        for row in payload[collection]:
            for program in PROGRAMS:
                for metric in CHAR_METRICS:
                    through = row["programThroughMonths"][program][metric]
                    sheet.append([row["year"], row["throughMonth"], program, metric, through, row["programs"][program][metric]])
    notes = book.create_sheet("Sources and method")
    notes.append(["Note"])
    for index, note in enumerate(payload["notes"]):
        notes.append([cell_value(note, f"notes/{index}")])
    for key, value in source.items():
        if key not in ("flows", "characteristics"):
            notes.append([f"{key}: " + str(cell_value(value, f"metadata/{key}"))])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.lstrip().startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1C3444")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 27
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def create_router(context, call, service=None):
    router = APIRouter(prefix="/api/v1/preferential-mortgage")
    service = service or PreferentialMortgageService(context)

    @router.get("/catalog")
    def catalog():
        return call(service.catalog, scope="Preferential mortgage")

    @router.get("")
    def report(region: str | None = None, month: int | None = None, characteristicMonth: int | None = None):
        return call(service.report, region, month, characteristicMonth, scope="Preferential mortgage")

    @router.get("/export")
    def export(required_version: str = Query(min_length=1, max_length=128), region: str | None = None,
               month: int | None = None, characteristicMonth: int | None = None):
        payload, source = call(service.export_snapshot, region, month, characteristicMonth, scope="Preferential mortgage")
        if required_version != payload["version"]:
            raise HTTPException(409, "Source changed; refresh the displayed preferential mortgage report before exporting",
                                headers={"X-Data-Version": payload["version"]})
        content = call(export_workbook, payload, source, scope="Preferential mortgage")
        return Response(content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="preferential-mortgage.xlsx"',
                                 "X-Data-Version": payload["version"]})

    return router
