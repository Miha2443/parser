from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.linear_objects import period_summary, read_linear_objects
from backend.linear_routes import create_router, export_workbook
from backend.linear_service import LinearService, records
from backend.profile_service import FilterUnavailable, SourceUnavailable
from pipeline.data_access import DataContext


class LinearTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        directory = self.root / "data/raw/realty/linear_objects"
        directory.mkdir(parents=True)
        self.path = directory / "linear_objects_2026-09-28.json"
        rows = [{"code": f"1.{code}", "indicator": f"Indicator {code}", "unit": "km",
                 "year": 2026, "quarter": q, "plan": 10, "fact": None if q == 3 else q * 2}
                for code in range(1, 5) for q in range(1, 4)]
        self.path.write_text(json.dumps(rows), encoding="utf-8")
        self.service = LinearService(DataContext(self.root), check_interval=0)

    def test_legacy_parity_and_defaults(self):
        self.assertEqual(self.service.catalog()["years"][0]["defaultQuarter"], 2)
        data = read_linear_objects(self.path)
        for q in (1, 2, 3):
            for cumulative in (False, True):
                with self.subTest(q=q, cumulative=cumulative):
                    payload = self.service.report(2026, q, cumulative, "1.2")
                    self.assertEqual(payload["summary"], records(period_summary(data, 2026, q, cumulative)))
                    self.assertEqual(len(payload["allPeriods"]), 12)
        self.assertIsNone(self.service.report(2026, 3, True)["summary"][0]["fact"])

    def test_filters_missing_and_immutable(self):
        for args in ((2025,), (2026, 4), (2026, 1, False, "bad")):
            with self.assertRaises(FilterUnavailable):
                self.service.report(*args)
        catalog = self.service.catalog()
        catalog["years"].clear()
        self.assertTrue(self.service.catalog()["years"])
        self.path.unlink()
        with self.assertRaises(SourceUnavailable):
            self.service.catalog()

    def test_export_complete_and_version_conflict(self):
        payload = self.service.report()
        book = load_workbook(BytesIO(export_workbook(payload)))
        self.assertEqual(book["All quarters"].max_row, 13)
        app = FastAPI()
        app.include_router(create_router(DataContext(self.root), lambda function, *args, **kwargs: function(*args), self.service))
        client = TestClient(app)
        self.assertEqual(client.get("/api/v1/linear/export").status_code, 422)
        self.assertEqual(client.get("/api/v1/linear/export", params={"required_version": "old"}).status_code, 409)
        response = client.get("/api/v1/linear/export", params={"required_version": payload["version"]})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"PK"))

    def test_source_change(self):
        before = self.service.catalog()["version"]
        rows = json.loads(self.path.read_text())
        rows[0]["plan"] = 25
        self.path.write_text(json.dumps(rows), encoding="utf-8")
        self.assertNotEqual(before, self.service.catalog()["version"])
        self.assertEqual(self.service.report(2026, 1)["summary"][0]["plan"], 25)

    def test_corrupt_workbook_is_source_failure(self):
        self.path.unlink()
        self.path.with_suffix(".xlsx").write_bytes(b"not a valid zip workbook")
        with self.assertRaises(SourceUnavailable):
            self.service.catalog()


if __name__ == "__main__":
    unittest.main()
