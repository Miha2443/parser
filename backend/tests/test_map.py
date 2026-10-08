from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import pandas as pd
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.map_routes import create_router, workbook
from backend.map_service import MapService
from pipeline.data_access import DataContext
from pipeline.map_calculations import build_objects
from pipeline.map_coordinates import apply_geocode_cache, load_geocode_cache


class MapTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        directory = self.root / "data/raw/realty/nashdom"
        directory.mkdir(parents=True)
        self.path = directory / "monitoring_2_0_20260101.xlsx"
        self.rv = pd.DataFrame([
            {"УИН": "rv1", "Группа компаний": "A", "Строительный адрес": "Test", "Год ввода по Мосстату": 2008},
            {"УИН": "rv2", "Группа компаний": "B", "Строительный адрес": "Other", "Год ввода по Мосстату": None},
        ])
        geometry = '{"geometry":{"coordinates":[[37.6,55.7],[37.7,55.8]]}}'
        self.oks = pd.DataFrame([
            {"УИН": "o1", "Группа компаний": "A", "Срок выдачи РС": pd.Timestamp("2011-01-01"),
             "Строительный адрес": "Test", "Геометрия": geometry, "Год ввода по графику": 2026},
            {"УИН": "o2", "Срок выдачи РС": pd.Timestamp("2010-12-31"), "Год ввода по графику": 2030},
            {"УИН": "o3", "Срок выдачи РС": None, "Год выдачи": 2020, "Год ввода по графику": 2030},
        ])
        with pd.ExcelWriter(self.path) as writer:
            self.rv.to_excel(writer, sheet_name="Реестр РВ", index=False)
            self.oks.to_excel(writer, sheet_name="Реестр ОКС", index=False)
            pd.DataFrame([{"УИН": "forbidden", "Строительный адрес": "Лист4"}]).to_excel(
                writer, sheet_name="Лист4", index=False)
        self.service = MapService(DataContext(self.root), check_interval=0)

    def test_scope_date_boundary_unknown_rv_year_and_no_sheet4(self):
        report = self.service.report(only_with_coords=False)
        self.assertEqual({row["id"] for row in report["rows"]}, {"rv:rv1", "rv:rv2", "oks:o1"})
        self.assertEqual(report["scope"]["excludedOksRows"], 2)
        self.assertEqual(report["scope"]["missingIssueDateRows"], 1)
        bounded = self.service.report(year_from=2025, year_to=2027, only_with_coords=False)
        self.assertEqual({row["id"] for row in bounded["rows"]}, {"rv:rv2", "oks:o1"})
        self.assertEqual(len(report["geojson"]["features"]), 2)
        for feature in report["geojson"]["features"]:
            self.assertEqual(feature["geometry"]["coordinates"], [37.650000000000006, 55.75])
        self.assertEqual(report["totals"]["missing"], 1)
        self.assertEqual(len(self.service.report(quality="missing", only_with_coords=False)["rows"]), 1)

    def test_coordinate_priorities_and_no_mutation(self):
        cache = pd.DataFrame([{"registry": "rv", "object_id": "rv1", "address_key": "test",
                               "lat": 55.8, "lon": 37.8, "coord_source": "manual_exact", "precision": "exact"}])
        before_rv, before_oks = self.rv.copy(deep=True), self.oks.copy(deep=True)
        objects, _ = build_objects(self.rv, self.oks, cache)
        row = objects[objects["id"].eq("rv:rv1")].iloc[0]
        self.assertEqual(row["lat"], 55.8)
        self.assertEqual(row["coord_source"], "manual_exact")
        pd.testing.assert_frame_equal(before_rv, self.rv)
        pd.testing.assert_frame_equal(before_oks, self.oks)

    def test_blank_address_never_matches_cache_and_ids_keep_leading_zero(self):
        cache_path = self.root / "cache.csv"
        cache_path.write_text("registry,object_id,address,lat,lon,coord_source\nrv,001,,55.8,37.8,manual_exact\n")
        cache = load_geocode_cache(cache_path)
        self.assertEqual(cache.iloc[0]["object_id"], "001")
        objects, _ = build_objects(self.rv, self.oks, pd.DataFrame())
        missing = objects[objects["id"].eq("rv:rv2")].copy()
        missing["address_key"] = ""
        missing["address_match_key"] = ""
        result = apply_geocode_cache(missing, cache)
        self.assertTrue(result.iloc[0]["lat"] != result.iloc[0]["lat"])

    def test_full_export_pin_copy_and_refresh(self):
        report = self.service.report(only_with_coords=False)
        report["rows"][0]["address"] = "=Test"
        book = load_workbook(BytesIO(workbook(report)))
        self.assertEqual(book["Objects"].max_row, 4)
        address_col = [cell.value for cell in book["Objects"][1]].index("address") + 1
        self.assertTrue(any(row[address_col - 1].value == "'=Test" for row in list(book["Objects"])[1:]))
        app = FastAPI()
        app.include_router(create_router(DataContext(self.root), lambda fn, *args, **kw: fn(*args), self.service))
        client = TestClient(app)
        self.assertEqual(client.get("/api/v1/map/export").status_code, 422)
        self.assertEqual(client.get("/api/v1/map/export", params={"required_version": "old"}).status_code, 409)
        report["rows"].clear()
        self.assertEqual(len(self.service.report(only_with_coords=False)["rows"]), 3)
        old = self.service.catalog()["version"]
        self.path.touch()
        self.assertNotEqual(old, self.service.catalog()["version"])


if __name__ == "__main__":
    unittest.main()
