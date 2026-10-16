"""Temporary ready-sheet fixtures; no real inputs or external loaders are used."""
import builtins
from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import runpy
import shutil
import struct
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import MagicMock, patch
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app import create_app
from backend.construction_tree_service import COLUMN, SHEET, ConstructionTreeService
from backend.profile_service import SourceUnavailable
from pipeline.data_access import DataAccess, DataContext
from pipeline import source_availability as availability


ROWS = (
    (2, "total", "\u0412\u0441\u0435\u0433\u043e", 51.65728716),
    (3, "housing", "\u0416\u0438\u043b\u044c\u0435", 37.92555967),
    (4, "embedded", "\u041d\u0435\u0436\u0438\u043b\u044c\u0435 \u0432 \u0436\u0438\u043b\u044c\u0435", 5.09855928),
    (5, "mop", "\u041c\u041e\u041f", 10.68700039),
    (6, "living", "\u0416\u0438\u043b\u0430\u044f \u043f\u043b\u043e\u0449\u0430\u0434\u044c", 22.14),
    (7, "investors", "\u0418\u043d\u0432\u0435\u0441\u0442\u043e\u0440\u044b", 15.74),
    (8, "renovation", "\u0420\u0435\u043d\u043e\u0432\u0430\u0446\u0438\u044f", 6.4),
    (9, "nonresidential", "\u041d\u0435\u0436\u0438\u043b\u044c\u0435", 13.73172749),
    (11, "budget", "\u0411\u044e\u0434\u0436\u0435\u0442", 11.20503593),
    (12, "nonbudget", "\u0412\u043d\u0435\u0431\u044e\u0434\u0436\u0435\u0442", 40.45225123),
)
XML_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class ConstructionTreeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / "releases/release-a"
        self.directory = self.root / "data/raw/realty/nashdom"
        self.directory.mkdir(parents=True)
        self.path = self.directory / "monitoring_2_0_20260928.xlsx"
        self.context = DataContext(self.root, downloads=Path(self.tmp.name) / "downloads",
                                   require_marts=True)
        self.mtime = 1_792_100_000
        self.write_book()
        self.service = ConstructionTreeService(self.context, check_interval=0)

    def write_book(self, path=None, values=None, row_offset=0, column=9, sheet=True):
        path = path or self.path
        path.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        ready = book.active
        ready.title = SHEET if sheet else "Not the ready sheet"
        ready.cell(1 + row_offset, column, COLUMN)
        for row, identifier, label, value in ROWS:
            ready.cell(row + row_offset, column - 1, label)
            ready.cell(row + row_offset, column, (values or {}).get(identifier, value))
            # A historical block with identical labels must never win selection.
            ready.cell(row + row_offset, 2, label)
            ready.cell(row + row_offset, 3, 99999)
        book.create_sheet("\u0420\u0435\u0435\u0441\u0442\u0440 \u041e\u041a\u0421").append(["Do not calculate", 999999])
        book.save(path)
        book.close()
        os.utime(path, (self.mtime, self.mtime))
        return path

    def edit_book(self, edits, path=None):
        path = path or self.path
        book = load_workbook(path)
        for cell, value in edits.items():
            book[SHEET][cell] = value
        book.save(path)
        book.close()
        self.mtime += 1
        os.utime(path, (self.mtime, self.mtime))

    def cached_xml(self, cell, value, data_type=None):
        # openpyxl writes formulas but not their cached results. Supply cached
        # XML in this temporary fixture to exercise data_only, NaN and infinity.
        buffer = BytesIO()
        with ZipFile(self.path) as source, ZipFile(buffer, "w") as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(content)
                    cell_node = root.find(f".//{XML_NS}c[@r='{cell}']")
                    if data_type is not None:
                        cell_node.set("t", data_type)
                    node = cell_node.find(f"{XML_NS}v")
                    node.text = value
                    content = ET.tostring(root, encoding="utf-8")
                target.writestr(item, content)
        self.path.write_bytes(buffer.getvalue())
        self.mtime += 1
        os.utime(self.path, (self.mtime, self.mtime))

    def client(self, context=None):
        context = context or self.context
        profile = MagicMock(context=context)
        with patch("pipeline.data_access.DataContext.from_environment", return_value=context):
            app = create_app(profile, MagicMock(), MagicMock())
        self.addCleanup(profile.catalog.assert_not_called)
        return TestClient(app)

    def test_exact_contract_values_cells_and_download_timestamp(self):
        payload = self.service.report()
        self.assertEqual(set(payload), {"schemaVersion", "version", "generatedAt", "unit", "source", "metrics"})
        self.assertEqual(payload["schemaVersion"], 1)
        self.assertEqual(payload["unit"], "million_m2")
        self.assertIsInstance(payload["version"], str)
        self.assertTrue(payload["version"])
        self.assertIsNotNone(datetime.fromisoformat(payload["generatedAt"]).tzinfo)
        self.assertEqual(payload["source"], {
            "file": "data/raw/realty/nashdom/monitoring_2_0_20260928.xlsx",
            "sheet": SHEET, "column": COLUMN, "fileDate": "2026-09-28",
            "downloadedAt": datetime.fromtimestamp(self.mtime, timezone.utc).isoformat(),
        })
        self.assertEqual(payload["metrics"], [
            {"id": identifier, "label": label, "value": value, "cell": f"I{row}"}
            for row, identifier, label, value in ROWS
        ])
        json.dumps(payload, allow_nan=False)

    def test_cached_formula_value_not_formula_or_recalculation(self):
        self.edit_book({"I2": "=999999"})
        self.cached_xml("I2", "51.65728716")
        self.assertEqual(self.service.report()["metrics"][0]["value"], 51.65728716)

    def test_missing_formula_cache_is_null_not_recalculated(self):
        self.edit_book({"I2": "=999999"})
        metric = self.service.report()["metrics"][0]
        self.assertIsNone(metric["value"])
        self.assertEqual(metric["cell"], "I2")

    def test_null_zero_and_whitespace_are_not_conflated(self):
        self.write_book(values={"total": None, "housing": 0, "embedded": " \t "})
        metrics = self.service.report()["metrics"]
        self.assertIsNone(metrics[0]["value"])
        self.assertEqual(metrics[0]["cell"], "I2")
        self.assertEqual(metrics[1]["value"], 0)
        self.assertIsNone(metrics[2]["value"])

    def test_all_blank_is_unavailable_but_all_zero_is_valid(self):
        self.write_book(values={identifier: None for _, identifier, _, _ in ROWS})
        with self.assertRaises(SourceUnavailable):
            self.service.report()
        self.write_book(values={identifier: 0 for _, identifier, _, _ in ROWS})
        self.assertTrue(all(metric["value"] == 0 for metric in self.service.report()["metrics"]))

    def test_malformed_negative_boolean_error_and_date_are_unavailable(self):
        for value in ("invalid", "51.65728716", "1,23", "NaN", "Infinity", -0.01, True,
                      "#DIV/0!", datetime(2026, 1, 1)):
            with self.subTest(value=value):
                self.write_book(values={"total": value})
                with self.assertRaises(SourceUnavailable):
                    self.service.report()

    def test_nonfinite_excel_numbers_are_unavailable(self):
        for value in ("1e309", "-1e309", "NaN"):
            with self.subTest(value=value):
                self.write_book()
                self.cached_xml("I2", value)
                with self.assertRaises(SourceUnavailable):
                    self.service.report()

    def test_invalid_workbook_shared_string_reference_is_503(self):
        self.cached_xml("I2", "123456", data_type="s")
        with self.assertLogs("backend.app", level="ERROR"):
            response = self.client().get("/api/v1/construction/tree")
        self.assertEqual(response.status_code, 503, response.text)

    def test_shifted_header_labels_and_rows_use_actual_cells(self):
        self.write_book(row_offset=5, column=12)
        self.edit_book({"K8": "  \u0436\u0438\u043b\u044c\u0451\u00a0\n", "L6": " \u0438\u0442\u043e\u0433 "})
        payload = self.service.report()
        self.assertEqual([metric["cell"] for metric in payload["metrics"]], [f"L{row + 5}" for row, _, _, _ in ROWS])
        self.assertEqual([metric["value"] for metric in payload["metrics"]], [value for _, _, _, value in ROWS])

    def test_reordered_metrics_are_identified_by_label_not_row_position(self):
        self.edit_book({"H2": ROWS[-1][2], "I2": ROWS[-1][3], "H12": ROWS[0][2], "I12": ROWS[0][3]})
        metrics = self.service.report()["metrics"]
        self.assertEqual(metrics[0]["id"], "total")
        self.assertEqual(metrics[0]["cell"], "I12")
        self.assertEqual(metrics[0]["value"], ROWS[0][3])
        self.assertEqual(metrics[-1]["cell"], "I2")

    def test_missing_renamed_and_duplicate_labels_fail_closed(self):
        for edits in ({"H2": None}, {"H2": "Unknown label"}, {"H20": ROWS[0][2], "I20": 1}):
            with self.subTest(edits=edits):
                self.write_book()
                self.edit_book(edits)
                with self.assertRaises(SourceUnavailable):
                    self.service.report()

    def test_missing_duplicate_and_unadjacent_headers_fail_closed(self):
        for edits in ({"I1": None}, {"I1": "Other"}, {"C1": COLUMN}, {"I20": COLUMN},
                      {"I1": None, "J1": COLUMN}, {"I1": None, "A1": COLUMN}):
            with self.subTest(edits=edits):
                self.write_book()
                self.edit_book(edits)
                with self.assertRaises(SourceUnavailable):
                    self.service.report()

    def test_missing_sheet_and_missing_source_are_unavailable(self):
        self.write_book(sheet=False)
        with self.assertRaises(SourceUnavailable):
            self.service.report()
        self.path.unlink()
        with self.assertRaises(SourceUnavailable):
            self.service.report()

    def test_no_snapshot_or_oks_fallback(self):
        self.path.unlink()
        marts = self.root / "data/marts/realty"
        marts.mkdir(parents=True)
        (marts / "monitoring_2_0.pkl").write_bytes(b"must not be read")
        self.write_book(path=self.directory / "unregistered.xlsx")
        with self.assertRaises(SourceUnavailable):
            self.service.report()

    def test_latest_is_mtime_not_filename_and_refreshes_in_same_process(self):
        before = self.service.report()
        self.mtime += 10
        newer = self.directory / "monitoring_2_0_20200101.xlsx"
        self.write_book(path=newer, values={"total": 123.456789})
        after = self.service.report()
        self.assertNotEqual(before["version"], after["version"])
        self.assertTrue(after["source"]["file"].endswith(newer.name))
        self.assertEqual(after["metrics"][0]["value"], 123.456789)
        self.edit_book({"I2": 42.123456789}, newer)
        updated = self.service.report()
        self.assertNotEqual(after["version"], updated["version"])
        self.assertEqual(updated["metrics"][0]["value"], 42.123456789)

    def test_copy_creation_order_with_preserved_mtime_does_not_select_older_download(self):
        old_mtime = self.mtime - 100
        source = self.write_book(path=Path(self.tmp.name) / "older-download/monitoring_2_0_20261016.xlsx",
                                 values={"total": 999})
        os.utime(source, (old_mtime, old_mtime))
        copied = self.directory / source.name
        shutil.copy2(source, copied)
        self.assertGreaterEqual(copied.stat().st_ctime_ns, self.path.stat().st_ctime_ns)
        self.assertEqual(copied.stat().st_mtime_ns, source.stat().st_mtime_ns)
        payload = self.service.report()
        self.assertTrue(payload["source"]["file"].endswith(self.path.name))
        self.assertEqual(payload["source"]["downloadedAt"], datetime.fromtimestamp(self.mtime, timezone.utc).isoformat())
        self.assertEqual(payload["metrics"][0]["value"], ROWS[0][3])

    def test_newer_invalid_never_falls_back_to_old_or_cached_payload(self):
        self.service.report()
        newer = self.directory / "monitoring_2_0_current.xlsx"
        for invalid in ("missing_sheet", "missing_label", "corrupt", "all_blank"):
            with self.subTest(invalid=invalid):
                self.mtime += 10
                self.write_book(path=newer, sheet=invalid != "missing_sheet",
                                values={identifier: None for _, identifier, _, _ in ROWS} if invalid == "all_blank" else None)
                if invalid == "missing_label":
                    self.edit_book({"H2": None}, newer)
                if invalid == "corrupt":
                    newer.write_bytes(b"not a workbook")
                    os.utime(newer, (self.mtime, self.mtime))
                for _ in range(2):
                    with self.assertRaises(SourceUnavailable):
                        self.service.report()
        self.mtime += 10
        self.write_book(path=newer, values={"total": 12})
        self.assertEqual(self.service.report()["metrics"][0]["value"], 12)

    def test_tied_latest_sources_are_ambiguous(self):
        self.write_book(path=self.directory / "monitoring_2_0_current.xlsx")
        with self.assertRaises(SourceUnavailable):
            self.service.report()

    def test_default_five_second_check_interval_and_error_invalidates_cache(self):
        service = ConstructionTreeService(self.context)
        with patch("backend.construction_tree_service.time.monotonic", return_value=10):
            before = service.report()
        self.edit_book({"I2": 100})
        with patch("backend.construction_tree_service.time.monotonic", return_value=14.99):
            self.assertEqual(before, service.report())
        with patch("backend.construction_tree_service.time.monotonic", return_value=15):
            self.assertEqual(service.report()["metrics"][0]["value"], 100)
        self.edit_book({"H2": None})
        with patch("backend.construction_tree_service.time.monotonic", return_value=20):
            with self.assertRaises(SourceUnavailable):
                service.report()
        with patch("backend.construction_tree_service.time.monotonic", return_value=20.01):
            with self.assertRaises(SourceUnavailable):
                service.report()

    def test_unchanged_generation_is_cached_and_returned_as_deep_copy(self):
        payload = self.service.report()
        with patch("backend.construction_tree_service._read_metrics", side_effect=AssertionError("Unexpected reread")):
            self.assertEqual(payload, self.service.report())
        payload["metrics"][0]["value"] = -1
        payload["source"].clear()
        self.assertEqual(self.service.report()["metrics"][0]["value"], ROWS[0][3])
        self.assertTrue(self.service.report()["source"])

    def test_modified_source_during_parse_is_not_published(self):
        from backend import construction_tree_service as module
        original = module._read_metrics

        def changing_source(path):
            metrics = original(path)
            self.edit_book({"I2": 100}, path)
            return metrics

        self.service.report()
        self.edit_book({"I2": 99})
        with patch.object(module, "_read_metrics", side_effect=changing_source):
            with self.assertRaises(SourceUnavailable):
                self.service.report()
        self.assertIsNone(self.service._payload)
        self.assertEqual(self.service.report()["metrics"][0]["value"], 100)

    def test_new_source_during_parse_is_not_published(self):
        from backend import construction_tree_service as module
        original = module._read_metrics

        def changing_source(path):
            metrics = original(path)
            self.mtime += 10
            self.write_book(path=self.directory / "monitoring_2_0_current.xlsx", values={"total": 100})
            return metrics

        with patch.object(module, "_read_metrics", side_effect=changing_source):
            with self.assertRaises(SourceUnavailable):
                self.service.report()
        self.assertIsNone(self.service._payload)
        self.assertEqual(self.service.report()["metrics"][0]["value"], 100)

    def test_file_date_unknown_or_invalid_is_null_never_modification_date(self):
        self.path.unlink()
        for name in ("monitoring_2_0_current.xlsx", "monitoring_2_0_20260230.xlsx"):
            with self.subTest(name=name):
                path = self.write_book(path=self.directory / name)
                payload = self.service.report()
                self.assertIsNone(payload["source"]["fileDate"])
                self.assertEqual(payload["source"]["downloadedAt"], datetime.fromtimestamp(self.mtime, timezone.utc).isoformat())
                path.unlink()

    def test_discovery_is_only_monitoring_registry_with_supplied_context(self):
        external = self.write_book(path=self.context.downloads / "monitoring_2_0_current.xlsx", values={"total": 456})
        with patch("backend.construction_tree_service.DataAccess", autospec=True) as factory:
            access = factory.return_value
            access.source_files.return_value = [external]
            payload = self.service.report()
            factory.assert_called_once_with(self.context)
            self.assertEqual(access.source_files.call_args_list, [unittest.mock.call("monitoring_2_0")] * 2)
            self.assertEqual(access.method_calls, [unittest.mock.call.source_files("monitoring_2_0")] * 2)
        self.assertEqual(payload["source"]["file"], external.name)
        self.assertEqual(payload["metrics"][0]["value"], 456)

    def test_runtime_environment_context_and_no_other_loaders(self):
        environment = {"PARSER_ROOT": str(self.root), "PARSER_DOWNLOADS": str(self.context.downloads),
                       "PARSER_REQUIRE_REALTY_MARTS": "1"}
        with patch.dict(os.environ, environment):
            client = TestClient(create_app())
        with patch.object(DataAccess, "load_monitoring_2_0", side_effect=AssertionError("No monitoring loader")), \
             patch.object(DataAccess, "load_construction_operational", side_effect=AssertionError("No OKS loader")), \
             patch.object(DataAccess, "load_erzrf_top", side_effect=AssertionError("No ERZ loader")), \
             patch.object(DataAccess, "load_erzrf_cards", side_effect=AssertionError("No ERZ cards loader")), \
             patch.object(DataAccess, "load_rasprodannost", side_effect=AssertionError("No sales loader")):
            response = client.get("/api/v1/construction/tree")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["metrics"][0]["value"], ROWS[0][3])

    def test_api_contract_no_filters_no_export_and_uses_app_context(self):
        other = DataContext(Path(self.tmp.name) / "not-the-injected-context")
        with patch("pipeline.data_access.DataContext.from_environment", return_value=other):
            client = TestClient(create_app(MagicMock(context=self.context), MagicMock(), MagicMock()))
        response = client.get("/api/v1/construction/tree")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.assertEqual(response.json()["metrics"], self.service.report()["metrics"])
        self.assertEqual(client.get("/api/v1/construction/tree", params={"region": "anything", "year": 1}).json(), response.json())
        spec = client.get("/openapi.json").json()["paths"]["/api/v1/construction/tree"]["get"]
        self.assertFalse(spec.get("parameters"))
        self.assertEqual(client.get("/api/v1/construction/tree/export").status_code, 404)

    def test_api_invalid_sources_are_503_not_500_or_previous_payload(self):
        for invalid in ("missing_source", "missing_sheet", "missing_header", "missing_label", "duplicate_label",
                        "duplicate_header", "bad_value", "negative", "all_blank", "nonfinite", "corrupt"):
            with self.subTest(invalid=invalid):
                self.write_book(sheet=invalid != "missing_sheet")
                edits = {"missing_header": {"I1": None}, "missing_label": {"H2": None},
                         "duplicate_label": {"H20": ROWS[0][2]}, "duplicate_header": {"C1": COLUMN},
                         "bad_value": {"I2": "bad"}, "negative": {"I2": -1},
                         "all_blank": {f"I{row}": None for row, _, _, _ in ROWS}}
                if invalid in edits:
                    self.edit_book(edits[invalid])
                if invalid == "missing_source":
                    self.path.unlink()
                if invalid == "nonfinite":
                    self.cached_xml("I2", "1e309")
                if invalid == "corrupt":
                    self.path.write_bytes(b"not a workbook")
                with self.assertLogs("backend.app", level="ERROR"):
                    response = self.client().get("/api/v1/construction/tree")
                self.assertEqual(response.status_code, 503, response.text)
                self.assertEqual(response.headers["Cache-Control"], "no-store")

    def test_api_same_process_refresh_failure_is_503_until_source_is_repaired(self):
        with patch("backend.construction_tree_routes.ConstructionTreeService", return_value=self.service):
            client = self.client()
        before = client.get("/api/v1/construction/tree")
        self.assertEqual(before.status_code, 200)
        self.edit_book({"I2": 100})
        updated = client.get("/api/v1/construction/tree")
        self.assertEqual(updated.status_code, 200)
        self.assertNotEqual(before.json()["version"], updated.json()["version"])
        self.edit_book({"H2": None})
        with self.assertLogs("backend.app", level="ERROR"):
            for _ in range(2):
                self.assertEqual(client.get("/api/v1/construction/tree").status_code, 503)
        self.edit_book({"H2": ROWS[0][2]})
        self.assertEqual(client.get("/api/v1/construction/tree").status_code, 200)

    def assert_corrupt_newer_source_invalidates_cache(self, corruption):
        service = ConstructionTreeService(self.context)
        with patch("backend.construction_tree_routes.ConstructionTreeService", return_value=service):
            client = self.client()
        with patch("backend.construction_tree_service.time.monotonic", return_value=10):
            before = client.get("/api/v1/construction/tree")
        self.assertEqual(before.status_code, 200, before.text)
        self.assertIsNotNone(service._payload)
        newer = self.directory / "monitoring_2_0_current.xlsx"
        shutil.copy2(self.path, newer)
        if corruption == "deflate":
            with ZipFile(newer) as archive:
                info = archive.getinfo("[Content_Types].xml")
            self.assertEqual(info.compress_type, ZIP_DEFLATED)
            content = bytearray(newer.read_bytes())
            name_length, extra_length = struct.unpack_from("<HH", content, info.header_offset + 26)
            # BTYPE=3 is a reserved, invalid raw DEFLATE block type.
            content[info.header_offset + 30 + name_length + extra_length] = 0x07
            newer.write_bytes(content)
        else:
            buffer = BytesIO()
            with ZipFile(newer) as source, ZipFile(buffer, "w") as target:
                for item in source.infolist():
                    content = b"<Types>" if item.filename == "[Content_Types].xml" else source.read(item.filename)
                    target.writestr(item, content)
            newer.write_bytes(buffer.getvalue())
        self.mtime += 10
        os.utime(newer, (self.mtime, self.mtime))
        for now in (15, 15.01):
            with patch("backend.construction_tree_service.time.monotonic", return_value=now), \
                 self.assertLogs("backend.app", level="ERROR"):
                response = client.get("/api/v1/construction/tree")
            self.assertEqual(response.status_code, 503, response.text)
            self.assertEqual(response.headers["Cache-Control"], "no-store")
            self.assertNotIn("metrics", response.json())
            self.assertIsNone(service._payload)
            self.assertIsNone(service._signature)
            self.assertEqual(service._checked, 0)
        self.mtime += 1
        self.write_book(path=newer, values={"total": 123})
        with patch("backend.construction_tree_service.time.monotonic", return_value=15.02):
            repaired = client.get("/api/v1/construction/tree")
        self.assertEqual(repaired.status_code, 200, repaired.text)
        self.assertEqual(repaired.json()["metrics"][0]["value"], 123)
        self.assertNotEqual(repaired.json()["version"], before.json()["version"])

    def test_corrupt_newer_deflate_is_503_and_invalidates_good_cached_payload(self):
        self.assert_corrupt_newer_source_invalidates_cache("deflate")

    def test_corrupt_newer_content_types_xml_is_503_and_invalidates_good_cached_payload(self):
        self.assert_corrupt_newer_source_invalidates_cache("xml")

    def test_service_import_and_standard_xml_errors_without_optional_lxml(self):
        original_import = builtins.__import__

        def without_lxml(name, *args, **kwargs):
            if name == "lxml" or name.startswith("lxml."):
                raise ImportError("Optional lxml is not installed")
            return original_import(name, *args, **kwargs)

        path = Path(__file__).resolve().parents[1] / "construction_tree_service.py"
        with patch("builtins.__import__", side_effect=without_lxml):
            module = runpy.run_path(str(path))
        service = module["ConstructionTreeService"](self.context, check_interval=0)
        self.assertEqual(service.report()["metrics"][0]["value"], ROWS[0][3])
        service_globals = service.report.__func__.__globals__
        with patch.dict(service_globals, load_workbook=MagicMock(side_effect=ET.ParseError("Invalid XML"))):
            self.edit_book({"I2": 100})
            with self.assertRaises(SourceUnavailable):
                service.report()
        self.assertIsNone(service._payload)

    def test_route_availability_requires_monitoring_not_construction_or_erz(self):
        for path in ("/api/v1/construction/tree", "/api/v1/construction/tree/"):
            self.assertEqual(availability.route_sources(path), {"monitoring"})
        self.assertEqual(availability.route_sources("/api/v1/construction"), {"construction"})
        self.assertEqual(availability.route_sources("/api/v1/construction/export"), {"construction"})
        with patch.object(availability, "unavailable", return_value={"construction", "erz-top", "erz-cards"}):
            client = self.client()
            self.assertEqual(client.get("/api/v1/construction/tree").status_code, 200)
        publication = {"sources": {"monitoring": {"reason": "Monitoring unavailable"}}}
        with patch.object(availability, "unavailable", return_value={"monitoring"}), \
             patch.object(availability, "read_publication", return_value=publication):
            client = self.client()
            response = client.get("/api/v1/construction/tree")
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()["detail"]["sources"], ["monitoring"])


if __name__ == "__main__":
    unittest.main()
