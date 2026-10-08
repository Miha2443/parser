"""Compare operational payloads with execution of the original page, not a reimplementation."""
import ast
from io import BytesIO
from pathlib import Path
import unittest

import pandas as pd
import plotly.graph_objects as go
from openpyxl import load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.vvod_operational import housing_ytd_table, nonres_ytd_table, quarter_tree
from backend.operational_routes import build_report, create_service, create_router, export_workbook
from backend.profile_service import FilterUnavailable
from pipeline.data_access import DataAccess, DataContext
from pipeline.operational_calculations import catalog
from pipeline.report_values import records
from scripts.check_apartments_parity import Capture


ROOT = Path(__file__).resolve().parents[2]


class OperationalCapture(Capture):
    def __init__(self, month, exclude_mkd, year, quarter, cumulative):
        super().__init__("msk")
        self.selection = {"Период с начала года": month, "Год": year, "Квартал": quarter,
                          "Состав показателя": "Без нежилых помещений в жилых объектах" if exclude_mkd else "Всё нежильё",
                          "Расчёт периода": "С начала года" if cumulative else "За квартал"}

    def selectbox(self, label, options, **kwargs):
        value = self.selection[label]
        assert value in options
        return value

    def radio(self, label, options, **kwargs):
        return self.selectbox(label, options)


class OperationalTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.context = DataContext(ROOT)
        access = DataAccess(cls.context)
        cls.loaders = ("load_vvod_static", "load_monitoring_2_0", "load_monitoring_operational_history")
        cls.before = tuple(access.data_version(name) for name in cls.loaders)
        cls.data = {name: getattr(access, name)() for name in cls.loaders}
        cls.options = catalog(cls.data)

    def capture(self, month, exclude, year, quarter, cumulative):
        path = ROOT / "app/pages/8_Ввод_недвижимости_оперативные.py"
        tree = ast.parse(path.read_text(encoding="utf-8"))
        tree.body = [node for node in tree.body if not isinstance(node, (ast.Import, ast.ImportFrom))]
        ui = OperationalCapture(month, exclude, year, quarter, cumulative)
        env = {"st": ui, "pd": pd, "go": go,
               "COLORS": {key: "#333333" for key in ("red", "green")},
               "apply_theme": lambda: None, "page_header": lambda *args: None,
               "style_plotly": lambda *args, **kwargs: None,
               "chart_download_button": lambda *args, **kwargs: None,
               "table_download_buttons": lambda *args, **kwargs: None,
               "housing_ytd_table": housing_ytd_table, "nonres_ytd_table": nonres_ytd_table,
               "quarter_tree": quarter_tree,
               **{name: lambda value=value: value for name, value in self.data.items()}}
        exec(compile(tree, str(path), "exec"), env)
        return ui, env

    def test_original_page_parity(self):
        months = sorted({self.options["months"][0]["id"], self.options["defaultMonth"], self.options["months"][-1]["id"]})
        year = self.options["currentYear"]
        for month in months:
            for exclude in (False, True):
                for cumulative in (False, True):
                    with self.subTest(month=month, exclude=exclude, cumulative=cumulative):
                        payload = build_report(self.data, self.options, month=month, exclude_mkd=exclude,
                                               year=year, quarter=2, cumulative=cumulative)
                        ui, env = self.capture(month, exclude, year, 2, cumulative)
                        for table, name, figure in zip(payload["tables"], ("housing", "nonres"), ui.figures):
                            self.assertEqual(table["rows"], records(env[name]))
                            self.assertEqual(table["chart"]["x"], list(figure.data[0].x))
                            for field, trace in (("period", 0), ("remainder", 1), ("totals", 2)):
                                self.assertEqual(table["chart"][field], records(pd.DataFrame({"value": figure.data[trace].y})))
                        self.assertEqual(payload["tree"], env["tree"])
        access = DataAccess(self.context)
        self.assertEqual(self.before, tuple(access.data_version(name) for name in self.loaders))

    def test_filters_and_exports(self):
        with self.assertRaises(FilterUnavailable):
            build_report(self.data, self.options, year=2000)
        service = create_service(self.context, check_interval=0)
        payload = service.report()
        book = load_workbook(BytesIO(export_workbook(payload)))
        self.assertEqual(book["housing"].max_row, len(payload["tables"][0]["rows"]) + 1)
        self.assertEqual(book["Structure"].max_row, len(payload["treeRows"]) + 1)
        copy = service.catalog()
        copy["months"].clear()
        self.assertTrue(service.catalog()["months"])
        app = FastAPI()
        app.include_router(create_router(self.context, lambda function, *args, **kwargs: function(*args), service))
        client = TestClient(app)
        self.assertEqual(client.get("/api/v1/commissioning/operational/export").status_code, 422)
        self.assertEqual(client.get("/api/v1/commissioning/operational/export", params={"required_version": "old"}).status_code, 409)


if __name__ == "__main__":
    unittest.main()
