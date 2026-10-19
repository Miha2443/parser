"""Independent Moscow permits generation, retaining the legacy main calculation."""
from __future__ import annotations

import copy
from dataclasses import replace
from datetime import datetime, timezone
import hashlib
import json
import re
import threading
import time
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
from zlib import error as DecompressionError

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from backend.profile_service import FilterUnavailable, SourceUnavailable
from pipeline.construction_calculations import catalog, report
from pipeline.data_access import DataAccess, DataContext, file_signature
from pipeline.permit_page_calculations import (OKS_SHEET, breakdowns, read_breakdown_sources,
                                               read_forecast, select_forecast)


_SOURCE_ERRORS = (OSError, ValueError, TypeError, LookupError, RuntimeError, OverflowError,
                  BadZipFile, InvalidFileException, ParseError, EOFError, DecompressionError)
try:
    from lxml.etree import XMLSyntaxError
except ImportError:
    pass
else:
    _SOURCE_ERRORS += (XMLSyntaxError,)


class _PermitsAccess(DataAccess):
    """Use the unchanged loader without its unrelated JSON or mart providers."""

    def __init__(self, context, monitoring):
        super().__init__(replace(context, use_marts=False, require_marts=False))
        self.CONSTRUCTION_OPERATIONAL_PATHS = []
        self.monitoring = monitoring

    def _realty_mart_manifest_entry(self, *args, **kwargs):
        return None

    def load_monitoring_2_0(self, **kwargs):
        # Only these five OKS columns feed the unchanged operational loader.
        # Match its legacy missing -> 0 normalization for the main table only;
        # supplemental calculations separately read the untouched source cells.
        frame = pd.read_excel(self.monitoring, sheet_name=OKS_SHEET,
                              usecols=lambda column: column in {
                                  "Назначение", "Срок выдачи РС", "Год выдачи", "Общая площадь", "Жилая площадь"})
        for column in ("Общая площадь", "Жилая площадь"):
            values = frame[column].map(lambda value: value.replace("\xa0", "").replace(" ", "").replace(",", ".")
                                       if isinstance(value, str) else value)
            frame[column] = pd.to_numeric(values, errors="coerce").fillna(0)
        return {"oks": frame}


def _date(path):
    match = re.search(r"_(\d{8})$", path.stem)
    if match:
        try:
            return datetime.strptime(match[1], "%Y%m%d").date().isoformat()
        except ValueError:
            pass
    return None


class PermitsPageService:
    def __init__(self, context: DataContext, check_interval: float = 5):
        self.context = context
        self.check_interval = check_interval
        self._lock = threading.RLock()
        self._checked = 0.0
        self._signature = None
        self._generation = None

    def _inputs(self):
        access = DataAccess(replace(self.context, use_marts=False, require_marts=False))
        signatures = [(path, file_signature(path)) for path in access.source_files("monitoring_2_0")]
        if not signatures or any(signature[1] is None for _, signature in signatures):
            raise SourceUnavailable("Monitoring workbook is missing")
        newest = max(signature[2] for _, signature in signatures)
        selected = [(path, signature) for path, signature in signatures if signature[2] == newest]
        if len(selected) != 1:
            raise SourceUnavailable("Latest monitoring workbook is ambiguous")
        base = access._vvod_dir()
        static = base / "static_vvod_rs_2011_2025.xlsx" if base else None
        used_static = [static] if static is not None and static.is_file() else []
        path, signature = selected[0]
        return path, tuple([signature] + [file_signature(item) for item in used_static]), [path] + used_static

    def _relative(self, path):
        try:
            return path.resolve().relative_to(self.context.root).as_posix()
        except ValueError:
            return path.name

    def _refresh(self):
        now = time.monotonic()
        if self._generation is not None and now - self._checked < self.check_interval:
            return
        path, signature, files = self._inputs()
        if signature == self._signature:
            self._checked = now
            return
        with path.open("rb") as stream:
            book = load_workbook(stream, read_only=True, data_only=True)
            try:
                if OKS_SHEET not in book.sheetnames:
                    raise SourceUnavailable("Primary permit registry is missing")
                headers = {cell.value for cell in next(book[OKS_SHEET].iter_rows(max_row=1))}
                if not {"Назначение", "Срок выдачи РС", "Год выдачи", "Общая площадь", "Жилая площадь"} <= headers:
                    raise SourceUnavailable("Primary permit columns are missing")
                forecast = read_forecast(book, self._relative(path))
                detailed = read_breakdown_sources(book)
            finally:
                book.close()
        access = _PermitsAccess(self.context, path)
        operational = access.load_construction_operational()
        if operational["permits"].empty:
            raise SourceUnavailable("Monitoring contains no permit observations")
        data = {"load_construction_operational": operational, "load_rasprodannost": {}}
        controls = catalog(data)
        controls = {"permitKinds": sorted(controls["permitKinds"], key=lambda row:
                                         ("housing", "nonresidential", "total").index(row["id"])),
                    "monthsByKind": controls["monthsByKind"]}
        issues = [{"level": level, "message": message} for level, message in access.issues]
        if any(row["level"] == "error" for row in issues):
            raise SourceUnavailable("Permit loader reported a source error")
        if not forecast["rows"]:
            issues.append({"level": "warning", "message": "Прогноз недоступен; см. forecast.notes."})
        payload = {"schemaVersion": 1, "version": hashlib.sha256(repr(signature).encode()).hexdigest(),
                   "generatedAt": datetime.now(timezone.utc).isoformat(),
                   "source": {"date": _date(path), "files": [self._relative(item) for item in files],
                              "fileEvidence": "selected", "issues": issues},
                   "controls": controls}
        # Build every selection before publishing so a partial generation is never visible.
        reports = {}
        for kind, months in controls["monthsByKind"].items():
            for month in months:
                reports[(kind, month)] = {
                    **payload, "selection": {"permitKind": kind, "month": month},
                    "permits": report(data, controls, permit_kind=kind, month=month)["permits"],
                    "forecast": select_forecast(forecast, kind), "breakdowns": breakdowns(detailed, kind, month),
                    "notes": ["Разрешения относятся только к Москве. Основная таблица и график сохраняют прежний "
                              "расчёт: жилая площадь жилья плюс общая площадь нежилья, без фильтра действующего статуса РС.",
                              "Прогноз и детальная структура не изменяют фактические годовые значения, остаток и рост.",
                              "Исторические структурные ряды имеют собственную исходную базу; расхождения с основной "
                              "таблицей не исправляются масштабированием."]}
        json.dumps(list(reports.values()), allow_nan=False)
        if self._inputs()[1] != signature:
            raise SourceUnavailable("Permit sources changed during loading")
        self._generation, self._signature, self._checked = reports, signature, now

    def report(self, permit_kind="total", month=None):
        if permit_kind not in ("housing", "nonresidential", "total") or (month is not None and
                (isinstance(month, bool) or not isinstance(month, int) or not 1 <= month <= 12)):
            raise FilterUnavailable("Permit filter is unavailable")
        with self._lock:
            try:
                self._refresh()
            except _SOURCE_ERRORS as exc:
                self._generation = self._signature = None
                self._checked = 0.0
                raise SourceUnavailable("Permit sources are unavailable") from exc
            months = next(iter(self._generation.values()))["controls"]["monthsByKind"][permit_kind]
            month = months[-1] if month is None else month
            if (permit_kind, month) not in self._generation:
                raise FilterUnavailable("Permit month is unavailable")
            return copy.deepcopy(self._generation[(permit_kind, month)])
