"""Focused workbook, calendar and weighted-population regression tests."""
import copy
from datetime import datetime
import math
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from pipeline.preferential_mortgage import (CHAR_METRICS, CHAR_SOURCE_METRICS, PROGRAMS, SOURCE_URL,
                                            build_report, parse_workbook, region_periods, validate_source)


def flow(period, count=2, volume=10, program="family", region="msk", basis="month"):
    return {"region": region, "program": program, "period": period, "basis": basis,
            "count": count, "volumeMln": volume}


def characteristic(period, program="family", region="msk", basis="month", **values):
    return {"region": region, "program": program, "period": period, "basis": basis,
            **dict.fromkeys(CHAR_SOURCE_METRICS), **values}


def source(flows=None, characteristics=None):
    return {"schemaVersion": 1, "sourceUrl": SOURCE_URL, "sourceDate": "2026-10-11",
            "collectedAt": "2026-10-23T12:00:00+03:00", "notes": [],
            "methodology": {"ratePct": {"weight": "Сумма кредита, срок кредита"},
                            "termMonths": {"weight": "Сумма кредита"},
                            "ownFundsPct": {"weight": "Стоимость квартиры"}},
            "flows": [flow("2026-01")] if flows is None else flows,
            "characteristics": [characteristic("2026-01", averageAmountMln=5, ratePct=6)]
            if characteristics is None else characteristics}


def weighted_source():
    flows, chars = [], []
    for month in range(1, 13):
        count, amount, rate, term, own, area, price = (2, 5, 2, 100, 50, 20, 10) if month == 1 else (3, 10, 8, 300, 20, 50, 20)
        volume = count * amount
        period = f"2025-{month:02d}"
        flows.append(flow(period, count, volume))
        chars.append(characteristic(period, averageAmountMln=amount, ratePct=rate, termMonths=term,
                                    ownFundsPct=own, areaM2=area, propertyAmountMln=price,
                                    priceThousandM2=price / area * 1000))
    annual_count = sum(r["count"] for r in flows)
    annual_volume = sum(r["volumeMln"] for r in flows)
    weights = {
        "averageAmountMln": [r["count"] for r in flows], "areaM2": [r["count"] for r in flows],
        "termMonths": [r["volumeMln"] for r in flows],
        "ratePct": [f["volumeMln"] * c["termMonths"] for f, c in zip(flows, chars)],
        "ownFundsPct": [f["count"] * c["propertyAmountMln"] for f, c in zip(flows, chars)],
        "priceThousandM2": [f["count"] * c["areaM2"] for f, c in zip(flows, chars)],
        "propertyAmountMln": [r["count"] for r in flows],
    }
    annual = {metric: sum(c[metric] * weight for c, weight in zip(chars, weight_list)) / sum(weight_list)
              for metric, weight_list in weights.items()}
    flows.append(flow("2025", annual_count, annual_volume, basis="annual"))
    chars.append(characteristic("2025", basis="annual", **annual))
    return source(flows, chars)


def partial_source():
    document = weighted_source()
    for row in document["flows"] + document["characteristics"]:
        row["period"] = row["period"].replace("2025", "2026")
    for collection in ("flows", "characteristics"):
        document[collection] = [row for row in document[collection]
                                if row["basis"] == "annual" or row["period"] <= "2026-10"]
    for row in document["characteristics"]:
        if row["basis"] == "month" and row["period"] > "2026-07":
            row.update(dict.fromkeys((*CHAR_METRICS[1:], "propertyAmountMln")))
        elif row["basis"] == "annual":
            row.update(ownFundsPct=30.0552221846817, ratePct=5.87903022732177)
    document["notes"] = [
        "Семейная ипотека: до июля 2026 г. – ежемесячные отчеты по возмещению, с августа 2026 г. – ежемесячные опросы и еженедельные отчеты",
        "IT ипотека: до июля 2026 г. – ежемесячные отчеты по возмещению, с августа 2026 г. – еженедельные отчеты",
        "Параметры кредитов приведены по отчетам о возмещении (в оперативной статистике отсутствует).",
    ]
    return document


def fixture_workbook(path):
    book = Workbook()
    overview = book.active
    overview.title = "01_02_00"
    overview.cell(13, 2, "1. Данные актуальны по состоянию на 11.10.2026.")
    for row, label, weight in ((33, "Процентная ставка", "Сумма кредита, срок кредита"),
                                (34, "Доля собственных средств", "Стоимость квартиры"),
                                (35, "Срок кредита", "Сумма кредита")):
        overview.cell(row, 2, label)
        overview.cell(row, 3, weight)
    flows = book.create_sheet("01_02_01")
    flows.append([])
    flows.append([])
    flows.append([])
    flows.append(["Программы", "Размерность величин", "ИТОГО (за весь период)",
                  *[f"Итого {year} год" for year in range(2018, 2027)],
                  datetime(2026, 1, 1), datetime(2026, 2, 1), datetime(2026, 10, 1)])
    for label, count in (("Российская Федерация", 4), ("Московская область", 999), ("г. Москва", 2)):
        flows.append([label])
        flows.append(["Все программы", "Всего, шт.", 9999])
        for program in [*PROGRAMS.values(), "Дальневосточная и арктическая ипотека"]:
            flows.append([" " + program, "Всего, шт.", 100, *([0] * 8), 42.5, count, 0, 0.5])
            flows.append([" " + program, "Всего, млн руб.", 500, *([0] * 8), 101.75, count * 5, 0, 2.5])
    chars = book.create_sheet("01_02_05")
    chars.append([])
    chars.append([])
    chars.append([])
    chars.append(["Программы", "ИТОГО (за весь период)", *[f"Итого {y} год" for y in range(2018, 2027)],
                  datetime(2026, 1, 1), datetime(2026, 2, 1), datetime(2026, 10, 1)])
    labels = ["Средняя сумма кредита, млн руб.", "Средневзвешенный размер текущей процентной ставки, %",
              "Размер собственных средств заемщика, %", "Средний срок кредита, мес.",
              "Средняя стоимость помещения, млн руб.", "Средняя площадь, м2", "Средняя стоимость 1 м2, тыс. руб."]
    for region in ("Российская Федерация", "Московская область", "г. Москва"):
        chars.append([region])
        for program in ["Все программы", *PROGRAMS.values(), "Дальневосточная и арктическая ипотека"]:
            chars.append([program])
            for index, label in enumerate(labels):
                value = (5, 6, 30, 240, 10, 50, 200)[index]
                chars.append(["   " + label, value, *([None] * 8), value, value, None,
                              value if index == 0 else None])
    book.save(path)


class PreferentialWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "source.xlsx"
        fixture_workbook(self.path)

    def test_extraction_headers_units_scoping_precision_and_provenance(self):
        before = self.path.read_bytes()
        document = parse_workbook(self.path)
        self.assertEqual(document["sourceDate"], "2026-10-11")
        self.assertEqual(set(r["region"] for r in document["flows"]), {"msk", "rf"})
        self.assertEqual(set(r["program"] for r in document["characteristics"]), set(PROGRAMS))
        monthly = next(r for r in document["flows"] if r["region"] == "msk" and r["program"] == "family" and r["period"] == "2026-10")
        self.assertEqual((monthly["count"], monthly["volumeMln"]), (0.5, 2.5))
        self.assertTrue(monthly["sourceCells"]["count"].startswith("O"))
        annual = next(r for r in document["flows"] if r["region"] == "msk" and r["program"] == "family" and r["period"] == "2026")
        self.assertEqual(annual["count"], 42.5)
        self.assertTrue(annual["sourceCells"]["count"].startswith("L"))
        january = next(r for r in document["characteristics"] if r["region"] == "msk" and r["program"] == "family" and r["period"] == "2026-01")
        self.assertEqual(january["propertyAmountMln"], 10)
        self.assertTrue(january["sourceCells"]["ratePct"].startswith("L"))
        february = next(r for r in document["flows"] if r["period"] == "2026-02")
        self.assertEqual(february["count"], 0)
        char_feb = next(r for r in document["characteristics"] if r["period"] == "2026-02")
        self.assertIsNone(char_feb["ratePct"])
        self.assertEqual(document["methodology"]["ratePct"]["weight"], "Сумма кредита, срок кредита")
        self.assertEqual(len(document["sourceWorkbook"]["sha256"]), 64)
        self.assertEqual(before, self.path.read_bytes())

    def test_missing_sheet_wrong_units_duplicate_rows_and_duplicate_period_fail_closed(self):
        from openpyxl import load_workbook
        for change in ("sheet", "unit", "row", "period"):
            with self.subTest(change=change):
                fixture_workbook(self.path)
                book = load_workbook(self.path)
                if change == "sheet":
                    del book["01_02_05"]
                elif change == "unit":
                    book["01_02_01"].cell(7, 2, "Всего, тыс. руб.")
                elif change == "row":
                    book["01_02_01"].insert_rows(8)
                    for col in range(1, 16):
                        book["01_02_01"].cell(8, col, book["01_02_01"].cell(7, col).value)
                else:
                    book["01_02_01"].cell(4, 14, datetime(2026, 1, 1))
                book.save(self.path)
                book.close()
                with self.assertRaises(ValueError):
                    parse_workbook(self.path)


class PreferentialCalculationTests(unittest.TestCase):
    def test_annual_is_source_not_monthly_sum_and_all_years_are_present(self):
        document = source([flow("2026-01", count=0, volume=0), flow("2026-10", count=0.5, volume=2.5),
                           flow("2026", count=42.5, volume=101.75, basis="annual")])
        report = build_report(document, "msk", 12, 12)
        self.assertEqual([r["year"] for r in report["annual"]], list(range(2018, 2027)))
        self.assertEqual(report["annual"][-1]["programs"]["family"]["count"], 42.5)
        self.assertEqual(report["annual"][-1]["throughMonth"], 10)
        self.assertEqual(report["cumulative"][-1]["throughMonth"], 10)
        self.assertIsNone(report["cumulative"][-1]["programs"]["family"]["count"])
        self.assertIsNone(report["annual"][-1]["total"]["count"])
        self.assertTrue(any("неполный" in note for note in report["notes"]))

    def test_totals_sum_only_three_programs_strictly_and_preserve_null_zero(self):
        document = source([flow("2026-01", count=0, volume=0, program=p) for p in PROGRAMS])
        report = build_report(document, "msk", 1, 1)
        self.assertEqual(report["cumulative"][-1]["total"], {"volumeMln": 0, "count": 0})
        document["flows"][0]["count"] = None
        report = build_report(document, "msk", 1, 1)
        self.assertIsNone(report["cumulative"][-1]["total"]["count"])
        self.assertEqual(report["cumulative"][-1]["total"]["volumeMln"], 0)

    def test_integer_flow_precision_and_overflow_remain_json_safe(self):
        import json
        integer = 2 ** 53 + 1
        document = source([flow(f"2026-{month:02d}", count=integer, volume=1e308) for month in (1, 2)])
        report = build_report(document, "msk", 2, 2)
        values = report["cumulative"][-1]["programs"]["family"]
        self.assertEqual(values["count"], integer * 2)
        self.assertIsInstance(values["count"], int)
        self.assertIsNone(values["volumeMln"])
        json.dumps(report, allow_nan=False)

    def test_partial_month_uses_actual_calendar_end_and_unknown_date_is_not_complete(self):
        document = source([flow("2026-12")])
        document["sourceDate"] = "2026-12-29"
        self.assertTrue(build_report(document, "msk", 12, 12)["latestPeriodPartial"])
        document["sourceDate"] = "2026-12-31"
        self.assertFalse(build_report(document, "msk", 12, 12)["latestPeriodPartial"])
        document["sourceDate"] = None
        self.assertIsNone(build_report(document, "msk", 12, 12)["latestPeriodPartial"])

    def test_null_only_characteristics_do_not_hide_observed_issuance(self):
        document = validate_source(source(characteristics=[characteristic("2026-01")]))
        self.assertEqual(region_periods(document, characteristics=True)["msk"], [])
        report = build_report(document, "msk", 1, 1)
        self.assertEqual(report["cumulative"][-1]["programs"]["family"]["count"], 2)
        self.assertTrue(all(v is None for v in report["characteristicsAnnual"][-1]["programs"]["family"].values()))

    def test_verified_weighted_formulas_are_not_unweighted_means(self):
        document = validate_source(weighted_source())
        values = build_report(document, "msk", 2, 2)["characteristicsCumulative"][-1]["programs"]["family"]
        self.assertEqual(values["averageAmountMln"], 8)
        self.assertAlmostEqual(values["ratePct"], (2 * 10 * 100 + 8 * 30 * 300) / (10 * 100 + 30 * 300))
        self.assertEqual(values["termMonths"], 250)
        self.assertEqual(values["ownFundsPct"], (50 * 2 * 10 + 20 * 3 * 20) / (2 * 10 + 3 * 20))
        self.assertEqual(values["areaM2"], 38)
        self.assertEqual(values["priceThousandM2"], 1000 * (2 * 10 + 3 * 20) / (2 * 20 + 3 * 50))
        self.assertNotEqual(values["ratePct"], 5)

    def test_failed_annual_reconciliation_withholds_only_affected_metric(self):
        document = weighted_source()
        document["characteristics"][-1]["ownFundsPct"] += 0.001
        report = build_report(document, "msk", 7, 7)
        values = report["characteristicsCumulative"][-1]["programs"]["family"]
        self.assertIsNone(values["ownFundsPct"])
        self.assertIsNotNone(values["ratePct"])
        self.assertTrue(any("ownFundsPct" in note and "2025" in note for note in report["notes"]))
        full = build_report(document, "msk", 12, 12)
        self.assertEqual(full["characteristicsCumulative"][-1]["programs"]["family"]["ownFundsPct"],
                         document["characteristics"][-1]["ownFundsPct"])
        self.assertEqual(report["characteristicsAnnual"][-1]["programs"]["family"]["ownFundsPct"],
                         document["characteristics"][-1]["ownFundsPct"])

    def test_exact_current_source_window_uses_published_values_without_reconstruction(self):
        document = partial_source()
        document["methodology"] = {}
        july = build_report(document, "msk", 7, 7)["characteristicsCumulative"][-1]
        self.assertEqual(july["programs"]["family"]["ownFundsPct"], 30.0552221846817)
        self.assertEqual(july["programs"]["family"]["ratePct"], 5.87903022732177)
        self.assertEqual(july["programThroughMonths"]["family"]["ownFundsPct"], 7)
        june = build_report(document, "msk", 6, 6)["characteristicsCumulative"][-1]
        self.assertIsNone(june["programs"]["family"]["ownFundsPct"])
        self.assertIsNone(june["programs"]["family"]["ratePct"])
        october = build_report(document, "msk", 10, 10)["characteristicsCumulative"][-1]
        self.assertEqual(october["programs"]["family"]["ownFundsPct"], 30.0552221846817)
        self.assertEqual(october["programThroughMonths"]["family"]["ownFundsPct"], 7)

    def test_individual_trailing_blank_does_not_confirm_a_shorter_source_window(self):
        document = partial_source()
        for row in document["characteristics"]:
            if row["period"] == "2026-07":
                row["ratePct"] = None
        values = build_report(document, "msk", 7, 7)["characteristicsCumulative"][-1]["programs"]["family"]
        self.assertIsNone(values["ratePct"])
        self.assertEqual(values["ownFundsPct"], 30.0552221846817)

    def test_removing_multiple_july_fields_cannot_reinterpret_published_aggregate_as_june(self):
        document = partial_source()
        for row in document["characteristics"]:
            if row["period"] == "2026-07":
                row.update(dict.fromkeys((*CHAR_METRICS[1:], "propertyAmountMln")))
        july_flow = next(row for row in document["flows"] if row["period"] == "2026-07")
        self.assertGreater(july_flow["volumeMln"], 0)
        for month in (6, 7):
            values = build_report(document, "msk", month, month)["characteristicsCumulative"][-1]
            self.assertIsNone(values["programs"]["family"]["ownFundsPct"])
            self.assertIsNone(values["programs"]["family"]["ratePct"])

    def test_lagged_source_reuse_requires_both_dated_transition_and_parameter_population_note(self):
        for keep in ([], [0], [2], [1, 2]):
            document = partial_source()
            document["notes"] = [document["notes"][index] for index in keep]
            values = build_report(document, "msk", 7, 7)["characteristicsCumulative"][-1]["programs"]["family"]
            self.assertIsNone(values["ownFundsPct"])
            self.assertIsNone(values["ratePct"])

    def test_historical_stopped_program_can_reuse_source_only_without_later_issuance(self):
        document = weighted_source()
        for row in document["characteristics"]:
            if row["basis"] == "month" and row["period"] > "2025-10":
                row.update(dict.fromkeys(CHAR_SOURCE_METRICS))
            elif row["basis"] == "annual":
                row["ownFundsPct"] = 31.23456789
        unknown = build_report(document, "msk", 10, 10)["characteristicsCumulative"][-1]
        self.assertIsNone(unknown["programs"]["family"]["ownFundsPct"])
        for row in document["flows"]:
            if row["basis"] == "month" and row["period"] > "2025-10":
                row.update(count=0, volumeMln=0)
        for month in (10, 11):
            values = build_report(document, "msk", month, month)["characteristicsCumulative"][-1]
            self.assertEqual(values["programs"]["family"]["ownFundsPct"], 31.23456789)
            self.assertEqual(values["programThroughMonths"]["family"]["ownFundsPct"], 10)
        september = build_report(document, "msk", 9, 9)["characteristicsCumulative"][-1]
        self.assertIsNone(september["programs"]["family"]["ownFundsPct"])
        for replacement in ({"count": 0, "volumeMln": None}, {"count": None, "volumeMln": 0}, None):
            incomplete = copy.deepcopy(document)
            december = next(row for row in incomplete["flows"] if row["period"] == "2025-12")
            if replacement is None:
                incomplete["flows"].remove(december)
            else:
                december.update(replacement)
            values = build_report(incomplete, "msk", 10, 10)["characteristicsCumulative"][-1]
            self.assertIsNone(values["programs"]["family"]["ownFundsPct"])

    def test_published_window_match_is_program_specific(self):
        document = partial_source()
        it_flows = [dict(row, program="it") for row in document["flows"]]
        it_chars = [dict(row, program="it") for row in document["characteristics"]]
        for row in it_chars:
            if row["basis"] == "annual":
                row["ownFundsPct"] = 45.123456789
        document["flows"].extend(it_flows)
        document["characteristics"].extend(it_chars)
        document["notes"][0] = document["notes"][0].replace("с августа", "с сентября")
        for row in document["characteristics"]:
            if row["program"] == "family" and row["period"] == "2026-08":
                row.update(dict(averageAmountMln=10, ratePct=8, termMonths=300, ownFundsPct=20,
                                areaM2=50, propertyAmountMln=20, priceThousandM2=400))
        values = build_report(document, "msk", 7, 7)["characteristicsCumulative"][-1]
        self.assertIsNone(values["programs"]["family"]["ownFundsPct"])
        self.assertEqual(values["programs"]["it"]["ownFundsPct"], 45.123456789)
        self.assertEqual(values["programThroughMonths"]["family"]["ownFundsPct"], 7)
        self.assertEqual(values["programThroughMonths"]["it"]["ownFundsPct"], 7)

    def test_no_methodology_no_rate_or_term_weighting_and_population_mismatch_is_not_silently_used(self):
        document = weighted_source()
        document["methodology"] = {}
        values = build_report(document, "msk", 2, 2)["characteristicsCumulative"][-1]["programs"]["family"]
        self.assertIsNone(values["ratePct"])
        self.assertIsNone(values["termMonths"])
        self.assertIsNotNone(values["areaM2"])
        document = weighted_source()
        document["flows"][1]["count"] += 1
        values = build_report(document, "msk", 2, 2)["characteristicsCumulative"][-1]["programs"]["family"]
        self.assertTrue(all(value is None for value in values.values()))

    def test_missing_positive_flow_characteristic_is_not_dropped_but_zero_flow_can_be_skipped(self):
        document = weighted_source()
        document["characteristics"][1]["ratePct"] = None
        values = build_report(document, "msk", 7, 7)["characteristicsCumulative"][-1]["programs"]["family"]
        self.assertIsNone(values["ratePct"])
        self.assertIsNotNone(values["termMonths"])
        document = weighted_source()
        document["flows"][1].update(count=0, volumeMln=0)
        document["characteristics"][1].update(dict.fromkeys(CHAR_SOURCE_METRICS))
        # Recompute the source aggregate after the genuinely empty month.
        for metric in CHAR_SOURCE_METRICS:
            if metric == "priceThousandM2":
                continue
            document["characteristics"][-1][metric] = None
        document["characteristics"][-1]["areaM2"] = (2 * 20 + 30 * 50) / 32
        values = build_report(document, "msk", 2, 2)["characteristicsCumulative"][-1]["programs"]["family"]
        self.assertEqual(values["areaM2"], 20)

    def test_coverage_per_metric_and_program_never_carries_forward(self):
        document = weighted_source()
        for row in document["flows"] + document["characteristics"]:
            row["period"] = row["period"].replace("2025", "2026")
        document["flows"] = [r for r in document["flows"] if r["basis"] == "annual" or r["period"] <= "2026-10"]
        for row in document["characteristics"]:
            if row["basis"] == "month" and row["period"] > "2026-07":
                for metric in CHAR_METRICS[1:]:
                    row[metric] = None
        document["characteristics"] = [r for r in document["characteristics"] if r["basis"] == "annual" or r["period"] <= "2026-10"]
        report = build_report(document, "msk", 12, 10)
        self.assertEqual(report["characteristicsAnnual"][-1]["throughMonths"]["ratePct"], 7)
        self.assertEqual(report["characteristicsAnnual"][-1]["throughMonths"]["averageAmountMln"], 10)
        self.assertEqual(report["characteristicsCumulative"][-1]["programThroughMonths"]["family"]["ratePct"], 7)
        self.assertIsNone(report["characteristicsCumulative"][-1]["programThroughMonths"]["it"]["ratePct"])
        self.assertEqual(report["characteristicLatestPeriod"], "2026-07")

    def test_validation_invalid_dimensions_nonfinite_values_duplicates_and_detachment(self):
        for changes in ({"schemaVersion": True}, {"sourceUrl": "wrong"}, {"collectedAt": "2026-01-01"},
                        {"sourceDate": "2026-02-30"}, {"flows": []}, {"notes": [None]}, {"methodology": []}):
            with self.subTest(changes=changes), self.assertRaises((ValueError, TypeError)):
                validate_source({**source(), **changes})
        for changes in ({"region": []}, {"region": "mo"}, {"program": "all"}, {"basis": "ytd"},
                        {"period": "2026-13"}, {"count": True}, {"count": "1"}, {"count": -1},
                        {"count": math.inf}, {"count": math.nan}, {"count": 10 ** 1000}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                validate_source(source([{**flow("2026-01"), **changes}]))
        document = source()
        document["flows"].append(copy.deepcopy(document["flows"][0]))
        with self.assertRaises(ValueError):
            validate_source(document)
        document = source()
        detached = validate_source(document)
        detached["flows"][0]["count"] = 999
        self.assertEqual(document["flows"][0]["count"], 2)
        self.assertEqual(region_periods(document)["msk"], ["2026-01"])


@unittest.skipUnless(os.environ.get("PREFERENTIAL_MORTGAGE_TEST_WORKBOOK"), "Set PREFERENTIAL_MORTGAGE_TEST_WORKBOOK to the attached 2026-10-11 XLSX")
class AttachedWorkbookRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.document = parse_workbook(os.environ["PREFERENTIAL_MORTGAGE_TEST_WORKBOOK"])

    def test_source_precision_and_published_annual_values(self):
        self.assertEqual(self.document["sourceDate"], "2026-10-11")
        row = next(r for r in self.document["flows"] if (r["region"], r["program"], r["period"]) == ("rf", "family", "2026-10"))
        self.assertEqual(row["count"], 2900.154)
        self.assertEqual(row["volumeMln"], 16946.9180810941)
        report = build_report(self.document, "msk", 10, 7)
        self.assertEqual(report["annual"][-1]["total"], {"volumeMln": 281274.694940557, "count": 25251.3941836885})
        self.assertEqual(report["latestPeriod"], "2026-10")
        self.assertTrue(report["latestPeriodPartial"])

    def test_exact_published_window_is_preserved_but_shorter_mismatched_windows_remain_unknown(self):
        rf = build_report(self.document, "rf", 7, 7)
        ytd = rf["characteristicsCumulative"][-1]["programs"]
        self.assertEqual(ytd["family"]["ownFundsPct"], 30.0552221846817)
        self.assertEqual(ytd["it"]["ratePct"], 5.87903022732177)
        self.assertAlmostEqual(ytd["family"]["ratePct"], 5.77399037750004, places=11)
        self.assertAlmostEqual(ytd["family"]["averageAmountMln"], 5.974323443, places=8)
        self.assertEqual(rf["characteristicsAnnual"][-1]["programs"]["family"]["averageAmountMln"], 6.02591754237987)
        june = build_report(self.document, "rf", 6, 6)["characteristicsCumulative"][-1]["programs"]
        self.assertIsNone(june["family"]["ownFundsPct"])
        self.assertIsNone(june["it"]["ratePct"])
        msk = build_report(self.document, "msk", 7, 7)
        row2019 = next(r for r in msk["characteristicsCumulative"] if r["year"] == 2019)
        self.assertIsNone(row2019["programs"]["family"]["termMonths"])
        full = build_report(self.document, "msk", 12, 12)
        annual2019 = next(r for r in full["characteristicsCumulative"] if r["year"] == 2019)
        self.assertEqual(annual2019["programs"]["family"]["termMonths"], 248.733453524856)

    def test_mixed_metric_and_program_horizons_and_undefined_zero_issuance(self):
        report = build_report(self.document, "msk", 10, 10)
        self.assertEqual(report["characteristicLatestPeriod"], "2026-07")
        current = report["characteristicsAnnual"][-1]
        self.assertEqual(current["throughMonths"]["averageAmountMln"], 10)
        self.assertEqual(current["throughMonths"]["ratePct"], 7)
        self.assertIsNone(current["programs"]["it"]["ratePct"])
        self.assertEqual(report["annual"][-1]["programs"]["it"]["count"], 0)
        last_it = next(r for r in report["characteristicsAnnual"] if r["year"] == 2024)
        self.assertEqual(last_it["programThroughMonths"]["it"]["ratePct"], 10)
        for collection in ("characteristicsAnnual", "characteristicsCumulative"):
            for row in report[collection]:
                for program in PROGRAMS:
                    self.assertEqual(set(row["programs"][program]), set(CHAR_METRICS))

    def test_missing_source_rows_cannot_move_confirmed_source_aggregate_to_shorter_window(self):
        missing_july = copy.deepcopy(self.document)
        for row in missing_july["characteristics"]:
            if (row["region"], row["program"], row["period"]) == ("rf", "family", "2026-07"):
                row.update(dict.fromkeys((*CHAR_METRICS[1:], "propertyAmountMln")))
        values = build_report(missing_july, "rf", 6, 6)["characteristicsCumulative"][-1]
        self.assertIsNone(values["programs"]["family"]["ownFundsPct"])
        self.assertIsNone(values["programs"]["family"]["ratePct"])
        missing_tail = copy.deepcopy(self.document)
        for collection in ("flows", "characteristics"):
            missing_tail[collection] = [row for row in missing_tail[collection]
                                       if not (row["region"] == "rf" and row["program"] == "family"
                                               and row["basis"] == "month" and "2025-08" <= row["period"] <= "2025-12")]
        values = build_report(missing_tail, "rf", 7, 7)["characteristicsCumulative"][-2]
        self.assertEqual(values["year"], 2025)
        self.assertIsNone(values["programs"]["family"]["ownFundsPct"])


if __name__ == "__main__":
    unittest.main()
