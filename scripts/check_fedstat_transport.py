"""Offline regressions for browser export fallback and forced-run state safety."""
from __future__ import annotations

import base64
from contextlib import redirect_stdout
from io import BytesIO, StringIO
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import fedstat_checker as fc
from openpyxl import Workbook, load_workbook


class TransportTests(unittest.TestCase):
    def test_ipc_short_selection_expands_both_parts(self):
        self.assertEqual(fc._parse_only_ids("31074"), ["31074_часть1", "31074_часть2"])

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.output = self.root / "result.xls"
        workbook = Workbook()
        workbook.active.append(["ИПЦ", 105.3])
        buffer = BytesIO()
        workbook.save(buffer)
        self.excel = buffer.getvalue()
        self.driver = Mock(current_url="https://www.fedstat.ru/indicator/31074")
        self.entries = [("id", "31074"), ("selectedFilterIds", "3_2025"),
                        ("selectedFilterIds", "3_2026")]

    def download(self, result, *, form_body=None):
        self.driver.execute_async_script.return_value = result

        def submit(script, *args):
            self.assertIn("form.submit()", script)
            self.assertEqual(args[1], [list(entry) for entry in self.entries])
            if form_body is not None:
                (self.root / "server.xls").write_bytes(form_body)

        self.driver.execute_script.side_effect = submit
        with patch.object(fc, "_set_driver_download_dir"), \
                patch.object(fc, "_export_post_with_token", return_value=self.entries), \
                patch.object(fc, "_wait_for_browser_download", side_effect=lambda *a, **kw:
                             self.root / "server.xls" if form_body is not None else None), \
                redirect_stdout(StringIO()):
            return fc._download_excel_via_browser(
                self.driver, "https://www.fedstat.ru/indicator/data.do?format=excel",
                self.entries, self.root, self.output)

    def test_html_and_http_errors_reach_native_form(self):
        for status in (200, 503):
            with self.subTest(status=status):
                self.driver.reset_mock()
                result = self.download({"ok": status == 200, "status": status,
                    "contentType": "text/html", "bodyBase64": base64.b64encode(b"<html>error</html>").decode()},
                    form_body=self.excel)
                self.assertEqual(result, self.output)
                self.driver.execute_script.assert_called_once()
                self.assertEqual(load_workbook(BytesIO(result.read_bytes())).active["B1"].value, 105.3)

    def test_successful_fetch_does_not_submit_again(self):
        result = self.download({"ok": True, "status": 200, "contentType": "application/octet-stream",
                                "bodyBase64": base64.b64encode(self.excel).decode()})
        self.assertEqual(result, self.output)
        self.driver.execute_script.assert_not_called()

    def test_error_status_cannot_save_even_excel_shaped_body(self):
        self.assertIsNone(self.download({"ok": False, "status": 503,
            "contentType": "application/octet-stream", "bodyBase64": base64.b64encode(self.excel).decode()}))
        self.assertFalse(self.output.exists())
        self.driver.execute_script.assert_called_once()

    def test_failed_form_preserves_previous_file(self):
        self.output.write_bytes(self.excel)
        self.assertIsNone(self.download({"error": "Failed to fetch"}, form_body=b"<html>503</html>"))
        self.assertEqual(self.output.read_bytes(), self.excel)

    def test_native_http_error_does_not_wait_full_download_timeout(self):
        self.driver.title = "503 Service Temporarily Unavailable"
        with patch.object(fc, "wait_for_download", return_value=None) as wait, redirect_stdout(StringIO()):
            self.assertIsNone(fc._wait_for_browser_download(self.driver, self.root, set()))
            wait.assert_called_once()
            self.assertLessEqual(wait.call_args.kwargs["timeout"], 5)

    def test_force_preserves_state_of_failed_and_unselected_sources(self):
        initial = {"31074": "01.09.2026", "other": "01.08.2026"}
        for succeeded in (False, True):
            with self.subTest(succeeded=succeeded), \
                    patch.object(fc, "DOWNLOAD_DIR", self.root), \
                    patch.object(fc, "INDICATORS", {"31074": "ИПЦ", "other": "Другое"}), \
                    patch.object(fc, "load_state", return_value=initial.copy()), \
                    patch.object(fc, "save_state") as save, \
                    patch.object(fc, "create_driver", return_value=self.driver), \
                    patch.object(fc, "get_last_update_date", return_value="01.09.2026"), \
                    patch.object(fc, "download_excel", return_value=self.output if succeeded else None) as download, \
                    redirect_stdout(StringIO()):
                files, ok = fc.run(force=True, only_ids=["31074"])
                download.assert_called_once()
                self.assertEqual(ok, succeeded)
                save.assert_called_once_with(initial)

    def test_unknown_selection_does_not_run_partial_request(self):
        with patch.object(fc, "DOWNLOAD_DIR", self.root), \
                patch.object(fc, "load_state", return_value={}), \
                patch.object(fc, "create_driver") as create, \
                patch.object(fc, "save_state") as save, redirect_stdout(StringIO()):
            self.assertEqual(fc.run(only_ids=["31074_часть1", "typo"]), ([], False))
            create.assert_not_called()
            save.assert_not_called()

    def test_force_advances_only_successful_selected_date(self):
        initial = {"31074_часть2": "01.08.2026", "other": "01.07.2026"}
        with patch.object(fc, "DOWNLOAD_DIR", self.root), \
                patch.object(fc, "load_state", return_value=initial.copy()), \
                patch.object(fc, "save_state") as save, \
                patch.object(fc, "create_driver", return_value=self.driver), \
                patch.object(fc, "get_last_update_date", return_value="01.09.2026"), \
                patch.object(fc, "download_excel", return_value=self.output), redirect_stdout(StringIO()):
            self.assertTrue(fc.run(force=True, only_ids=["31074_часть2"])[1])
            save.assert_called_once_with({"31074_часть2": "01.09.2026", "other": "01.07.2026"})

    def test_diagnostic_worker_keeps_output_and_receipt_isolated(self):
        import diagnose_fedstat_ipc as diag
        import pandas as pd
        from pipeline.download_provenance import receipt_path
        folder = self.root / "diagnostic"
        folder.mkdir()

        @fc.record_successful_download
        def fake_download(source, destination, **kwargs):
            path = destination / "fixture.xls"
            path.write_bytes(self.excel)
            return path

        with patch.object(diag, "ROOT", self.root), \
                patch.object(fc, "create_driver", return_value=self.driver), \
                patch.object(fc, "get_last_update_date", return_value="01.09.2026"), \
                patch.object(fc, "download_excel", side_effect=fake_download), \
                patch.object(fc, "save_state") as save, \
                patch("pipeline.parsers.fedstat_ipc.parse", return_value=pd.DataFrame({"year": [2026], "month": [8]})), \
                redirect_stdout(StringIO()):
            self.assertEqual(diag.main(["--worker-dir", str(folder)]), 0)
            save.assert_not_called()
        self.assertTrue(receipt_path(folder / "fixture.xls").is_file())
        log = (folder / "diagnostic.log").read_text(encoding="utf-8")
        self.assertIn("latest period: 2026-08", log)
        self.assertIn(str(self.root), log)
        self.driver.quit.assert_called_once()

    def test_browser_csrf_refresh_preserves_repeated_filters_and_session(self):
        session = Mock()
        self.driver.execute_script.return_value = [["struts.token.name", "token"], ["token", "fresh-fixture"]]
        self.driver.get_cookies.return_value = [{"name": "JSESSIONID", "value": "session-fixture",
                                               "domain": "www.fedstat.ru", "path": "/"}]
        original = self.entries + [("struts.token.name", "token"), ("token", "expired-fixture")]
        result = fc._export_post_with_token(session, self.driver, "31074", original, refresh=True)
        self.driver.get.assert_called_once_with("https://www.fedstat.ru/indicator/31074")
        self.assertEqual(result, self.entries + [("struts.token.name", "token"), ("token", "fresh-fixture")])
        session.cookies.set.assert_called_once_with("JSESSIONID", "session-fixture", domain="www.fedstat.ru", path="/")

    def test_export_contract_current_endpoint_title_and_csrf(self):
        session = Mock()
        session.get.return_value = Mock(text='''<div id="downloadTokenHolder">
            <input type="hidden" name="struts.token.name" value="token">
            <input type="hidden" name="token" value="fixture-csrf"></div>''')
        response = Mock(headers={"Content-Type": "application/vnd.ms-excel"})
        response.iter_content.return_value = [self.excel]
        session.post.return_value = response
        template = {"id": "31074", "title": "ИПЦ", "selectedFilterIds": ["3_2025", "3_2026"]}
        with patch.object(fc.requests, "Session", return_value=session), redirect_stdout(StringIO()):
            result = fc.download_excel("31074_часть2", self.root, remote_date="19.09.2026",
                                       payload_template_override=template, save_path_override=self.output)
        self.assertEqual(result, self.output)
        session.get.assert_called_once_with("https://www.fedstat.ru/indicator/31074", timeout=30)
        self.assertEqual(session.post.call_args.args[0], "https://www.fedstat.ru/indicator/downloadData.do?format=excel")
        data = session.post.call_args.kwargs["data"]
        self.assertIn(("title", "ИПЦ"), data)
        self.assertNotIn("indicator_title", dict(data))
        self.assertIn(("token", "fixture-csrf"), data)
        self.assertEqual([v for k, v in data if k == "selectedFilterIds"], ["3_2025", "3_2026"])

    def test_missing_csrf_stops_before_any_export_post(self):
        session = Mock()
        session.get.return_value = Mock(text="<html>unavailable</html>")
        with patch.object(fc.requests, "Session", return_value=session), redirect_stdout(StringIO()):
            result = fc.download_excel("31074_часть2", self.root)
        self.assertIsNone(result)
        session.post.assert_not_called()


if __name__ == "__main__":
    unittest.main()
