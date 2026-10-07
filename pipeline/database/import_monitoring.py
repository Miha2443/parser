"""Adapt the shared raw loader to typed rows; no second normalization pipeline."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, localcontext
import hashlib
from itertools import zip_longest
import re
from pathlib import Path

import openpyxl
import pandas as pd

from pipeline.data_access import DataAccess, DataContext
from pipeline.source_records import file_sha256
from .repository import AREA_FIELDS, canonical_json, decimal_text


AREA_COLUMNS = {
    "total_area_m2": "Общая площадь",
    "residential_area_m2": "Жилая площадь",
    "category_residential_m2": "category_жилое",
    "category_common_m2": "category_моп",
    "category_nonres_in_res_m2": "category_нежилое_в_жилом",
    "category_nonres_other_m2": "category_нежилое_отдельное",
}
TEXT_COLUMNS = {
    "uin": "УИН", "developer": "Группа компаний", "object_name": "Наименование объекта",
    "address": "Адрес", "okrug": "Округ", "district": "Район",
    "grouping": "Группировка", "purpose": "Назначение", "funding_source": "Источник финансирования",
}


def _missing(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip()) or bool(pd.isna(value))


def _text(value) -> str | None:
    return None if _missing(value) else str(value)


def _json_cell(value):
    """Preserve cell types without storing binary floats in JSON."""
    if value is None or bool(pd.isna(value)):
        return None
    if isinstance(value, (datetime, date)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float, Decimal)) or hasattr(value, "dtype"):
        return {"type": "number", "value": decimal_text(value)}
    return str(value)


def _fingerprint() -> str:
    base = Path(__file__).resolve().parents[1]
    paths = [base / "data_access.py", base / "dev_name_utils.py", base / "source_records.py", Path(__file__).resolve()]
    parts = [(path.name, file_sha256(path)) for path in paths]
    return "monitoring-decimal-v1:" + hashlib.sha256(canonical_json(parts).encode("utf-8")).hexdigest()


@dataclass
class PreparedMonitoring:
    raw_file: dict
    normalizer_version: str
    records: list[dict]
    validation: dict


def decimal_area_sums(records: list[dict]) -> dict[str, str]:
    """Use the repository's exact aggregate precision, not ambient context."""
    with localcontext() as context:
        context.prec = 60
        return {
            field: format(sum((Decimal(r[field]) for r in records if r[field] is not None), Decimal(0)), "f")
            for field in AREA_FIELDS
        }


def prepare_monitoring(source_path: Path, *, root: Path) -> PreparedMonitoring:
    source_path = source_path.resolve(strict=True)
    normalizer_version = _fingerprint()
    raw_hash = file_sha256(source_path)
    access = DataAccess(DataContext(root=root, use_marts=False))
    payload = access.load_monitoring_2_0(source_path=source_path, include_provenance=True)
    if not payload["rv"].size or not payload["oks"].size:
        raise ValueError("Pilot requires nonempty RV and OKS from the raw loader")
    needed = {}
    records = []
    for registry in ("rv", "oks"):
        frame = payload[registry]
        required = {"source_sheet", "source_row", "source_file_sha256", "source_record_id", *AREA_COLUMNS.values()}
        if not required <= set(frame.columns):
            raise ValueError(f"Loader missing provenance/area columns: {required - set(frame.columns)}")
        for item in frame.to_dict("records"):
            sheet, row = str(item["source_sheet"]), int(item["source_row"])
            if str(item["source_file_sha256"]) != raw_hash:
                raise ValueError("Loader provenance hash mismatch")
            record = {
                "source_sheet": sheet, "source_row": row,
                "source_record_id": str(item["source_record_id"]), "registry": registry,
                **{key: _text(item.get(column)) for key, column in TEXT_COLUMNS.items()},
                "permit_number": _text(item.get("№ РВ" if registry == "rv" else "№РС")),
            }
            year_value = item.get("Год ввода по Мосстату" if registry == "rv" else "Год ввода по графику")
            year_number = pd.to_numeric(year_value, errors="coerce")
            record["year"] = int(year_number) if pd.notna(year_number) and float(year_number).is_integer() else None
            for field, column in AREA_COLUMNS.items():
                record[field] = decimal_text(item[column])
            record["normalized_payload_json"] = canonical_json({str(k): _json_cell(v) for k, v in item.items()})
            identity = (sheet, row)
            if identity in needed:
                raise ValueError(f"Loader repeats source identity {identity}")
            needed[identity] = record
            records.append(record)
    # Read physical rows only for provenance/preservation and null counts.
    # No area/developer/category normalization is duplicated here.
    workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=False)
    cached_workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
    source_rows = {}
    workbook_metadata = {}
    try:
        for sheet in sorted({sheet for sheet, _ in needed}):
            iterator = workbook[sheet].iter_rows(values_only=True)
            cached_iterator = cached_workbook[sheet].iter_rows(values_only=True)
            formula_columns = list(next(iterator))
            columns = list(next(cached_iterator))
            workbook_metadata[sheet] = {
                "columns": [_json_cell(c) for c in columns],
                "formula_columns": [_json_cell(c) for c in formula_columns],
            }
            total_index = columns.index("Общая площадь")
            residential_index = columns.index("Жилая площадь")
            source_rows[sheet] = 0
            for row_number, (cells, cached_cells) in enumerate(zip_longest(iterator, cached_iterator), start=2):
                if cells is None or cached_cells is None:
                    raise ValueError("Formula/cached row streams disagree")
                source_rows[sheet] += 1
                record = needed.get((sheet, row_number))
                if record is None:
                    continue
                for index, prefix, field in (
                    (total_index, "raw_total_area", "total_area_m2"),
                    (residential_index, "raw_residential_area", "residential_area_m2"),
                ):
                    missing = _missing(cached_cells[index])
                    invalid = not missing and pd.isna(pd.to_numeric(cached_cells[index], errors="coerce"))
                    record[prefix + "_missing"] = int(missing)
                    record[prefix + "_invalid"] = int(invalid)
                    if missing or invalid:
                        record[field] = None
                record["raw_payload_json"] = canonical_json({
                    "cells": [_json_cell(c) for c in cells],
                    "cached_cells": [_json_cell(c) for c in cached_cells],
                    "cell_mode": "formula/value and cached values; styles remain in the original XLSX only",
                })
    finally:
        workbook.close()
        cached_workbook.close()
    if any("raw_payload_json" not in record for record in records):
        raise ValueError("Some loader records cannot be traced to a physical Excel row")
    if file_sha256(source_path) != raw_hash or normalizer_version != _fingerprint():
        raise ValueError("Source file or normalization code changed during preparation")
    validation = {"source": "raw_xlsx_shared_loader", "provenance": "sha256+sheet+1_based_excel_row",
                  "row_counts": {}, "decimal_sums": {}, "missing_uin": {}, "raw_area_missing": {},
                  "physical_source_rows": source_rows, "typed_rejected_rows": 0,
                  "source_dates": {
                      "source_file_mtime_utc": datetime.fromtimestamp(source_path.stat().st_mtime, timezone.utc).isoformat(),
                      "filename_date": next(iter(re.findall(r"\d{8}", source_path.stem)), None),
                      "source_published_at": None,
                      "observation_year_min": payload["min_year"], "observation_year_max": payload["max_year"],
                  },
                  "quality_warnings": [
                      "Source publication date is unknown; file-name date/mtime are not publication evidence.",
                      "UIN is nonunique. Sums are record sums, not verified distinct-object area.",
                      "Category areas retain the shared loader's zero-fill/clip policy; raw main area NULLs are preserved.",
                      "Decimal strings preserve the loader's values without rounding; Excel/pandas float artifacts may remain.",
                  ]}
    validation["shared_loader_filtered_rows"] = sum(source_rows.values()) - len(records)
    for registry in ("rv", "oks"):
        selected = [r for r in records if r["registry"] == registry]
        validation["row_counts"][registry] = len(selected)
        validation["missing_uin"][registry] = sum(r["uin"] is None for r in selected)
        validation["decimal_sums"][registry] = decimal_area_sums(selected)
        validation["raw_area_missing"][registry] = {
            key: sum(r[key] for r in selected) for key in (
                "raw_total_area_missing", "raw_residential_area_missing",
                "raw_total_area_invalid", "raw_residential_area_invalid",
            )
        }
    return PreparedMonitoring(
        raw_file={"sha256": raw_hash, "path": str(source_path), "size_bytes": source_path.stat().st_size,
                  "workbook_metadata": workbook_metadata},
        normalizer_version=normalizer_version, records=records, validation=validation,
    )
