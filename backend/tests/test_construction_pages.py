"""Construction page contracts against temporary books/JSON, never live data."""
import copy
from datetime import datetime
from io import BytesIO
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from zlib import error as DecompressionError

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.construction_pages_export import export_workbook
from backend.construction_pages_routes import create_router
from backend.domrf_service import DomrfService
from backend.permits_page_service import PermitsPageService
from backend.profile_service import FilterUnavailable, SourceUnavailable
from pipeline.construction_calculations import catalog, report
from pipeline.data_access import DataAccess, DataContext, MONTH_SHORT_RU, file_signature
from pipeline.permit_page_calculations import (FORECAST_FIELDS, FORECAST_SHEET, HISTORY_SHEET,
                                               HOUSING_SHEET, OKS_SHEET)


class ConstructionPagesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / "releases/pinned"
        self.raw = self.root / "data/raw/realty/nashdom"
        self.raw.mkdir(parents=True)
        self.context = DataContext(self.root, require_marts=True)
        self.book_path = self.raw / "monitoring_2_0_20261016.xlsx"
        self.json_path = self.raw / "construction_operational_20261001.json"
        self.timestamp = 1_800_000_000
        self.write_book()
        self.write_json()
        self.permits = PermitsPageService(self.context, check_interval=0)
        self.domrf = DomrfService(self.context, check_interval=0)

    def touch(self, path):
        self.timestamp += 1
        os.utime(path, (self.timestamp, self.timestamp))

    def write_book(self, path=None, missing=()):
        path = path or self.book_path
        book = Workbook()
        book.remove(book.active)
        registry = book.create_sheet(OKS_SHEET)
        headers = {1: "УИН", 2: "Назначение", 4: "Срок выдачи РС", 6: "Год выдачи", 8: "Статус РС",
                   17: "Общая площадь", 18: "Жилая площадь", 37: "Группа компаний",
                   44: "ОСМ код", 45: "ФНО 2", 46: "ФНО для Минстроя России", 47: "ФНО для дашборда"}
        for column in range(1, 48):
            registry.cell(1, column, headers.get(column, f"unused {column}"))
        for year in (2025, 2026):
            for month in range(1, 13):
                for kind in ("Жилье", "Нежилье"):
                    row = registry.max_row + 1
                    values = {1: f"id-{row}", 2: kind, 4: datetime(year, month, 2), 6: year,
                              8: "Не действует", 17: 100000 + month if kind == "Жилье" else 40000 + month,
                              18: 60000 + month if kind == "Жилье" else None, 37: "Developer",
                              44: "Do not classify by OSM", 46: "Коммерческие" if month % 2 else "Административные"}
                    for column, value in values.items():
                        registry.cell(row, column, value)
        registry.cell(50, 2, "Нежилье")
        registry.cell(50, 6, 2026)
        registry.cell(50, 17, 1000)
        registry.cell(50, 46, "Другие")
        # Original DataAccess also expects RV, although permit calculations do not use it.
        rv = book.create_sheet("Реестр РВ")
        rv.append(["Год ввода по Мосстату", "Группа компаний", "Общая площадь", "Жилая площадь", "Отрасли", "Группировка"])
        rv.append([2026, "Developer", 1, 1, "Жилые объекты", "Жилье"])
        history = book.create_sheet(HOUSING_SHEET)
        history["C3"], history["D3"], history["G3"] = "ОБЩАЯ", "ЖИЛАЯ", "ОБЩАЯ"
        for year in range(2011, 2022):
            start = 4 + (year - 2011) * 13
            history.cell(start, 2, year)
            history.cell(start, 6, year)
            for month in range(1, 13):
                row = start + month
                history.cell(row, 2, MONTH_SHORT_RU[month - 1])
                history.cell(row, 6, MONTH_SHORT_RU[month - 1])
                history.cell(row, 3, 61021.2 + month - 1)
                history.cell(row, 4, 41463.8 + month - 1)
                history.cell(row, 7, 348192.6 + month - 1)
        detailed = book.create_sheet(HISTORY_SHEET)
        detailed["B2"], detailed["B3"], detailed["C3"], detailed["B4"], detailed["C4"] = (
            "выдача рс жилье по годам", 2011, 9999.123456789, 2026, 12000.123456789)
        detailed["B7"], detailed["B8"], detailed["C8"], detailed["B9"], detailed["C9"] = (
            "выдача рс нежилье по годам", 2011, 7777.123456789, 2026, 8000.123456789)
        labels = ("Год", "№ месяца", "Месяц", "Нежильё,\nвсего", "Административно-\nделовые",
                  "Промышленные", "Торговые", "Социальные", "Гостиницы", "Прочие")
        for column, value in enumerate(labels, 6):
            detailed.cell(114, column, value)
        for index in range(132):
            year, month = 2011 + index // 12, index % 12 + 1
            for column, value in enumerate((year, month, MONTH_SHORT_RU[month - 1], 327.442, 234.535, 0, 0, 13.096, 0, 79.811), 6):
                detailed.cell(115 + index, column, value)
        forecast = book.create_sheet(FORECAST_SHEET)
        for column in range(4, 14):
            forecast.cell(2, column, 2026 + column - 4)
        for row, value in enumerate((5810.123456789, 10017.24138, 4207.241379, 6850.987654321,
                                     2329, 822, 1438.5, 205.5, 548, 1507), 3):
            forecast.cell(row, 3, FORECAST_FIELDS[row - 3])
            forecast.cell(row, 4, value)
        for sheet in missing:
            del book[sheet]
        book.save(path)
        book.close()
        self.touch(path)
        return path

    @staticmethod
    def json_payload():
        return {"source": "construction_operational", "construction": [
            {"region_key": region, "report_period": "Октябрь 2026", "report_date": "01.10.2026",
             "area_kind": area, "area_thousand_m2": value}
            for region in ("msk", "rf") for area, value in (("living", "15 640"), ("total", "26 964"))],
            "sales": [{"region_key": region, "report_period": "Август 2026", "report_date": "11.09.2026",
                       "metrics": {"sold_thousand_m2": "7 366 тыс. м²", "sold_pct": "47%",
                                   "price_per_m2_rub": "441 306 руб.", "funds_million_rub": "3 250 869 млн руб.",
                                   "blank": "", "zero": 0, "precise_pct": "42,123456789%"}}
                      for region in ("msk", "rf")]}

    def write_json(self, payload=None, path=None):
        path = path or self.json_path
        path.write_text(json.dumps(self.json_payload() if payload is None else payload, ensure_ascii=False), encoding="utf-8")
        self.touch(path)
        return path

    def edit_book(self, sheet, edits):
        book = load_workbook(self.book_path)
        for address, value in edits.items():
            book[sheet][address] = value
        book.save(self.book_path)
        book.close()
        self.touch(self.book_path)

    def client(self):
        app = FastAPI()

        def call(function, *args, scope=None):
            try:
                return function(*args)
            except FilterUnavailable as exc:
                raise HTTPException(404, str(exc)) from exc
            except SourceUnavailable as exc:
                raise HTTPException(503, str(exc)) from exc

        app.include_router(create_router(self.context, call, self.permits, self.domrf))
        return TestClient(app)

    @staticmethod
    def by_id(payload):
        return {row["id"]: row for row in payload["breakdowns"]}

    def test_main_exact_legacy_parity_all_kinds_and_months(self):
        self.edit_book(OKS_SHEET, {"Q26": "100 000,125", "R26": "60\xa0000,125", "Q27": None,
                                   "R28": None, "D26": None, "F26": 2026})
        raw_context = DataContext(self.root, use_marts=False, require_marts=False)
        access = DataAccess(raw_context)
        data = {"load_construction_operational": access.load_construction_operational(), "load_rasprodannost": {}}
        original = copy.deepcopy(data)
        controls = catalog(data)
        for kind in ("housing", "nonresidential", "total"):
            for month in range(1, 13):
                with self.subTest(kind=kind, month=month):
                    actual = self.permits.report(kind, month)
                    self.assertEqual({key: value for key, value in actual["permits"].items() if key != "presentation"},
                                     report(data, controls, permit_kind=kind, month=month)["permits"])
                    for row in actual["permits"]["presentation"]["rows"]:
                        source = next((item for item in actual["permits"]["rows"] if item["Год"] == row["year"]), {})
                        self.assertEqual(row["selected"], source.get("С начала года, тыс. м²"))
                        self.assertEqual(row["growth"], source.get("Изменение, %"))
                    self.assertEqual(actual["controls"]["monthsByKind"], controls["monthsByKind"])
        for key, value in original["load_construction_operational"].items():
            if hasattr(value, "equals"):
                self.assertTrue(data["load_construction_operational"][key].equals(value))

    def test_permit_contract_forecast_basis_full_precision_and_no_actual_inflation(self):
        payload = self.permits.report("total", 1)
        self.assertEqual(set(payload), {"schemaVersion", "version", "generatedAt", "source", "selection", "controls",
                                       "permits", "forecast", "breakdowns", "notes"})
        self.assertEqual([row["id"] for row in payload["controls"]["permitKinds"]], ["housing", "nonresidential", "total"])
        self.assertEqual(payload["source"]["fileEvidence"], "selected")
        self.assertEqual(payload["source"]["date"], "2026-10-16")
        forecast = payload["forecast"]
        self.assertEqual(forecast["unitBasis"], "assumed_thousand_m2")
        self.assertEqual(forecast["source"]["sheet"], FORECAST_SHEET)
        self.assertEqual(len(forecast["rows"]), 10)
        self.assertEqual(forecast["rows"][0]["value"], 5810.123456789 + 6850.987654321)
        self.assertEqual(forecast["rows"][0]["cells"]["housing_living"], "D3")
        self.assertEqual(forecast["rows"][1]["cells"]["other"], "E12")
        self.assertTrue(all(row[field] is None for row in forecast["rows"][1:] for field in (*FORECAST_FIELDS, "value")))
        before = {key: value for key, value in payload["permits"].items() if key != "presentation"}
        self.edit_book(FORECAST_SHEET, {"D3": 9999999.123, "D6": None})
        updated = self.permits.report("total", 1)
        self.assertEqual(before, {key: value for key, value in updated["permits"].items() if key != "presentation"})
        current = next(row for row in updated["permits"]["presentation"]["rows"] if row["year"] == 2026)
        self.assertEqual(current["basis"], "none")
        self.assertIsNone(current["remainder"])
        self.assertIsNone(updated["forecast"]["rows"][0]["value"])
        self.assertNotEqual(payload["version"], updated["version"])
        self.assertEqual(self.permits.report("housing", 1)["forecast"]["rows"][0]["value"], 9999999.123)

    def test_forecast_cached_formula_and_out_of_period_year(self):
        self.edit_book(FORECAST_SHEET, {"D2": 2040, "D3": "=999999999", "E3": 0})
        buffer = BytesIO()
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        with ZipFile(self.book_path) as source, ZipFile(buffer, "w") as target:
            sheet_name = "xl/worksheets/sheet5.xml"
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == sheet_name:
                    root = ET.fromstring(content)
                    root.find(f".//{namespace}c[@r='D3']/{namespace}v").text = "123.123456789"
                    content = ET.tostring(root)
                target.writestr(item, content)
        self.book_path.write_bytes(buffer.getvalue())
        self.touch(self.book_path)
        payload = self.permits.report("housing", 1)
        self.assertEqual(payload["forecast"]["rows"][0]["year"], 2040)
        self.assertEqual(payload["forecast"]["rows"][0]["value"], 123.123456789)
        self.assertEqual(payload["forecast"]["rows"][1]["value"], 0)
        self.assertNotIn(2040, [row["Год"] for row in payload["permits"]["rows"]])
        self.edit_book(FORECAST_SHEET, {"D3": "=1+1"})
        self.assertIsNone(self.permits.report("housing", 1)["forecast"]["rows"][0]["value"])

    def test_source_specific_history_not_rescaled_and_ytd_everywhere(self):
        january = self.permits.report("total", 1, structure_month_to=1)
        blocks = self.by_id(january)
        house = blocks["housing_components"]["rows"][0]
        self.assertAlmostEqual(house["living"], 41.4638)
        self.assertAlmostEqual(house["housing_total"], 61.0212)
        self.assertAlmostEqual(house["mop_embedded"], 19.5574)
        self.assertEqual(blocks["nonresidential_historical"]["rows"][0]["nonresidential"], 327.442)
        self.assertAlmostEqual(january["permits"]["rows"][0]["С начала года, тыс. м²"], 41.4638 + 348.1926)
        composition = blocks["area_composition"]
        self.assertEqual(composition["title"], "Структура общей площади")
        self.assertAlmostEqual(composition["rows"][0]["total"], 61.0212 + 327.442)
        self.assertNotEqual(composition["rows"][0]["total"], january["permits"]["rows"][0]["С начала года, тыс. м²"])
        feb = self.by_id(self.permits.report("total", 2, structure_month_to=2))
        self.assertEqual(feb["nonresidential_historical"]["rows"][0]["nonresidential"], 327.442 * 2)
        self.assertAlmostEqual(feb["housing_components"]["rows"][0]["living"], (41463.8 + 41464.8) / 1000)
        self.assertTrue(all(block["periodLabel"] == january["permits"]["periodLabel"] for block in blocks.values()))
        self.assertTrue(any("I115:O115" in note for note in blocks["nonresidential_historical"]["notes"]))

    def test_missing_raw_values_zero_negative_gap_native_taxonomy_no_status_filter(self):
        self.edit_book(OKS_SHEET, {"Q26": None, "R26": 0, "Q28": 10, "R28": 100,
                                   "AT27": None, "AU27": "Не использовать AU", "Q27": 0, "D50": datetime(2026, 1, 3)})
        jan = self.by_id(self.permits.report("total", 1, structure_month_to=1))
        row = next(row for row in jan["housing_components"]["rows"] if row["year"] == 2026)
        self.assertEqual(row["living"], 0)
        self.assertIsNone(row["housing_total"])
        self.assertIsNone(row["mop_embedded"])
        native = jan["nonresidential_native"]
        self.assertIn({"id": "unknown", "label": "Не указано в источнике"}, native["columns"])
        nrow = next(row for row in native["rows"] if row["year"] == 2026)
        self.assertEqual(nrow["unknown"], 0)
        self.assertEqual(nrow["nonresidential"], 1)
        self.assertNotIn("offices", nrow)
        feb = self.by_id(self.permits.report("housing", 2, structure_month_to=2))["housing_components"]
        self.assertIsNone(next(row for row in feb["rows"] if row["year"] == 2026)["mop_embedded"])
        self.assertTrue(any("общая площадь меньше жилой" in note for note in feb["notes"]))

    def test_breakdown_stacks_include_only_additive_components(self):
        aggregates = {"housing_components": "housing_total", "area_composition": "total",
                      "nonresidential_historical": "nonresidential", "nonresidential_native": "nonresidential"}
        for month in (1, 2, 12):
            for block in self.permits.report("total", month)["breakdowns"]:
                aggregate = aggregates[block["id"]]
                self.assertIn(aggregate, [column["id"] for column in block["columns"]])
                self.assertNotIn(aggregate, [series["id"] for series in block["series"]])
                for index, row in enumerate(block["rows"]):
                    points = [series["points"][index]["y"] for series in block["series"]]
                    if row[aggregate] is not None and all(value is not None for value in points):
                        self.assertAlmostEqual(sum(points), row[aggregate], places=8)

    def test_strict_historical_coverage_missing_supplements_do_not_block_primary(self):
        self.edit_book(HOUSING_SHEET, {"B5": None})
        blocks = self.by_id(self.permits.report("total", 2, structure_month_to=2))
        incomplete = blocks["housing_components"]["rows"][0]
        self.assertEqual(incomplete["year"], 2011)
        self.assertIsNone(incomplete["living"])
        self.assertFalse(incomplete["periodComplete"])
        self.assertIsNone(blocks["area_composition"]["rows"][0]["living"])
        self.write_book(missing=(HOUSING_SHEET, HISTORY_SHEET, FORECAST_SHEET))
        payload = self.permits.report("total", 1)
        self.assertTrue(payload["permits"]["rows"])
        self.assertEqual(payload["forecast"]["rows"], [])
        self.assertTrue(payload["forecast"]["notes"])
        self.assertEqual(self.by_id(payload)["nonresidential_historical"]["rows"], [])

    @staticmethod
    def presentation(payload, year):
        return next(row for row in payload["permits"]["presentation"]["rows"] if row["year"] == year)

    def partial_current_book(self):
        edits = {"B50": None}
        for month in (11, 12):
            edits[f"B{24 + month * 2}"] = None
            edits[f"B{25 + month * 2}"] = None
        self.edit_book(OKS_SHEET, edits)

    def test_permit_presentation_forecast_same_year_basis_not_latest_month_total(self):
        self.partial_current_book()
        payload = self.permits.report("total", 5)
        current = self.presentation(payload, 2026)
        raw = next(row for row in payload["permits"]["rows"] if row["Год"] == 2026)
        self.assertEqual(current["selected"], raw["С начала года, тыс. м²"])
        self.assertEqual(current["growth"], raw["Изменение, %"])
        self.assertEqual(current["annualForecast"], 5810.123456789 + 6850.987654321)
        self.assertEqual(current["difference"], current["annualForecast"] - current["selected"])
        self.assertEqual(current["remainder"], current["difference"])
        self.assertEqual(current["basis"], "forecast")
        self.assertIsNone(current["annualActual"])
        self.assertFalse(current["annualComplete"])
        self.assertTrue(current["periodComplete"])
        self.assertEqual(current["observedThrough"], "2026-10-16")
        self.assertNotEqual(current["remainder"], payload["permits"]["chart"]["remainder"][-1]["value"])
        self.assertEqual([row["year"] for row in payload["permits"]["presentation"]["rows"]].count(2026), 1)
        self.assertTrue(any("D3" in ref for ref in current["evidence"]))
        self.assertTrue(any("assumed_thousand_m2" in ref for ref in current["evidence"]))
        self.edit_book(FORECAST_SHEET, {"D3": None, "D6": None})
        unplanned = self.permits.report("total", 5)
        current = self.presentation(unplanned, 2026)
        self.assertEqual(current["basis"], "none")
        for field in ("annualActual", "annualForecast", "difference", "remainder"):
            self.assertIsNone(current[field])
        self.assertEqual(current["selected"], raw["С начала года, тыс. м²"])
        self.assertGreater(unplanned["permits"]["chart"]["remainder"][-1]["value"], 0)
        october = self.presentation(self.permits.report("total", 10), 2026)
        self.assertFalse(october["periodComplete"])
        self.assertIsNone(october["remainder"])

    def test_permit_confirmed_annual_controls_and_full_literal_months_only(self):
        self.edit_book(OKS_SHEET, {"B50": None})
        payload = self.permits.report("total", 5)
        historic = self.presentation(payload, 2011)
        self.assertEqual(historic["annualActual"], 9999.123456789 + 7777.123456789)
        self.assertEqual(historic["basis"], "annual_observed")
        self.assertTrue(historic["annualComplete"])
        self.assertTrue(any("годовой контроль" in ref for ref in historic["evidence"]))
        previous = self.presentation(payload, 2025)
        self.assertEqual(previous["basis"], "annual_observed")
        self.assertTrue(previous["annualComplete"])
        self.assertAlmostEqual(previous["annualActual"], sum((100000 + 2 * month) / 1000 for month in range(1, 13)))
        self.edit_book(OKS_SHEET, {"B10": None, "B11": None})
        missing = self.presentation(self.permits.report("total", 5), 2025)
        self.assertIsNone(missing["annualActual"])
        self.assertIsNone(missing["remainder"])
        self.assertFalse(missing["periodComplete"])
        self.assertEqual(missing["basis"], "none")

    def test_permit_forecast_only_zero_and_overplan_are_not_phantom_actuals(self):
        self.edit_book(FORECAST_SHEET, {"D3": 1, "E3": 0, "E2": 2040})
        payload = self.permits.report("housing", 5)
        current = self.presentation(payload, 2026)
        self.assertLess(current["difference"], 0)
        self.assertEqual(current["remainder"], 0)
        future = self.presentation(payload, 2040)
        self.assertEqual(future["annualForecast"], 0)
        self.assertEqual(future["basis"], "forecast")
        for field in ("selected", "annualActual", "difference", "remainder", "growth", "observedThrough"):
            self.assertIsNone(future[field])
        self.assertFalse(future["annualComplete"])
        self.assertFalse(future["periodComplete"])
        self.assertNotIn(2028, [row["year"] for row in payload["permits"]["presentation"]["rows"]])
        self.assertNotIn(2040, [row["Год"] for row in payload["permits"]["rows"]])

    def test_permit_forecast_overflow_unknown_snapshot_and_raw_null_coverage(self):
        self.edit_book(FORECAST_SHEET, {"D3": 1e308, "D6": 1e308})
        payload = self.permits.report("total", 5)
        self.assertIsNone(payload["forecast"]["rows"][0]["value"])
        self.assertIsNone(self.presentation(payload, 2026)["remainder"])
        self.edit_book(OKS_SHEET, {"R2": None})
        previous = self.presentation(self.permits.report("housing", 5), 2025)
        self.assertFalse(previous["annualComplete"])
        self.assertIsNone(previous["annualActual"])
        self.assertIsNone(previous["remainder"])
        renamed = self.book_path.with_name("monitoring_2_0_unknown.xlsx")
        self.book_path.rename(renamed)
        unknown = self.permits.report("housing", 5, structure_mode="annual")
        self.assertIsNone(unknown["source"]["date"])
        self.assertIsNone(self.presentation(unknown, 2025)["annualActual"])
        modern = next(row for row in self.by_id(unknown)["housing_components"]["rows"] if row["year"] == 2025)
        self.assertFalse(modern["annualComplete"])
        self.assertIsNone(modern["housing_total"])

    def test_structure_filters_independent_arbitrary_start_month_and_source_basis(self):
        options = {"structure_month_from": 3, "structure_month_to": 5,
                   "structure_year_from": 2012, "structure_year_to": 2013}
        first, second = self.permits.report("total", 1, **options), self.permits.report("total", 2, **options)
        self.assertEqual(first["selection"]["structure"], {"mode": "period", "monthFrom": 3, "monthTo": 5,
                                                            "yearFrom": 2012, "yearTo": 2013})
        self.assertEqual(first["breakdowns"], second["breakdowns"])
        blocks = self.by_id(first)
        housing = blocks["housing_components"]
        self.assertEqual([row["year"] for row in housing["rows"]], [2012, 2013])
        self.assertAlmostEqual(housing["rows"][0]["living"], sum(41463.8 + month - 1 for month in (3, 4, 5)) / 1000)
        self.assertEqual(blocks["nonresidential_historical"]["rows"][0]["nonresidential"], 327.442 * 3)
        self.assertAlmostEqual(blocks["area_composition"]["rows"][0]["total"],
                               sum(61021.2 + month - 1 for month in (3, 4, 5)) / 1000 + 327.442 * 3)
        self.assertTrue(all(row["periodComplete"] for block in first["breakdowns"] for row in block["rows"]))
        for block in first["breakdowns"]:
            self.assertNotIn(block["aggregateField"], [series["id"] for series in block["series"]])
        self.partial_current_book()
        default = self.permits.report("housing", 1)
        self.assertEqual(default["selection"]["structure"]["monthTo"], 10)
        self.assertEqual(default["selection"]["month"], 1)
        self.assertEqual(default["selection"]["structure"]["yearFrom"], 2011)
        self.assertEqual(default["selection"]["structure"]["yearTo"], 2026)

    def test_structure_annual_null_incomplete_current_and_partial_month_coverage(self):
        self.edit_book(OKS_SHEET, {"B50": None})
        payload = self.permits.report("total", 5, structure_mode="annual")
        blocks = self.by_id(payload)
        for identifier in ("housing_components", "nonresidential_native", "area_composition"):
            block = blocks[identifier]
            current = next(row for row in block["rows"] if row["year"] == 2026)
            self.assertFalse(current["annualComplete"])
            self.assertFalse(current["periodComplete"])
            self.assertTrue(all(current[column["id"]] is None for column in block["columns"] if column["id"] != "year"))
            past = next(row for row in block["rows"] if row["year"] == 2025)
            self.assertTrue(past["annualComplete"])
            self.assertTrue(past["periodComplete"])
        old = blocks["housing_components"]["rows"][0]
        self.assertTrue(old["annualComplete"])
        self.assertAlmostEqual(old["living"], sum(41463.8 + month - 1 for month in range(1, 13)) / 1000)
        coverage = next(row for row in payload["controls"]["structure"]["coverage"]
                        if row["blockId"] == "housing_components" and row["year"] == 2026)
        self.assertEqual(coverage["months"], list(range(1, 13)))
        self.assertEqual(coverage["completeMonths"], list(range(1, 10)))
        self.assertEqual(coverage["observedThrough"], "2026-10-16")
        self.assertFalse(coverage["annualComplete"])

    def test_structure_unknown_month_not_january_and_missing_raw_not_zero(self):
        payload = self.permits.report("nonresidential", 1, structure_month_to=1)
        native = self.by_id(payload)["nonresidential_native"]
        row = next(row for row in native["rows"] if row["year"] == 2026)
        self.assertIsNone(row["nonresidential"])
        self.assertFalse(row["periodComplete"])
        coverage = next(row for row in payload["controls"]["structure"]["coverage"]
                        if row["blockId"] == "nonresidential_native" and row["year"] == 2026)
        self.assertEqual(coverage["completeMonths"], [])
        self.assertTrue(any("месяц неизвестен" in ref for ref in coverage["evidence"]))
        # The old raw main table deliberately retains its legacy year -> January fallback.
        raw = next(row for row in payload["permits"]["rows"] if row["Год"] == 2026)
        self.assertEqual(raw["С начала года, тыс. м²"], 41.001)
        self.edit_book(OKS_SHEET, {"Q26": None, "R26": 0})
        house = next(row for row in self.by_id(self.permits.report("housing", 1, structure_month_to=1))
                     ["housing_components"]["rows"] if row["year"] == 2026)
        self.assertEqual(house["living"], 0)
        self.assertIsNone(house["housing_total"])
        self.assertIsNone(house["mop_embedded"])
        self.assertFalse(house["periodComplete"])

    def test_permit_views_use_one_parsed_generation_no_filter_grid(self):
        with patch("backend.permits_page_service.report", wraps=report) as reports:
            first = self.permits.report("total", 1, structure_month_to=3)
            self.assertEqual(reports.call_count, 1)
            generation = self.permits._generation
            original = copy.deepcopy(generation)
            for month in (2, 3, 4):
                self.permits.report("housing", month, structure_month_from=2, structure_month_to=5,
                                    structure_year_from=2012, structure_year_to=2025)
            self.assertEqual(reports.call_count, 4)
            self.assertIs(generation, self.permits._generation)
            self.assertEqual(generation["structures"], original["structures"])
            self.assertEqual(generation["main_sources"], original["main_sources"])
            first["controls"]["structure"]["coverage"][0]["months"].clear()
            first["permits"]["presentation"]["rows"].clear()
            restored = self.permits.report("total", 1, structure_month_to=3)
            self.assertTrue(restored["controls"]["structure"]["coverage"][0]["months"])
            self.assertTrue(restored["permits"]["presentation"]["rows"])

    def test_permit_structure_http_validation_and_all_export_selection_controls(self):
        client = self.client()
        for path in ("permits", "permits/export"):
            for query in ({"structure_mode": "other"}, {"structure_month_from": 0}, {"structure_month_to": 13},
                          {"structure_month_from": 5, "structure_month_to": 3},
                          {"structure_year_from": 2025, "structure_year_to": 2011}):
                params = {"required_version": "not-important", **query}
                self.assertEqual(client.get(f"/api/v1/construction/{path}", params=params).status_code, 422)
            for query in ({"structure_year_from": 2022}, {"structure_year_to": 2040}):
                self.assertEqual(client.get(f"/api/v1/construction/{path}",
                                            params={"required_version": "not-important", **query}).status_code, 404)
        params = {"permit_kind": "total", "month": 5, "structure_mode": "period", "structure_month_from": 3,
                  "structure_month_to": 5, "structure_year_from": 2012, "structure_year_to": 2026}
        payload = client.get("/api/v1/construction/permits", params=params).json()
        response = client.get("/api/v1/construction/permits/export", params={**params, "required_version": payload["version"]})
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        self.addCleanup(book.close)
        selected = dict(book["Selection"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(selected, {"permitKind": "total", "month": 5, "structure.mode": "period",
                                    "structure.monthFrom": 3, "structure.monthTo": 5,
                                    "structure.yearFrom": 2012, "structure.yearTo": 2026})
        self.assertIn("Permit Presentation", book.sheetnames)
        self.assertIn("Structure Coverage", book.sheetnames)
        self.assertIn("Controls", book.sheetnames)
        full = {row[0]: (row[1], row[3]) for row in book["Full Payload"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(full["/selection/structure/monthFrom"], ("int", 3))
        index = next(i for i, row in enumerate(payload["permits"]["presentation"]["rows"]) if row["year"] == 2026)
        self.assertEqual(full[f"/permits/presentation/rows/{index}/annualActual"], ("null", None))
        self.assertEqual(full[f"/permits/presentation/rows/{index}/annualForecast"][1], 5810.123456789 + 6850.987654321)

    def test_domrf_literal_units_distinct_periods_and_only_observed_history(self):
        payload = self.domrf.report()
        self.assertEqual(set(payload), {"schemaVersion", "version", "generatedAt", "source", "region", "construction", "sales", "notes"})
        self.assertEqual(payload["construction"]["period"], {"id": "2026-10", "label": "Октябрь 2026", "reportDate": "2026-10-01"})
        self.assertEqual(payload["sales"]["period"], {"id": "2026-08", "label": "Август 2026", "reportDate": "2026-09-11"})
        self.assertEqual(payload["construction"]["metrics"], {"living": 15640, "total": 26964})
        metrics = payload["sales"]["metrics"]
        self.assertEqual(metrics["sold_thousand_m2"], 7366)
        self.assertEqual(metrics["sold_pct"], 47)
        self.assertEqual(metrics["price_per_m2_rub"], 441306)
        self.assertEqual(metrics["funds_million_rub"], 3250869)
        self.assertEqual(metrics["precise_pct"], 42.123456789)
        self.assertIsNone(metrics["blank"])
        self.assertEqual(metrics["zero"], 0)
        self.assertEqual(len(payload["construction"]["history"]), 1)
        self.assertEqual(len(payload["sales"]["history"]), 1)
        self.assertEqual(payload["sales"]["history"][0]["period"], "2026-08")
        json.dumps(payload, allow_nan=False)

    def test_both_services_five_second_cache_and_retry_after_expiry(self):
        for cls in (PermitsPageService, DomrfService):
            with self.subTest(service=cls.__name__):
                service = cls(self.context)
                first = service.report()
                with patch.object(service, "_inputs", side_effect=SourceUnavailable("new broken source")):
                    self.assertEqual(first, service.report())
                    service._checked -= 6
                    for _ in range(2):
                        with self.assertRaises(SourceUnavailable):
                            service.report()
                self.assertTrue(service.report()["version"])

    def test_only_used_static_control_changes_permit_version(self):
        original = self.permits.report("housing", 1)
        directory = self.root / "data/raw/realty/vvod"
        directory.mkdir(parents=True)
        path = directory / "static_vvod_rs_2011_2025.xlsx"
        book = Workbook()
        book.active.append(["year", "permits", "input", "", "year", "permits", "input"])
        book.active.append([2011, 1234.123456789, 0, "", 2011, 4567.987654321, 0])
        book.save(path)
        book.close()
        self.touch(path)
        updated = self.permits.report("housing", 1)
        self.assertNotEqual(original["version"], updated["version"])
        self.assertEqual(updated["permits"]["rows"][0]["За год, тыс. м²"], 1234.123456789)
        self.assertEqual(self.presentation(updated, 2011)["annualActual"], 1234.123456789)
        self.assertEqual(self.presentation(updated, 2011)["basis"], "annual_observed")
        self.assertTrue(any("static_vvod_rs_2011_2025.xlsx" in ref for ref in self.presentation(updated, 2011)["evidence"]))
        self.assertEqual(len(updated["source"]["files"]), 2)
        (directory / "static_vvod_rs_unused.xlsx").write_bytes(b"corrupt unused file")
        self.assertEqual(updated["version"], self.permits.report("housing", 1)["version"])
        old_b = self.domrf.report()["version"]
        self.touch(path)
        self.assertEqual(old_b, self.domrf.report()["version"])

    def test_domrf_rejects_invalid_shapes_and_nonfinite_json_no_cached_fallback(self):
        self.domrf.report()
        invalid = [{"construction": False}, {"sales": "bad"}, {"construction": [1]},
                   {"sales": [{"region_key": "msk", "metrics": []}]},
                   {"construction": [{"region_key": "msk", "area_kind": "living", "report_period": 2026}]}]
        for payload in invalid:
            with self.subTest(payload=payload):
                self.write_json(payload)
                self.assertEqual(self.client().get("/api/v1/construction/domrf").status_code, 503)
        self.json_path.write_text('{"sales": NaN}', encoding="utf-8")
        self.touch(self.json_path)
        self.assertEqual(self.client().get("/api/v1/construction/domrf").status_code, 503)
        self.write_json()
        self.assertEqual(self.client().get("/api/v1/construction/domrf").status_code, 200)

    def test_domrf_unknown_area_is_not_a_fabricated_history_bar(self):
        payload = {"construction": [{"region_key": "msk", "area_kind": "unknown", "report_period": "Январь 2026",
                                     "report_date": "01.02.2026", "area_thousand_m2": 100}]}
        self.write_json(payload)
        self.assertEqual(self.domrf.report()["construction"]["history"], [])

    def test_http_old_version_rejected_after_real_source_replacement(self):
        client = self.client()
        for page in ("permits", "domrf"):
            before = client.get(f"/api/v1/construction/{page}").json()
            if page == "permits":
                self.edit_book(FORECAST_SHEET, {"D3": 0})
            else:
                data = self.json_payload()
                data["sales"][0]["metrics"]["sold_pct"] = 0
                self.write_json(data)
            response = client.get(f"/api/v1/construction/{page}/export", params={"required_version": before["version"]})
            self.assertEqual(response.status_code, 409)
            self.assertNotEqual(response.headers["x-data-version"], before["version"])

    def test_all_registered_history_duplicate_period_full_snapshot_latest_mtime(self):
        older = self.json_payload()
        for row in older["construction"]:
            row["report_period"] = "Июнь 2026"
            row["report_date"] = "01.12.2030"
        for row in older["sales"]:
            row["report_period"] = "Май 2026"
        older_path = self.raw / "construction_operational_20260601.json"
        self.write_json(older, older_path)
        replacement = self.json_payload()
        replacement["construction"] = [replacement["construction"][0]]
        replacement["construction"][0]["area_thousand_m2"] = 123.123456789
        replacement["sales"][0]["metrics"] = {"sold_pct": 0}
        new_path = self.raw / "construction_operational_20260901.json"
        self.write_json(replacement, new_path)
        # Archive and downloads are outside the active registered source directories.
        archive = self.raw / "archive"
        archive.mkdir()
        (archive / "construction_operational_99999999.json").write_text("corrupt", encoding="utf-8")
        payload = self.domrf.report()
        self.assertEqual([row["period"] for row in payload["construction"]["history"]], ["2026-06", "2026-10"])
        self.assertEqual(payload["construction"]["metrics"], {"living": 123.123456789, "total": None})
        self.assertEqual(payload["sales"]["metrics"], {"sold_pct": 0})
        self.assertEqual(payload["construction"]["history"][-1]["sourceFile"], "data/raw/realty/nashdom/" + new_path.name)
        self.assertEqual(len(payload["source"]["files"]), 3)

    def test_missing_block_period_and_blank_values_no_fake_chronology(self):
        payload = self.json_payload()
        payload.pop("sales")
        for row in payload["construction"]:
            row.pop("report_period")
            row["area_thousand_m2"] = None if row["area_kind"] == "living" else 0
        self.write_json(payload)
        actual = self.domrf.report()
        self.assertEqual(actual["sales"], {"period": None, "metrics": None, "history": []})
        self.assertEqual(actual["construction"]["period"], {"id": None, "label": "", "reportDate": "2026-10-01"})
        self.assertEqual(actual["construction"]["history"], [])
        self.assertEqual(actual["construction"]["metrics"], {"living": None, "total": 0})
        self.assertTrue(all(set(issue) == {"level", "message"} for issue in actual["source"]["issues"]))
        self.write_json({"sales": self.json_payload()["sales"]})
        actual = self.domrf.report()
        self.assertEqual(actual["construction"], {"period": None, "metrics": {"living": None, "total": None}, "history": []})
        self.assertTrue(actual["sales"]["metrics"])

    def test_services_are_independent_and_signatures_only_used_sources(self):
        a, b = self.permits.report(), self.domrf.report()
        self.json_path.write_text("corrupt", encoding="utf-8")
        self.touch(self.json_path)
        for name in ("rasprodannost_20261019.xlsx", "not-a-source.json"):
            (self.raw / name).write_text("corrupt", encoding="utf-8")
        marts = self.root / "data/marts/realty"
        marts.mkdir(parents=True)
        (marts / "manifest.json").write_text("corrupt", encoding="utf-8")
        self.assertEqual(a, self.permits.report())
        with self.assertRaises(SourceUnavailable):
            self.domrf.report()
        self.write_json()
        b = self.domrf.report()
        self.book_path.write_bytes(b"broken workbook")
        self.touch(self.book_path)
        self.assertEqual(b, self.domrf.report())
        with self.assertRaises(SourceUnavailable):
            self.permits.report()

    def test_latest_monitoring_uses_signature_mtime_not_size_or_filename(self):
        older = self.raw / "monitoring_2_0_20250101.xlsx"
        self.write_book(older)
        os.utime(older, (self.timestamp - 100, self.timestamp - 100))
        initial = self.permits.report()
        self.assertTrue(initial["source"]["files"][0].endswith("20261016.xlsx"))
        # A larger file must not win selection by signature[1] (size).
        with older.open("ab") as stream:
            stream.write(b" " * 50000)
        os.utime(older, (self.timestamp - 100, self.timestamp - 100))
        self.assertEqual(initial["version"], self.permits.report()["version"])
        self.touch(older)
        self.assertTrue(self.permits.report()["source"]["files"][0].endswith("20250101.xlsx"))
        older.write_bytes(b"corrupt latest")
        self.touch(older)
        self.assertEqual(self.client().get("/api/v1/construction/permits").status_code, 503)

    def test_same_process_refresh_deepcopy_cache_interval_and_failed_refresh(self):
        before = self.permits.report("housing", 1)
        before["controls"]["monthsByKind"]["housing"].clear()
        before["breakdowns"][0]["rows"].clear()
        self.assertTrue(self.permits.report("housing", 1)["breakdowns"][0]["rows"])
        b = self.domrf.report()
        b["sales"]["metrics"]["sold_pct"] = -999
        self.assertEqual(self.domrf.report()["sales"]["metrics"]["sold_pct"], 47)
        cached = DomrfService(self.context)
        first = cached.report()
        with patch.object(cached, "_inputs", side_effect=AssertionError("cache should throttle source checks")):
            self.assertEqual(first, cached.report())
        self.write_json({"construction": self.json_payload()["construction"]})
        changed = self.domrf.report()
        self.assertNotEqual(first["version"], changed["version"])
        self.assertIsNone(changed["sales"]["metrics"])
        self.json_path.write_text("{", encoding="utf-8")
        self.touch(self.json_path)
        for _ in range(2):
            self.assertEqual(self.client().get("/api/v1/construction/domrf").status_code, 503)
        self.write_json()
        self.assertEqual(self.client().get("/api/v1/construction/domrf").status_code, 200)

    def test_atomic_source_guard_rejects_changed_generation(self):
        original = self.permits._inputs
        calls = 0

        def changed():
            nonlocal calls
            calls += 1
            value = original()
            return value if calls == 1 else (value[0], ("changed",), value[2])

        with patch.object(self.permits, "_inputs", side_effect=changed):
            with self.assertRaises(SourceUnavailable):
                self.permits.report()
        self.assertIsNone(self.permits._generation)
        self.assertTrue(self.permits.report()["permits"]["rows"])
        original = self.domrf._inputs
        with patch.object(self.domrf, "_inputs", side_effect=[original(), ([], ("changed",))]):
            with self.assertRaises(SourceUnavailable):
                self.domrf.report()

    def test_workbook_xml_and_zlib_corruption_become_503(self):
        client = self.client()
        for error in (DecompressionError("bad compression"), ET.ParseError("bad XML")):
            with self.subTest(error=type(error).__name__):
                with patch("backend.permits_page_service.load_workbook", side_effect=error):
                    self.assertEqual(client.get("/api/v1/construction/permits").status_code, 503)
        self.assertEqual(client.get("/api/v1/construction/permits").status_code, 200)

    def test_http_filters_required_version_409_and_full_precision_exports(self):
        client = self.client()
        for path in ("permits", "domrf"):
            self.assertEqual(client.get(f"/api/v1/construction/{path}/export").status_code, 422)
            payload = client.get(f"/api/v1/construction/{path}").json()
            self.assertEqual(client.get(f"/api/v1/construction/{path}/export", params={"required_version": "stale"}).status_code, 409)
            response = client.get(f"/api/v1/construction/{path}/export", params={"required_version": payload["version"]})
            self.assertEqual(response.status_code, 200, response.text[:500] if response.status_code != 200 else "")
            self.assertEqual(response.headers["x-data-version"], payload["version"])
            book = load_workbook(BytesIO(response.content))
            full = {row[0]: row[3] for row in book["Full Payload"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(full["/version"], payload["version"])
            self.assertEqual(full["/notes/0"], payload["notes"][0])
            if path == "permits":
                self.assertEqual(book["Forecast"]["C2"].value, 5810.123456789)
                self.assertEqual(full["/forecast/rows/0/cells/housing_living"], "D3")
                self.assertEqual(book["Permits"].max_row, len(payload["permits"]["rows"]) + 1)
                self.assertEqual(full["/controls/monthsByKind/housing/11"], 12)
            else:
                self.assertEqual(full["/sales/metrics/precise_pct"], 42.123456789)
                self.assertEqual(book["Sales History"].max_row, len(payload["sales"]["history"]) + 1)
                self.assertEqual(full["/construction/metrics/living"], 15640)
            book.close()
        for query in ("permit_kind=invalid", "month=0", "month=13"):
            self.assertEqual(client.get("/api/v1/construction/permits?" + query).status_code, 422)
        self.assertEqual(client.get("/api/v1/construction/domrf?region=other").status_code, 422)
        selected = client.get("/api/v1/construction/permits?permit_kind=housing&month=2").json()
        response = client.get("/api/v1/construction/permits/export", params={"permit_kind": "housing", "month": 2,
                                                                            "required_version": selected["version"]})
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book["Selection"]["B2"].value, "housing")
        self.assertEqual(book["Selection"]["B3"].value, 2)
        book.close()

    def test_exports_are_formula_injection_safe_and_inputs_immutable(self):
        payload = self.permits.report()
        payload["notes"].extend(["=HYPERLINK(\"bad\")", " +SUM(1,1)", "\t@bad", "-bad"])
        payload["forecast"]["notes"].append("x" * 35000)
        before = copy.deepcopy(payload)
        signatures = [file_signature(path) for path in (self.book_path, self.json_path)]
        book = load_workbook(BytesIO(export_workbook(payload)))
        self.assertTrue(all(cell.data_type != "f" for sheet in book for row in sheet for cell in row))
        parts = [row[3] for row in book["Full Payload"].iter_rows(min_row=2, values_only=True)
                 if row[0] == "/forecast/notes/2"]
        self.assertEqual("".join(parts), "x" * 35000)
        self.assertEqual(payload, before)
        self.assertEqual(signatures, [file_signature(path) for path in (self.book_path, self.json_path)])
        book.close()


if __name__ == "__main__":
    unittest.main()
