"""Mortgage endpoints and version-pinned, unrounded source exports."""
from io import BytesIO
import json

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from backend.mortgage_service import MortgageService
from pipeline.mortgage_calculations import METRICS


def export_workbook(payload, source):
    book = Workbook()
    overflow = None

    def cell_value(value, location):
        nonlocal overflow
        text = json.dumps(value, ensure_ascii=False, allow_nan=False) if isinstance(value, (dict, list)) else value
        if not isinstance(text, str) or len(text) <= 30000:
            return text
        if overflow is None:
            overflow = book.create_sheet("Long source fields")
            overflow.append(["Field", "Part", "Text"])
        for start in range(0, len(text), 30000):
            overflow.append([location, start // 30000 + 1, text[start:start + 30000]])
        return "Long source fields: " + location

    metadata = book.active
    metadata.title = "Report"
    metadata.append(["Field", "Value"])
    for key, value in {"version": payload["version"], "generatedAt": payload["generatedAt"],
                       **payload["selection"], "sourceUrl": payload["source"]["url"],
                       "sourceDate": payload["source"]["date"], "collectedAt": payload["source"]["collectedAt"],
                       "sourceFile": "; ".join(payload["source"]["files"])}.items():
        metadata.append([key, value])
    raw = book.create_sheet("Source observations")
    headers = ["region", "period", "segment", "basis", *METRICS]
    extra = sorted({key for row in source["records"] for key in row} - set(headers))
    raw.append(headers + extra)
    for index, row in enumerate(source["records"]):
        raw.append([cell_value(row.get(key), f"records/{index}/{key}") for key in headers + extra])
    annual = book.create_sheet("Annual calculations")
    fields = ("annual", "selected", "actualThroughMonth", "annualYoYPct", "selectedYoYPct")
    annual.append(["year", "metric", *fields])
    for row in payload["annual"]:
        for metric in ("volume", "count"):
            annual.append([row["year"], metric, *[row[metric][key] for key in fields]])
    monthly = book.create_sheet("Monthly selected region")
    monthly.append(["period", "segment", *METRICS])
    for row in payload["monthly"]:
        for segment in ("all", "primary", "secondary"):
            monthly.append([row["period"], segment, *[row[segment][key] for key in METRICS]])
    notes = book.create_sheet("Sources and method")
    notes.append(["Note"])
    for note in ["Volume is in rubles; count is in contracts; rate is in percent; term is in months.",
                 "Annual totals require twelve observed monthly flows. Missing values are not zero.",
                 "Selected values sum January through the selected month; the latest year uses the available interval.",
                 "Year-on-year changes compare matching calendar intervals only.",
                 "Rates, terms and average loan sizes are source values, not averages calculated by the API.",
                 "Primary and secondary series are independent source segments, not inferred residuals.",
                 *payload["notes"]]:
        notes.append([cell_value(note, f"notes/{notes.max_row}")])
    for key, value in source.items():
        if key != "records":
            notes.append([f"{key}: " + str(cell_value(value, f"metadata/{key}"))])
    for tab in book:
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for row in tab:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.lstrip().startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
        for cell in tab[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1C3444")
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = 27
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def create_router(context, call, service=None):
    router = APIRouter(prefix="/api/v1/mortgage")
    service = service or MortgageService(context)

    @router.get("/catalog")
    def catalog():
        return call(service.catalog, scope="Mortgage")

    @router.get("")
    def report(region: str | None = None, month: int = 12):
        return call(service.report, region, month, scope="Mortgage")

    @router.get("/export")
    def export(required_version: str = Query(min_length=1, max_length=128), region: str | None = None, month: int = 12):
        payload, source = call(service.export_snapshot, region, month, scope="Mortgage")
        if required_version != payload["version"]:
            raise HTTPException(409, "Source changed; refresh the displayed mortgage report before exporting",
                                headers={"X-Data-Version": payload["version"]})
        content = call(export_workbook, payload, source, scope="Mortgage")
        return Response(content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="mortgage.xlsx"',
                                 "X-Data-Version": payload["version"]})

    return router
