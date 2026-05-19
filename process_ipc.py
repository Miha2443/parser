"""
Обработка данных Росстата для DataLens / Excel.

Архитектура:
  - Лист 'data' разбит на блоки — по одному на каждый показатель (ИПЦ, ЗП, ...).
  - Каждый блок имеет свою шапку и оформлен как «умная таблица» с уникальным
    именем (tbl_IPC, tbl_ZP, ...).
  - Между блоками — пустая строка.
  - Имена умных таблиц детерминированы (один и тот же показатель → всегда то
    же имя), поэтому сводные на других листах после обновления данных не
    теряют связь с источником.
  - Скрипт обновляет ТОЛЬКО лист 'data'. Прочие листы (ваши сводные)
    остаются нетронутыми.
  - Перед каждым обновлением создаётся бэкап (хранится 10 последних).
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# ── Настройки путей ──────────────────────────────────────────────────────────
DOWNLOADS_DIR = Path(r"C:\cloud\downloads")
PROCESSED_DIR = Path(r"C:\cloud\processed")
BACKUPS_DIR = PROCESSED_DIR / "backups"
OUTPUT_FILE = PROCESSED_DIR / "ipc_latest.xlsx"

BACKUPS_KEEP = 10

# Паттерны для поиска свежих файлов (берём самый поздний по имени)
IPC_PART1_PATTERN = "*Индексы потребительских цен*часть1*.xls"
IPC_PART2_PATTERN = "*Индексы потребительских цен*часть2*.xls"
ZP_PATTERN = "*Среднемесячная номинальная начисленная заработная плата*.xls"

# Имя листа, которым управляет скрипт. Все остальные листы не трогаются.
DATA_SHEET = "data"

# Стиль умной таблицы по умолчанию
DEFAULT_TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2", showRowStripes=True
)

# Колонки итоговых данных (общая шапка для всех блоков)
DATA_COLUMNS = [
    "показатель", "вид", "регион",
    "год", "месяц", "период_тип",
    "значение", "единица",
]

# ── Общие справочники ───────────────────────────────────────────────────────
MONTHS = {
    "январь": 1, "февраль": 2, "март": 3, "апрель": 4, "май": 5, "июнь": 6,
    "июль": 7, "август": 8, "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12,
}
MONTH_NAME_BY_NUM = {v: k for k, v in MONTHS.items()}

REGION_CLEAN = {
    "Российская Федерация": "Российская Федерация",
    "Российская Федерация без учета новых субъектов (с 01.01.2023)": "Российская Федерация без учета новых субъектов",
    "Город Москва столица Российской Федерации город федерального значения": "Москва",
}

# ── Настройки источника ИПЦ ──────────────────────────────────────────────────
IPC_PERIOD_LABEL = {
    "К предыдущему месяцу": "месяц к месяцу",
    "Период с начала года к соответствующему периоду предыдущего года": "Период с начала года к АППГ",
}

# ── Настройки источника по зарплате ──────────────────────────────────────────
ZP_INDICATOR = "Среднемесячная номинальная начисленная заработная плата"
ZP_VIEW_MAP = {
    "Всего по обследуемым видам экономической деятельности": "Всего",
    "    СТРОИТЕЛЬСТВО": "Строительство",
}
ZP_REGIONS_KEEP = {
    "Российская Федерация",
    "Российская Федерация без учета новых субъектов",
    "Москва",
}


# ── Поиск свежих файлов ──────────────────────────────────────────────────────
def find_latest(pattern: str, required: bool = True) -> Path | None:
    files = sorted(DOWNLOADS_DIR.glob(pattern))
    if not files:
        if required:
            raise FileNotFoundError(f"Не найдены файлы по паттерну {pattern} в {DOWNLOADS_DIR}")
        return None
    return files[-1]


# ── Парсинг ИПЦ (широкая таблица) ────────────────────────────────────────────
def parse_ipc(xls_path: Path) -> pd.DataFrame:
    df = pd.read_excel(xls_path, sheet_name="Данные", header=None)
    years = df.iloc[2].ffill()
    months = df.iloc[3].ffill()
    idx_types = df.iloc[4]

    records = []
    for row_idx in range(5, len(df)):
        indicator = df.iat[row_idx, 0]
        unit = df.iat[row_idx, 1]
        region_raw = df.iat[row_idx, 2]
        view_raw = df.iat[row_idx, 3]

        if pd.isna(region_raw):
            continue

        region = REGION_CLEAN.get(str(region_raw).strip(), str(region_raw).strip())

        for col in range(4, df.shape[1]):
            value = df.iat[row_idx, col]
            if pd.isna(value):
                continue
            year = years.iat[col]
            month_name = months.iat[col]
            idx_type_raw = idx_types.iat[col]
            if pd.isna(year) or pd.isna(month_name) or pd.isna(idx_type_raw):
                continue
            month_num = MONTHS.get(str(month_name).strip().lower())
            if month_num is None:
                continue
            period_label = IPC_PERIOD_LABEL.get(str(idx_type_raw).strip())
            if period_label is None:
                continue
            try:
                value_num = float(value)
            except (TypeError, ValueError):
                continue

            records.append({
                "показатель": str(indicator).strip() if pd.notna(indicator) else "",
                "вид": str(view_raw).strip() if pd.notna(view_raw) else "",
                "регион": region,
                "год": int(year),
                "месяц": MONTH_NAME_BY_NUM[month_num],
                "период_тип": period_label,
                "значение": round(value_num, 4),
                "единица": str(unit).strip() if pd.notna(unit) else "процент",
            })

    return pd.DataFrame.from_records(records)


# ── Парсинг зарплаты ─────────────────────────────────────────────────────────
def parse_zp(xls_path: Path) -> pd.DataFrame:
    df = pd.read_excel(xls_path, sheet_name="Данные", header=None)
    regions = df.iloc[2].ffill()
    years = df.iloc[3]

    records = []
    for row_idx in range(4, len(df)):
        view_raw = df.iat[row_idx, 0]
        period_raw = df.iat[row_idx, 1]
        if pd.isna(view_raw) or pd.isna(period_raw):
            continue

        view = ZP_VIEW_MAP.get(str(view_raw))
        if view is None:
            continue

        period_str = str(period_raw).strip().lower()
        if "-" in period_str:
            last_month_name = period_str.split("-")[-1].strip()
            period_label = "период с начала года"
        else:
            last_month_name = period_str
            period_label = "месяц"
        month_num = MONTHS.get(last_month_name)
        if month_num is None:
            continue

        for col in range(2, df.shape[1]):
            value = df.iat[row_idx, col]
            if pd.isna(value):
                continue
            region_raw = regions.iat[col]
            year_raw = years.iat[col]
            if pd.isna(region_raw) or pd.isna(year_raw):
                continue
            region = REGION_CLEAN.get(str(region_raw).strip(), str(region_raw).strip())
            if region not in ZP_REGIONS_KEEP:
                continue
            try:
                value_num = float(value)
                year = int(float(year_raw))
            except (TypeError, ValueError):
                continue

            records.append({
                "показатель": ZP_INDICATOR,
                "вид": view,
                "регион": region,
                "год": year,
                "месяц": MONTH_NAME_BY_NUM[month_num],
                "период_тип": period_label,
                "значение": round(value_num, 4),
                "единица": "рубль",
            })

    return pd.DataFrame.from_records(records)


# ── Сортировка строк блока ───────────────────────────────────────────────────
def sort_block(df: pd.DataFrame) -> pd.DataFrame:
    """Сортировка внутри блока: показатель, регион, вид, год, месяц, тип."""
    df = df.copy()
    df["_месяц_номер"] = df["месяц"].map(MONTHS)
    df = (
        df.drop_duplicates(
            subset=["показатель", "вид", "регион", "год", "месяц", "период_тип"],
            keep="last",
        )
        .sort_values(["показатель", "регион", "вид", "год", "_месяц_номер", "период_тип"])
        .drop(columns=["_месяц_номер"])
        .reset_index(drop=True)
    )
    # Гарантируем порядок колонок
    return df[DATA_COLUMNS]


# ── Бэкап ────────────────────────────────────────────────────────────────────
def make_backup() -> None:
    """Копирует текущий ipc_latest.xlsx в backups/ с номером NNN."""
    if not OUTPUT_FILE.exists():
        print("  ℹ️  Бэкап: исходный файл не существует (первый запуск)")
        return

    BACKUPS_DIR.mkdir(parents=True, exist_ok=True)
    today = datetime.now().strftime("%Y%m%d")
    existing = list(BACKUPS_DIR.glob(f"ipc_latest_{today}_*.xlsx"))
    nums = []
    for f in existing:
        try:
            n = int(f.stem.split("_")[-1])
            nums.append(n)
        except ValueError:
            continue
    next_num = max(nums, default=0) + 1
    backup_path = BACKUPS_DIR / f"ipc_latest_{today}_{next_num:03d}.xlsx"
    shutil.copy2(OUTPUT_FILE, backup_path)
    print(f"  💾 Бэкап: {backup_path}")

    # Чистим старые: оставляем только BACKUPS_KEEP последних по имени
    # (имя содержит дату и номер, сортировка по имени = по времени создания)
    all_backups = sorted(
        BACKUPS_DIR.glob("ipc_latest_*.xlsx"),
        reverse=True,
    )
    for old in all_backups[BACKUPS_KEEP:]:
        old.unlink()
        print(f"  🗑️  Удалён старый бэкап: {old.name}")


# ── Запись листа data с блоками ──────────────────────────────────────────────
def write_data_sheet(wb, blocks: list[tuple[str, pd.DataFrame]]) -> None:
    """
    Создаёт лист data с блоками. blocks = [(table_name, dataframe), ...].
    Каждый блок: шапка + данные, оформлен как умная таблица.
    Между блоками — одна пустая строка.
    """
    # Удаляем старый лист data, если есть
    if DATA_SHEET in wb.sheetnames:
        del wb[DATA_SHEET]
    ws = wb.create_sheet(DATA_SHEET, index=0)  # сделаем data первым листом

    current_row = 1  # текущая строка для записи

    for table_name, df in blocks:
        if df.empty:
            continue

        # Шапка блока
        header_row = current_row
        for col_idx, col_name in enumerate(DATA_COLUMNS, start=1):
            ws.cell(row=header_row, column=col_idx, value=col_name)

        # Данные
        for r_offset, row in enumerate(df.itertuples(index=False), start=1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=header_row + r_offset, column=c_idx, value=value)

        last_row = header_row + len(df)
        # Создаём умную таблицу на диапазон шапка+данные
        ref = f"A{header_row}:{_col_letter(len(DATA_COLUMNS))}{last_row}"
        tbl = Table(displayName=table_name, ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name=DEFAULT_TABLE_STYLE.name,
            showRowStripes=DEFAULT_TABLE_STYLE.showRowStripes,
        )
        ws.add_table(tbl)

        # Сдвигаемся на 2 строки вниз: одну пустую + начало следующего блока
        current_row = last_row + 2


def _col_letter(col_num: int) -> str:
    """Номер колонки (1=A, 2=B...) → буква."""
    letters = ""
    while col_num > 0:
        col_num, rem = divmod(col_num - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


# ── Запись итогового файла ───────────────────────────────────────────────────
def write_output(blocks: list[tuple[str, pd.DataFrame]]) -> None:
    """
    Если файл существует — открывает его, заменяет лист 'data'.
    Прочие листы остаются нетронутыми.
    Если файла нет — создаёт новый.
    Если файл повреждён — переименовывает в BROKEN и создаёт новый.
    """
    if OUTPUT_FILE.exists():
        try:
            wb = load_workbook(OUTPUT_FILE)
        except Exception as e:
            broken_path = OUTPUT_FILE.parent / (
                f"ipc_latest_BROKEN_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
            )
            OUTPUT_FILE.rename(broken_path)
            print(f"  ⚠️  Существующий файл повреждён ({type(e).__name__}): {e}")
            print(f"  📦 Переименован в: {broken_path}")
            print(f"  🆕 Создаю новый файл с нуля")
            wb = _new_workbook()
    else:
        wb = _new_workbook()

    write_data_sheet(wb, blocks)
    wb.save(OUTPUT_FILE)
    wb.close()


def _new_workbook():
    """Создаёт новую пустую книгу."""
    from openpyxl import Workbook
    wb = Workbook()
    # Удаляем дефолтный пустой лист — write_data_sheet создаст свой
    wb.remove(wb.active)
    return wb


# ── Главный процесс ──────────────────────────────────────────────────────────
def run():
    PROCESSED_DIR.mkdir(exist_ok=True)

    # Блоки в порядке обработки: (имя_умной_таблицы, dataframe)
    blocks: list[tuple[str, pd.DataFrame]] = []

    # ИПЦ
    ipc_p1 = find_latest(IPC_PART1_PATTERN, required=False)
    ipc_p2 = find_latest(IPC_PART2_PATTERN, required=False)
    if ipc_p1 and ipc_p2:
        print(f"📂 ИПЦ ч.1: {ipc_p1.name}")
        print(f"📂 ИПЦ ч.2: {ipc_p2.name}")
        df1 = parse_ipc(ipc_p1)
        df2 = parse_ipc(ipc_p2)
        ipc = sort_block(pd.concat([df1, df2], ignore_index=True))
        print(f"  ✓ ИПЦ: {len(ipc)} строк, {ipc['год'].min()}–{ipc['год'].max()}")
        blocks.append(("tbl_IPC", ipc))
    else:
        print("⚠️  Файлы ИПЦ не найдены, пропускаю")

    # Зарплата
    zp = find_latest(ZP_PATTERN, required=False)
    if zp:
        print(f"📂 Зарплата: {zp.name}")
        zp_df = sort_block(parse_zp(zp))
        print(f"  ✓ Зарплата: {len(zp_df)} строк, {zp_df['год'].min()}–{zp_df['год'].max()}")
        blocks.append(("tbl_ZP", zp_df))
    else:
        print("⚠️  Файл зарплаты не найден, пропускаю")

    if not blocks:
        raise RuntimeError("Не найдено ни одного источника")

    total = sum(len(df) for _, df in blocks)
    print(f"\n  ✓ Всего: {total} строк в {len(blocks)} блоках")
    for name, df in blocks:
        print(f"      {name}: {len(df)} строк")

    # Бэкап → запись
    make_backup()
    write_output(blocks)

    print(f"\n💾 Обновлён файл: {OUTPUT_FILE}")
    print(f"\nГотово ✓")


if __name__ == "__main__":
    try:
        run()
    except Exception as e:
        print(f"❌ Ошибка: {e}", file=sys.stderr)
        sys.exit(1)