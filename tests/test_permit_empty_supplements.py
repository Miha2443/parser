"""Absent optional worksheet ranges stay unavailable instead of breaking permits."""
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from pipeline.permit_page_calculations import (FORECAST_SHEET, HISTORY_SHEET, HOUSING_SHEET,
                                               OKS_SHEET, read_breakdown_sources, read_forecast)


class EmptySupplementTests(unittest.TestCase):
    def test_short_forecast_preserves_known_value_and_leaves_missing_measures_null(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "monitoring.xlsx"
            book = Workbook()
            book.active.title = FORECAST_SHEET
            book.active["D2"], book.active["D3"] = 2026, 5810
            book.save(path)
            book.close()
            for read_only in (False, True):
                with self.subTest(read_only=read_only):
                    saved = load_workbook(path, data_only=True, read_only=read_only)
                    try:
                        result = read_forecast(saved, path.name)
                    finally:
                        saved.close()
                    self.assertEqual(len(result["rows"]), 1)
                    row = result["rows"][0]
                    self.assertEqual((row["year"], row["housing_living"]), (2026, 5810))
                    self.assertIsNone(row["nonresidential"])
                    self.assertEqual(row["cells"]["nonresidential"], "D6")

    def test_short_and_empty_optional_ranges_remain_empty_in_read_only_workbooks(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "monitoring.xlsx"
            for short in (False, True):
                book = Workbook()
                book.remove(book.active)
                for name in (FORECAST_SHEET, HISTORY_SHEET, HOUSING_SHEET, OKS_SHEET):
                    sheet = book.create_sheet(name)
                    if short:
                        sheet.append(["Unrelated source header"])
                book.save(path)
                book.close()
                for read_only in (False, True):
                    with self.subTest(short=short, read_only=read_only):
                        saved = load_workbook(path, data_only=True, read_only=read_only)
                        try:
                            forecast = read_forecast(saved, path.name)
                            structure = read_breakdown_sources(saved)
                        finally:
                            saved.close()
                        self.assertEqual(forecast["rows"], [])
                        self.assertTrue(forecast["notes"])
                        for field in ("housing", "nonresidential", "groups"):
                            self.assertEqual(structure[field], {})
                        self.assertEqual(structure["modern"], [])
                        for field in ("housing_notes", "historical_notes", "modern_notes"):
                            self.assertTrue(structure[field])


if __name__ == "__main__":
    unittest.main()
