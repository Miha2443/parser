"""Full pinned sales payload, including every chart point and all six raw tables."""
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def export_workbook(payload):
    book = Workbook()
    book.remove(book.active)

    def safe(value):
        if isinstance(value, (list, dict)):
            value = json.dumps(value, ensure_ascii=False, allow_nan=False)
        return value

    def sheet(name, headers, rows):
        tab = book.create_sheet(name[:31])
        tab.append([safe(value) for value in headers])
        for row in rows:
            tab.append([safe(value) for value in row])
        # Explicit strings prevent formula execution without altering raw text.
        for row in tab:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1C3444")
        for index, header in enumerate(headers, 1):
            tab.column_dimensions[get_column_letter(index)].width = min(55, max(18, len(header) + 2))

    sheet("Отчёт", ["Показатель", "Значение"], [
        ["Регион", payload["region"]["label"]], ["Период", payload["period"]["label"]],
        ["Период ID", payload["period"]["id"]], ["Версия источников", payload["version"]],
        ["Сформировано (UTC)", payload["generatedAt"]], ["Версия схемы", payload["schemaVersion"]],
    ])
    sheet("KPI", ["ID", "Показатель", "Единица", "Значение", "Знаков после запятой"],
          [[metric[key] for key in ("id", "label", "unit", "value", "digits")] for metric in payload["metrics"]])
    for chart in payload["charts"]:
        sheet(chart["id"], ["Серия ID", "Показатель", "Единица серии", "Период / год", "Значение"],
              [[series["id"], series["name"], series["unit"], point["x"], point["y"]]
               for series in chart["series"] for point in series["points"]])
    sheet("Графики", ["ID", "Название", "Тип", "Единица оси"],
          [[chart[key] for key in ("id", "title", "kind", "unit")] for chart in payload["charts"]])
    for table in payload["tables"]:
        columns = table["columns"]
        sheet(table["title"], [column["label"] for column in columns],
              [[row[column["id"]] for column in columns] for row in table["rows"]])
    sheet("Колонки", ["Таблица ID", "Исходная колонка ID", "Подпись"],
          [[table["id"], column["id"], column["label"]] for table in payload["tables"] for column in table["columns"]])
    source = payload["source"]
    sheet("Источники", ["Показатель", "Значение"], [
        ["Дата источников", source["date"]], ["Файлы-кандидаты", source["files"]],
        ["Диагностика загрузчика", source["issues"]],
        ["Примечание", "Кандидаты DataAccess, не доказательство участия; выбор mart/raw и fallback не меняется."],
        ["Прогноз", "Общая legacy-ось без пересчёта: тыс. м² для объёма, % для остальных KPI. Единица указана для каждой серии."],
    ])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
