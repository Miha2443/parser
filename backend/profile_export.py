"""Excel export of the exact API profile, retaining nulls and provenance."""
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_workbook(payload):
    profile = payload["profiles"][0]
    book = Workbook()
    book.remove(book.active)

    def cell(value):
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False, allow_nan=False)
        # Imported addresses/company names must remain text, never Excel formulas.
        return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value

    def sheet(name, headers, rows):
        tab = book.create_sheet(name)
        tab.append(headers)
        for row in rows:
            tab.append([cell(value) for value in row])
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for c in tab[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1C3444")
        for cells in tab.iter_rows(min_row=2):
            for c in cells:
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(c.value, (int, float)):
                    c.number_format = "#,##0.00"
        for index, title in enumerate(headers, 1):
            tab.column_dimensions[get_column_letter(index)].width = min(55, max(18, len(str(title)) + 2))
        return tab

    sheet("Профиль", ["Показатель", "Значение"], [
        ["Компания", profile["developer"]], ["Регион квартирографии", profile["region"]],
        ["Сформировано (UTC)", payload["generatedAt"]], ["Версия источников", payload["version"]],
        ["Область фильтра региона", "Только квартирография; другие разделы сохраняют собственную географию"],
    ])
    categories = [("residential", "Жилое, м²"), ("common", "МОП, м²"),
                  ("nonresidentialInHousing", "Нежилое в жилье, м²"), ("standaloneNonresidential", "Нежилое отдельное, м²")]
    sheet("Ввод по годам", ["Год", *[label for _, label in categories], "Всего, м²"],
          [[r["year"], *[r["valuesM2"].get(key) for key, _ in categories], r["totalM2"]] for r in profile["annual"]["allYears"]])
    sheet("Структура площадей", ["Показатель", "Годы", "Категория", "Площадь, м²", "Статус"],
          [[d["title"], d["years"], s["label"], s["valueM2"], d["status"]]
           for d in profile["categoryDonuts"] for s in (d.get("segments") or [{"label": None, "valueM2": None}])])
    sheet("Квартирография", ["Тип", "Количество", "Площадь, тыс. м²", "Доля, %"],
          [[r["type"], r["count"], r["areaThousandM2"], r["sharePercent"]] for r in profile["apartments"].get("rows") or []])
    sheet("Переносы", ["Показатель", "География", "Перенос, м²", "База, м²", "Доля, %", "Примечание"],
          [[r["title"], r["region"], r["valueM2"], r["baseM2"], r["percent"], r["sourceNote"]] for r in profile["delays"]["cards"]])
    sheet("Источники", ["Семейство", "Дата источника", "Файлы-кандидаты"],
          [[name, value, payload["provenance"]["sources"][name]["candidateRawFiles"]] for name, value in payload["sourceDates"].items()])
    for key, title in (("commissioned", "Реестр РВ"), ("permitted", "Реестр ОКС")):
        table = profile["objects"][key]
        sheet(title, table["columns"], [[r.get(c) for c in table["columns"]] for r in table["rows"]])
    # Complete machine-readable profile includes ratings, escrow, sales, source
    # notes and availability flags without introducing new calculation rules.
    sheet("Все показатели", ["Раздел", "Поле", "JSON / значение"],
          [[section, key, value] for section, data in profile.items()
           if section != "objects" for key, value in (data.items() if isinstance(data, dict) else [("value", data)])])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
