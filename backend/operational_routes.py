"""Operational commissioning endpoints preserving Moscow RV/history scope."""
from io import BytesIO
import json

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook

from backend.generation_service import GenerationService
from backend.profile_service import FilterUnavailable
from pipeline.operational_calculations import catalog, report


def build_report(*args, **kwargs):
    try:
        return report(*args, **kwargs)
    except LookupError as exc:
        raise FilterUnavailable(str(exc)) from exc


def create_service(context, check_interval=5):
    return GenerationService(context, ("load_vvod_static", "load_monitoring_2_0", "load_monitoring_operational_history"),
                             ("vvod_static", "monitoring_2_0", "monitoring_operational_history"),
                             catalog, build_report, check_interval)


def export_workbook(payload):
    def safe(value):
        return "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value

    book = Workbook()
    meta = book.active
    meta.title = "Report"
    meta.append(["Field", "Value"])
    for key in ("version", "generatedAt", "region", "periodLabel", "selection", "source", "sourceDetails", "notes"):
        value = payload[key]
        meta.append([key, safe(json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value)])
    for table in payload["tables"]:
        tab = book.create_sheet(table["id"])
        tab.append([column["label"] for column in table["columns"]])
        for row in table["rows"]:
            tab.append([safe(row[column["id"]]) for column in table["columns"]])
    tab = book.create_sheet("Structure")
    tab.append(["Indicator", "million m2"])
    for row in payload["treeRows"]:
        tab.append([safe(row["label"]), row["value"]])
    for tab in book:
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = 40
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def create_router(context, call, service=None):
    router = APIRouter(prefix="/api/v1/commissioning/operational")
    service = service or create_service(context)

    @router.get("/catalog")
    def options():
        return call(service.catalog, scope="Operational commissioning")

    @router.get("")
    def view(month: int | None = Query(default=None, ge=1, le=12), exclude_mkd: bool = False,
             year: int | None = None, quarter: int = Query(default=1, ge=1, le=4), cumulative: bool = False):
        return call(lambda: service.report(month=month, exclude_mkd=exclude_mkd, year=year,
                                          quarter=quarter, cumulative=cumulative), scope="Operational commissioning")

    @router.get("/export")
    def export(month: int | None = Query(default=None, ge=1, le=12), exclude_mkd: bool = False,
               year: int | None = None, quarter: int = Query(default=1, ge=1, le=4), cumulative: bool = False,
               required_version: str = Query(min_length=1, max_length=128)):
        payload = view(month, exclude_mkd, year, quarter, cumulative)
        if required_version is not None and required_version != payload["version"]:
            raise HTTPException(409, "Source changed; refresh the displayed report before exporting")
        return Response(export_workbook(payload), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="operational-commissioning.xlsx"',
                                 "X-Data-Version": payload["version"]})

    return router
