from datetime import date
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pandas as pd

from backend.linear_routes import create_router, export_dashboard
from backend.linear_service import LinearService
from backend.profile_service import SourceUnavailable
from pipeline.commissioning_forecasts import FORECAST_FIELDS
from pipeline.data_access import DataContext
from pipeline.linear_dashboard import (HISTORY_SHEET, METRICS, STATIC_DEPOT_HISTORY,
                                       STATIC_DEPOT_SOURCE, build_dashboard, number, read_history)


class LinearDashboardTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.year = date.today().year
        self.service = LinearService(DataContext(self.root), check_interval=0)
        self.monitoring = self.root / f"data/raw/realty/nashdom/monitoring_2_0_{self.year}1016.xlsx"
        self.dsti = self.root / f"data/raw/realty/linear_objects/linear_objects_{self.year}-09-28.json"

    def write_monitoring(self, horizon=2, history=True, forecast=True):
        self.monitoring.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        book.remove(book.active)
        if history:
            tab = book.create_sheet(HISTORY_SHEET)
            tab.cell(254, 4, "Ввод линейные объекты")
            tab.cell(255, 4, "Показатель")
            tab.cell(255, 5, "ед. изм")
            for col, year in enumerate([self.year - 2, self.year - 1, self.year, self.year + 1], 6):
                tab.cell(255, col, year)
            for row, (key, label, unit, source_label, code) in enumerate(METRICS, 256):
                if key == "depots":
                    continue
                tab.cell(row, 4, source_label)
                tab.cell(row, 5, unit)
                for col, value in enumerate([20, 30, 0 if key == "metro_stations" else 7, 999], 6):
                    tab.cell(row, col, value)
        if forecast:
            tab = book.create_sheet("Прогнозы")
            for offset in range(horizon + 1):
                tab.cell(2, offset + 4, self.year + offset)
            for row, (key, label) in FORECAST_FIELDS.items():
                tab.cell(row, 3, label)
                for offset in range(horizon + 1):
                    if not (key == "depots" and offset == 0):
                        tab.cell(row, offset + 4, 0 if key == "metro_stations" else 50 + offset)
        if not book.sheetnames:
            book.create_sheet("Unrelated")
        book.save(self.monitoring)
        book.close()

    def write_dsti(self, quarters=4, fact=2, plan=10):
        self.dsti.parent.mkdir(parents=True, exist_ok=True)
        rows = [{"code": f"1.{code}", "indicator": f"Indicator {code}", "unit": "км" if code == 1 else "ед",
                 "year": year, "quarter": q, "fact": fact if q < 4 else None, "plan": plan}
                for code in range(1, 5) for year in [self.year - 1, self.year, self.year + 1]
                for q in range(1, quarters + 1)]
        self.dsti.write_text(json.dumps(rows), encoding="utf-8")
        return rows

    def client(self):
        app = FastAPI()
        app.include_router(create_router(DataContext(self.root), lambda fn, *args, **kwargs: fn(*args), self.service))
        return TestClient(app)

    def test_absent_dynamic_sources_keep_only_agreed_static_depot_history(self):
        payload = self.service.dashboard()
        self.assertEqual([row["id"] for row in payload["metrics"]], [row[0] for row in METRICS])
        for row in payload["metrics"]:
            if row["id"] == "depots":
                self.assertEqual([(year["year"], year["fact"]) for year in row["years"]], list(STATIC_DEPOT_HISTORY))
            else:
                self.assertEqual(row["years"], [])
            for key in ("fact", "annualPlan", "forecast", "target", "percent", "remaining"):
                self.assertIsNone(row[key])
        self.assertTrue(all(row["file"] is None for row in payload["sources"] if row["id"] != "static_depots"))
        self.assertEqual(payload["sources"][-1], STATIC_DEPOT_SOURCE)
        json.dumps(payload, allow_nan=False)

    def test_real_named_history_discovered_axis_zero_and_no_future_facts(self):
        self.write_monitoring()
        payload = self.service.dashboard()
        self.assertEqual(payload["metrics"][0]["years"][0]["year"], self.year - 2)
        station = payload["metrics"][3]
        self.assertEqual(station["fact"], 0)
        self.assertEqual(station["forecast"], 0)
        self.assertIsNone(station["percent"])
        self.assertEqual(station["remaining"], 0)
        self.assertIsNone(payload["metrics"][5]["fact"])
        self.assertEqual(payload["metrics"][5]["years"][0]["fact"], 0)
        self.assertEqual(payload["metrics"][5]["years"][0]["factSource"], "static_depots")
        for metric in payload["metrics"]:
            for row in metric["years"]:
                if row["year"] > self.year:
                    self.assertIsNone(row["fact"])
                    self.assertIsNone(row["factKind"])
                if row["year"] == self.year and row["fact"] is not None:
                    self.assertEqual(row["factKind"], "partial")
                    self.assertIn("неполный", row["factPeriod"])

    def test_dsti_partial_not_annual_no_missing_quarter_plan(self):
        self.write_dsti(quarters=3)
        payload = self.service.dashboard()
        roads = payload["metrics"][0]
        self.assertEqual(roads["fact"], min(3, (date.today().month - 1) // 3 + 1) * 2)
        self.assertIsNone(roads["annualPlan"])
        self.assertIsNone(roads["percent"])
        self.assertIsNone(roads["years"][0]["fact"])
        self.assertTrue(all(row["year"] <= self.year for row in roads["years"]))

    def test_dsti_complete_plan_zero_fact_and_monitoring_precedence(self):
        self.write_dsti(fact=0, plan=0)
        roads = self.service.dashboard()["metrics"][0]
        self.assertEqual(roads["annualPlan"], 0)
        self.assertEqual(roads["fact"], 0)
        self.assertIsNone(roads["remaining"])
        self.assertIsNone(roads["percent"])
        self.write_monitoring()
        roads = self.service.dashboard()["metrics"][0]
        self.assertEqual(roads["fact"], 7)
        self.assertEqual(roads["factSource"], "monitoring")
        self.assertEqual(roads["annualPlan"], 0)
        self.assertEqual(roads["forecast"], 50)
        self.assertEqual(roads["targetKind"], "forecast")
        self.assertEqual(roads["target"], 50)
        self.assertAlmostEqual(roads["percent"], 14)
        self.assertEqual(roads["remaining"], 43)

    def test_only_forecast_defines_target_with_missing_zero_or_exceeded_values(self):
        self.write_dsti(fact=2, plan=100)
        roads = self.service.dashboard()["metrics"][0]
        self.assertEqual(roads["annualPlan"], 400)
        for key in ("target", "targetKind", "percent", "remaining"):
            self.assertIsNone(roads[key])
        self.write_monitoring()
        book = load_workbook(self.monitoring)
        book[HISTORY_SHEET].cell(256, 8, 70)
        book.save(self.monitoring)
        book.close()
        roads = self.service.dashboard()["metrics"][0]
        self.assertEqual(roads["target"], 50)
        self.assertEqual(roads["percent"], 140)
        self.assertEqual(roads["remaining"], 0)
        book = load_workbook(self.monitoring)
        book["Прогнозы"].cell(40, 4, 0)
        book.save(self.monitoring)
        book.close()
        roads = self.service.dashboard()["metrics"][0]
        self.assertEqual(roads["target"], 0)
        self.assertIsNone(roads["percent"])
        self.assertEqual(roads["remaining"], 0)

    def test_complete_historical_dsti_only_when_all_four_measured(self):
        rows = self.write_dsti()
        for row in rows:
            if row["year"] < self.year:
                row["fact"] = 1
        self.dsti.write_text(json.dumps(rows), encoding="utf-8")
        historic = self.service.dashboard()["metrics"][0]["years"][0]
        self.assertEqual(historic["fact"], 4)
        self.assertEqual(historic["factKind"], "annual")

    def test_forecast_horizons_change_version_and_source_removal(self):
        self.write_monitoring(horizon=1)
        before = self.service.dashboard()
        self.write_monitoring(horizon=4)
        after = self.service.dashboard()
        self.assertNotEqual(before["version"], after["version"])
        self.assertEqual(after["metrics"][0]["years"][-1]["year"], self.year + 4)
        self.monitoring.unlink()
        self.assertEqual(self.service.dashboard()["metrics"][0]["years"], [])

    def test_forecast_axis_uses_each_metric_horizon_and_retains_zero(self):
        self.write_monitoring(horizon=9)
        book = load_workbook(self.monitoring)
        forecasts = book["Прогнозы"]
        for offset in range(10):
            if offset > 4:
                forecasts.cell(40, offset + 4).value = None
            if offset > 2:
                forecasts.cell(41, offset + 4).value = None
            forecasts.cell(43, offset + 4).value = None
            forecasts.cell(42, offset + 4).value = 0 if offset == 2 else None
        book.save(self.monitoring)
        book.close()
        metrics = {metric["id"]: metric for metric in self.service.dashboard()["metrics"]}
        self.assertEqual(metrics["roads"]["years"][-1]["year"], self.year + 4)
        self.assertEqual(metrics["bridges"]["years"][-1]["year"], self.year + 2)
        self.assertEqual(metrics["metro_stations"]["years"][-1]["year"], self.year + 9)
        self.assertTrue(all(row["forecast"] == 0 for row in metrics["metro_stations"]["years"] if row["year"] >= self.year))
        self.assertEqual([row["year"] for row in metrics["crossings"]["years"]],
                         [self.year - 2, self.year - 1, self.year])
        self.assertEqual([(row["year"], row["forecast"]) for row in metrics["depots"]["years"]],
                         [(year, None) for year, _ in STATIC_DEPOT_HISTORY] + [(self.year + 2, 0)])

    def test_static_depots_exact_counts_with_zeros_and_annual_provenance(self):
        self.assertEqual(STATIC_DEPOT_HISTORY, tuple(zip(range(2011, 2026),
                         (0, 0, 0, 1, 0, 1, 0, 2, 1, 0, 0, 0, 1, 1, 2))))
        depot = self.service.dashboard()["metrics"][5]
        self.assertEqual(sum(row["fact"] for row in depot["years"]), 9)
        for row in depot["years"]:
            self.assertEqual(row["factKind"], "annual")
            self.assertEqual(row["factPeriod"], f"За {row['year']} год")
            self.assertEqual(row["factSource"], "static_depots")
            self.assertIsNone(row["forecast"])

    def test_static_depot_counts_override_only_covered_past_years(self):
        sources = [{"id": "dsti", "date": None}, {"id": "monitoring", "date": "2026-10-01"}]
        history = [{"id": "depots", "year": 2025, "value": 999},
                   {"id": "depots", "year": 2026, "value": 3},
                   {"id": "roads", "year": 2025, "value": 41}]
        report = build_dashboard(None, history, {"rows": [{"year": 2026, "depots": 5}]}, sources,
                                 date(2026, 10, 23), static_depots=STATIC_DEPOT_HISTORY)
        depots = report["metrics"][5]
        self.assertEqual(next(row for row in depots["years"] if row["year"] == 2025)["fact"], 2)
        self.assertEqual(depots["fact"], 3)
        self.assertEqual(depots["factSource"], "monitoring")
        self.assertEqual(depots["forecast"], 5)
        self.assertEqual(depots["percent"], 60)
        self.assertEqual(depots["remaining"], 2)
        self.assertEqual(report["metrics"][0]["years"][0]["fact"], 41)

    def test_static_values_change_version_without_changing_mutable_sources(self):
        before = self.service.dashboard()
        changed = (*STATIC_DEPOT_HISTORY[:-1], (2025, 3))
        with patch("backend.linear_service.STATIC_DEPOT_HISTORY", changed):
            after = self.service.dashboard()
        self.assertNotEqual(before["version"], after["version"])
        self.assertEqual(after["metrics"][5]["years"][-1]["fact"], 3)
        self.assertEqual(self.service.dashboard()["metrics"][5]["years"][-1]["fact"], 2)

    def test_static_depots_exported_with_source_and_no_current_year_fabrication(self):
        payload = self.service.dashboard()
        response = self.client().get("/api/v1/linear/dashboard/export", params={"required_version": payload["version"]})
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content), data_only=True)
        rows = [row for row in book["Yearly data"].iter_rows(min_row=2, values_only=True) if row[0] == "depots"]
        self.assertEqual([(row[3], row[4]) for row in rows], list(STATIC_DEPOT_HISTORY))
        self.assertTrue(all(row[5] == "annual" and row[7] == "static_depots" for row in rows))
        self.assertEqual(list(book["Sources"].values)[-1], tuple(STATIC_DEPOT_SOURCE.values()))
        self.assertIsNone(payload["metrics"][5]["fact"])
        book.close()

    def test_missing_sheets_do_not_fallback_to_demo(self):
        self.write_monitoring(history=False, forecast=False)
        payload = self.service.dashboard()
        self.assertTrue(all(metric["fact"] is None and metric["forecast"] is None for metric in payload["metrics"]))
        self.assertTrue(payload["notes"])

    def test_version_pinned_export_and_complete_yearly_and_forecast_sheets(self):
        self.write_monitoring(horizon=3)
        payload = self.service.dashboard()
        client = self.client()
        self.assertEqual(client.get("/api/v1/linear/dashboard").status_code, 200)
        self.assertEqual(client.get("/api/v1/linear/dashboard/export").status_code, 422)
        self.assertEqual(client.get("/api/v1/linear/dashboard/export", params={"required_version": "old"}).status_code, 409)
        response = client.get("/api/v1/linear/dashboard/export", params={"required_version": payload["version"]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["x-data-version"], payload["version"])
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book["Yearly data"].max_row, 1 + sum(len(metric["years"]) for metric in payload["metrics"]))
        self.assertEqual(book["Forecasts"].max_row, 24)
        self.assertEqual(book["Current year"].max_row, 7)
        for sheet in ("Current year", "Yearly data"):
            self.assertNotIn("annualPlan", [cell.value for cell in book[sheet][1]])
            self.assertNotIn("plan", [cell.value for cell in book[sheet][1]])
        self.write_monitoring(horizon=1)
        self.assertEqual(client.get("/api/v1/linear/dashboard/export", params={"required_version": payload["version"]}).status_code, 409)

    def test_snapshot_immutable_and_export_formula_safety(self):
        payload = self.service.dashboard()
        payload["notes"].append("=1+2")
        book = load_workbook(BytesIO(export_dashboard(payload)))
        self.assertEqual(book["Notes"].cell(book["Notes"].max_row, 1).value, "'=1+2")
        payload["metrics"].clear()
        self.assertEqual(len(self.service.dashboard()["metrics"]), 6)

    def test_source_replacement_while_loading_is_not_published(self):
        self.write_monitoring()
        original = self.service._dashboard_inputs
        calls = 0
        def inputs(access):
            nonlocal calls
            calls += 1
            if calls == 2:
                self.write_monitoring(horizon=4)
            return original(access)
        with patch.object(self.service, "_dashboard_inputs", side_effect=inputs):
            with self.assertRaises(SourceUnavailable):
                self.service.dashboard()
        self.assertIsNone(self.service._dashboard)

    def test_reader_ignores_other_blocks_and_rejects_ambiguous_or_units(self):
        book = Workbook()
        tab = book.active
        tab.title = HISTORY_SHEET
        tab.append(["Ввод жилья", 2011, 888])
        self.assertEqual(read_history(book)[0], [])
        tab.cell(10, 4, "Ввод линейные объекты")
        tab.cell(11, 4, "Показатель")
        tab.cell(11, 6, 2020)
        tab.cell(12, 4, "дороги")
        tab.cell(12, 5, "тыс. м²")
        with self.assertRaises(ValueError):
            read_history(book)
        tab.cell(12, 5, "км")
        tab.cell(30, 4, "Ввод линейные объекты")
        with self.assertRaises(ValueError):
            read_history(book)

    def test_old_partial_snapshot_does_not_become_annual_and_number_nulls(self):
        sources = [{"id": "dsti", "date": None}, {"id": "monitoring", "date": "2024-10-01"}]
        payload = build_dashboard(None, [{"id": "roads", "year": 2024, "value": 10}], {"rows": []}, sources, date(2026, 10, 21))
        self.assertIsNone(payload["metrics"][0]["years"][0]["fact"])
        for value in (None, True, float("nan"), float("inf"), "<1", -1):
            self.assertIsNone(number(value))
        self.assertEqual(number(0), 0)


if __name__ == "__main__":
    unittest.main()
