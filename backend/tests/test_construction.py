import ast
from io import BytesIO
from pathlib import Path
import unittest

import pandas as pd
import plotly.graph_objects as go
from openpyxl import load_workbook

from backend.construction_routes import create_service, export_workbook
from pipeline.construction_calculations import catalog, report
from pipeline.data_access import DataAccess, DataContext
from pipeline.report_values import records
from scripts.check_apartments_parity import Capture

ROOT = Path(__file__).resolve().parents[2]


class ConstructionCapture(Capture):
    def __init__(self, region, kind, month):
        super().__init__(region)
        self.kind, self.month = kind, month

    def container(self, *args, **kwargs):
        return self

    def radio(self, label, options, **kwargs):
        return self.region if label == "Регион" else self.kind

    def selectbox(self, label, options, **kwargs):
        assert self.month in options
        return self.month


class ConstructionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.context = DataContext(ROOT)
        access = DataAccess(cls.context)
        cls.loaders = ("load_construction_operational", "load_rasprodannost")
        cls.before = tuple(access.data_version(name) for name in cls.loaders)
        cls.data = {name: getattr(access, name)() for name in cls.loaders}
        cls.options = catalog(cls.data)

    def capture(self, region, kind, month):
        path = ROOT / "app/pages/0_Текущее_строительство.py"
        tree = ast.parse(path.read_text(encoding="utf-8"))
        tree.body = [node for node in tree.body if not isinstance(node, (ast.Import, ast.ImportFrom))]
        ui = ConstructionCapture(region, kind, month)
        env = {"st": ui, "pd": pd, "go": go,
               "COLORS": {key: "#333333" for key in ("red", "green", "blue")},
               "apply_theme": lambda: None, "page_header": lambda *args: None,
               "style_plotly": lambda *args, **kwargs: None, "chart_data_expander": lambda *args, **kwargs: None,
               **{name: lambda value=value: value for name, value in self.data.items()}}
        exec(compile(tree, str(path), "exec"), env)
        return ui, env

    def test_original_page_parity(self):
        for region in ("msk", "rf"):
            for kind, months in self.options["monthsByKind"].items():
                for month in sorted({months[0], months[-1]}):
                    with self.subTest(region=region, kind=kind, month=month):
                        payload = report(self.data, self.options, region=region, permit_kind=kind, month=month)
                        ui, env = self.capture(region, kind, month)
                        values = {row["id"]: row["value"] for row in payload["metrics"]}
                        self.assertEqual(values["living"], env["in_millions"](env["living_value"]))
                        self.assertEqual(values["total"], env["in_millions"](env["total_value"]))
                        self.assertEqual(payload["permits"]["rows"], records(env["table"]))
                        self.assertEqual(payload["sales"], records(env["selected_sales"].head(1))[0] if not env["selected_sales"].empty else None)
                        figure = ui.figures[-1]
                        for field, trace in (("period", 0), ("remainder", 1), ("totals", 2)):
                            self.assertEqual(payload["permits"]["chart"][field], records(pd.DataFrame({"value": figure.data[trace].y})))
        access = DataAccess(self.context)
        self.assertEqual(self.before, tuple(access.data_version(name) for name in self.loaders))

    def test_full_export_and_region_independent_permits(self):
        service = create_service(self.context, check_interval=0)
        msk, rf = service.report(region="msk"), service.report(region="rf")
        self.assertEqual(msk["permits"], rf["permits"])
        book = load_workbook(BytesIO(export_workbook(rf)))
        self.assertEqual(book["Permits Moscow"].max_row, len(rf["permits"]["rows"]) + 1)


if __name__ == "__main__":
    unittest.main()
