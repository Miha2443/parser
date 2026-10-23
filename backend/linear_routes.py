"""Linear-object endpoints and complete version-pinned Excel export."""
from io import BytesIO

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from backend.linear_service import LinearService


def export_dashboard(payload):
    book = Workbook()
    meta = book.active
    meta.title = "Dashboard"
    meta.append(["Field", "Value"])
    meta.append(["version", payload["version"]])
    meta.append(["currentYear", payload["currentYear"]])
    meta.append(["generatedAt", payload["generatedAt"]])
    sources = book.create_sheet("Sources")
    sources.append(["id", "file", "date", "sheet"])
    for source in payload["sources"]:
        sources.append([source[key] for key in ("id", "file", "date", "sheet")])
    notes = book.create_sheet("Notes")
    notes.append(["Note"])
    for note in payload["notes"]:
        notes.append([note])
    yearly = book.create_sheet("Yearly data")
    yearly.append(["id", "indicator", "unit", "year", "fact", "factKind", "factPeriod", "factSource", "forecast"])
    forecasts = book.create_sheet("Forecasts")
    forecasts.append(["id", "indicator", "unit", "year", "forecast"])
    kpis = book.create_sheet("Current year")
    kpis.append(["id", "indicator", "unit", "fact", "factPeriod", "forecast", "targetKind", "percent", "remaining"])
    for metric in payload["metrics"]:
        base = [metric[key] for key in ("id", "label", "unit")]
        kpis.append(base + [metric[key] for key in ("fact", "factPeriod", "forecast", "targetKind", "percent", "remaining")])
        for row in metric["years"]:
            yearly.append(base + [row[key] for key in ("year", "fact", "factKind", "factPeriod", "factSource", "forecast")])
            if row["year"] >= payload["currentYear"]:
                forecasts.append(base + [row["year"], row["forecast"]])
    for tab in book:
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for row in tab:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.lstrip().startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
        for cell in tab[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1C3444")
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = 26
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def export_workbook(payload):
    book = Workbook()
    meta = book.active
    meta.title = "Report"
    meta.append(["Field", "Value"])
    for key, value in {**payload["selection"], "version": payload["version"],
                       "sourceDate": payload["source"]["date"], "sourceFile": payload["source"]["files"][0]}.items():
        meta.append([key, value])
    headers = ["code", "indicator", "unit", "year", "quarter", "plan", "fact", "percent"]
    for name, rows in [("Selected period", payload["summary"]), ("All quarters", payload["allPeriods"])]:
        tab = book.create_sheet(name)
        tab.append(headers)
        for row in rows:
            tab.append(["'" + row[key] if isinstance(row[key], str) and row[key].lstrip().startswith(("=", "+", "-", "@"))
                        else row[key] for key in headers])
    for tab in book:
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1C3444")
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = 22 if column[0].value != "indicator" else 65
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def create_router(context, call, service=None):
    router = APIRouter(prefix="/api/v1/linear")
    service = service or LinearService(context)

    @router.get("/dashboard")
    def dashboard():
        return call(service.dashboard, scope="Linear")

    @router.get("/dashboard/export")
    def dashboard_export(required_version: str = Query(min_length=1, max_length=128)):
        payload = call(service.dashboard, scope="Linear")
        if required_version != payload["version"]:
            raise HTTPException(409, "Source changed; refresh the displayed dashboard before exporting")
        return Response(export_dashboard(payload), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="linear-dashboard.xlsx"',
                                 "X-Data-Version": payload["version"]})

    @router.get("/catalog")
    def catalog():
        return call(service.catalog, scope="Linear")

    @router.get("")
    def report(year: int | None = None, quarter: int | None = Query(default=None, ge=1, le=4),
               cumulative: bool = False, indicator: str | None = None):
        return call(service.report, year, quarter, cumulative, indicator, scope="Linear")

    @router.get("/export")
    def export(year: int | None = None, quarter: int | None = Query(default=None, ge=1, le=4),
               cumulative: bool = False, indicator: str | None = None,
               required_version: str = Query(min_length=1, max_length=128)):
        payload = call(service.report, year, quarter, cumulative, indicator, scope="Linear")
        if required_version is not None and required_version != payload["version"]:
            raise HTTPException(409, "Source changed; refresh the displayed report before exporting")
        return Response(export_workbook(payload), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="linear-objects.xlsx"',
                                 "X-Data-Version": payload["version"]})

    return router
