"""Mortgage history semantics, immutable snapshots and pinned Excel exports."""
import copy
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.mortgage_routes import create_router, export_workbook
from backend.mortgage_service import MortgageService
from backend.profile_service import FilterUnavailable, SourceUnavailable
from pipeline.data_access import DataContext, file_signature
from pipeline.mortgage_calculations import (METRICS, SOURCE_FILE, SOURCE_URL, build_report,
                                             validate_source)


def observation(period, region="msk", segment="all", volume=100.25, count=10, **metrics):
    return {"region": region, "period": period, "segment": segment, "basis": "month",
            "volumeRub": volume, "count": count, "ratePct": None, "termMonths": None,
            "averageAmountRub": None, **metrics}


def source(rows=None):
    return {"schemaVersion": 1, "sourceUrl": SOURCE_URL, "collectedAt": "2026-10-21T10:00:00Z",
            "sourceDate": None, "records": rows if rows is not None else [observation("2026-01")]}


def year_rows(year, **kwargs):
    return [observation(f"{year}-{month:02d}", **kwargs) for month in range(1, 13)]


class MortgageCalculationsTests(unittest.TestCase):
    def test_year_totals_and_comparable_partial_year_yoy(self):
        document = source(year_rows(2024, volume=100, count=10) + year_rows(2025, volume=200, count=20) +
                          [observation(f"2026-{month:02d}", volume=300, count=30) for month in range(1, 6)])
        payload = build_report(document, "msk", 12)
        old, full, latest = payload["annual"]
        self.assertEqual(old["volume"]["annual"], 1200)
        self.assertEqual(full["volume"]["annual"], 2400)
        self.assertEqual(full["volume"]["selected"], 2400)
        self.assertEqual(full["volume"]["annualYoYPct"], 100)
        self.assertEqual(latest["volume"], {"annual": None, "selected": 1500, "actualThroughMonth": 5,
                                               "annualYoYPct": None, "selectedYoYPct": 50})
        self.assertEqual(latest["count"]["selected"], 150)
        self.assertEqual(latest["count"]["selectedYoYPct"], 50)

    def test_selection_all_twelve_months_and_old_year_full_annual_base(self):
        document = source(year_rows(2025, volume=1, count=0))
        for month in range(1, 13):
            with self.subTest(month=month):
                row = build_report(document, "msk", month)["annual"][0]
                self.assertEqual(row["volume"]["annual"], 12)
                self.assertEqual(row["volume"]["selected"], month)
                self.assertEqual(row["volume"]["actualThroughMonth"], month)
                self.assertEqual(row["count"]["annual"], 0)

    def test_internal_gap_does_not_become_zero_or_valid_ytd(self):
        rows = year_rows(2025) + [observation("2026-01"), observation("2026-03")]
        latest = build_report(source(rows), "msk", 12)["annual"][-1]
        self.assertIsNone(latest["volume"]["annual"])
        self.assertIsNone(latest["volume"]["selected"])
        self.assertEqual(latest["volume"]["actualThroughMonth"], 3)
        self.assertIsNone(latest["volume"]["selectedYoYPct"])
        january = build_report(source(rows), "msk", 1)["annual"][-1]
        self.assertEqual(january["volume"]["selected"], 100.25)
        missing_february = build_report(source(rows), "msk", 2)["annual"][-1]
        self.assertIsNone(missing_february["volume"]["selected"])
        self.assertEqual(missing_february["volume"]["actualThroughMonth"], 1)

    def test_independent_metric_coverage_does_not_cap_to_other_metric(self):
        document = source([observation("2026-01", volume=100, count=10),
                           observation("2026-02", volume=200, count=None)])
        row = build_report(document, "msk", 12)["annual"][0]
        self.assertEqual(row["volume"]["selected"], 300)
        self.assertEqual(row["volume"]["actualThroughMonth"], 2)
        self.assertEqual(row["count"]["selected"], 10)
        self.assertEqual(row["count"]["actualThroughMonth"], 1)

    def test_old_partial_year_not_presented_as_complete_selected_interval(self):
        document = source([observation(f"2024-{month:02d}") for month in range(1, 11)] + year_rows(2025))
        old = build_report(document, "msk", 12)["annual"][0]
        self.assertIsNone(old["volume"]["annual"])
        self.assertIsNone(old["volume"]["selected"])
        self.assertEqual(old["volume"]["actualThroughMonth"], 10)
        self.assertIsNone(build_report(document, "msk", 12)["annual"][1]["volume"]["annualYoYPct"])

    def test_null_source_metric_is_not_derived_and_independent_segments_are_preserved(self):
        document = source([observation("2026-01", volume=100, count=5, ratePct=8, termMonths=200),
                           observation("2026-01", segment="primary", volume=90, count=4, ratePct=6),
                           observation("2026-01", segment="secondary", volume=30, count=2, ratePct=15)])
        before = copy.deepcopy(document)
        row = build_report(document, "msk", 12)["monthly"][0]
        self.assertEqual(row["all"]["volumeRub"], 100)
        self.assertEqual(row["primary"]["volumeRub"], 90)
        self.assertEqual(row["secondary"]["volumeRub"], 30)
        self.assertEqual(row["all"]["ratePct"], 8)
        self.assertEqual(row["all"]["termMonths"], 200)
        self.assertIsNone(row["all"]["averageAmountRub"])
        self.assertEqual(document, before)

    def test_year_and_month_axis_gaps_retained_without_cross_year_yoy(self):
        rows = year_rows(2023) + [observation("2025-01")]
        report = build_report(source(rows), "msk", 12)
        self.assertEqual([row["year"] for row in report["annual"]], [2023, 2024, 2025])
        self.assertTrue(all(value is None for value in report["annual"][1]["volume"].values()))
        self.assertIsNone(report["annual"][2]["volume"]["selectedYoYPct"])
        self.assertEqual(len(report["monthly"]), 25)
        missing = next(row for row in report["monthly"] if row["period"] == "2024-06")
        for segment in ("all", "primary", "secondary"):
            self.assertEqual(missing[segment], dict.fromkeys(METRICS))

    def test_no_total_inferred_from_segments_and_zero_baseline_has_no_yoy(self):
        document = source(year_rows(2025, volume=0, count=0) + [observation("2026-01", volume=0, count=0),
                          observation("2026-02", segment="primary", volume=25, count=1)])
        report = build_report(document, "msk", 12)
        self.assertIsNone(report["annual"][-1]["volume"]["selectedYoYPct"])
        self.assertEqual(report["annual"][-1]["volume"]["selected"], 0)
        self.assertIsNone(report["monthly"][-1]["all"]["volumeRub"])
        self.assertIsNone(report["monthly"][-1]["secondary"]["volumeRub"])

    def test_source_validation_accepts_zero_nullable_date_and_detaches_input(self):
        document = source([observation("2026-01", volume=0, count=0, ratePct=0, termMonths=0, averageAmountRub=0)])
        parsed = validate_source(document)
        parsed["records"][0]["count"] = 99
        self.assertEqual(document["records"][0]["count"], 0)

    def test_all_null_source_is_invalid_and_empty_observations_do_not_extend_axis(self):
        blank = observation("9999-12", volume=None, count=None)
        with self.assertRaises(ValueError):
            validate_source(source([blank]))
        document = source([observation("2026-01"), blank, observation("1000-01", volume=None, count=None)])
        report = build_report(validate_source(document), "msk", 12)
        self.assertEqual([row["year"] for row in report["annual"]], [2026])
        self.assertEqual([row["period"] for row in report["monthly"]], ["2026-01"])

    def test_calculated_overflow_becomes_null_without_losing_finite_raw_values(self):
        rows = year_rows(2025, volume=1e-308) + year_rows(2026, volume=1e308)
        document = validate_source(source(rows))
        report = build_report(document, "msk", 1)
        latest = report["annual"][-1]["volume"]
        self.assertIsNone(latest["annual"])
        self.assertEqual(latest["selected"], 1e308)
        self.assertIsNone(latest["selectedYoYPct"])
        self.assertEqual(report["monthly"][-1]["all"]["volumeRub"], 1e308)
        json.dumps(report, allow_nan=False)

    def test_integer_flows_are_not_rounded_to_float_during_annual_summing(self):
        value = 2 ** 53 + 1
        report = build_report(source(year_rows(2025, volume=value, count=value)), "msk", 12)
        for metric in ("volume", "count"):
            self.assertEqual(report["annual"][0][metric]["annual"], value * 12)
            self.assertIsInstance(report["annual"][0][metric]["annual"], int)

    def test_invalid_source_dimensions_and_numeric_values_are_rejected(self):
        invalids = [
            {"schemaVersion": 2}, {"schemaVersion": True}, {"sourceUrl": "https://example.test/"},
            {"sourceDate": "2026-02-30"}, {"sourceDate": "2026-10"}, {"collectedAt": None},
            {"collectedAt": "2026-10-21T12:00:00"}, {"notes": "bad"}, {"notes": [None]},
            {"coverage": []}, {"records": []}, {"records": {}},
        ]
        for replacement in invalids:
            with self.subTest(replacement=replacement), self.assertRaises(ValueError):
                validate_source({**source(), **replacement})
        invalid_rows = [{"region": "xx"}, {"region": []}, {"region": {}}, {"segment": "other"},
                        {"segment": []}, {"period": "2026-13"},
                        {"period": "2026-1"}, {"period": "2026-01-01"}, {"basis": "ytd"},
                        *[{metric: value} for metric in METRICS for value in
                          (-1, float("nan"), float("inf"), 10 ** 1000, True, "12")]]
        for replacement in invalid_rows:
            with self.subTest(replacement=replacement), self.assertRaises(ValueError):
                validate_source(source([observation("2026-01", **replacement)] if "period" not in replacement
                                       else [{**observation("2026-01"), **replacement}]))
        missing = observation("2026-01")
        del missing["ratePct"]
        with self.assertRaises(ValueError):
            validate_source(source([missing]))
        with self.assertRaises(ValueError):
            validate_source(source([observation("2026-01"), observation("2026-01")]))


class MortgageServiceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.context = DataContext(Path(self.tmp.name))
        self.path = self.context.root / SOURCE_FILE
        self.path.parent.mkdir(parents=True)
        self.service = MortgageService(self.context)

    def write(self, document=None):
        document = source() if document is None else document
        self.path.write_text(json.dumps(document, ensure_ascii=False), encoding="utf-8")
        return document

    def client(self):
        def call(function, *args, **kwargs):
            try:
                return function(*args)
            except SourceUnavailable as exc:
                raise HTTPException(503, "Mortgage source unavailable") from exc
            except FilterUnavailable as exc:
                raise HTTPException(404, "Mortgage filter unavailable") from exc
        app = FastAPI()
        app.include_router(create_router(self.context, call, self.service))
        return TestClient(app)

    def test_catalog_both_regions_all_months_and_null_source_date(self):
        self.write(source([observation("2026-02", region="rf"), observation("2025-12"), observation("2026-01")]))
        catalog = self.service.catalog()
        self.assertEqual(catalog["defaultRegion"], "msk")
        self.assertEqual(catalog["regions"][0]["periods"], ["2025-12", "2026-01"])
        self.assertEqual(catalog["regions"][1]["periods"], ["2026-02"])
        self.assertEqual([month["id"] for month in catalog["months"]], list(range(1, 13)))
        self.assertIsNone(catalog["source"]["date"])
        self.assertEqual(catalog["source"]["files"], [SOURCE_FILE])
        self.assertEqual(catalog["version"], self.service.report()["version"])

    def test_only_available_region_is_default_and_empty_region_is_404(self):
        self.write(source([observation("2026-01", region="rf")]))
        self.assertEqual(self.service.catalog()["defaultRegion"], "rf")
        self.assertEqual(self.service.report()["selection"]["region"], "rf")
        with self.assertRaises(FilterUnavailable):
            self.service.report("msk")
        for region in ("", "unknown", "msk"):
            self.assertEqual(self.client().get("/api/v1/mortgage", params={"region": region}).status_code, 404)

    def test_null_only_population_is_not_advertised_but_observed_zero_is_available(self):
        self.write(source([observation("2026-01", volume=None, count=None),
                           observation("2026-01", region="rf", volume=0, count=None),
                           observation("2026-02", region="rf", volume=None, count=None)]))
        catalog = self.service.catalog()
        self.assertEqual(catalog["defaultRegion"], "rf")
        self.assertEqual(catalog["regions"], [{"id": "rf", "label": "Российская Федерация", "periods": ["2026-01"]}])
        self.assertEqual(self.service.report()["annual"][0]["volume"]["selected"], 0)
        self.assertEqual(self.client().get("/api/v1/mortgage", params={"region": "msk"}).status_code, 404)
        self.write(source([observation("2026-01", volume=None, count=None)]))
        self.assertEqual(self.client().get("/api/v1/mortgage/catalog").status_code, 503)
        self.assertIsNone(self.service._generation)

    def test_missing_or_malformed_source_is_503_without_unrelated_file_fallback(self):
        self.path.with_name("other_history.json").write_text(json.dumps(source()), encoding="utf-8")
        for value in (None, "{", "[]", json.dumps(source([])), json.dumps({**source(), "sourceUrl": "wrong"})):
            with self.subTest(value=value):
                if value is not None:
                    self.path.write_text(value, encoding="utf-8")
                for endpoint in ("/api/v1/mortgage/catalog", "/api/v1/mortgage",
                                 "/api/v1/mortgage/export?required_version=old"):
                    self.assertEqual(self.client().get(endpoint).status_code, 503)

    def test_changed_source_refreshes_and_failure_never_returns_previous_generation(self):
        self.write()
        before = self.service.report()
        self.write(source([observation("2026-01", volume=200)]))
        after = self.service.report()
        self.assertNotEqual(before["version"], after["version"])
        self.assertEqual(after["annual"][0]["volume"]["selected"], 200)
        self.path.write_text("bad", encoding="utf-8")
        with self.assertRaises(SourceUnavailable):
            self.service.report()
        self.assertIsNone(self.service._generation)
        self.write()
        self.assertEqual(self.service.report()["annual"][0]["volume"]["selected"], 100.25)
        self.path.unlink()
        with self.assertRaises(SourceUnavailable):
            self.service.catalog()

    def test_content_digest_detects_changes_even_with_unchanged_stat_signature(self):
        self.write()
        signature = file_signature(self.path)
        with patch("backend.mortgage_service.file_signature", return_value=signature):
            before = self.service.catalog()
            self.write(source([observation("2026-01", volume=999)]))
            after = self.service.report()
        self.assertNotEqual(before["version"], after["version"])
        self.assertEqual(after["annual"][0]["volume"]["selected"], 999)

    def test_stat_change_without_content_change_also_changes_generation(self):
        self.write()
        signature = file_signature(self.path)
        with patch("backend.mortgage_service.file_signature", return_value=signature):
            before = self.service.catalog()
        with patch("backend.mortgage_service.file_signature", return_value=(*signature[:-1], "replacement")):
            after = self.service.catalog()
        self.assertNotEqual(before["version"], after["version"])

    def test_mid_read_or_mid_validation_replacement_is_not_published(self):
        self.write()
        self.service.catalog()
        signature = file_signature(self.path)
        for sequence in ((signature, ("different",)), (signature, signature, ("different",))):
            with self.subTest(sequence=sequence), patch("backend.mortgage_service.file_signature", side_effect=sequence):
                self.service._signature = None
                with self.assertRaises(SourceUnavailable):
                    self.service.catalog()
                self.assertIsNone(self.service._generation)

    def test_cached_generation_rechecks_publication_before_returning(self):
        self.write()
        self.service.catalog()
        signature = file_signature(self.path)
        with patch("backend.mortgage_service.file_signature", side_effect=(signature, signature, ("different",))):
            with self.assertRaises(SourceUnavailable):
                self.service.catalog()
        self.assertIsNone(self.service._generation)

    def test_payloads_and_export_source_are_detached(self):
        original = self.write()
        before_bytes = self.path.read_bytes()
        before_signature = file_signature(self.path)
        catalog = self.service.catalog()
        catalog["source"]["files"].clear()
        catalog["regions"].clear()
        payload, raw = self.service.export_snapshot()
        payload["monthly"][0]["all"]["count"] = 999
        raw["records"].clear()
        self.assertEqual(self.service.report()["monthly"][0]["all"]["count"], 10)
        self.assertEqual(self.service.export_snapshot()[1], original)
        self.assertEqual(self.path.read_bytes(), before_bytes)
        self.assertEqual(file_signature(self.path), before_signature)

    def test_excel_keeps_long_history_evidence_without_cell_truncation(self):
        original = source([observation("2026-01", sourceEvidence={"literal": "a" * 80000})])
        original["yearSources"] = {"evidence": "b" * 90000}
        self.write(original)
        payload, raw = self.service.export_snapshot()
        book = load_workbook(BytesIO(export_workbook(payload, raw)))
        pieces = {}
        for location, part, text in book["Long source fields"].iter_rows(min_row=2, values_only=True):
            pieces.setdefault(location, []).append((part, text))
        for location, expected in (("records/0/sourceEvidence", original["records"][0]["sourceEvidence"]),
                                   ("metadata/yearSources", original["yearSources"])):
            restored = "".join(text for _, text in sorted(pieces[location]))
            self.assertEqual(json.loads(restored), expected)

    def test_api_filter_validation_and_export_version_pin(self):
        self.write(source(year_rows(2025) + [observation("2026-01")]))
        client = self.client()
        catalog = client.get("/api/v1/mortgage/catalog").json()
        for month in range(1, 13):
            self.assertEqual(client.get("/api/v1/mortgage", params={"region": "msk", "month": month}).status_code, 200)
        for month in (0, 13, -1):
            self.assertEqual(client.get("/api/v1/mortgage", params={"month": month}).status_code, 404)
        self.assertEqual(client.get("/api/v1/mortgage/export").status_code, 422)
        self.assertEqual(client.get("/api/v1/mortgage/export", params={"required_version": "old"}).status_code, 409)
        response = client.get("/api/v1/mortgage/export", params={"required_version": catalog["version"], "month": 5})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["x-data-version"], catalog["version"])
        self.assertIn("mortgage.xlsx", response.headers["content-disposition"])
        self.write(source([observation("2026-01", volume=200)]))
        stale = client.get("/api/v1/mortgage/export", params={"required_version": catalog["version"]})
        self.assertEqual(stale.status_code, 409)
        self.assertNotEqual(stale.headers["x-data-version"], catalog["version"])

    def test_excel_exact_raw_values_provenance_blanks_zero_and_formula_safety(self):
        raw = self.write(source([observation("2026-01", volume=1234567.89123, count=0, ratePct=8.120123,
                                             provenance={"sheet": "source", "row": 3}, comment="=1+2"),
                                observation("2026-01", region="rf", segment="secondary", volume=None)]))
        raw["notes"] = ["=HYPERLINK(\"bad\")", " =2+2", "normal"]
        self.write(raw)
        payload, snapshot = self.service.export_snapshot()
        book = load_workbook(BytesIO(export_workbook(payload, snapshot)), data_only=False)
        self.assertEqual(book.sheetnames, ["Report", "Source observations", "Annual calculations",
                                          "Monthly selected region", "Sources and method"])
        tab = book["Source observations"]
        headers = [cell.value for cell in tab[1]]
        observed = dict(zip(headers, [cell.value for cell in tab[2]]))
        self.assertEqual(observed["volumeRub"], 1234567.89123)
        self.assertEqual(observed["count"], 0)
        self.assertEqual(observed["ratePct"], 8.120123)
        self.assertIsNone(observed["averageAmountRub"])
        self.assertEqual(json.loads(observed["provenance"]), {"sheet": "source", "row": 3})
        self.assertEqual(observed["comment"], "'=1+2")
        self.assertEqual(tab.max_row, 3)
        notes = [cell.value for row in book["Sources and method"] for cell in row]
        self.assertIn("'=HYPERLINK(\"bad\")", notes)
        self.assertIn("' =2+2", notes)
        self.assertFalse(any(cell.data_type == "f" for tab in book for row in tab for cell in row))

    def test_export_uses_one_pinned_snapshot_even_if_source_changes_during_rendering(self):
        original = self.write()
        version = self.service.catalog()["version"]
        def replaced(payload, snapshot):
            self.write(source([observation("2026-01", volume=999)]))
            self.assertEqual(snapshot, original)
            self.assertEqual(payload["version"], version)
            return export_workbook(payload, snapshot)
        with patch("backend.mortgage_routes.export_workbook", side_effect=replaced):
            response = self.client().get("/api/v1/mortgage/export", params={"required_version": version})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["x-data-version"], version)
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book["Source observations"].cell(2, 5).value, 100.25)
        self.assertEqual(self.service.report()["monthly"][0]["all"]["volumeRub"], 999)
        self.assertEqual(self.client().get("/api/v1/mortgage/export", params={"required_version": version}).status_code, 409)

    def test_router_registered_with_supplied_pinned_root(self):
        self.write()
        from backend.app import create_app
        class ProfileStub:
            context = self.context
        client = TestClient(create_app(service=ProfileStub()))
        response = client.get("/api/v1/mortgage/catalog")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertEqual(response.json()["source"]["files"], [SOURCE_FILE])


if __name__ == "__main__":
    unittest.main()
