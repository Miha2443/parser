from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook

from pipeline.commissioning_forecasts import FORECAST_FIELDS, read_commissioning_forecasts


class CommissioningForecastTests(unittest.TestCase):
    def book(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Прогнозы"
        for row, (_, label) in FORECAST_FIELDS.items():
            sheet.cell(row, 3, label)
        sheet["D2"], sheet["E2"], sheet["F2"] = 2026, 2027, 2028
        self.addCleanup(book.close)
        return book, sheet

    def test_literal_units_horizons_zero_and_missing_are_preserved(self):
        book, sheet = self.book()
        sheet["D13"], sheet["E13"], sheet["D42"], sheet["D38"] = 7500, "7 600,5", 0, 3
        result = read_commissioning_forecasts(book, "monitoring.xlsx", "2026-10-16")
        self.assertEqual([row["year"] for row in result["rows"]], [2026, 2027])
        self.assertEqual(result["rows"][0]["housing_living"], 7500)
        self.assertEqual(result["rows"][1]["housing_living"], 7600.5)
        self.assertEqual(result["rows"][0]["depots"], 0)
        self.assertIsNone(result["rows"][1]["depots"])
        self.assertIsNone(result["rows"][0]["housing_total"])
        self.assertEqual(result["source_date"], "2026-10-16")

    def test_bad_label_does_not_shift_to_another_metric_and_duplicate_years_skip(self):
        book, sheet = self.book()
        sheet["C13"], sheet["D13"], sheet["D14"], sheet["E14"] = "РС жилье", 9999, 10000, 11000
        result = read_commissioning_forecasts(book)
        self.assertIsNone(result["rows"][0]["housing_living"])
        self.assertEqual(result["rows"][0]["housing_total"], 10000)
        sheet["E2"] = 2026
        self.assertEqual(read_commissioning_forecasts(book)["rows"], [])

    def test_formula_without_cache_is_not_computed_or_read_as_fact(self):
        book, sheet = self.book()
        sheet["D13"], sheet["D38"] = "=7500", 0
        stream = BytesIO()
        book.save(stream)
        stream.seek(0)
        cached = load_workbook(stream, data_only=True, read_only=True)
        self.addCleanup(cached.close)
        result = read_commissioning_forecasts(cached)
        self.assertIsNone(result["rows"][0]["housing_living"])
        self.assertEqual(result["rows"][0]["metro_stations"], 0)

    def test_missing_sheet_and_short_sheet_do_not_fabricate_forecasts(self):
        book = Workbook()
        self.addCleanup(book.close)
        self.assertEqual(read_commissioning_forecasts(book)["rows"], [])
        book.active.title = "Прогнозы"
        self.assertEqual(read_commissioning_forecasts(book)["rows"], [])

    def test_future_columns_follow_explicit_year_headers_not_old_fixed_horizon(self):
        book, sheet = self.book()
        sheet["N2"], sheet["N40"], sheet["O40"] = 2036, 100, 999999
        result = read_commissioning_forecasts(book)
        self.assertEqual([row["year"] for row in result["rows"]], [2036])
        self.assertEqual(result["rows"][0]["roads"], 100)
        self.assertIsNone(result["rows"][0]["metro_stations"])


if __name__ == "__main__":
    unittest.main()
