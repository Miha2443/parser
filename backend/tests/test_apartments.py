"""Independent legacy DataFrame parity, real inputs, HTTP and read-only generations."""
import copy
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import tempfile
import unittest
from unittest.mock import MagicMock, patch

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import create_app
from backend.apartments_export import export_workbook
from backend.apartments_service import ApartmentsService
from pipeline.apartment_calculations import developer_detail, overview
from pipeline.data_access import DataAccess, DataContext

ROOT = Path(__file__).resolve().parents[2]


def finite(value):
    return None if value is None or pd.isna(value) or not math.isfinite(float(value)) else float(value)


def legacy_frames(data, region):
    # Literal DataFrame operations from pages 4/5, independently of new helpers.
    devs = data["developers"][data["developers"]["region_key"] == region].copy()
    devs = devs.sort_values("площадь_тыс_м²_num", ascending=False).reset_index(drop=True)
    devs["place"] = devs.index + 1
    regs = data["regions"]
    df = regs[regs["region_key"] == region].copy()
    moscow_mask = regs["наименование"].astype(str).str.strip().str.casefold().isin(["город москва", "г. москва", "москва"])
    if region == "msk":
        moscow = regs[regs["region_key"].eq("rf") & moscow_mask]
        df = pd.concat([moscow, df], ignore_index=True).drop_duplicates("наименование", keep="first")
    df = df.sort_values("площадь_тыс_м²_num", ascending=False)
    first = df["наименование"].astype(str).str.strip().str.casefold().isin(["город москва", "г. москва", "москва"])
    return devs, pd.concat([df[first], df[~first]])


def legacy_row(row):
    out = {"id": row["наименование"], "name": row["наименование"],
           "apartmentThousandCount": finite(row["квартиры_тыс_шт_num"]),
           "areaThousandM2": finite(row["площадь_тыс_м²_num"]),
           "rooms": [{"type": label, "sharePercent": finite(row[column])} for label, column in [
               ("1 комн", "доля_1комн_%_num"), ("2 комн", "доля_2комн_%_num"),
               ("3 комн", "доля_3комн_%_num"), ("4+ комн", "доля_4+комн_%_num")]]}
    if "place" in row:
        out["place"] = int(row["place"])
    return out


def legacy_detail(data, region, developer):
    devs, _ = legacy_frames(data, region)
    row = devs[devs["наименование"] == developer].iloc[0]
    qty = row["квартиры_тыс_шт_num"]
    area = row["площадь_тыс_м²_num"]
    avg_area = None
    references = []
    if qty and area and qty > 0:
        avg_area = (area * 1000) / (qty * 1000)
        reference_rows = data["apartments"]
        for region_key in ["msk", "rf"]:
            reference = reference_rows[reference_rows["region_key"].eq(region_key) & reference_rows["тип"].eq("Все квартиры")]
            if not reference.empty:
                item = reference.iloc[0]
                reference_qty = pd.to_numeric(item["количество_шт_num"], errors="coerce")
                reference_area = pd.to_numeric(item["площадь_тыс_м²_num"], errors="coerce")
                if pd.notna(reference_qty) and reference_qty > 0 and pd.notna(reference_area):
                    references.append({"region": region_key, "averageAreaM2": finite(reference_area * 1000 / reference_qty)})
    top10 = devs.head(10).copy()
    if developer not in top10["наименование"].values:
        extra = devs[devs["наименование"] == developer]
        top10 = pd.concat([top10, extra], ignore_index=True)
    base = devs["площадь_тыс_м²_num"].sum()
    share = row["площадь_тыс_м²_num"] / base * 100 if base != 0 else None
    return {"region": region, "developer": legacy_row(row),
            "summary": {"countThousand": finite(qty), "areaThousandM2": finite(area),
                        "averageAreaM2": finite(avg_area), "marketSharePercent": finite(share),
                        "marketBaseAreaThousandM2": finite(base), "place": int(row["place"]),
                        "totalDevelopers": len(devs)},
            "referenceAverages": references, "rooms": legacy_row(row)["rooms"],
            "comparison": [legacy_row(item) for _, item in top10.iterrows()]}


def sample_data(count=65):
    apartments, developers, regions = [], [], []
    for region in ("rf", "msk"):
        apartments.extend([
            {"region_key": region, "тип": "Все квартиры", "количество_шт_num": 200000.0, "площадь_тыс_м²_num": 10000.0},
            {"region_key": region, "тип": "1-комн", "количество_шт_num": 100000.0, "площадь_тыс_м²_num": 3000.0},
        ])
        for i in range(count):
            developers.append({"region_key": region, "наименование": f" Dev {i:02d} ",
                               "квартиры_тыс_шт_num": 2.0, "площадь_тыс_м²_num": float(count - i),
                               "доля_1комн_%_num": 20.0, "доля_2комн_%_num": 30.0,
                               "доля_3комн_%_num": 40.0, "доля_4+комн_%_num": 10.0})
        for name, area in [("Город Москва", 10.0 if region == "rf" else 999.0),
                           ("Б область", 500.0), ("А область", 600.0)]:
            regions.append({**developers[-1], "region_key": region, "наименование": name, "площадь_тыс_м²_num": area})
    return {"apartments": pd.DataFrame(apartments), "developers": pd.DataFrame(developers),
            "regions": pd.DataFrame(regions),
            "distribution": pd.DataFrame([{"region_key": r, "диапазон": "< 30", "доля_num": 0.5} for r in ("rf", "msk")]),
            "report_date": "01.01.2026", "regions_available": ["rf", "msk"]}


def fake_access(data, issues=None):
    access = MagicMock()
    access.load_kvartirografia.return_value = data
    access.data_version.return_value = ("generation-one",)
    access.latest_realty_mart_source_date.return_value = "01.01.2026"
    access.latest_raw_source_date.return_value = ""
    access.source_files.return_value = [ROOT / "data/raw/realty/nashdom/kvartirografia_candidate.json"]
    access.issues = issues or []
    return access


def client_for(service):
    # Apartment tests deliberately do not depend on profile/monitoring inputs.
    return TestClient(create_app(service=MagicMock(), apartments_service=service))


class CalculationTests(unittest.TestCase):
    def setUp(self):
        self.data = sample_data()

    def test_overview_full_rows_units_and_moscow_first_both_regions(self):
        for region in ("msk", "rf"):
            with self.subTest(region=region):
                original = copy.deepcopy(self.data)
                result = overview(self.data, region)
                devs, regs = legacy_frames(self.data, region)
                self.assertEqual(result["developers"], [legacy_row(row) for _, row in devs.iterrows()])
                self.assertEqual(result["regions"], [legacy_row(row) for _, row in regs.iterrows()])
                self.assertEqual(result["developerCount"], 65)
                self.assertEqual(result["regionCount"], 3)
                self.assertEqual(result["regions"][0]["name"], "Город Москва")
                self.assertEqual(result["regions"][0]["areaThousandM2"], 10.0)
                self.assertEqual(result["apartments"][0]["count"], 200000.0)
                self.assertEqual(result["apartments"][0]["areaThousandM2"], 10000.0)
                self.assertEqual(result["distribution"], [{"range": "< 30", "sharePercent": 0.5}])
                for key in ("apartments", "developers", "regions", "distribution"):
                    pd.testing.assert_frame_equal(original[key], self.data[key])

    def test_developer_independent_legacy_all_ranks_and_market_denominator(self):
        for region in ("msk", "rf"):
            for name in [" Dev 00 ", " Dev 09 ", " Dev 13 ", " Dev 64 "]:
                with self.subTest(region=region, name=name):
                    actual = developer_detail(self.data, region, name)
                    self.assertEqual(actual, legacy_detail(self.data, region, name))
                    summary = actual["summary"]
                    self.assertEqual(summary["marketBaseAreaThousandM2"], 2145.0)
                    self.assertNotEqual(summary["marketBaseAreaThousandM2"], 10000.0)
                    self.assertEqual(summary["averageAreaM2"], summary["areaThousandM2"] / 2)
                    self.assertEqual(actual["referenceAverages"], [
                        {"region": "msk", "averageAreaM2": 50.0}, {"region": "rf", "averageAreaM2": 50.0}])
                    self.assertEqual(len(actual["comparison"]), 10 if summary["place"] <= 10 else 11)
                    if summary["place"] > 10:
                        self.assertEqual(actual["comparison"][-1]["id"], name)

    def test_zero_missing_negative_and_reference_condition(self):
        for qty, area in [(0.0, 20.0), (2.0, 0.0), (float("nan"), 20.0),
                          (2.0, float("nan")), (2.0, -20.0)]:
            with self.subTest(qty=qty, area=area):
                data = sample_data(1)
                mask = data["developers"]["region_key"].eq("msk")
                data["developers"].loc[mask, "квартиры_тыс_шт_num"] = qty
                data["developers"].loc[mask, "площадь_тыс_м²_num"] = area
                data["developers"].loc[mask, "доля_2комн_%_num"] = float("nan")
                data["developers"].loc[mask, "доля_4+комн_%_num"] = 0.0
                result = developer_detail(data, "msk", " Dev 00 ")
                self.assertEqual(result, legacy_detail(data, "msk", " Dev 00 "))
                self.assertIsNone(result["rooms"][1]["sharePercent"])
                self.assertEqual(result["rooms"][3]["sharePercent"], 0.0)
                json.dumps(result, allow_nan=False)
        data["apartments"].loc[:, "количество_шт_num"] = 0.0
        self.assertEqual(developer_detail(data, "msk", " Dev 00 ")["referenceAverages"], [])

    def test_ties_duplicates_and_moscow_aliases_follow_exact_page_sort(self):
        data = self.data
        data["developers"].loc[:11, "площадь_тыс_м²_num"] = 100.0
        data["developers"].loc[64, "площадь_тыс_м²_num"] = float("nan")
        extra = data["regions"].iloc[0].copy()
        extra["наименование"] = " г. МОСКВА "
        extra["площадь_тыс_м²_num"] = 900.0
        data["regions"] = pd.concat([data["regions"], pd.DataFrame([extra])], ignore_index=True)
        for region in ("rf", "msk"):
            devs, regs = legacy_frames(data, region)
            actual = overview(data, region)
            self.assertEqual(actual["developers"], [legacy_row(r) for _, r in devs.iterrows()])
            self.assertEqual(actual["regions"], [legacy_row(r) for _, r in regs.iterrows()])


class ServiceAndHttpTests(unittest.TestCase):
    def setUp(self):
        self.data = sample_data()
        self.access = fake_access(self.data, [("warning", "test source warning")])
        self.service = ApartmentsService(DataContext(ROOT))
        self.factory = patch("backend.apartments_service.DataAccess", return_value=self.access)
        self.factory.start()
        self.addCleanup(self.factory.stop)
        self.log = patch("backend.app.logger")
        self.log.start()
        self.addCleanup(self.log.stop)
        self.client = client_for(self.service)

    def test_catalog_contract_defaults_raw_ids_and_shared_generation(self):
        response = self.client.get("/api/v1/apartments/catalog")
        self.assertEqual(response.status_code, 200)
        catalog = response.json()
        self.assertEqual(set(catalog), {"schemaVersion", "version", "generatedAt", "reportDate", "source", "regions", "developersByRegion"})
        self.assertEqual(catalog["schemaVersion"], 1)
        self.assertEqual(len(catalog["version"]), 64)
        self.assertEqual([r["id"] for r in catalog["regions"]], ["msk", "rf"])
        self.assertEqual(catalog["developersByRegion"]["msk"][0], {"id": " Dev 00 ", "name": " Dev 00 ", "place": 1})
        self.assertEqual(catalog["source"]["issues"], [["warning", "test source warning"]])
        self.assertEqual(catalog["source"]["files"], ["data/raw/realty/nashdom/kvartirografia_candidate.json"])
        for path, params in [("/api/v1/apartments", {}),
                             ("/api/v1/apartments/developer", {"developer": " Dev 13 "})]:
            result = self.client.get(path, params=params)
            self.assertEqual(result.status_code, 200)
            self.assertEqual(result.json()["region"], "msk")
            for key in ("version", "generatedAt", "source", "reportDate"):
                self.assertEqual(result.json()[key], catalog[key])
            self.assertEqual(result.headers["cache-control"], "no-store")
        self.access.load_kvartirografia.assert_called_once_with()

    def test_filters_and_exact_name_not_normalized(self):
        for path, params, status in [
            ("/api/v1/apartments", {"region": "bad"}, 422),
            ("/api/v1/apartments/developer", {}, 422),
            ("/api/v1/apartments/developer", {"developer": ""}, 422),
            ("/api/v1/apartments/developer", {"developer": "dev 00"}, 404),
            ("/api/v1/apartments/developer", {"developer": " Dev 00 "}, 200),
            ("/api/v1/apartments/export", {"developer": "missing"}, 404),
            ("/api/v1/apartments/export", {"required_version": ""}, 422),
        ]:
            with self.subTest(path=path, params=params):
                self.assertEqual(self.client.get(path, params=params).status_code, status)
        self.data["regions_available"] = ["rf"]
        self.access.data_version.return_value = ("rf-only",)
        self.service.check_interval = 0
        self.assertEqual(self.client.get("/api/v1/apartments").json()["region"], "rf")
        self.assertEqual(self.client.get("/api/v1/apartments", params={"region": "msk"}).status_code, 404)
        self.assertEqual(set(self.service.catalog()["developersByRegion"]), {"rf"})

    def test_cache_freshness_five_seconds_and_refresh_invalidation(self):
        with patch("backend.apartments_service.time.monotonic", return_value=100.0):
            first = self.service.developer(" Dev 13 ", "msk")
        self.data["developers"].loc[self.data["developers"]["наименование"].eq(" Dev 13 "), "площадь_тыс_м²_num"] = 500.0
        self.access.data_version.return_value = ("generation-two",)
        with patch("backend.apartments_service.time.monotonic", return_value=104.99):
            self.assertEqual(self.service.developer(" Dev 13 ", "msk"), first)
        self.access.load_kvartirografia.assert_called_once()
        with patch("backend.apartments_service.time.monotonic", return_value=105.0):
            updated = self.service.developer(" Dev 13 ", "msk")
        self.assertNotEqual(updated["version"], first["version"])
        self.assertNotEqual(updated["summary"], first["summary"])
        self.assertEqual(updated["summary"]["place"], 1)
        self.assertEqual(self.access.load_kvartirografia.call_count, 2)
        with patch("backend.apartments_service.time.monotonic", return_value=110.0):
            self.assertEqual(self.service.developer(" Dev 13 ", "msk"), updated)
        self.assertEqual(self.access.load_kvartirografia.call_count, 2)

    def test_failed_refresh_never_serves_previous_generation_and_recovers(self):
        self.service.check_interval = 0
        first = self.service.overview("msk")
        self.access.data_version.return_value = ("broken-new-generation",)
        self.access.issues = [("error", "fatal load failure")]
        for path in ("/api/v1/apartments/catalog", "/api/v1/apartments", "/api/v1/apartments/export"):
            self.assertEqual(self.client.get(path).status_code, 503)
        self.assertEqual(self.service._metadata["version"], first["version"])
        self.access.issues = []
        self.assertNotEqual(self.service.overview("msk")["version"], first["version"])

    def test_inconsistent_load_cannot_publish(self):
        self.access.data_version.side_effect = [("before",), ("after",)]
        self.assertEqual(self.client.get("/api/v1/apartments/catalog").status_code, 503)
        self.assertIsNone(self.service._metadata)

    def test_fatal_load_exceptions_and_schema_failure_are_503_not_404(self):
        for exception in (OSError("input"), ValueError("input"), RuntimeError("input"), KeyError("input"), TypeError("input")):
            with self.subTest(exception=exception):
                self.access.load_kvartirografia.side_effect = exception
                self.assertEqual(self.client.get("/api/v1/apartments/catalog").status_code, 503)
        self.access.load_kvartirografia.side_effect = None
        for bad_data in [{}, {**self.data, "developers": self.data["developers"].drop(columns=["площадь_тыс_м²_num"])},
                         {**self.data, "apartments": pd.DataFrame(), "developers": pd.DataFrame()},
                         {**self.data, "regions_available": []}]:
            with self.subTest(keys=list(bad_data)):
                self.access.load_kvartirografia.return_value = bad_data
                self.assertEqual(self.client.get("/api/v1/apartments/catalog").status_code, 503)
                self.assertIsNone(self.service._metadata)

    def test_empty_optional_tables_and_absent_developers_for_region(self):
        self.data["distribution"] = pd.DataFrame()
        self.data["regions"] = pd.DataFrame()
        self.data["developers"] = self.data["developers"][self.data["developers"]["region_key"].eq("rf")]
        result = self.client.get("/api/v1/apartments", params={"region": "msk"}).json()
        self.assertEqual(result["distribution"], [])
        self.assertEqual(result["regions"], [])
        self.assertEqual(result["developers"], [])
        self.assertEqual(result["developerCount"], 0)
        self.assertEqual(self.client.get("/api/v1/apartments/developer", params={"region": "msk", "developer": " Dev 00 "}).status_code, 404)

    def test_partial_sources_only_require_endpoint_primary_table(self):
        self.service.check_interval = 0
        self.data["developers"] = pd.DataFrame()
        self.access.data_version.return_value = ("apartments-only",)
        self.assertEqual(self.client.get("/api/v1/apartments/catalog").status_code, 200)
        self.assertEqual(self.service.catalog()["developersByRegion"], {"msk": [], "rf": []})
        result = self.client.get("/api/v1/apartments")
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()["developers"], [])
        self.assertEqual(self.client.get("/api/v1/apartments/export").status_code, 200)
        self.assertEqual(self.client.get("/api/v1/apartments/developer", params={"developer": "unknown"}).status_code, 503)
        self.access.load_kvartirografia.return_value = {**sample_data(), "apartments": pd.DataFrame()}
        self.access.data_version.return_value = ("developers-only",)
        self.assertEqual(self.client.get("/api/v1/apartments/catalog").status_code, 200)
        for path in ("/api/v1/apartments/developer", "/api/v1/apartments/export"):
            result = self.client.get(path, params={"developer": " Dev 13 "})
            self.assertEqual(result.status_code, 200)
            if path.endswith("developer"):
                self.assertEqual(result.json()["referenceAverages"], [])
                self.assertGreater(result.json()["summary"]["averageAreaM2"], 0)
        self.assertEqual(self.client.get("/api/v1/apartments").status_code, 503)
        self.assertEqual(self.client.get("/api/v1/apartments/export").status_code, 503)

    def test_exact_duplicate_ids_dedup_catalog_only_first_detail_all_extra_matches(self):
        duplicate = self.data["developers"].iloc[13].copy()
        duplicate["площадь_тыс_м²_num"] = 1.5
        self.data["developers"] = pd.concat([self.data["developers"], pd.DataFrame([duplicate])], ignore_index=True)
        overview_payload = self.service.overview("rf")
        self.assertEqual(overview_payload["developerCount"], 66)
        self.assertEqual(sum(row["id"] == " Dev 13 " for row in overview_payload["developers"]), 2)
        catalog = self.service.catalog()["developersByRegion"]["rf"]
        self.assertEqual(len(catalog), 65)
        self.assertEqual([row["place"] for row in catalog if row["id"] == " Dev 13 "], [14])
        self.assertEqual([row["place"] for row in catalog], sorted(row["place"] for row in catalog))
        actual = self.service.developer(" Dev 13 ", "rf")
        expected = legacy_detail(self.data, "rf", " Dev 13 ")
        self.assertEqual({key: actual[key] for key in expected}, expected)
        self.assertEqual(actual["developer"]["place"], 14)
        self.assertEqual(len(actual["comparison"]), 12)
        self.assertEqual(sum(row["id"] == " Dev 13 " for row in actual["comparison"]), 2)

    def test_response_mutation_does_not_change_cached_generation(self):
        result = self.service.overview()
        result["developers"][0]["name"] = "tampered"
        result["source"]["issues"].append(["error", "tampered"])
        self.assertEqual(self.service.overview()["developers"][0]["name"], " Dev 00 ")
        self.assertEqual(len(self.service.catalog()["source"]["issues"]), 1)
        result = self.service.developer(" Dev 00 ")
        result["rooms"].clear()
        self.assertEqual(len(self.service.developer(" Dev 00 ")["rooms"]), 4)

    def test_excel_full_rows_and_developer_comparison_match_payload(self):
        for region in ("msk", "rf"):
            for developer in (None, " Dev 13 "):
                with self.subTest(region=region, developer=developer):
                    params = {"region": region}
                    payload = self.service.overview(region)
                    if developer is not None:
                        params["developer"] = developer
                        payload = self.service.developer(developer, region)
                    params["required_version"] = payload["version"]
                    response = self.client.get("/api/v1/apartments/export", params=params)
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(response.headers["x-data-version"], payload["version"])
                    self.assertIn("spreadsheetml.sheet", response.headers["content-type"])
                    book = load_workbook(BytesIO(response.content), data_only=False)
                    keys = [("developers", "Девелоперы"), ("regions", "Регионы")] if developer is None else [("comparison", "Сравнение")]
                    for key, sheet in keys:
                        expected = [(r["name"], r["apartmentThousandCount"], r["areaThousandM2"], r.get("place"),
                                     *[room["sharePercent"] for room in r["rooms"]]) for r in payload[key]]
                        self.assertEqual(list(book[sheet].values)[1:], expected)
                        self.assertEqual(book[sheet].freeze_panes, "A2")
                    self.assertEqual(book["Отчёт"]["B5"].value, payload["version"])
                    if developer is None:
                        self.assertEqual(book["Девелоперы"].max_row - 1, 65)
                        self.assertEqual(book["Типы квартир"]["B2"].value, 200000)
                    else:
                        self.assertEqual(book["Сравнение"].max_row - 1, 11)
                        self.assertEqual(book["Показатели"]["B6"].value, payload["summary"]["marketBaseAreaThousandM2"])

    def test_excel_version_conflict_header_and_no_export(self):
        for developer in (None, " Dev 00 "):
            params = {"region": "rf", "required_version": "old"}
            if developer:
                params["developer"] = developer
            with patch("backend.apartments_export.export_workbook") as exporter:
                response = self.client.get("/api/v1/apartments/export", params=params)
                self.assertEqual(response.status_code, 409)
                self.assertEqual(response.headers["x-data-version"], self.service.catalog()["version"])
                exporter.assert_not_called()

    def test_excel_injection_and_nulls_all_imported_text_columns(self):
        payload = self.service.overview()
        for i, prefix in enumerate(["=", "+", "-", "@", " \t="]):
            payload["developers"][i]["name"] = prefix + 'HYPERLINK("evil")'
            payload["developers"][i]["rooms"][0]["sharePercent"] = None
        payload["apartments"][0]["type"] = "=evil"
        payload["distribution"][0]["range"] = "@evil"
        payload["regions"][0]["name"] = "+evil"
        payload["reportDate"] = "-evil"
        book = load_workbook(BytesIO(export_workbook(payload)), data_only=False)
        for i in range(2, 7):
            self.assertEqual(book["Девелоперы"].cell(i, 1).data_type, "s")
            self.assertTrue(book["Девелоперы"].cell(i, 1).value.startswith("'"))
            self.assertIsNone(book["Девелоперы"].cell(i, 5).value)
        for tab in book:
            for row in tab:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")


class RealSourceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.context = DataContext(ROOT)
        cls.access = DataAccess(cls.context)
        cls.signature = cls.access.data_version("load_kvartirografia")
        cls.hashes = {item[0]: hashlib.sha256(Path(item[0]).read_bytes()).hexdigest()
                      for item in cls.signature if Path(item[0]).is_file()}
        cls.data = cls.access.load_kvartirografia()
        if cls.data["apartments"].empty or cls.data["developers"].empty:
            raise AssertionError("Existing apartment sources are required for this integration check")
        cls.service = ApartmentsService(cls.context, check_interval=0)
        cls.client = client_for(cls.service)

    @classmethod
    def tearDownClass(cls):
        if cls.access.data_version("load_kvartirografia") != cls.signature:
            raise AssertionError("Apartment input signatures changed during tests")
        for path, expected in cls.hashes.items():
            if hashlib.sha256(Path(path).read_bytes()).hexdigest() != expected:
                raise AssertionError(f"Input modified during tests: {path}")

    def test_real_catalog_and_complete_overviews_legacy_dataframe_parity(self):
        catalog = self.client.get("/api/v1/apartments/catalog").json()
        self.assertEqual(catalog["reportDate"], self.data["report_date"])
        self.assertEqual(catalog["source"]["issues"], [list(i) for i in self.access.issues])
        for region in ("msk", "rf"):
            with self.subTest(region=region):
                response = self.client.get("/api/v1/apartments", params={"region": region})
                self.assertEqual(response.status_code, 200)
                actual = response.json()
                devs, regs = legacy_frames(self.data, region)
                self.assertEqual(actual["developers"], [legacy_row(row) for _, row in devs.iterrows()])
                self.assertEqual(actual["regions"], [legacy_row(row) for _, row in regs.iterrows()])
                apartments = self.data["apartments"][self.data["apartments"]["region_key"] == region]
                self.assertEqual(actual["apartments"], [
                    {"type": r["тип"], "count": finite(r["количество_шт_num"]), "areaThousandM2": finite(r["площадь_тыс_м²_num"])}
                    for _, r in apartments.iterrows()])
                distribution = self.data["distribution"][self.data["distribution"]["region_key"] == region]
                self.assertEqual(actual["distribution"], [
                    {"range": r["диапазон"], "sharePercent": finite(r["доля_num"])} for _, r in distribution.iterrows()])
                self.assertEqual(actual["developerCount"], len(devs))
                self.assertEqual(actual["regionCount"], len(regs))
                self.assertGreater(len(devs), 50)
                self.assertEqual(catalog["developersByRegion"][region], [
                    {"id": r["наименование"], "name": r["наименование"], "place": int(r["place"])} for _, r in devs.iterrows()])

    def test_real_developer_kpis_rooms_and_top10_selected_against_legacy(self):
        for region in ("msk", "rf"):
            devs, _ = legacy_frames(self.data, region)
            for index in (0, 9, 13, len(devs) - 1):
                developer = devs.iloc[index]["наименование"]
                with self.subTest(region=region, developer=developer):
                    response = self.client.get("/api/v1/apartments/developer", params={"region": region, "developer": developer})
                    self.assertEqual(response.status_code, 200)
                    actual = response.json()
                    expected = legacy_detail(self.data, region, developer)
                    self.assertEqual({k: actual[k] for k in expected}, expected)

    def test_existing_raw_mode_matches_legacy_loader_without_marts_writes(self):
        context = DataContext(ROOT, use_marts=False)
        data = DataAccess(context).load_kvartirografia()
        service = ApartmentsService(context)
        for region in ("msk", "rf"):
            devs, regs = legacy_frames(data, region)
            actual = service.overview(region)
            self.assertEqual(actual["developers"], [legacy_row(r) for _, r in devs.iterrows()])
            self.assertEqual(actual["regions"], [legacy_row(r) for _, r in regs.iterrows()])
            self.assertEqual(actual["reportDate"], data["report_date"])

    def test_missing_actual_sources_and_required_mart_errors_are_503(self):
        with tempfile.TemporaryDirectory() as directory, patch("backend.app.logger"):
            for context in (DataContext(Path(directory)), DataContext(Path(directory), require_marts=True)):
                client = client_for(ApartmentsService(context))
                for path in ("/api/v1/apartments/catalog", "/api/v1/apartments", "/api/v1/apartments/export"):
                    self.assertEqual(client.get(path).status_code, 503)


if __name__ == "__main__":
    unittest.main()
