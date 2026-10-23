"""Approved permit controls, complete month selection and honest compositions."""
from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook

from backend.tests import test_construction_pages as fixtures
from pipeline.permit_page_calculations import (FORECAST_SHEET, HISTORICAL_FIELDS, HISTORY_SHEET,
                                               HOUSING_SHEET, OKS_SHEET, _osm_group)


class CanonicalPermitTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.ConstructionPagesTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.service = self.fixture.permits
        self.fixture.partial_current_book()

    @staticmethod
    def main(payload, year):
        return next(row for row in payload["permits"]["presentation"]["rows"] if row["year"] == year)

    @staticmethod
    def block(payload, identifier, year):
        block = next(block for block in payload["breakdowns"] if block["id"] == identifier)
        return next(row for row in block["rows"] if row["year"] == year)

    def test_december_available_with_current_october_and_approved_full_previous_years(self):
        payload = self.service.report("housing", 12, structure_mode="annual")
        self.assertEqual(payload["controls"]["monthsByKind"]["housing"], list(range(1, 13)))
        self.assertEqual(payload["controls"]["structure"]["months"], list(range(1, 13)))
        self.assertEqual(payload["selection"]["month"], 12)
        old = self.main(payload, 2011)
        self.assertEqual(old["selected"], 9999.123456789)
        self.assertEqual(old["actualMonth"], 12)
        self.assertTrue(old["periodComplete"])
        current = self.main(payload, 2026)
        self.assertAlmostEqual(current["selected"], sum((60000 + month) / 1000 for month in range(1, 11)))
        self.assertEqual(current["actualMonth"], 10)
        self.assertFalse(current["periodComplete"])
        self.assertEqual(current["observedThrough"], "2026-10-16")
        self.assertEqual(self.block(payload, "housing_components", 2026)["actualMonth"], 10)

    def test_latest_measurement_not_snapshot_month_bounds_selected_year(self):
        self.fixture.edit_book(OKS_SHEET, {"B44": None, "B45": None})
        payload = self.service.report("total", 12, structure_mode="annual")
        current = self.main(payload, 2026)
        self.assertEqual(current["actualMonth"], 9)
        self.assertAlmostEqual(current["selected"], sum((140000 + 2 * month) / 1000 for month in range(1, 10)))
        for block in payload["breakdowns"]:
            row = next(row for row in block["rows"] if row["year"] == 2026)
            self.assertEqual(row["actualMonth"], 9)
            self.assertIsNotNone(row[block["aggregateField"]])

    def test_default_main_and_structure_months_use_source_observations_per_kind(self):
        self.fixture.edit_book(OKS_SHEET, {"B44": None})
        for kind, expected in (("housing", 9), ("nonresidential", 10), ("total", 10)):
            with self.subTest(kind=kind):
                payload = self.service.report(kind)
                self.assertEqual(payload["selection"]["month"], expected)
                self.assertEqual(payload["selection"]["structure"]["monthTo"], expected)
                for month in range(1, 13):
                    chosen = self.service.report(kind, month, structure_month_to=month)
                    self.assertEqual(chosen["selection"]["month"], month)
                    self.assertEqual(chosen["selection"]["structure"]["monthTo"], month)

    def test_default_structure_ignores_future_registry_months_but_keeps_explicit_december(self):
        self.fixture.write_book()
        payload = self.service.report("total")
        self.assertEqual(payload["selection"]["month"], 10)
        self.assertEqual(payload["selection"]["structure"]["monthTo"], 10)
        self.assertEqual(self.service.report("total", 12, structure_month_to=12)["selection"]["structure"]["monthTo"], 12)

    def test_default_months_distinguish_missing_measurements_from_observed_zero(self):
        self.fixture.edit_book(OKS_SHEET, {"B44": None, "Q45": None})
        for kind in ("housing", "nonresidential", "total"):
            selection = self.service.report(kind)["selection"]
            self.assertEqual(selection["month"], 9)
            self.assertEqual(selection["structure"]["monthTo"], 9)
        self.fixture.edit_book(OKS_SHEET, {"Q45": 0})
        selection = self.service.report("nonresidential")["selection"]
        self.assertEqual(selection["month"], 10)
        self.assertEqual(selection["structure"]["monthTo"], 10)

    def test_api_and_export_share_default_months_and_preserve_explicit_periods(self):
        self.fixture.edit_book(OKS_SHEET, {"B44": None})
        client = self.fixture.client()
        url = "/api/v1/construction/permits"
        for month in (None, 12):
            params = {"permit_kind": "housing"}
            if month is not None:
                params.update(month=month, structure_month_to=month)
            payload = client.get(url, params=params).json()
            expected = 9 if month is None else month
            self.assertEqual(payload["selection"]["month"], expected)
            self.assertEqual(payload["selection"]["structure"]["monthTo"], expected)
            response = client.get(url + "/export", params={**params, "required_version": payload["version"]})
            self.assertEqual(response.status_code, 200)
            book = load_workbook(BytesIO(response.content))
            selected = dict(book["Selection"].values)
            self.assertEqual(selected["month"], expected)
            self.assertEqual(selected["structure.monthTo"], expected)

    def test_default_uses_observed_month_without_filename_date_and_after_year_rollover(self):
        renamed = self.fixture.book_path.with_name("monitoring_2_0_unknown.xlsx")
        self.fixture.book_path.rename(renamed)
        self.assertEqual(self.service.report("housing")["selection"]["month"], 10)
        self.assertEqual(self.service.report("housing")["selection"]["structure"]["monthTo"], 10)
        renamed.rename(renamed.with_name("monitoring_2_0_20270105.xlsx"))
        payload = self.service.report("housing")
        self.assertEqual(payload["selection"]["month"], 10)
        self.assertEqual(payload["selection"]["structure"]["monthTo"], 10)

    def test_main_and_structure_share_approved_controls_without_rescaling_mop(self):
        for kind, field, block in (("housing", "living", "housing_components"),
                                   ("nonresidential", "nonresidential", "nonresidential_components"),
                                   ("total", "total", "area_composition")):
            payload = self.service.report(kind, 12, structure_mode="annual")
            for year in (2011, 2018, 2021, 2025, 2026):
                self.assertAlmostEqual(self.main(payload, year)["selected"], self.block(payload, block, year)[field])
        payload = self.service.report("total", 12, structure_mode="annual")
        old = self.block(payload, "housing_components", 2011)
        self.assertEqual(old["living"], 9999.123456789)
        self.assertAlmostEqual(old["mop_embedded"], 12 * (61021.2 - 41463.8) / 1000)
        self.assertAlmostEqual(old["source_living"], sum((41463.8 + month - 1) / 1000 for month in range(1, 13)))
        self.assertNotEqual(old["housing_total"], old["source_housing_total"])

    def test_approved_static_precedes_workbook_in_both_views(self):
        directory = self.fixture.root / "data/raw/realty/vvod"
        directory.mkdir(parents=True)
        path = directory / "static_vvod_rs_2011_2025.xlsx"
        book = Workbook()
        book.active.append(["year", "permits", "input", "", "year", "permits", "input"])
        book.active.append([2011, 1234.5, 0, "", 2011, 4567.8, 0])
        book.save(path)
        book.close()
        self.fixture.touch(path)
        for kind, block, field, value in (("housing", "housing_components", "living", 1234.5),
                                           ("nonresidential", "nonresidential_components", "nonresidential", 4567.8)):
            payload = self.service.report(kind, 12, structure_mode="annual")
            self.assertEqual(self.main(payload, 2011)["selected"], value)
            self.assertEqual(self.block(payload, block, 2011)[field], value)

    def test_total_contains_mop_in_selected_fact_and_forecast(self):
        payload = self.service.report("total", 5, structure_month_to=5)
        self.assertAlmostEqual(self.main(payload, 2026)["selected"], sum((140000 + 2 * month) / 1000 for month in range(1, 6)))
        self.assertEqual(self.main(payload, 2026)["annualForecast"], 10017.24138 + 6850.987654321)
        self.fixture.edit_book(FORECAST_SHEET, {"D4": None, "D5": None})
        self.assertIsNone(self.main(self.service.report("total", 5), 2026)["annualForecast"])

    def test_missing_historic_mop_never_invented_from_living_annual_control(self):
        self.fixture.edit_book(HOUSING_SHEET, {"C5": None})
        housing = self.service.report("housing", 12, structure_mode="annual")
        self.assertEqual(self.main(housing, 2011)["selected"], 9999.123456789)
        self.assertIsNone(self.block(housing, "housing_components", 2011)["mop_embedded"])
        total = self.service.report("total", 12, structure_mode="annual")
        self.assertIsNone(self.main(total, 2011)["selected"])
        self.assertIsNone(self.block(total, "area_composition", 2011)["total"])

    def test_latest_living_month_survives_missing_general_area_in_both_structure_modes(self):
        self.fixture.edit_book(OKS_SHEET, {"Q44": None})
        for mode in ("annual", "period"):
            with self.subTest(mode=mode):
                payload = self.service.report("housing", 12, structure_mode=mode, structure_month_to=12)
                main = self.main(payload, 2026)
                structure = self.block(payload, "housing_components", 2026)
                self.assertEqual(main["actualMonth"], 10)
                self.assertEqual(structure["actualMonth"], 10)
                self.assertEqual(structure["living"], main["selected"])
                self.assertAlmostEqual(structure["living"], sum((60000 + month) / 1000 for month in range(1, 11)))
                self.assertIsNone(structure["mop_embedded"])
                self.assertIsNone(structure["housing_total"])

    def test_one_nonresidential_chart_reconciles_approved_controls_through_other(self):
        payload = self.service.report("nonresidential", 12, structure_mode="annual")
        self.assertEqual([block["id"] for block in payload["breakdowns"]], ["nonresidential_components"])
        block = payload["breakdowns"][0]
        self.assertEqual([row["year"] for row in block["rows"]], [*range(2011, 2022), 2025, 2026])
        old, modern = self.block(payload, block["id"], 2011), self.block(payload, block["id"], 2026)
        self.assertIsNone(old["undetailed"])
        self.assertAlmostEqual(old["offices"], 234.535 * 12)
        self.assertAlmostEqual(old["source_offices"], 234.535 * 12)
        self.assertAlmostEqual(old["source_other"], 79.811 * 12)
        self.assertAlmostEqual(old["other_adjustment"], 7777.123456789 - 327.442 * 12)
        self.assertAlmostEqual(old["other"], old["source_other"] + old["other_adjustment"])
        self.assertFalse(any(column["label"].startswith("Минстрой:") for column in block["columns"]))
        self.assertEqual(old["trade"], 0)
        self.assertAlmostEqual(modern["trade"], sum((40000 + month) / 1000 for month in range(1, 11, 2)))
        self.assertAlmostEqual(modern["offices"], sum((40000 + month) / 1000 for month in range(2, 11, 2)))
        for index, row in enumerate(block["rows"]):
            values = [series["points"][index]["y"] for series in block["series"]]
            self.assertAlmostEqual(sum(value for value in values if value is not None), row["nonresidential"])

    def test_annual_other_adjustment_is_signed_and_does_not_change_periods_or_other_categories(self):
        source_total, source_other = 327.442 * 12, 79.811 * 12
        for difference in (100.25, -100.25, -1000.25, 0):
            with self.subTest(difference=difference):
                self.fixture.edit_book(HISTORY_SHEET, {"C8": source_total + difference})
                for mode in ("annual", "period"):
                    payload = self.service.report("nonresidential", 12, structure_mode=mode, structure_month_to=12)
                    row = self.block(payload, "nonresidential_components", 2011)
                    self.assertAlmostEqual(row["nonresidential"], source_total + difference)
                    self.assertAlmostEqual(row["other"], source_other + difference)
                    self.assertAlmostEqual(row["other_adjustment"], difference)
                    self.assertAlmostEqual(row["source_other"], source_other)
                    self.assertAlmostEqual(row["offices"], row["source_offices"])
                    self.assertIsNone(row["undetailed"])
                    self.assertAlmostEqual(sum(series["points"][0]["y"] or 0 for series in payload["breakdowns"][0]["series"]), row["nonresidential"])
                    columns = {column["id"] for column in payload["breakdowns"][0]["columns"]}
                    self.assertIn("other_adjustment", columns)
                    self.assertIn("source_other", columns)
                period = self.service.report("nonresidential", 10, structure_month_to=10)
                row = self.block(period, "nonresidential_components", 2011)
                self.assertAlmostEqual(row["other"], 79.811 * 10)
                self.assertIsNone(row["other_adjustment"])

    def test_modern_osm_other_and_signed_difference_are_retained(self):
        self.fixture.edit_book(OKS_SHEET, {"AR3": "12 Коммунально-складские", "AR5": "17 Специальные объекты",
                                           "AT3": "Коммерческие", "AT5": "Административные"})
        source_total = sum((40000 + month) / 1000 for month in range(1, 13))
        self.fixture.edit_book(HISTORY_SHEET, {"B10": 2025, "C10": source_total - 5})
        payload = self.service.report("nonresidential", 12, structure_mode="annual")
        block = payload["breakdowns"][0]
        row = self.block(payload, block["id"], 2025)
        self.assertAlmostEqual(row["source_other"], (40001 + 40002) / 1000)
        self.assertAlmostEqual(row["other"], row["source_other"] - 5)
        self.assertAlmostEqual(row["other_adjustment"], -5)
        self.assertEqual([column["id"] for column in block["columns"] if column["label"] == "Прочие"], ["other"])
        self.assertFalse(any(column["label"].startswith("Минстрой:") for column in block["columns"]))

    def test_osm_codes_are_exact_and_all_other_codes_are_other(self):
        expected = {**{code: "social" for code in range(2, 10)},
                    10: "trade", 11: "hotels", 13: "industry", 14: "offices"}
        for code in range(21):
            for value in (code, float(code), str(code), f"{code:02}", f"  {code:02} Название из источника  "):
                with self.subTest(value=value):
                    self.assertEqual(_osm_group(value), expected.get(code, "other"))
        for value in (None, "", "Социальные", "ОСМ 02", "140", "02.5", "02x", 2.5,
                      True, False, float("nan"), float("inf")):
            with self.subTest(value=value):
                self.assertEqual(_osm_group(value), "other")

    def test_modern_categories_use_only_ar_and_preserve_totals_and_historical_rows(self):
        before = self.service.report("nonresidential", 12, structure_mode="annual")
        codes = [*range(2, 12), 13, 14]
        edits = {"AT1": "Неприменимый заголовок ФНО", "AU1": "Другой заголовок"}
        for month, code in enumerate(codes, 1):
            edits.update({f"AR{month * 2 + 1}": f"{code:02} Противоречащее названию описание",
                          f"AS{month * 2 + 1}": "Коммерческие", f"AT{month * 2 + 1}": "Другие",
                          f"AU{month * 2 + 1}": "Административно-деловые"})
        self.fixture.edit_book(OKS_SHEET, edits)
        after = self.service.report("nonresidential", 12, structure_mode="annual")
        row = self.block(after, "nonresidential_components", 2025)
        for field, _ in HISTORICAL_FIELDS:
            expected = sum((40000 + month) / 1000 for month, code in enumerate(codes, 1) if _osm_group(code) == field)
            self.assertAlmostEqual(row[field], expected)
        self.assertEqual(row["nonresidential"], self.main(before, 2025)["selected"])
        self.assertEqual(self.main(after, 2025), self.main(before, 2025))
        self.assertEqual(self.block(after, "nonresidential_components", 2011),
                         self.block(before, "nonresidential_components", 2011))
        block = after["breakdowns"][0]
        self.assertEqual([series["id"] for series in block["series"]], [field for field, _ in HISTORICAL_FIELDS] + ["undetailed"])
        coverage = next(item for item in after["controls"]["structure"]["coverage"] if item["year"] == 2025)
        self.assertTrue(any("D/F/Q/R/AR" in note for note in coverage["evidence"]))
        self.assertFalse(any("/AT" in note for note in coverage["evidence"]))

    def test_unavailable_osm_header_uses_other_without_fno_fallback(self):
        self.fixture.edit_book(OKS_SHEET, {"AR1": "Не ОСМ код"})
        payload = self.service.report("nonresidential", 12, structure_mode="annual")
        row = self.block(payload, "nonresidential_components", 2025)
        self.assertEqual(row["other"], row["nonresidential"])
        self.assertTrue(all(row[field] == 0 for field, _ in HISTORICAL_FIELDS if field != "other"))
        self.assertTrue(any("Заголовок AR" in note for note in payload["breakdowns"][0]["notes"]))

    def test_incomplete_annual_categories_remain_undetailed_instead_of_fabricating_other(self):
        self.fixture.edit_book(HISTORY_SHEET, {"J115": None})
        payload = self.service.report("nonresidential", 12, structure_mode="annual")
        row = self.block(payload, "nonresidential_components", 2011)
        self.assertEqual(row["undetailed"], 7777.123456789)
        self.assertIsNone(row["other_adjustment"])
        self.assertIsNone(row["other"])

    def test_periods_use_category_total_not_incompatible_auxiliary_nonresidential(self):
        payload = self.service.report("nonresidential", 5, structure_month_to=5)
        old = self.main(payload, 2011)
        self.assertAlmostEqual(old["selected"], 327.442 * 5)
        self.assertEqual(old["selected"], self.block(payload, "nonresidential_components", 2011)["nonresidential"])
        self.assertNotAlmostEqual(old["selected"], sum((348192.6 + month - 1) / 1000 for month in range(1, 6)))


if __name__ == "__main__":
    unittest.main()
