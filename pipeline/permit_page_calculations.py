"""Supplemental permit series from cached, explicitly identified source cells."""
from __future__ import annotations

from calendar import monthrange
from datetime import date, datetime
import math
import re

import pandas as pd
from openpyxl.utils import get_column_letter

from pipeline.operational_calculations import MONTH_LABELS


HOUSING_SHEET = "Вспомогат1"
HISTORY_SHEET = "Данные с 2011 года"
OKS_SHEET = "Реестр ОКС"
FORECAST_SHEET = "Прогнозы"
FORECAST_FIELDS = ("housing_living", "housing_total", "mop_embedded", "nonresidential",
                   "offices", "industry", "social", "trade", "hotels", "other")
HISTORICAL_FIELDS = (("offices", "Административно-деловые"), ("industry", "Промышленные"),
                     ("trade", "Торговые"), ("social", "Социальные"),
                     ("hotels", "Гостиницы"), ("other", "Прочие"))
OSM_GROUPS = {**{code: "social" for code in range(2, 10)},
              10: "trade", 11: "hotels", 13: "industry", 14: "offices"}


def normalized(value):
    return " ".join(str(value or "").split()).casefold().replace("ё", "е")


def number(value):
    if isinstance(value, str):
        value = re.sub(r"\s", "", value).replace(",", ".")
    if isinstance(value, bool) or value is None or value == "":
        return None
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return result if math.isfinite(result) else None


def strict_sum(values):
    values = list(values)
    return number(sum(values)) if values and all(value is not None for value in values) else None


def _osm_group(value):
    if isinstance(value, str):
        match = re.fullmatch(r"(\d{1,2})(?:\s+.*)?", value.strip())
        code = int(match[1]) if match else None
    else:
        numeric = number(value)
        code = int(numeric) if numeric is not None and numeric == int(numeric) else None
    return OSM_GROUPS.get(code, "other")


def _cell_number(cell, notes, scale=1):
    value = number(cell.value)
    if value is None and cell.value is not None and str(cell.value).strip():
        notes.append(f"{cell.coordinate}: значение не является конечным числом; показано null.")
    return value / scale if value is not None else None


def read_forecast(book, source_file):
    notes = ["Единица в строках 1–12 источника не указана. Рабочее предположение: тыс. м², "
             "согласованное с масштабом прежнего графика; единица не подтверждена источником.",
             "Годовой прогноз задаёт годовой ориентир, но не увеличивает факт и не входит в темпы роста. "
             "Использованы только сохранённые значения ячеек, без пересчёта формул и коэффициентов."]
    result = {"unit": "тыс. м²", "unitBasis": "assumed_thousand_m2",
              "source": {"file": source_file, "sheet": FORECAST_SHEET}, "rows": [], "notes": notes}
    if FORECAST_SHEET not in book.sheetnames:
        notes.append(f"Лист «{FORECAST_SHEET}» отсутствует; прогноз недоступен.")
        return result
    cells = list(book[FORECAST_SHEET].iter_rows(min_row=2, max_row=12, min_col=4, max_col=13))
    if not cells:
        notes.append("Прогнозы!D2:M12: диапазон отсутствует; прогноз недоступен.")
        return result
    if len(cells) < 11:
        notes.append("Прогнозы!D2:M12: диапазон неполон; отсутствующие значения оставлены null.")
    seen = set()
    for column in range(10):
        letter = get_column_letter(column + 4)
        year = number(cells[0][column].value)
        if year is None or year != int(year) or not 1900 <= year <= 2200 or year in seen:
            notes.append(f"{letter}2: отсутствующий или неоднозначный год; столбец пропущен.")
            continue
        seen.add(year)
        row = {"year": int(year), "cells": {"year": f"{letter}2"}}
        for offset, field in enumerate(FORECAST_FIELDS, start=1):
            cell = cells[offset][column] if offset < len(cells) else None
            row[field] = _cell_number(cell, notes) if cell is not None else None
            row["cells"][field] = f"{letter}{offset + 2}"
        result["rows"].append(row)
    if not result["rows"]:
        notes.append("Нет столбцов прогноза с подтверждённым годом.")
    return result


def select_forecast(forecast, kind):
    # The caller owns this generation; selection never mutates the cached rows.
    return {**forecast, "rows": [{**row, "value": row["housing_living"] if kind == "housing" else
                                  row["nonresidential"] if kind == "nonresidential" else
                                  strict_sum((row["housing_total"] if row["housing_total"] is not None else
                                              strict_sum((row["housing_living"], row["mop_embedded"])),
                                              row["nonresidential"]))}
                                 for row in forecast["rows"]]}


def _gap(total, living, notes, source):
    if total is None or living is None:
        return None
    if total < living:
        notes.append(f"{source}: общая площадь меньше жилой; МОП и встроенное нежильё = null, не ноль.")
        return None
    return total - living


def _housing_history(book, notes):
    if HOUSING_SHEET not in book.sheetnames:
        notes.append(f"Лист «{HOUSING_SHEET}» отсутствует; историческая структура жилья недоступна.")
        return {}
    rows = list(book[HOUSING_SHEET].iter_rows(min_row=1, max_row=146, min_col=2, max_col=4))
    if len(rows) < 3 or normalized(rows[2][1].value) != "общая" or normalized(rows[2][2].value) != "жилая":
        notes.append("Вспомогат1!C3:D3: заголовки ОБЩАЯ/ЖИЛАЯ не подтверждены; историческая структура недоступна.")
        return {}
    result, year = {}, None
    stems = ("янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек")
    for label, total_cell, living_cell in rows[3:]:
        numeric_year = number(label.value)
        if numeric_year is not None and numeric_year == int(numeric_year) and 2011 <= numeric_year <= 2021:
            year = int(numeric_year)
            continue
        month = next((index for index, stem in enumerate(stems, 1) if normalized(label.value).startswith(stem)), None)
        if year is None or month is None:
            continue
        key = (year, month)
        if key in result:
            notes.append(f"{HOUSING_SHEET}!{label.coordinate}: повтор месяца; историческая структура недоступна.")
            return {}
        total = _cell_number(total_cell, notes, 1000)
        living = _cell_number(living_cell, notes, 1000)
        source = f"{HOUSING_SHEET}!C{label.row}:D{label.row}"
        result[key] = {"housing_total": total, "living": living,
                       "mop_embedded": _gap(total, living, notes, source), "source": source}
    return result


def _nonres_history(book, notes):
    if HISTORY_SHEET not in book.sheetnames:
        notes.append(f"Лист «{HISTORY_SHEET}» отсутствует; исторические категории нежилья недоступны.")
        return {}
    rows = list(book[HISTORY_SHEET].iter_rows(min_row=114, max_row=246, min_col=6, max_col=15))
    expected = ("год", "№ месяца", "месяц", "нежилье, всего", "административно- деловые",
                "промышленные", "торговые", "социальные", "гостиницы", "прочие")
    # Whitespace in wrapped labels is presentation, not a different taxonomy.
    clean = lambda value: normalized(value).replace("- ", "-").replace(", ", ",")
    if not rows or [clean(cell.value) for cell in rows[0]] != [clean(value) for value in expected]:
        notes.append("Данные с 2011 года!F114:O114: заголовки исторических категорий не подтверждены; разрез недоступен.")
        return {}
    result = {}
    for cells in rows[1:]:
        year, month = number(cells[0].value), number(cells[1].value)
        if year is None or month is None or year != int(year) or month != int(month):
            continue
        if not 2011 <= year <= 2021 or not 1 <= month <= 12:
            continue
        key = (int(year), int(month))
        if key in result:
            notes.append("Повтор месяца в исторических категориях нежилья; разрез недоступен.")
            return {}
        item = {"nonresidential": _cell_number(cells[3], notes)}
        item.update({field: _cell_number(cell, notes) for (field, _), cell in zip(HISTORICAL_FIELDS, cells[4:])})
        item["source"] = f"{HISTORY_SHEET}!I{cells[0].row}:O{cells[0].row}"
        subtotal = strict_sum(item[field] for field, _ in HISTORICAL_FIELDS)
        if subtotal is not None and item["nonresidential"] is not None and not math.isclose(subtotal, item["nonresidential"], abs_tol=.000001):
            notes.append(f"{item['source']}: сумма категорий отличается от итога; исходные значения не изменены.")
        result[key] = item
    return result


def read_breakdown_sources(book):
    housing_notes, historical_notes, modern_notes = [], [], []
    housing = _housing_history(book, housing_notes)
    nonresidential = _nonres_history(book, historical_notes)
    modern, groups = [], {}
    if OKS_SHEET not in book.sheetnames:
        modern_notes.append("Реестр ОКС отсутствует; структура с 2022 года недоступна.")
    else:
        rows = book[OKS_SHEET].iter_rows(min_row=1, max_col=44)
        headers = next(rows, ())
        expected = {1: "Назначение", 3: "Срок выдачи РС", 5: "Год выдачи",
                    16: "Общая площадь", 17: "Жилая площадь"}
        valid = len(headers) >= 18 and all(normalized(headers[col].value) == normalized(label) for col, label in expected.items())
        classification = len(headers) >= 44 and normalized(headers[43].value) == normalized("ОСМ код")
        if not valid:
            modern_notes.append("Заголовки B/D/F/Q/R реестра ОКС не подтверждены; современная структура недоступна.")
        else:
            for cells in rows:
                purpose = normalized(cells[1].value)
                if purpose not in ("жилье", "нежилье"):
                    continue
                date = pd.to_datetime(cells[3].value, errors="coerce")
                fallback = number(cells[5].value)
                fallback_used = pd.isna(date) and fallback is not None and fallback == int(fallback)
                if fallback_used:
                    date = pd.to_datetime(f"{int(fallback)}-01-01", errors="coerce")
                if pd.isna(date) or not 2022 <= date.year <= datetime.now().year:
                    continue
                source_row = cells[1].row
                general = _cell_number(cells[16], modern_notes, 1000)
                living = _cell_number(cells[17], modern_notes, 1000) if purpose == "жилье" else None
                group = _osm_group(cells[43].value if classification else None)
                if purpose == "нежилье" and not groups:
                    groups.update(HISTORICAL_FIELDS)
                modern.append({"year": date.year, "month": None if fallback_used else date.month,
                               "kind": "housing" if purpose == "жилье" else "nonresidential",
                               "housing_total": general if purpose == "жилье" else None,
                               "living": living, "mop_embedded": _gap(general, living, modern_notes,
                                                                        f"{OKS_SHEET}!Q{source_row}:R{source_row}") if purpose == "жилье" else None,
                               "nonresidential": general if purpose == "нежилье" else None,
                               "group": group, "source_row": source_row, "fallback": fallback_used})
            if not classification:
                modern_notes.append("Заголовок AR «ОСМ код» не подтверждён; нежилые объекты отнесены в «Прочие». "
                                    "Классификация ФНО вместо ОСМ не используется.")
    return {"housing": housing, "nonresidential": nonresidential, "modern": modern, "groups": groups,
            "housing_notes": housing_notes, "historical_notes": historical_notes, "modern_notes": modern_notes}


def _as_date(value):
    try:
        return date.fromisoformat(value) if value else None
    except (TypeError, ValueError):
        return None


def _month_end(year, month):
    return date(year, month, monthrange(year, month)[1])


def _row_refs(rows):
    numbers = sorted({row["source_row"] for row in rows})
    ranges = []
    for value in numbers:
        if ranges and value == ranges[-1][1] + 1:
            ranges[-1][1] = value
        else:
            ranges.append([value, value])
    return ", ".join(str(start) if start == end else f"{start}-{end}" for start, end in ranges)


def _coverage(block_id, monthly, fields, unknown, source_date, evidence):
    snapshot = _as_date(source_date)
    result = []
    for year in sorted({key[0] for key in monthly} | set(unknown)):
        months = sorted(month for y, month in monthly if y == year)
        complete = [month for month in months if all(monthly[(year, month)].get(field) is not None for field in fields)]
        if year in unknown:
            complete = []
        if year >= 2022:
            # A dated registry snapshot is not evidence for the rest of its month/year.
            complete = [month for month in complete if snapshot is not None and _month_end(year, month) <= snapshot]
        through = _month_end(year, months[-1]) if months else None
        if year >= 2022:
            through = min(through, snapshot) if through is not None and snapshot is not None and snapshot.year >= year else None
        refs = list(evidence.get(year, []))
        if year in unknown:
            refs.append("Есть строки с годом F без даты D: месяц неизвестен; интервальные суммы недоступны.")
        if source_date and year >= 2022:
            refs.append(f"Дата сохранённого источника: {source_date}; незавершённый месяц не подтверждён как полный.")
        result.append({"blockId": block_id, "year": year, "months": months, "completeMonths": complete,
                       "annualComplete": complete == list(range(1, 13)),
                       "observedThrough": through.isoformat() if through else None, "evidence": refs})
    return result


def structure_base(data, source_date, annual_sources=None):
    """Index literal months once; request views only select/sum this immutable base."""
    housing_fields = (("living", "Жилая площадь"), ("mop_embedded", "МОП и встроенное нежильё"),
                      ("housing_total", "Общая площадь жилья"))
    historical_fields = (("nonresidential", "Нежильё, всего"),) + HISTORICAL_FIELDS
    native_fields = historical_fields
    composition_fields = (("living", "Жилая площадь"), ("mop_embedded", "МОП и встроенное нежильё"),
                          ("nonresidential", "Общая площадь нежилья"), ("total", "Общая площадь, всего"))
    housing, historical, native = dict(data["housing"]), dict(data["nonresidential"]), {}
    housing_evidence, history_evidence, native_evidence = {}, {}, {}
    for target, source in ((housing_evidence, housing), (history_evidence, historical)):
        for (year, _), row in source.items():
            target.setdefault(year, []).append(row["source"])
    unknown_h, unknown_n = set(), set()
    grouped = {}
    for row in data["modern"]:
        if row["month"] is None:
            (unknown_h if row["kind"] == "housing" else unknown_n).add(row["year"])
        else:
            grouped.setdefault((row["year"], row["month"], row["kind"]), []).append(row)
    for (year, month, kind), rows in grouped.items():
        if kind == "housing":
            housing[(year, month)] = {key: strict_sum(row[key] for row in rows) for key, _ in housing_fields}
        else:
            native[(year, month)] = {"nonresidential": strict_sum(row["nonresidential"] for row in rows),
                                    **{group: strict_sum(row["nonresidential"] for row in rows if row["group"] == group)
                                       if any(row["group"] == group for row in rows) else 0 for group in data["groups"]}}
    for year in sorted({row["year"] for row in data["modern"]}):
        for kind, target in (("housing", housing_evidence), ("nonresidential", native_evidence)):
            rows = [row for row in data["modern"] if row["year"] == year and row["kind"] == kind]
            if rows:
                target[year] = [f"{OKS_SHEET}!D/F/Q/R/AR, строки {_row_refs(rows)}; м² / 1000. Статус РС не фильтруется."]
    nvalues = {**historical, **native}
    composition = {}
    for key in housing.keys() | nvalues.keys():
        h, n = housing.get(key, {}), nvalues.get(key, {})
        row = {"living": h.get("living"), "mop_embedded": h.get("mop_embedded"),
               "nonresidential": n.get("nonresidential")}
        row["total"] = strict_sum(row.values())
        composition[key] = row
    combined_evidence = {year: housing_evidence.get(year, []) + history_evidence.get(year, []) + native_evidence.get(year, [])
                         for year in set(housing_evidence) | set(history_evidence) | set(native_evidence)}
    basis = ["Для полных прошедших лет жилая площадь и нежильё берутся из утверждённых годовых контролей, "
             "как на основном графике. МОП и встроенное нежильё сохраняются по исходной разности общей и жилой площади; "
             "общий итог рассчитывается из этих компонентов. Исходные суммы сохраняются отдельно, без масштабирования.",
             "Суммируются только буквальные месяцы выбранного интервала. Пустые значения и неизвестное "
             "покрытие не заменяются нулями. Годовой режим требует подтверждения всех 12 месяцев."]
    specifications = [
        ("area_composition", "Структура общей площади", composition_fields, "total", composition, unknown_h | unknown_n,
         combined_evidence, data["housing_notes"] + data["historical_notes"] + data["modern_notes"] +
         ["Общая структура = жилая площадь + МОП и встроенное нежильё + общая площадь нежилья. "
          "В режиме «Всего» основного графика используется та же база."]),
        ("housing_components", "Структура площади жилья", housing_fields, "housing_total", housing, unknown_h,
         housing_evidence, data["housing_notes"] + data["modern_notes"] +
         ["МОП и встроенное нежильё = общая минус жилая площадь; отдельного источника МОП нет."]),
        ("nonresidential_historical", "Нежильё: исторические категории", historical_fields, "nonresidential", historical,
         set(), history_evidence, data["historical_notes"]),
        ("nonresidential_native", "Нежильё: группы ОСМ", native_fields, "nonresidential", native, unknown_n,
         native_evidence, data["modern_notes"]),
        ("nonresidential_components", "Структура площади нежилья", historical_fields,
         "nonresidential", nvalues, unknown_n, {**history_evidence, **native_evidence},
         data["historical_notes"] + data["modern_notes"] + [
             "До 2022 года используются исторические категории I:O, с 2022 года — ОСМ код из AR реестра ОКС: "
             "02–09 — социальные, 10 — торговые, 14 — административно-деловые, 13 — промышленные, "
             "11 — гостиницы, остальные и отсутствующие коды — прочие. Столбцы ФНО не используются. "
             "Для согласованного полного года разница между итогом и суммой категорий относится в «Прочие» со знаком, "
             "без изменения остальных категорий. Исходные значения и корректировка сохранены отдельно. "
             "Если исходная разбивка неполна, итог остаётся без детализации."])]
    result = {identifier: {"id": identifier, "title": title, "fields": fields, "aggregateField": aggregate,
                         "monthly": monthly, "unknown": unknown, "notes": basis + notes,
                         "coverage": _coverage(identifier, monthly, [key for key, _ in fields], unknown, source_date, evidence)}
            for identifier, title, fields, aggregate, monthly, unknown, evidence, notes in specifications}
    unified_block = result["nonresidential_components"]
    unified_block["fields"] += (("undetailed", "Без детализации"),)
    unified_block["coverage"] = [{**row, "blockId": "nonresidential_components"}
                                  for key in ("nonresidential_historical", "nonresidential_native")
                                  for row in result[key]["coverage"]]
    unified_block["applicable"] = {
        year: {key for key, _ in HISTORICAL_FIELDS} for year in {key[0] for key in nvalues}}
    snapshot = _as_date(source_date)
    if annual_sources is not None and snapshot is not None:
        for identifier, block in result.items():
            kinds = ("housing",) if identifier == "housing_components" else ("housing", "nonresidential") if identifier == "area_composition" else ("nonresidential",)
            years = {year for year, kind in annual_sources["annual"] if kind in kinds and year < snapshot.year}
            present = {row["year"] for row in block["coverage"]}
            for year in years - present:
                refs = [annual_sources["annual"][(year, kind)]["evidence"] for kind in kinds if (year, kind) in annual_sources["annual"]]
                block["coverage"].append({"blockId": identifier, "year": year, "months": [], "completeMonths": [],
                                          "annualComplete": False, "observedThrough": f"{year}-12-31", "evidence": refs})
            block["coverage"].sort(key=lambda row: row["year"])
    return result


def structure_controls(base, kind):
    identifiers = _block_ids(kind)
    coverage = [row for identifier in identifiers for row in base[identifier]["coverage"]]
    return {"modes": [{"id": "period", "label": "По периодам"}, {"id": "annual", "label": "По годам"}],
            "years": sorted({row["year"] for row in coverage}),
            "months": list(range(1, 13)), "coverage": coverage}


def _block_ids(kind):
    return (["area_composition", "housing_components", "nonresidential_components"] if kind == "total"
            else ["housing_components"] if kind == "housing" else ["nonresidential_components"])


def _available_month(year, end, source_date, observed_months=None):
    snapshot = _as_date(source_date)
    if snapshot is None or year != snapshot.year:
        return end
    available = [month for month in observed_months or [] if month <= snapshot.month]
    return min(end, max(available)) if available else None


def _monthly_value(sources, kind, year, start, end):
    if end is None or year in sources["unknown"].get(kind, set()):
        return None
    return strict_sum(sources["monthly"].get(kind, {}).get((year, month), {}).get("value")
                      for month in range(start, end + 1))


def _annual_values(sources, year, source_date):
    snapshot = _as_date(source_date)
    if snapshot is None or year >= snapshot.year:
        return {}, []
    result, evidence = {}, []
    for kind in ("housing", "nonresidential"):
        control = sources["annual"].get((year, kind))
        if control is not None:
            result[kind] = control["value"]
            evidence.append(control["evidence"])
        else:
            result[kind] = _monthly_value(sources, kind, year, 1, 12)
    # A living-area control says nothing about missing common/embedded areas.
    result["mop_embedded"] = _monthly_value(sources, "mop_embedded", year, 1, 12)
    result["housing_total"] = strict_sum((result["housing"], result["mop_embedded"]))
    result["total"] = strict_sum((result["housing_total"], result["nonresidential"]))
    return result, evidence


def breakdowns(base, kind, selection, sources=None, source_date=None):
    start, end = (1, 12) if selection["mode"] == "annual" else (selection["monthFrom"], selection["monthTo"])
    label = "За год" if selection["mode"] == "annual" else MONTH_LABELS[end - 1] if start == 1 else f"Месяцы {start}–{end}"
    result = []
    for identifier in _block_ids(kind):
        block = base[identifier]
        rows, evidence = [], []
        for coverage in block["coverage"]:
            year = coverage["year"]
            if not selection["yearFrom"] <= year <= selection["yearTo"]:
                continue
            # Housing must share the main living-area period even when its
            # common/total component is missing for the latest month.
            period_field = "living" if identifier == "housing_components" else block["aggregateField"]
            actual_month = _available_month(year, end, source_date, [month for (y, month), item in block["monthly"].items()
                                                                    if y == year and item.get(period_field) is not None])
            months = [block["monthly"].get((year, month)) for month in range(start, actual_month + 1)] if actual_month else []
            known = year not in block["unknown"] and all(row is not None for row in months)
            row = {"year": year,
                   **{field: strict_sum(item.get(field) for item in months) if known else None for field, _ in block["fields"]},
                   "periodComplete": set(range(start, end + 1)) <= set(coverage["completeMonths"]),
                   "annualComplete": coverage["annualComplete"], "actualMonth": actual_month}
            approved = {}
            if sources is not None:
                values = {key: _monthly_value(sources, key, year, start, actual_month)
                          for key in ("housing", "mop_embedded", "nonresidential")}
                approved, annual_evidence = _annual_values(sources, year, source_date) if start == 1 and end == 12 else ({}, [])
                if approved:
                    values.update({key: approved.get(key) for key in values})
                    evidence.extend(annual_evidence)
                if identifier in ("housing_components", "area_composition"):
                    row["source_living"] = row["living"]
                    row["living"], row["mop_embedded"] = values["housing"], values["mop_embedded"]
                    if identifier == "housing_components":
                        row["source_housing_total"] = row["housing_total"]
                        row["housing_total"] = strict_sum((row["living"], row["mop_embedded"]))
                if identifier in ("nonresidential_components", "area_composition"):
                    row["source_nonresidential"] = row["nonresidential"]
                    row["nonresidential"] = values["nonresidential"]
                if identifier == "area_composition":
                    row["total"] = strict_sum((row["living"], row["mop_embedded"], row["nonresidential"]))
                if approved:
                    row["annualComplete"] = row[block["aggregateField"]] is not None
                    row["periodComplete"] = row["annualComplete"]
            if identifier == "nonresidential_components":
                applicable = block["applicable"].get(year, set())
                subtotal = strict_sum(row[field] for field in applicable)
                total = row["nonresidential"]
                for field in applicable:
                    row[f"source_{field}"] = row[field]
                row["other_adjustment"] = None
                if total is not None and subtotal is not None and approved.get("nonresidential") is not None:
                    # The agreed annual difference belongs only to the residual
                    # category; preserve source amounts and never clip its sign.
                    row.setdefault("source_other", row.get("other"))
                    row["other_adjustment"] = total - subtotal
                    row["other"] = (row.get("other") or 0) + row["other_adjustment"]
                    subtotal = total
                if total is not None and (subtotal is None or not math.isclose(subtotal, total, abs_tol=1e-6)):
                    for field in applicable:
                        row[field] = None
                    row["undetailed"] = total
                else:
                    row["undetailed"] = None
            rows.append(row)
            evidence.extend(coverage["evidence"])
        raw_fields = sorted({key for row in rows for key in row if key.startswith("source_")})
        source_labels = dict(block["fields"])
        for row in rows:
            for key in raw_fields:
                row.setdefault(key, None)
        adjustment_columns = ([{"id": "other_adjustment", "label": "Годовая корректировка «Прочие»"}]
                              if identifier == "nonresidential_components" else [])
        result.append({"id": identifier, "title": block["title"], "unit": "тыс. м²", "periodLabel": label,
                       "aggregateField": block["aggregateField"],
                       "columns": [{"id": "year", "label": "Год"}] + [{"id": key, "label": name} for key, name in block["fields"]] +
                                  [{"id": key, "label": f"Исходное значение: {source_labels.get(key[7:], key[7:])}"} for key in raw_fields] + adjustment_columns,
                       "rows": rows, "series": [{"id": key, "name": name,
                                                  "points": [{"x": str(row["year"]), "y": row[key]} for row in rows]}
                                                 for key, name in block["fields"] if key != block["aggregateField"]],
                       "notes": block["notes"] + evidence +
                       (["Нет данных с подтверждённым покрытием выбранного периода."] if not rows else [])})
    return result


def read_main_sources(book, detailed, source_file):
    """Shared literal components and approved annual controls for every permit view."""
    monthly = {"housing": {key: {"value": row["living"]} for key, row in detailed["housing"].items()},
               "mop_embedded": {key: {"value": row["mop_embedded"]} for key, row in detailed["housing"].items()},
               "nonresidential": {key: {"value": row["nonresidential"]} for key, row in detailed["nonresidential"].items()}}
    evidence = {"housing": {}, "mop_embedded": {}, "nonresidential": {}}
    for (year, _), row in detailed["housing"].items():
        evidence["housing"].setdefault(year, []).append(f"{source_file}:{row['source']}")
        evidence["mop_embedded"].setdefault(year, []).append(f"{source_file}:{row['source']}")
    for (year, _), row in detailed["nonresidential"].items():
        evidence["nonresidential"].setdefault(year, []).append(f"{source_file}:{row['source']}")
    if HOUSING_SHEET in book.sheetnames:
        year = None
        stems = ("янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек")
        for label, cell in book[HOUSING_SHEET].iter_rows(min_row=1, min_col=6, max_col=7):
            numeric_year = number(label.value)
            if numeric_year is not None and numeric_year == int(numeric_year) and 2011 <= numeric_year <= 2021:
                year = int(numeric_year)
                continue
            month = next((m for m, stem in enumerate(stems, 1) if normalized(label.value).startswith(stem)), None)
            if year is not None and month is not None:
                key = (year, month)
                value = number(cell.value)
                # I:O has the total matching the categories. F:G is a fallback
                # only when that historical series is entirely unavailable.
                if not detailed["nonresidential"]:
                    monthly["nonresidential"][key] = {"value": value / 1000 if value is not None and key not in monthly["nonresidential"] else None}
                    evidence["nonresidential"].setdefault(year, []).append(f"{source_file}:{HOUSING_SHEET}!G{label.row}")
    unknown = {"housing": set(), "mop_embedded": set(), "nonresidential": set()}
    grouped = {}
    for row in detailed["modern"]:
        kind, year, month = row["kind"], row["year"], row["month"]
        if month is None:
            unknown[kind].add(year)
            if kind == "housing":
                unknown["mop_embedded"].add(year)
        else:
            grouped.setdefault((year, month, kind), []).append(row)
    for (year, month, kind), rows in grouped.items():
        field = "living" if kind == "housing" else "nonresidential"
        monthly[kind][(year, month)] = {"value": strict_sum(row[field] for row in rows)}
        if kind == "housing":
            monthly["mop_embedded"][(year, month)] = {"value": strict_sum(row["mop_embedded"] for row in rows)}
    for kind in monthly:
        source_kind = "housing" if kind == "mop_embedded" else kind
        for year in {row["year"] for row in detailed["modern"] if row["kind"] == source_kind}:
            rows = [row for row in detailed["modern"] if row["kind"] == source_kind and row["year"] == year]
            evidence[kind][year] = [f"{source_file}:{OKS_SHEET}!D/F/Q/R, строки {_row_refs(rows)}; "
                                    "исходные площади жилья/МОП/нежилья, без фильтра статуса."]
    annual = {}
    if HISTORY_SHEET in book.sheetnames:
        rows = list(book[HISTORY_SHEET].iter_rows(min_col=2, max_col=3))
        markers = {"выдача рс жилье по годам": "housing", "выдача рс нежилье по годам": "nonresidential"}
        kind = None
        for year_cell, value_cell in rows:
            marker = markers.get(normalized(year_cell.value))
            if marker:
                kind = marker
                continue
            year, value = number(year_cell.value), number(value_cell.value)
            if year is None:
                kind = None
            elif kind and year == int(year) and 2011 <= year <= datetime.now().year and value is not None:
                annual[(int(year), kind)] = {"value": value, "evidence":
                                           f"{source_file}:{HISTORY_SHEET}!{value_cell.coordinate}; годовой контроль."}
    return {"monthly": monthly, "unknown": unknown, "evidence": evidence, "annual": annual}


def read_static_annual(book, source_file):
    result = {}
    for cells in book.worksheets[0].iter_rows(min_row=2, max_col=7):
        for year_col, value_col, kind in ((0, 1, "housing"), (4, 5, "nonresidential")):
            year, value = number(cells[year_col].value), number(cells[value_col].value)
            if year is not None and year == int(year) and 2011 <= year <= 2025 and value is not None:
                result[(int(year), kind)] = {"value": value, "evidence":
                                             f"{source_file}:{book.worksheets[0].title}!{cells[value_col].coordinate}; утверждённый годовой контроль."}
    return result


def main_presentation(table, sources, forecast, kind, month, source_date):
    kinds = ("housing", "mop_embedded", "nonresidential") if kind == "total" else (kind,)
    monthly, evidence = {}, {}
    keys = set().union(*(sources["monthly"][k].keys() for k in kinds))
    for key in keys:
        monthly[key] = {"value": strict_sum(sources["monthly"][k].get(key, {}).get("value") for k in kinds)}
    unknown = set().union(*(sources["unknown"][k] for k in kinds))
    for year in {key[0] for key in keys} | unknown:
        evidence[year] = [ref for k in kinds for ref in sources["evidence"][k].get(year, [])]
    coverage = {row["year"]: row for row in _coverage("main", monthly, ["value"], unknown, source_date, evidence)}
    actual = {int(row["Год"]): row for row in table["rows"]}
    forecasts = {row["year"]: row for row in forecast["rows"] if row["value"] is not None}
    rows = []
    for year in sorted(actual.keys() | forecasts.keys()):
        current = coverage.get(year, {})
        refs = list(current.get("evidence", []))
        annual_values, annual_evidence = _annual_values(sources, year, source_date)
        annual = annual_values.get(kind)
        refs.extend(annual_evidence)
        actual_month = _available_month(year, month, source_date,
                                        [m for (y, m), item in monthly.items() if y == year and item["value"] is not None]) if year in actual else None
        selected = strict_sum(_monthly_value(sources, k, year, 1, actual_month) for k in kinds) if actual_month else None
        if month == 12 and annual is not None:
            selected = annual
        predicted = forecasts.get(year)
        annual_forecast = predicted["value"] if predicted else None
        basis = "forecast" if annual_forecast is not None else "annual_observed" if annual is not None else "none"
        reference = annual_forecast if basis == "forecast" else annual
        difference = number(reference - selected) if reference is not None and selected is not None else None
        if predicted:
            fields = (("housing_total", "nonresidential") if predicted["housing_total"] is not None else
                      ("housing_living", "mop_embedded", "nonresidential")) if kind == "total" else ("housing_living",) if kind == "housing" else ("nonresidential",)
            refs.extend(f"{forecast['source']['file']}:{forecast['source']['sheet']}!{predicted['cells'][field]}; "
                        f"годовой прогноз, {forecast['unitBasis']}." for field in fields)
        rows.append({"year": year, "selected": selected, "annualActual": annual, "annualForecast": annual_forecast,
                     "basis": basis, "difference": difference, "remainder": max(difference, 0) if difference is not None else None,
                     "growth": None,
                     "periodComplete": (month == 12 and annual is not None) or
                                       (selected is not None and set(range(1, month + 1)) <= set(current.get("completeMonths", []))),
                     "annualComplete": annual is not None,
                     "actualMonth": actual_month,
                     "observedThrough": current.get("observedThrough") or (f"{year}-12-31" if annual is not None else None),
                     "evidence": refs})
    by_year = {row["year"]: row for row in rows}
    for row in rows:
        previous = by_year.get(row["year"] - 1, {})
        if row["selected"] is not None and previous.get("selected") not in (None, 0) and row["actualMonth"] == previous.get("actualMonth"):
            row["growth"] = number((row["selected"] / previous["selected"] - 1) * 100)
    return {"rows": rows}


def apply_main_presentation(table, presentation):
    """Keep downloads/legacy chart arrays on the same facts as the visible chart."""
    indexed = {row["year"]: row for row in presentation["rows"]}
    for row in table["rows"]:
        selected = indexed[int(row["Год"])]
        row["С начала года, тыс. м²"] = selected["selected"]
        row["За год, тыс. м²"] = selected["annualActual"]
        row["Изменение, %"] = selected["growth"]
    if table["chart"]:
        rows = [indexed[int(year)] for year in table["chart"]["x"]]
        for field, key in (("period", "selected"), ("totals", "annualActual"), ("growth", "growth"), ("remainder", "remainder")):
            table["chart"][field] = [{"value": row[key]} for row in rows]
    table["presentation"] = presentation
    return table
