"""Read the ready construction tree, without recalculating registry data."""
from __future__ import annotations

import copy
import hashlib
import json
import math
import re
import threading
import time
from datetime import datetime, timezone
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
from zlib import error as DecompressionError

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from backend.profile_service import SourceUnavailable
from pipeline.data_access import DataAccess, DataContext, file_signature


_SOURCE_ERRORS = (OSError, ValueError, TypeError, LookupError, RuntimeError, OverflowError,
                  BadZipFile, InvalidFileException, ParseError, EOFError, DecompressionError)
try:
    from lxml.etree import XMLSyntaxError
except ImportError:
    pass
else:
    _SOURCE_ERRORS += (XMLSyntaxError,)


SHEET = "\u0414\u0435\u0440\u0435\u0432\u043e \u0442\u0435\u043a\u0443\u0449 \u0441\u0442\u0440 \u0432\u043e"
COLUMN = "\u0418\u0442\u043e\u0433"
METRICS = (
    ("total", "\u0412\u0441\u0435\u0433\u043e"),
    ("housing", "\u0416\u0438\u043b\u044c\u0435"),
    ("embedded", "\u041d\u0435\u0436\u0438\u043b\u044c\u0435 \u0432 \u0436\u0438\u043b\u044c\u0435"),
    ("mop", "\u041c\u041e\u041f"),
    ("living", "\u0416\u0438\u043b\u0430\u044f \u043f\u043b\u043e\u0449\u0430\u0434\u044c"),
    ("investors", "\u0418\u043d\u0432\u0435\u0441\u0442\u043e\u0440\u044b"),
    ("renovation", "\u0420\u0435\u043d\u043e\u0432\u0430\u0446\u0438\u044f"),
    ("nonresidential", "\u041d\u0435\u0436\u0438\u043b\u044c\u0435"),
    ("budget", "\u0411\u044e\u0434\u0436\u0435\u0442"),
    ("nonbudget", "\u0412\u043d\u0435\u0431\u044e\u0434\u0436\u0435\u0442"),
)


def _normalize(value):
    return " ".join(value.split()).casefold().replace("\u0451", "\u0435") if isinstance(value, str) else ""


def _metric_value(cell):
    value = cell.value
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if (cell.data_type != "n" or isinstance(value, bool) or not isinstance(value, (int, float))
            or not math.isfinite(value) or value < 0):
        raise SourceUnavailable(f"Invalid construction tree value at {cell.coordinate}")
    return value


def _read_metrics(path):
    # Own the stream even when openpyxl fails before returning a workbook.
    with path.open("rb") as source:
        book = load_workbook(source, read_only=True, data_only=True)
        try:
            if SHEET not in book.sheetnames:
                raise SourceUnavailable("Ready construction tree worksheet is missing")
            sheet = book[SHEET]
            headers = [(cell.row, cell.column) for row in sheet.iter_rows() for cell in row
                       if _normalize(cell.value) == _normalize(COLUMN)]
            if len(headers) != 1 or headers[0][1] < 2:
                raise SourceUnavailable("Construction tree total column is missing or ambiguous")
            header_row, column = headers[0]
            labels = {_normalize(label): identifier for identifier, label in METRICS}
            found = {}
            # Other blocks repeat these labels. Only the column adjacent to the
            # unique total header owns the ready tree, including shifted layouts.
            for label_cell, value_cell in sheet.iter_rows(min_row=header_row + 1,
                                                          min_col=column - 1, max_col=column):
                identifier = labels.get(_normalize(label_cell.value))
                if identifier is None:
                    continue
                if identifier in found:
                    raise SourceUnavailable(f"Duplicate construction tree label: {identifier}")
                found[identifier] = {"id": identifier, "label": " ".join(label_cell.value.split()),
                                     "value": _metric_value(value_cell),
                                     "cell": f"{get_column_letter(column)}{label_cell.row}"}
            if len(found) != len(METRICS):
                raise SourceUnavailable("Construction tree metric labels are incomplete")
            if all(metric["value"] is None for metric in found.values()):
                raise SourceUnavailable("Construction tree has no cached values")
            return [found[identifier] for identifier, _ in METRICS]
        finally:
            book.close()


def _file_date(path):
    match = re.fullmatch(r"monitoring_2_0_(\d{8})", path.stem)
    if match:
        try:
            return datetime.strptime(match[1], "%Y%m%d").date().isoformat()
        except ValueError:
            pass
    return None


class ConstructionTreeService:
    def __init__(self, context: DataContext, check_interval: float = 5):
        self.context = context
        self.check_interval = check_interval
        self._lock = threading.RLock()
        self._checked = 0.0
        self._signature = None
        self._payload = None

    def _inputs(self, access):
        files = access.source_files("monitoring_2_0")
        if not files:
            raise SourceUnavailable("Monitoring workbook is missing")
        signatures = [(path, file_signature(path)) for path in files]
        if any(signature[1] is None for _, signature in signatures):
            raise SourceUnavailable("Monitoring source disappeared during discovery")
        latest = max(signature[2] for _, signature in signatures)
        selected = [(path, signature) for path, signature in signatures if signature[2] == latest]
        if len(selected) != 1:
            raise SourceUnavailable("Latest monitoring workbook is ambiguous")
        return selected[0]

    def _refresh(self):
        now = time.monotonic()
        if self._payload is not None and now - self._checked < self.check_interval:
            return
        access = DataAccess(self.context)
        path, signature = self._inputs(access)
        if signature == self._signature:
            self._checked = now
            return
        metrics = _read_metrics(path)
        try:
            name = path.resolve().relative_to(self.context.root).as_posix()
        except ValueError:
            name = path.name
        payload = {
            "schemaVersion": 1,
            "version": hashlib.sha256(repr(signature).encode()).hexdigest(),
            "generatedAt": datetime.now(timezone.utc).isoformat(),
            "unit": "million_m2",
            "source": {"file": name, "sheet": SHEET, "column": COLUMN,
                       "downloadedAt": datetime.fromtimestamp(signature[2] / 1_000_000_000, timezone.utc).isoformat(),
                       "fileDate": _file_date(path)},
            "metrics": metrics,
        }
        json.dumps(payload, allow_nan=False)
        if self._inputs(access) != (path, signature):
            raise SourceUnavailable("Monitoring source changed while reading the tree")
        self._payload, self._signature, self._checked = payload, signature, now

    def report(self):
        with self._lock:
            try:
                self._refresh()
            except _SOURCE_ERRORS as exc:
                # Invalidate throttling as well: subsequent requests must retry
                # the source, never return the preceding successful generation.
                self._payload = self._signature = None
                self._checked = 0.0
                raise SourceUnavailable("Ready construction tree source is unavailable") from exc
            return copy.deepcopy(self._payload)
