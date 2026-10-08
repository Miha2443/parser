"""Actual page-6 execution parity, read-only inputs, HTTP and pinned XLSX."""
import ast
import copy
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import tempfile
import unittest
from unittest.mock import MagicMock, patch

import pandas as pd
import plotly.graph_objects as go
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import create_app
from backend.profile_service import SourceUnavailable
from backend.sales_export import export_workbook
from backend.sales_service import SalesService
from pipeline.data_access import DataAccess, DataContext, MONTH_NAMES_RU, MONTH_SHORT_RU
from pipeline.sales_calculations import KPI_DEFINITIONS, SECTION_TABS, sales_detail

ROOT = Path(__file__).resolve().parents[2]
PAGE = ROOT / "app/pages/6_Распроданность.py"


def finite(value):
    return None if value is None or pd.isna(value) or not math.isfinite(float(value)) else float(value)


class Capture:
    def __init__(self, region, period):
        self.region, self.period = region, period
        self.metrics, self.figures, self.frames = [], [], []
        self.column_config = self

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def columns(self, widths):
        return [self] * (widths if isinstance(widths, int) else len(widths))

    def tabs(self, titles):
        return [self] * len(titles)

    def radio(self, label, options, **kwargs):
        assert self.region in options
        return self.region

    def select_slider(self, label, options, **kwargs):
        return self.period

    def metric(self, label, value, **kwargs):
        self.metrics.append((label, value))

    def plotly_chart(self, figure, **kwargs):
        self.figures.append(figure)

    def dataframe(self, frame, **kwargs):
        self.frames.append(frame.copy())

    def stop(self):
        raise AssertionError("Legacy page stopped")

    def __getattr__(self, name):
        return lambda *args, **kwargs: None


def capture_legacy(data, region, period_index):
    tree = ast.parse(PAGE.read_text(encoding="utf-8"))
    tree.body = [node for node in tree.body if not isinstance(node, (ast.Import, ast.ImportFrom))]
    ui = Capture(region, period_index)
    env = {"st": ui, "pd": pd, "go": go, "load_rasprodannost": lambda: data,
           "MONTH_NAMES_RU": MONTH_NAMES_RU, "MONTH_SHORT_RU": MONTH_SHORT_RU,
           "COLORS": {key: "#333333" for key in ("blue", "red", "green", "teal")},
           "apply_theme": lambda: None, "page_header": lambda *a, **kw: None,
           "style_plotly": lambda *a, **kw: None, "chart_data_expander": lambda *a, **kw: None}
    exec(compile(tree, str(PAGE), "exec"), env)
    return ui, env


def sample_data(count=75):
    kpi, sections = [], {key: [] for key, _ in SECTION_TABS}
    periods = {"msk": [(2020, 1), (2026, 5)], "rf": [(2020, 1), (2025, 6), (2026, 5)]}
    for region, choices in periods.items():
        for year, month in choices:
            # Literal source row order differs from metric order.
            for name, value in [("Объем жилищного строительства", 15680.0), ("Распроданность", 44.0),
                                ("Отношение распроданности к стройготовности ", 105.0), ("Стройготовность", 42.0)]:
                kpi.append({"region_key": region, "year": year, "month": month, "название": name,
                            "значение_num": value, "единица": "6\u00a0953 тыс. м²",
                            "прогноз_2026": "17%", "прогноз_2031+": None, "прогноз_2020": "1 234",
                            "прогноз_2026_num": 17.0, "прогноз_2031+_num": float("nan"), "прогноз_2020_num": 1234.0})
            for key in sections:
                for index in range(count):
                    sections[key].append({"region_key": region, "year": year, "month": month,
                                          "month_name": "raw", "report_period": "raw", "section": key,
                                          "наименование": f" Сегмент {index} ", "Объем жил. строительства": "1 234",
                                          "Распроданность": "7\u00a0%", "Стройготовность": "—",
                                          "Отношение распроданности  к стройготовности": "105%",
                                          "Объем жил. строительства_num": 1234.0})
    return {"kpi": pd.DataFrame(kpi), **{key: pd.DataFrame(rows) for key, rows in sections.items()},
            "regions_available": ["rf", "msk"], "periods_by_region": periods}


def fake_access(data, issues=None):
    access = MagicMock()
    access.load_rasprodannost.return_value = data
    access.data_version.return_value = ("generation-one",)
    access.latest_realty_mart_source_date.return_value = "06.07.2026"
    access.latest_raw_source_date.return_value = ""
    access.source_files.return_value = [ROOT / "data/raw/realty/nashdom/rasprodannost_candidate.xlsx"]
    access.issues = issues or []
    return access


def client_for(service):
    return TestClient(create_app(service=MagicMock(), apartments_service=MagicMock(), sales_service=service))


class LegacyParityMixin:
    def assert_page_parity(self, data, region, period_index, actual):
        ui, env = capture_legacy(data, region, period_index)
        self.assertEqual(len(actual["metrics"]), 4)
        for metric, definition, (label, value) in zip(actual["metrics"], KPI_DEFINITIONS, ui.metrics):
            self.assertEqual(metric["label"], label)
            row = env["kpi_for_period"](definition[1])
            rendered = (f"{env['ru_num'](metric['value'], metric['digits'])} {metric['unit']}"
                        if row is not None else "—")
            self.assertEqual(rendered, value)
        self.assertEqual(len(actual["charts"]), len(ui.figures))
        for chart, figure in zip(actual["charts"], ui.figures):
            self.assertEqual(len(chart["series"]), len(figure.data))
            for series, trace in zip(chart["series"], figure.data):
                self.assertEqual(series["name"], trace.name)
                expected_x = [str(x) if chart["id"] == "forecast" else pd.Timestamp(x).strftime("%Y-%m") for x in trace.x]
                self.assertEqual([point["x"] for point in series["points"]], expected_x)
                self.assertEqual([point["y"] for point in series["points"]], [finite(value) for value in trace.y])
                self.assertIsNone(trace.yaxis)  # Original traces all use the default shared axis.
        frames = iter(ui.frames)
        self.assertEqual(len(actual["tables"]), 6)
        for table in actual["tables"]:
            if not table["rows"]:
                continue
            frame = next(frames)
            self.assertEqual([column["label"] for column in table["columns"]], frame.columns.tolist())
            expected = [[None if pd.isna(value) else value for value in row] for row in frame.itertuples(index=False, name=None)]
            actual_rows = [[row[column["id"]] for column in table["columns"]] for row in table["rows"]]
            self.assertEqual(actual_rows, expected)
        self.assertIsNone(next(frames, None))


class CalculationTests(LegacyParityMixin, unittest.TestCase):
    def test_exact_page_execution_all_regions_all_fixture_periods_and_no_mutation(self):
        data = sample_data()
        original = copy.deepcopy(data)
        for region, periods in data["periods_by_region"].items():
            for index, (year, month) in enumerate(periods):
                with self.subTest(region=region, year=year, month=month):
                    actual = sales_detail(data, region, year, month)
                    self.assert_page_parity(data, region, index, actual)
                    self.assertEqual([item["id"] for item in actual["metrics"]], ["volume", "sold", "ready", "ratio"])
                    self.assertEqual(actual["charts"][0]["unit"], "")
                    self.assertEqual([s["unit"] for s in actual["charts"][0]["series"]], ["тыс. м²", "%", "%", "%"])
                    self.assertEqual([p["x"] for p in actual["charts"][0]["series"][0]["points"]], ["2026", "2031+", "2020"])
                    self.assertTrue(all(len(table["rows"]) == 75 for table in actual["tables"]))
        for key in ("kpi", *[key for key, _ in SECTION_TABS]):
            pd.testing.assert_frame_equal(data[key], original[key])

    def test_missing_and_nonfinite_values_do_not_become_zero(self):
        data = sample_data()
        data["kpi"].loc[0, "значение_num"] = float("inf")
        data["kpi"].loc[1, "значение_num"] = float("nan")
        data["kpi"].loc[0, "прогноз_2026_num"] = float("-inf")
        data["developers"].loc[0, "Распроданность"] = None
        data["developers"].loc[0, "наименование"] = None
        data["developers"].loc[0, "Стройготовность"] = "—"
        actual = sales_detail(data, "msk", 2020, 1)
        self.assertIsNone(actual["metrics"][0]["value"])
        self.assertIsNone(actual["metrics"][1]["value"])
        self.assertIsNone(actual["charts"][0]["series"][0]["points"][0]["y"])
        self.assertIsNone(actual["tables"][2]["rows"][0]["Распроданность"])
        self.assertIsNone(actual["tables"][2]["rows"][0]["наименование"])
        self.assertEqual(actual["tables"][2]["rows"][0]["Стройготовность"], "—")
        json.dumps(actual, allow_nan=False)

    def test_partial_sections_and_absent_kpi_match_legacy_selection(self):
        data = sample_data()
        data["kpi"] = data["kpi"][~data["kpi"]["название"].str.contains("Стройготовность")]
        data["fed_okruga"] = pd.DataFrame()
        del data["by_class"]
        actual = sales_detail(data, "msk", 2020, 1)
        self.assertIsNone(actual["metrics"][2]["value"])
        self.assertEqual(actual["tables"][0]["rows"], [])
        self.assertEqual(actual["tables"][5]["rows"], [])
        self.assert_page_parity(data, "msk", 0, actual)


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.data = sample_data()
        self.access = fake_access(self.data)
        self.loader = patch("backend.sales_service.DataAccess", return_value=self.access)
        self.loader.start()
        self.addCleanup(self.loader.stop)
        self.service = SalesService(DataContext(ROOT), check_interval=0)
        self.client = client_for(self.service)

    def test_catalog_loads_only_data_and_valid_region_periods(self):
        with patch("backend.sales_service.sales_detail", side_effect=AssertionError("Catalog must not calculate detail")):
            response = self.client.get("/api/v1/sales/catalog")
        self.assertEqual(response.status_code, 200)
        result = response.json()
        self.assertEqual(set(result), {"schemaVersion", "version", "generatedAt", "source", "regions"})
        self.assertEqual([r["id"] for r in result["regions"]], ["msk", "rf"])
        self.assertEqual([p["id"] for p in result["regions"][0]["periods"]], ["2020-01", "2026-05"])
        self.assertEqual(result["source"]["files"], ["data/raw/realty/nashdom/rasprodannost_candidate.xlsx"])
        self.assertEqual(result["source"]["date"], "06.07.2026")
        self.assertEqual(response.headers["cache-control"], "no-store")

    def test_defaults_filters_and_required_export_version(self):
        actual = self.client.get("/api/v1/sales").json()
        self.assertEqual(actual["region"]["id"], "msk")
        self.assertEqual(actual["period"]["id"], "2026-05")
        for params in ({"region": "unknown"}, {"period": "bad"}, {"period": "2026-13"},
                       {"region": "msk", "period": "2025-06"}, {"period": "2030-01"}):
            self.assertEqual(self.client.get("/api/v1/sales", params=params).status_code, 404)
            self.assertEqual(self.client.get("/api/v1/sales/export", params={**params, "required_version": actual["version"]}).status_code, 404)
        self.assertEqual(self.client.get("/api/v1/sales/export").status_code, 422)

    def test_candidate_date_never_prefers_stale_manifest_over_raw(self):
        self.access.latest_raw_source_date.return_value = "08.10.2026"
        self.access.latest_realty_mart_source_date.return_value = "06.07.2026"
        self.assertEqual(self.service.catalog()["source"]["date"], "08.10.2026")
        self.access.latest_realty_mart_source_date.assert_not_called()

    def test_generation_and_response_deepcopy_and_calculation_cache(self):
        with patch("backend.sales_service.sales_detail", wraps=sales_detail) as calculate:
            first = self.service.detail("msk", "2020-01")
            expected = copy.deepcopy(first)
            first["tables"][0]["rows"][0]["наименование"] = "changed"
            first["source"]["files"].clear()
            catalog = self.service.catalog()
            catalog["regions"].clear()
            self.data["kpi"].loc[0, "значение_num"] = 999.0
            self.assertEqual(self.service.detail("msk", "2020-01"), expected)
            self.service.detail("rf", "2020-01")
            self.assertEqual(calculate.call_count, 2)
            self.access.load_rasprodannost.assert_called_once()

    def test_five_second_signature_check_and_version_change(self):
        service = SalesService(DataContext(ROOT))
        with patch("backend.sales_service.time.monotonic", return_value=10):
            old = service.detail()
        checked = self.access.data_version.call_count
        self.access.data_version.return_value = ("generation-two",)
        with patch("backend.sales_service.time.monotonic", return_value=14.99):
            self.assertEqual(service.detail()["version"], old["version"])
        self.assertEqual(self.access.data_version.call_count, checked)
        with patch("backend.sales_service.time.monotonic", return_value=15):
            self.assertNotEqual(service.detail()["version"], old["version"])
        self.assertEqual(self.access.load_rasprodannost.call_count, 2)

    def test_changed_signature_during_load_does_not_publish_and_retries(self):
        self.access.data_version.side_effect = [("one",), ("two",)]
        with self.assertRaises(SourceUnavailable):
            self.service.catalog()
        self.assertIsNone(self.service._metadata)
        self.access.data_version.side_effect = None
        self.assertEqual(self.service.catalog()["schemaVersion"], 1)

    def test_calculation_cache_is_bounded_and_source_null_raw_fields_remain_available(self):
        self.data["developers"].loc[0, "наименование"] = None
        self.data["developers"].loc[0, "Распроданность"] = None
        self.assertIsNone(self.service.detail("msk", "2020-01")["tables"][2]["rows"][0]["наименование"])
        for index in range(31):
            self.service._details[("old", str(index))] = {}
        self.service.detail("rf", "2020-01")
        self.assertEqual(len(self.service._details), 32)

    def test_failed_refresh_does_not_return_stale_payload(self):
        old = self.service.detail()
        self.access.data_version.return_value = ("generation-two",)
        self.access.issues = [("error", "failure")]
        with patch("backend.app.logger"):
            self.assertEqual(self.client.get("/api/v1/sales").status_code, 503)
        self.assertEqual(self.service._metadata["version"], old["version"])
        self.access.issues = []
        self.assertNotEqual(self.service.detail()["version"], old["version"])

    def test_warnings_preserved_errors_and_invalid_schemas_503(self):
        self.access.issues = [("warning", "fallback")]
        self.assertEqual(self.service.catalog()["source"]["issues"], [["warning", "fallback"]])
        malformed = []
        for mutate in (
            lambda d: d.update(kpi=pd.DataFrame()),
            lambda d: d.update(kpi=d["kpi"].drop(columns="значение_num")),
            lambda d: d.update(kpi=d["kpi"].assign(**{"значение_num": "bad"})),
            lambda d: d.update(kpi=d["kpi"].assign(month=13)),
            lambda d: d.update(developers=d["developers"].drop(columns="наименование")),
            lambda d: d.update(regions_available=["unknown"]),
            lambda d: d["periods_by_region"].update(msk=[(2020, 13)]),
            lambda d: d.update(developers=d["developers"].assign(**{"Стройготовность": [[1, 2]] * len(d["developers"])})),
        ):
            item = sample_data()
            mutate(item)
            malformed.append(item)
        for data in malformed:
            self.access.load_rasprodannost.return_value = data
            with self.subTest(columns=data["kpi"].columns.tolist()), patch("backend.app.logger"):
                client = client_for(SalesService(DataContext(ROOT)))
                for path in ("catalog", "", "export"):
                    response = client.get("/api/v1/sales" + ("/" + path if path else ""), params={"required_version": "old"})
                    self.assertEqual(response.status_code, 503)

    def test_excel_full_scope_pinned_version_nulls_and_original_units(self):
        payload = self.service.detail("msk", "2020-01")
        response = self.client.get("/api/v1/sales/export", params={"region": "msk", "period": "2020-01", "required_version": payload["version"]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["x-data-version"], payload["version"])
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book["Отчёт"]["B5"].value, payload["version"])
        self.assertEqual(book["KPI"]["D2"].value, 15680)
        for table in payload["tables"]:
            tab = book[table["title"][:31]]
            self.assertEqual(tab.max_row - 1, len(table["rows"]))
            self.assertEqual([cell.value for cell in tab[1]], [c["label"] for c in table["columns"]])
            self.assertEqual(tab["B2"].value, "1 234")
            self.assertEqual(tab["C2"].value, "7\u00a0%")
        for chart in payload["charts"]:
            tab = book[chart["id"]]
            expected = [[s["id"], s["name"], s["unit"], p["x"], p["y"]] for s in chart["series"] for p in s["points"]]
            self.assertEqual([list(row) for row in tab.iter_rows(min_row=2, values_only=True)], expected)
        self.assertTrue(all(tab.freeze_panes == "A2" for tab in book))

    def test_excel_conflict_prevents_export_and_provides_current_version(self):
        current = self.service.catalog()["version"]
        with patch("backend.sales_export.export_workbook") as exporter:
            response = self.client.get("/api/v1/sales/export", params={"required_version": "old"})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.headers["x-data-version"], current)
        exporter.assert_not_called()

    def test_excel_neutralizes_formula_in_all_text_surfaces(self):
        payload = self.service.detail()
        for table in payload["tables"]:
            for index, prefix in enumerate(("=", "+", "-", "@", " \t=")):
                table["rows"][index]["наименование"] = prefix + 'HYPERLINK("evil")'
                table["rows"][index]["Распроданность"] = None
            table["columns"][0]["label"] = "=evil"
        payload["charts"][0]["series"][0]["name"] = "=evil"
        payload["charts"][0]["series"][0]["points"][0]["x"] = "@evil"
        payload["metrics"][0]["label"] = "+evil"
        payload["source"]["date"] = "-evil"
        book = load_workbook(BytesIO(export_workbook(payload)), data_only=False)
        for tab in book:
            for row in tab:
                self.assertTrue(all(cell.data_type != "f" for cell in row))
        self.assertEqual(book["Девелоперы"]["A2"].value, '=HYPERLINK("evil")')
        self.assertEqual(book["Девелоперы"]["A2"].data_type, "s")
        self.assertIsNone(book["Девелоперы"]["C2"].value)


class RealSourceTests(LegacyParityMixin, unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.context = DataContext(ROOT)
        cls.access = DataAccess(cls.context)
        cls.signature = cls.access.data_version("load_rasprodannost")
        cls.hashes = {entry[0]: hashlib.sha256(Path(entry[0]).read_bytes()).hexdigest()
                      for entry in cls.signature if Path(entry[0]).is_file()}
        cls.data = cls.access.load_rasprodannost()
        if cls.data["kpi"].empty:
            raise AssertionError("Real sales inputs required for integration parity")
        cls.service = SalesService(cls.context)
        cls.client = client_for(cls.service)

    @classmethod
    def tearDownClass(cls):
        if cls.access.data_version("load_rasprodannost") != cls.signature:
            raise AssertionError("Sales source signatures changed")
        for path, expected in cls.hashes.items():
            if hashlib.sha256(Path(path).read_bytes()).hexdigest() != expected:
                raise AssertionError(f"Sales source modified: {path}")

    def test_real_both_regions_multiple_periods_page_execution_exact(self):
        catalog = self.client.get("/api/v1/sales/catalog").json()
        self.assertEqual(catalog["source"]["issues"], [list(issue) for issue in self.access.issues])
        self.assertEqual([r["id"] for r in catalog["regions"]], ["msk", "rf"])
        for region in catalog["regions"]:
            periods = self.data["periods_by_region"][region["id"]]
            self.assertEqual([p["id"] for p in region["periods"]], [f"{year:04d}-{month:02d}" for year, month in periods])
            for index in sorted({0, len(periods) // 2, len(periods) - 13, len(periods) - 1}):
                period = region["periods"][index]["id"]
                with self.subTest(region=region["id"], period=period):
                    response = self.client.get("/api/v1/sales", params={"region": region["id"], "period": period})
                    self.assertEqual(response.status_code, 200)
                    actual = response.json()
                    self.assert_page_parity(self.data, region["id"], index, actual)
                    self.assertEqual(actual["version"], catalog["version"])
                    self.assertEqual(actual["generatedAt"], catalog["generatedAt"])

    def test_real_full_excel_rows_and_series(self):
        for region in ("msk", "rf"):
            payload = self.service.detail(region)
            response = self.client.get("/api/v1/sales/export", params={"region": region, "required_version": payload["version"]})
            self.assertEqual(response.status_code, 200)
            book = load_workbook(BytesIO(response.content), read_only=True)
            for table in payload["tables"]:
                tab = book[table["title"][:31]]
                if table["rows"]:
                    expected = [[row[c["id"]] for c in table["columns"]] for row in table["rows"]]
                    self.assertEqual([list(row) for row in tab.iter_rows(min_row=2, values_only=True)], expected)
            for chart in payload["charts"]:
                self.assertEqual(book[chart["id"]].max_row - 1, sum(len(s["points"]) for s in chart["series"]))
            book.close()

    def test_raw_mode_legacy_page_parity_no_mart_generation(self):
        service = SalesService(DataContext(ROOT, use_marts=False))
        service.catalog()
        for region in ("msk", "rf"):
            periods = service._data["periods_by_region"][region]
            for index in (0, len(periods) - 1):
                year, month = periods[index]
                self.assert_page_parity(service._data, region, index, service.detail(region, f"{year:04d}-{month:02d}"))

    def test_missing_sources_and_required_marts_are_503(self):
        with tempfile.TemporaryDirectory() as directory, patch("backend.app.logger"):
            for context in (DataContext(Path(directory)), DataContext(Path(directory), require_marts=True)):
                client = client_for(SalesService(context))
                for path in ("/api/v1/sales/catalog", "/api/v1/sales", "/api/v1/sales/export?required_version=old"):
                    self.assertEqual(client.get(path).status_code, 503)


if __name__ == "__main__":
    unittest.main()
